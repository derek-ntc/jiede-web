"""Delivery-note layout regression checks, independent of the database."""

import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Table

import delivery_exports as exports


def sample_payload(*, assembly=True, count=17):
    return {
        "document_no": "DN-20260912-001", "customer": "宁波示例机械有限公司",
        "recipient_name": "张师傅", "recipient_phone": "13800000000",
        "address": "浙江省宁波市北仑区示例路 168 号 2 栋收货仓库",
        "updated_at": "2026-09-12", "is_assembly": assembly,
        "assembly_drawing_no": "ASM-100-2026", "assembly_set_quantity": 2,
        "items": [
            {"index": index, "drawing_no": f"JD-2026-{index:03d}",
             "product_name": "传动轴支撑座", "specification": "Φ20 × 120 / 304",
             "unit": "件", "quantity_per_set": 2, "order_quantity": 4,
             "shipped_quantity": 0 if index == 1 else 4,
             "order_no": "SO-20260912", "shipped_at": "2026-09-12",
             "remark": "暂不发货" if index == 1 else "按图加工"}
            for index in range(1, count + 1)
        ],
    }


class DeliveryExportLayoutTests(unittest.TestCase):
    def workbook(self, **kwargs):
        return load_workbook(exports.build_delivery_note_xlsx(sample_payload(**kwargs)))

    def pdf_story(self, **kwargs):
        story = []
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        with patch.object(exports.SimpleDocTemplate, "build", lambda _doc, values: story.extend(values)):
            exports.build_delivery_note_pdf(sample_payload(**kwargs), "STSong-Light")
        return story

    def test_excel_labels_fit_without_wrapping_and_principal_columns_are_wider(self):
        sheet = self.workbook().active
        self.assertGreaterEqual(sheet.column_dimensions["A"].width, 13)
        for column, old_width in (("B", 18), ("C", 18), ("D", 20)):
            self.assertGreater(sheet.column_dimensions[column].width, old_width)
        for coordinate in ("A4", "A6", "A7", "E4", "E6", "E7"):
            self.assertFalse(sheet[coordinate].alignment.wrap_text)
        self.assertEqual(sheet.page_setup.orientation, "portrait")
        self.assertEqual(str(sheet.page_setup.paperSize), sheet.PAPERSIZE_A4)
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.page_setup.fitToHeight, 0)
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)

    def test_excel_wrapped_content_gets_enough_row_height(self):
        payload = sample_payload()
        payload["items"][0]["remark"] = "首行备注\n第二行备注\n第三行备注"
        payload["address"] = "长地址用于检查自动换行后仍可完整阅读" * 5
        sheet = load_workbook(exports.build_delivery_note_xlsx(payload)).active
        self.assertGreaterEqual(sheet.row_dimensions[9].height, 50)
        self.assertGreater(sheet.row_dimensions[6].height, sheet.row_dimensions[4].height)
        self.assertEqual(sheet["H9"].value, "首行备注\n第二行备注\n第三行备注")

    def test_excel_removes_green_and_keeps_quantity_headers_readable(self):
        for assembly in (True, False):
            sheet = self.workbook(assembly=assembly).active
            header = 8 if assembly else 7
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.fill.fgColor.rgb, "0092D050")
            for column in (6, 7):
                cell = sheet.cell(header, column)
                self.assertEqual(cell.fill.fgColor.rgb, "00FFFFFF")
                self.assertEqual(cell.font.color.rgb, "00000000")
                self.assertEqual(cell.border.left.style, "thin")
            self.assertEqual(sheet.cell(header, 2).fill.fgColor.rgb, "005C7280")
            self.assertEqual(sheet.cell(header, 2).font.color.rgb, "00FFFFFF")
            self.assertEqual(sheet.cell(header + 1, 7).value, 0)

    def test_pdf_zero_is_not_blank_and_headers_have_actual_paragraph_colors(self):
        for assembly in (True, False):
            tables = [item for item in self.pdf_story(assembly=assembly) if isinstance(item, Table)]
            detail = next(table for table in tables if len(table._cellvalues[0]) == 8)
            self.assertEqual(detail._cellvalues[1][6].getPlainText(), "0")
            self.assertEqual(detail._cellvalues[0][1].style.textColor.hexval(), "0xffffff")
            self.assertEqual(detail._cellvalues[0][5].style.textColor.hexval(), "0x000000")

    def test_pdf_tables_match_excel_proportions_and_fit_a4(self):
        tables = [item for item in self.pdf_story() if isinstance(item, Table)]
        detail = next(table for table in tables if len(table._cellvalues[0]) == 8)
        metadata = next(table for table in tables if len(table._cellvalues[0]) == 4)
        self.assertLessEqual(sum(detail._colWidths), A4[0] - 16 * 72 / 25.4 + 0.1)
        self.assertAlmostEqual(metadata._colWidths[0], detail._colWidths[0], places=3)
        self.assertAlmostEqual(metadata._colWidths[1], sum(detail._colWidths[1:4]), places=3)
        self.assertAlmostEqual(metadata._colWidths[2], detail._colWidths[4], places=3)
        sheet = self.workbook().active
        excel_widths = [sheet.column_dimensions[column].width * 7 + 5 for column in "ABCDEFGH"]
        for actual, width in zip(detail._colWidths, excel_widths):
            self.assertAlmostEqual(actual / sum(detail._colWidths), width / sum(excel_widths), places=6)
        for table in tables:
            for command in table._bkgrndcmds:
                self.assertNotEqual(command[-1].hexval(), "0x92d050")


if __name__ == "__main__":
    unittest.main()

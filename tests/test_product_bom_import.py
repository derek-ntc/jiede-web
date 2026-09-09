import re
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import app

from openpyxl import Workbook
from werkzeug.datastructures import FileStorage


CUSTOMER = "宁波知了智能科技有限公司"
NOW = "2026-09-07T12:00:00"


def workbook_upload(rows, filename="组装清单.xlsx"):
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(tuple(row))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return FileStorage(
        stream=output,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def legacy_workbook_upload(filename="旧版组装清单.xls"):
    fixture = Path(__file__).parent / "fixtures" / "product_bom_legacy.xls"
    return FileStorage(
        stream=BytesIO(fixture.read_bytes()),
        filename=filename,
        content_type="application/vnd.ms-excel",
    )


def bom_row(row_number, assembly, sku, product_name, drawing_no, unit, quantity):
    return {
        "row_number": row_number,
        "assembly_drawing_no": assembly,
        "sku": sku,
        "product_name": product_name,
        "drawing_no": drawing_no,
        "unit": unit,
        "quantity_per_set": quantity,
    }


SCREENSHOT_ROWS = [
    ("组装件DZ-30",),
    ("物料编码", "物料名称", "规格型号", "单位", "每套数量"),
    ("21.12.015.00074", "后盖", "DZ-30-01-002T2S2", "PCS", 1, None, "组装件DZ-30", "组装件DZ-31"),
    ("21.12.015.00139", "右前面板", "DZ-30-01-006T5", "PCS", 1, None, "组装件DZ-30"),
    (None,),
    ("组装件DZ-31",),
    ("物料编码", "物料名称", "规格型号", "单位", "每套数量"),
    ("21.12.015.00074", "后盖", "DZ-30-01-002T2S2", "PCS", 1, None, "组装件DZ-30", "组装件DZ-31"),
    ("21.12.015.00172", "右前面板", "DZ-30-01-006T6", "PCS", 1, None, None, "组装件DZ-31"),
]


class ProductBomParserTests(unittest.TestCase):
    def test_parser_accepts_a_real_legacy_xls_workbook(self):
        rows, errors = app.parse_product_bom_workbook(legacy_workbook_upload())

        self.assertEqual(errors, [])
        self.assertEqual(
            rows,
            [
                {
                    "row_number": 3,
                    "assembly_drawing_no": "DZ-30",
                    "sku": "MAT-XLS-1",
                    "product_name": "旧版后盖",
                    "drawing_no": "PART-XLS-1",
                    "unit": "PCS",
                    "quantity_per_set": 2,
                }
            ],
        )

    def test_parser_accepts_standalone_products_without_assembly_or_quantity(self):
        rows, errors = app.parse_product_bom_workbook(
            workbook_upload(
                [
                    ("物料编码", "物料名称", "规格型号", "单位"),
                    ("MAT-SINGLE", "单件面板", "PART-SINGLE", "PCS"),
                ],
                filename="单件产品.xlsx",
            )
        )

        self.assertEqual(errors, [])
        self.assertEqual(
            rows,
            [
                {
                    "row_number": 2,
                    "assembly_drawing_no": "",
                    "sku": "MAT-SINGLE",
                    "product_name": "单件面板",
                    "drawing_no": "PART-SINGLE",
                    "unit": "PCS",
                    "quantity_per_set": None,
                }
            ],
        )

    def test_parser_can_switch_between_standalone_and_assembly_sections(self):
        rows, errors = app.parse_product_bom_workbook(
            workbook_upload(
                [
                    ("单件产品",),
                    ("物料编码", "物料名称", "规格型号", "单位"),
                    ("MAT-SHARED", "共用后盖", "PART-SHARED", "PCS"),
                    ("组装件DZ-30",),
                    ("物料编码", "物料名称", "规格型号", "单位", "每套数量"),
                    ("MAT-SHARED", "共用后盖", "PART-SHARED", "PCS", 1),
                    ("单件产品",),
                    ("物料编码", "物料名称", "规格型号", "单位"),
                    ("MAT-SINGLE", "单件面板", "PART-SINGLE", "PCS"),
                ]
            )
        )

        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 3)
        self.assertEqual(
            [row["assembly_drawing_no"] for row in rows],
            ["", "DZ-30", ""],
        )
        self.assertEqual(
            [row["quantity_per_set"] for row in rows],
            [None, 1, None],
        )

    def test_parser_does_not_treat_product_or_auxiliary_text_as_a_new_header(self):
        rows, errors = app.parse_product_bom_workbook(
            workbook_upload(
                [
                    ("组装件DZ-30",),
                    ("物料编码", "物料名称", "规格型号", "单位", "每套数量"),
                    (
                        "MAT-HEADER-WORD",
                        "物料名称",
                        "PART-HEADER-WORD",
                        "PCS",
                        1,
                        "辅助说明：单位",
                    ),
                ]
            )
        )

        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["drawing_no"], "PART-HEADER-WORD")
        self.assertEqual(rows[0]["product_name"], "物料名称")

    def test_parser_reads_repeated_blocks_and_matches_shared_parts_without_using_colors(self):
        rows, errors = app.parse_product_bom_workbook(
            workbook_upload(SCREENSHOT_ROWS)
        )

        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 4)
        self.assertEqual(
            [row["assembly_drawing_no"] for row in rows],
            ["DZ-30", "DZ-30", "DZ-31", "DZ-31"],
        )
        shared = [
            row
            for row in rows
            if row["drawing_no"] == "DZ-30-01-002T2S2"
        ]
        self.assertEqual(
            [(row["assembly_drawing_no"], row["quantity_per_set"]) for row in shared],
            [("DZ-30", 1), ("DZ-31", 1)],
        )

    def test_parser_reports_the_assembly_block_when_required_headers_are_missing(self):
        rows, errors = app.parse_product_bom_workbook(
            workbook_upload(
                [
                    ("组装件DZ-30",),
                    ("物料编码", "物料名称", "规格型号", "每套数量"),
                    ("21.12.015.00074", "后盖", "DZ-30-01-002T2S2", 1),
                ]
            )
        )

        self.assertEqual(rows, [])
        self.assertEqual(errors, ["组装件 DZ-30 缺少必要表头：单位"])

    def test_parser_rejects_non_positive_and_oversized_quantities(self):
        rows, errors = app.parse_product_bom_workbook(
            workbook_upload(
                [
                    ("组装件DZ-30",),
                    ("物料编码", "物料名称", "规格型号", "单位", "每套数量"),
                    ("MAT-1", "零件1", "PART-1", "PCS", 0),
                    ("MAT-2", "零件2", "PART-2", "PCS", 2147483648),
                ]
            )
        )

        self.assertEqual(rows, [])
        self.assertEqual(
            errors,
            [
                "第 3 行每套数量必须是大于 0 的整数",
                "第 4 行每套数量不能超过 2147483647",
            ],
        )

    def test_parser_preserves_visible_zero_padding_for_numeric_product_codes(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(("组装件DZ-30",))
        worksheet.append(("物料编码", "物料名称", "规格型号", "单位", "每套数量"))
        worksheet.append((123, "数字编码零件", 456, "PCS", 1))
        worksheet.append((1200, "科学计数编码零件", 789, "PCS", 1))
        worksheet["A3"].number_format = "000000"
        worksheet["C3"].number_format = "000000"
        worksheet["A4"].number_format = "0.00E+00"
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        upload = FileStorage(stream=output, filename="数字编码.xlsx")

        rows, errors = app.parse_product_bom_workbook(upload)

        self.assertEqual(errors, [])
        self.assertEqual(rows[0]["sku"], "000123")
        self.assertEqual(rows[0]["drawing_no"], "000456")
        self.assertEqual(rows[1]["sku"], "1.20E+03")

    def test_parser_returns_a_validation_error_for_corrupt_xlsx(self):
        upload = FileStorage(
            stream=BytesIO(b"this is not a zip workbook"),
            filename="损坏文件.xlsx",
        )

        rows, errors = app.parse_product_bom_workbook(upload)

        self.assertEqual(rows, [])
        self.assertEqual(
            errors,
            ["Excel 文件无法读取，请确认文件未损坏且格式为 .xlsx 或 .xlsm"],
        )

    def test_parser_returns_a_validation_error_for_lazily_corrupt_worksheet_xml(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(("物料编码", "物料名称", "规格型号", "单位"))
        valid_output = BytesIO()
        workbook.save(valid_output)

        damaged_output = BytesIO()
        with ZipFile(BytesIO(valid_output.getvalue()), "r") as source_archive:
            with ZipFile(damaged_output, "w", ZIP_DEFLATED) as target_archive:
                for member in source_archive.infolist():
                    contents = source_archive.read(member.filename)
                    if member.filename == "xl/worksheets/sheet1.xml":
                        contents = b"<worksheet><sheetData><row>"
                    target_archive.writestr(member, contents)
        damaged_output.seek(0)
        upload = FileStorage(
            stream=damaged_output,
            filename="工作表损坏.xlsx",
        )

        rows, errors = app.parse_product_bom_workbook(upload)

        self.assertEqual(rows, [])
        self.assertEqual(
            errors,
            ["Excel 文件无法读取，请确认文件未损坏且格式为 .xlsx 或 .xlsm"],
        )

    def test_parser_returns_a_validation_error_for_corrupt_xls(self):
        upload = FileStorage(
            stream=BytesIO(b"this is not a legacy Excel workbook"),
            filename="损坏文件.xls",
        )

        rows, errors = app.parse_product_bom_workbook(upload)

        self.assertEqual(rows, [])
        self.assertEqual(
            errors,
            ["Excel 文件无法读取，请确认文件未损坏且格式为 .xls、.xlsx 或 .xlsm"],
        )


class ProductBomDatabaseTestCase(unittest.TestCase):
    def setUp(self):
        self.original_db_path = app.DB_PATH
        self.original_database_ready = app.DATABASE_READY
        self.tmpdir = tempfile.TemporaryDirectory()
        app.DB_PATH = Path(self.tmpdir.name) / "manuals.db"
        app.DATABASE_READY = False
        app.init_db()
        with app.get_db() as conn:
            app.ensure_customer_exists(conn, CUSTOMER, NOW)

    def tearDown(self):
        app.DB_PATH = self.original_db_path
        app.DATABASE_READY = self.original_database_ready
        self.tmpdir.cleanup()


class ProductBomPlanTests(ProductBomDatabaseTestCase):
    def test_database_migration_adds_product_unit_field(self):
        with app.get_db() as conn:
            columns = {
                row["name"] for row in conn.execute("PRAGMA table_info(manuals)")
            }

        self.assertIn("unit", columns)

    def test_plan_deduplicates_shared_part_and_keeps_two_assembly_relationships(self):
        rows = [
            bom_row(
                3,
                "DZ-30",
                "21.12.015.00074",
                "后盖",
                "DZ-30-01-002T2S2",
                "PCS",
                1,
            ),
            bom_row(
                8,
                "DZ-31",
                "21.12.015.00074",
                "后盖",
                "DZ-30-01-002T2S2",
                "PCS",
                1,
            ),
        ]

        with app.get_db() as conn:
            plan = app.build_product_bom_import_plan(conn, CUSTOMER, rows)
            product_count = conn.execute(
                "SELECT COUNT(*) AS count FROM manuals"
            ).fetchone()["count"]

        self.assertEqual(plan["errors"], [])
        self.assertEqual(plan["new_product_count"], 1)
        self.assertEqual(plan["matched_product_count"], 0)
        self.assertEqual(len(plan["products"]), 1)
        self.assertEqual(
            [
                (item["assembly_drawing_no"], item["quantity_per_set"])
                for item in plan["relationships"]
            ],
            [("DZ-30", 1), ("DZ-31", 1)],
        )
        self.assertEqual(product_count, 0)

    def test_plan_keeps_standalone_product_without_creating_a_relationship(self):
        rows = [
            bom_row(2, "", "MAT-SINGLE", "单件面板", "PART-SINGLE", "PCS", None)
        ]

        with app.get_db() as conn:
            plan = app.build_product_bom_import_plan(conn, CUSTOMER, rows)

        self.assertEqual(plan["errors"], [])
        self.assertEqual(plan["new_product_count"], 1)
        self.assertEqual(plan["assembly_count"], 0)
        self.assertEqual(plan["relationships"], [])

    def test_plan_exposes_excel_specification_for_preview(self):
        rows = [
            bom_row(2, "", "MAT-SINGLE", "单件面板", "PART-SINGLE", "PCS", None)
        ]

        with app.get_db() as conn:
            plan = app.build_product_bom_import_plan(conn, CUSTOMER, rows)

        self.assertEqual(plan["errors"], [])
        self.assertEqual(plan["products"][0]["specification"], "PART-SINGLE")

    def test_plan_rejects_conflicting_shared_product_basics(self):
        rows = [
            bom_row(3, "DZ-30", "MAT-1", "后盖", "PART-1", "PCS", 1),
            bom_row(8, "DZ-31", "MAT-2", "后盖", "PART-1", "PCS", 1),
        ]

        with app.get_db() as conn:
            plan = app.build_product_bom_import_plan(conn, CUSTOMER, rows)

        self.assertEqual(
            plan["errors"],
            ["产品图号 PART-1 的物料编码不一致（第 3、8 行）"],
        )

    def test_plan_treats_case_only_basic_field_changes_as_conflicts(self):
        rows = [
            bom_row(3, "DZ-30", "MAT-1", "后盖", "PART-1", "PCS", 1),
            bom_row(8, "DZ-31", "mat-1", "后盖", "part-1", "PCS", 1),
        ]

        with app.get_db() as conn:
            plan = app.build_product_bom_import_plan(conn, CUSTOMER, rows)

        self.assertEqual(
            plan["errors"],
            ["产品图号 part-1 的物料编码不一致（第 3、8 行）"],
        )

    def test_plan_merges_duplicate_relationship_rows_and_reports_sources(self):
        rows = [
            bom_row(3, "DZ-30", "MAT-1", "后盖", "PART-1", "PCS", 1),
            bom_row(4, "dz-30", "MAT-1", "后盖", "part-1", "PCS", 1),
        ]

        with app.get_db() as conn:
            plan = app.build_product_bom_import_plan(conn, CUSTOMER, rows)

        self.assertEqual(plan["errors"], [])
        self.assertEqual(len(plan["relationships"]), 1)
        self.assertEqual(plan["relationships"][0]["source_rows"], [3, 4])
        self.assertEqual(plan["assembly_count"], 1)

    def test_plan_rejects_duplicate_relationship_rows_with_different_quantities(self):
        rows = [
            bom_row(3, "DZ-30", "MAT-1", "后盖", "PART-1", "PCS", 1),
            bom_row(4, "DZ-30", "MAT-1", "后盖", "PART-1", "PCS", 2),
        ]

        with app.get_db() as conn:
            plan = app.build_product_bom_import_plan(conn, CUSTOMER, rows)

        self.assertEqual(
            plan["errors"],
            ["产品图号 PART-1 在组装件 DZ-30 的每套数量不一致（第 3、4 行）"],
        )

    def test_plan_rejects_duplicate_existing_products_for_the_selected_customer(self):
        with app.get_db() as conn:
            for name in ("后盖 A", "后盖 B"):
                conn.execute(
                    """
                    INSERT INTO manuals (
                        drawing_no, product_name, customer, model, category,
                        version, remark, filename, original_filename,
                        created_at, updated_at
                    ) VALUES ('PART-1', ?, ?, '', '', '', '', '', '', ?, ?)
                    """,
                    (name, CUSTOMER, NOW, NOW),
                )
            plan = app.build_product_bom_import_plan(
                conn,
                CUSTOMER,
                [bom_row(3, "DZ-30", "MAT-1", "后盖", "PART-1", "PCS", 1)],
            )

        self.assertEqual(
            plan["errors"],
            [f"产品图号 PART-1 在客户 {CUSTOMER} 下存在多条产品记录，请先处理重复产品"],
        )

    def test_plan_matches_only_the_selected_customers_existing_product_and_link(self):
        with app.get_db() as conn:
            existing_id = conn.execute(
                """
                INSERT INTO manuals (
                    drawing_no, product_name, customer, model, category,
                    version, remark, filename, original_filename,
                    created_at, updated_at
                ) VALUES ('PART-1', '旧名称', ?, '', '', '', '', '', '', ?, ?)
                """,
                (CUSTOMER, NOW, NOW),
            ).lastrowid
            conn.execute(
                """
                INSERT INTO product_assembly_components (
                    manual_id, assembly_drawing_no, quantity_per_set,
                    sort_order, created_at, updated_at
                ) VALUES (?, 'DZ-30', 1, 0, ?, ?)
                """,
                (existing_id, NOW, NOW),
            )
            conn.execute(
                """
                INSERT INTO manuals (
                    drawing_no, product_name, customer, model, category,
                    version, remark, filename, original_filename,
                    created_at, updated_at
                ) VALUES ('PART-1', '其他客户产品', '其他客户', '', '', '', '', '', '', ?, ?)
                """,
                (NOW, NOW),
            )
            plan = app.build_product_bom_import_plan(
                conn,
                CUSTOMER,
                [bom_row(3, "DZ-30", "MAT-1", "新名称", "PART-1", "PCS", 2)],
            )

        self.assertEqual(plan["errors"], [])
        self.assertEqual(plan["new_product_count"], 0)
        self.assertEqual(plan["matched_product_count"], 1)
        self.assertEqual(plan["new_relationship_count"], 0)
        self.assertEqual(plan["updated_relationship_count"], 1)
        self.assertEqual(plan["products"][0]["manual_id"], existing_id)

    def test_plan_rejects_a_customer_that_is_not_in_the_customer_list(self):
        rows = [bom_row(3, "DZ-30", "MAT-1", "后盖", "PART-1", "PCS", 1)]

        with app.get_db() as conn:
            plan = app.build_product_bom_import_plan(conn, "不存在的客户", rows)

        self.assertEqual(plan["errors"], ["客户 不存在的客户 不存在，请先在客户信息中新增"])


class ProductBomApplyTests(ProductBomDatabaseTestCase):
    def test_apply_creates_one_shared_product_with_two_assembly_relationships(self):
        rows = [
            bom_row(3, "DZ-30", "MAT-1", "后盖", "PART-1", "PCS", 1),
            bom_row(8, "DZ-31", "MAT-1", "后盖", "PART-1", "PCS", 1),
        ]

        with app.get_db() as conn:
            result = app.apply_product_bom_import(
                conn, CUSTOMER, rows, "admin", NOW
            )
            products = conn.execute(
                """
                SELECT drawing_no, product_name, customer, sku, unit, supplier,
                       filename, unit_price_minor
                FROM manuals
                """
            ).fetchall()
            relationships = conn.execute(
                """
                SELECT assembly_drawing_no, quantity_per_set
                FROM product_assembly_components
                ORDER BY assembly_drawing_no
                """
            ).fetchall()

        self.assertEqual(result["new_product_count"], 1)
        self.assertEqual(
            [tuple(row) for row in products],
            [("PART-1", "后盖", CUSTOMER, "MAT-1", "PCS", "PART-1", "", None)],
        )
        self.assertEqual(
            [tuple(row) for row in relationships],
            [("DZ-30", 1), ("DZ-31", 1)],
        )

    def test_apply_updates_only_basic_import_fields_and_preserves_existing_data(self):
        with app.get_db() as conn:
            existing_id = conn.execute(
                """
                INSERT INTO manuals (
                    drawing_no, product_name, customer, sku, unit, supplier,
                    model, category, version, remark,
                    filename, original_filename, unit_price_minor, currency,
                    created_at, updated_at
                )
                VALUES (
                    'PART-1', '旧名称', ?, 'OLD-MAT', 'EA', '旧规格',
                    '', '', '', '保留备注',
                    'drawing.pdf', '原图纸.pdf', 1234, 'CNY', ?, ?
                )
                """,
                (CUSTOMER, NOW, NOW),
            ).lastrowid
            conn.execute(
                """
                INSERT INTO product_assembly_components (
                    manual_id, assembly_drawing_no, quantity_per_set,
                    sort_order, created_at, updated_at
                ) VALUES (?, 'OLD-ASM', 7, 0, ?, ?)
                """,
                (existing_id, NOW, NOW),
            )
            result = app.apply_product_bom_import(
                conn,
                CUSTOMER,
                [bom_row(3, "DZ-30", "MAT-NEW", "新名称", "PART-1", "PCS", 2)],
                "admin",
                "2026-09-07T13:00:00",
            )
            product = conn.execute(
                "SELECT * FROM manuals WHERE id = ?", (existing_id,)
            ).fetchone()
            relationships = conn.execute(
                """
                SELECT assembly_drawing_no, quantity_per_set
                FROM product_assembly_components
                WHERE manual_id = ?
                ORDER BY sort_order, id
                """,
                (existing_id,),
            ).fetchall()

        self.assertEqual(result["matched_product_count"], 1)
        self.assertEqual(product["product_name"], "新名称")
        self.assertEqual(product["sku"], "MAT-NEW")
        self.assertEqual(product["unit"], "PCS")
        self.assertEqual(product["supplier"], "PART-1")
        self.assertEqual(product["filename"], "drawing.pdf")
        self.assertEqual(product["original_filename"], "原图纸.pdf")
        self.assertEqual(product["unit_price_minor"], 1234)
        self.assertEqual(product["remark"], "保留备注")
        self.assertEqual(
            [tuple(row) for row in relationships],
            [("OLD-ASM", 7), ("DZ-30", 2)],
        )

    def test_apply_explicit_blank_specification_clears_only_the_imported_field(self):
        with app.get_db() as conn:
            existing_id = conn.execute(
                """
                INSERT INTO manuals (
                    drawing_no, product_name, customer, sku, unit, supplier,
                    model, category, version, remark,
                    filename, original_filename, unit_price_minor, currency,
                    created_at, updated_at
                )
                VALUES (
                    'PART-1', '旧名称', ?, 'OLD-MAT', 'EA', '旧规格',
                    '', '', '', '保留备注',
                    'drawing.pdf', '原图纸.pdf', 1234, 'CNY', ?, ?
                )
                """,
                (CUSTOMER, NOW, NOW),
            ).lastrowid
            conn.execute(
                """
                INSERT INTO manual_files (
                    manual_id, filename, original_filename, file_type, created_at
                ) VALUES (?, 'attachment.pdf', '附件.pdf', 'file', ?)
                """,
                (existing_id, NOW),
            )
            conn.execute(
                """
                INSERT INTO product_materials (
                    manual_id, material, thickness, surface_type, supplier,
                    sort_order, created_at, updated_at
                ) VALUES (?, '铝板', '2mm', '喷涂', '材料供应商', 0, ?, ?)
                """,
                (existing_id, NOW, NOW),
            )
            conn.execute(
                """
                INSERT INTO product_assembly_components (
                    manual_id, assembly_drawing_no, quantity_per_set,
                    sort_order, created_at, updated_at
                ) VALUES (?, 'OLD-ASM', 7, 0, ?, ?)
                """,
                (existing_id, NOW, NOW),
            )
            row = bom_row(3, "", "MAT-NEW", "新名称", "PART-1", "PCS", None)
            row["specification"] = ""

            app.apply_product_bom_import(
                conn, CUSTOMER, [row], "admin", "2026-09-07T13:00:00"
            )
            product = conn.execute(
                "SELECT * FROM manuals WHERE id = ?", (existing_id,)
            ).fetchone()
            file_count = conn.execute(
                "SELECT COUNT(*) FROM manual_files WHERE manual_id = ?", (existing_id,)
            ).fetchone()[0]
            material_count = conn.execute(
                "SELECT COUNT(*) FROM product_materials WHERE manual_id = ?", (existing_id,)
            ).fetchone()[0]
            relationship_count = conn.execute(
                "SELECT COUNT(*) FROM product_assembly_components WHERE manual_id = ?",
                (existing_id,),
            ).fetchone()[0]

        self.assertEqual(product["supplier"], "")
        self.assertEqual(product["filename"], "drawing.pdf")
        self.assertEqual(product["original_filename"], "原图纸.pdf")
        self.assertEqual(product["unit_price_minor"], 1234)
        self.assertEqual(product["remark"], "保留备注")
        self.assertEqual(file_count, 1)
        self.assertEqual(material_count, 1)
        self.assertEqual(relationship_count, 1)


class ProductBomRouteTests(ProductBomDatabaseTestCase):
    def setUp(self):
        super().setUp()
        self.client = app.app.test_client()
        with self.client.session_transaction() as session:
            session["admin_logged_in"] = True
            session["admin_username"] = "admin"
            session["admin_role"] = "admin"
        with app.get_db() as conn:
            app.ensure_customer_exists(conn, CUSTOMER, NOW)

    def test_upload_previews_without_writing_then_confirm_imports(self):
        preview = self.client.post(
            "/admin/products/import",
            data={
                "customer": CUSTOMER,
                "file": workbook_upload(SCREENSHOT_ROWS),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(preview.status_code, 200)
        preview_html = preview.get_data(as_text=True)
        self.assertIn("导入预览", preview_html)
        self.assertIn("新增产品 3 个", preview_html)
        self.assertIn("组装关系 4 条", preview_html)
        self.assertIn("组装件 2 个", preview_html)
        self.assertIn("新增组装关系 4 条", preview_html)
        self.assertIn("更新组装关系 0 条", preview_html)
        self.assertIn("第 3 行", preview_html)
        token_match = re.search(
            r'name="import_token"\s+value="([^"]+)"', preview_html
        )
        self.assertIsNotNone(token_match)
        with app.get_db() as conn:
            self.assertEqual(
                conn.execute("SELECT COUNT(*) AS count FROM manuals").fetchone()[
                    "count"
                ],
                0,
            )

        confirmed = self.client.post(
            "/admin/products/import/confirm",
            data={"import_token": token_match.group(1)},
        )

        self.assertEqual(confirmed.status_code, 302)
        self.assertTrue(confirmed.headers["Location"].endswith("/admin/products"))
        with app.get_db() as conn:
            products = conn.execute(
                "SELECT drawing_no, customer FROM manuals ORDER BY drawing_no"
            ).fetchall()
            relationships = conn.execute(
                """
                SELECT assembly_drawing_no, quantity_per_set
                FROM product_assembly_components
                ORDER BY assembly_drawing_no, manual_id
                """
            ).fetchall()
        self.assertEqual(len(products), 3)
        self.assertTrue(all(row["customer"] == CUSTOMER for row in products))
        self.assertEqual(len(relationships), 4)

    def test_standalone_product_is_visible_in_preview_and_confirmed_without_a_relationship(self):
        preview = self.client.post(
            "/admin/products/import",
            data={
                "customer": CUSTOMER,
                "file": workbook_upload(
                    [
                        ("物料编码", "物料名称", "规格型号", "单位"),
                        ("MAT-SINGLE", "单件面板", "PART-SINGLE", "PCS"),
                    ]
                ),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(preview.status_code, 200)
        preview_html = preview.get_data(as_text=True)
        self.assertIn("产品清单", preview_html)
        self.assertIn("PART-SINGLE", preview_html)
        self.assertIn("组装关系 0 条", preview_html)
        token = re.search(
            r'name="import_token"\s+value="([^"]+)"', preview_html
        ).group(1)

        confirmed = self.client.post(
            "/admin/products/import/confirm",
            data={"import_token": token},
        )

        self.assertEqual(confirmed.status_code, 302)
        with app.get_db() as conn:
            product = conn.execute(
                "SELECT drawing_no FROM manuals WHERE drawing_no = 'PART-SINGLE'"
            ).fetchone()
            relationship_count = conn.execute(
                "SELECT COUNT(*) AS count FROM product_assembly_components"
            ).fetchone()["count"]
        self.assertIsNotNone(product)
        self.assertEqual(relationship_count, 0)

    def test_template_download_contains_standalone_and_shared_assembly_examples(self):
        response = self.client.get("/admin/products/import-template")

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "filename=product-bom-import-template.xlsx",
            response.headers["Content-Disposition"],
        )
        workbook = app.load_workbook(BytesIO(response.data), data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["产品导入模板", "填写说明"])
            values = list(workbook["产品导入模板"].values)
            self.assertEqual(values[0][0], "单件产品")
            self.assertIn(("组装件DZ-30", None, None, None, None), values)
            self.assertIn(("组装件DZ-31", None, None, None, None), values)
            shared_rows = [row for row in values if row[0] == "MAT-SHARED"]
            self.assertEqual(len(shared_rows), 2)
            self.assertTrue(
                all(row[2] == "PART-SHARED" and row[4] == 1 for row in shared_rows)
            )
        finally:
            workbook.close()

    def test_confirm_rejects_tampered_token_without_writing(self):
        response = self.client.post(
            "/admin/products/import/confirm",
            data={"import_token": "tampered"},
            follow_redirects=True,
        )

        self.assertIn("导入确认已失效", response.get_data(as_text=True))
        with app.get_db() as conn:
            product_count = conn.execute(
                "SELECT COUNT(*) AS count FROM manuals"
            ).fetchone()["count"]
        self.assertEqual(product_count, 0)

    def test_confirm_revalidates_database_state_before_writing(self):
        preview = self.client.post(
            "/admin/products/import",
            data={
                "customer": CUSTOMER,
                "file": workbook_upload(
                    [
                        ("组装件DZ-30",),
                        ("物料编码", "物料名称", "规格型号", "单位", "每套数量"),
                        ("MAT-1", "后盖", "PART-1", "PCS", 1),
                    ]
                ),
            },
            content_type="multipart/form-data",
        )
        token = re.search(
            r'name="import_token"\s+value="([^"]+)"',
            preview.get_data(as_text=True),
        ).group(1)
        with app.get_db() as conn:
            for name in ("后盖 A", "后盖 B"):
                conn.execute(
                    """
                    INSERT INTO manuals (
                        drawing_no, product_name, customer, model, category,
                        version, remark, filename, original_filename,
                        created_at, updated_at
                    ) VALUES ('PART-1', ?, ?, '', '', '', '', '', '', ?, ?)
                    """,
                    (name, CUSTOMER, NOW, NOW),
                )

        response = self.client.post(
            "/admin/products/import/confirm",
            data={"import_token": token},
            follow_redirects=True,
        )

        self.assertIn("导入前数据状态已变化", response.get_data(as_text=True))
        with app.get_db() as conn:
            self.assertEqual(
                conn.execute(
                    "SELECT COUNT(*) AS count FROM manuals"
                ).fetchone()["count"],
                2,
            )
            self.assertEqual(
                conn.execute(
                    "SELECT COUNT(*) AS count FROM product_assembly_components"
                ).fetchone()["count"],
                0,
            )


class ProductBomPageTests(ProductBomDatabaseTestCase):
    def setUp(self):
        super().setUp()
        self.client = app.app.test_client()
        with self.client.session_transaction() as session:
            session["admin_logged_in"] = True
            session["admin_username"] = "admin"
            session["admin_role"] = "admin"
        with app.get_db() as conn:
            self.product_id = conn.execute(
                """
                INSERT INTO manuals (
                    drawing_no, product_name, customer, sku, unit,
                    model, category, version, remark, filename,
                    original_filename, created_at, updated_at
                ) VALUES (
                    'PART-UNIT', '单位测试产品', ?, 'MAT-UNIT', 'PCS',
                    '', '', '', '', '', '', ?, ?
                )
                """,
                (CUSTOMER, NOW, NOW),
            ).lastrowid

    def test_product_list_offers_import_and_displays_unit(self):
        response = self.client.get("/admin/products")
        html = response.get_data(as_text=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn("Excel批量导入产品", html)
        self.assertIn("/admin/products/import", html)
        self.assertIn("单位", html)
        self.assertIn("PCS", html)

    def test_import_page_offers_template_and_accepts_all_supported_excel_formats(self):
        response = self.client.get("/admin/products/import")
        html = response.get_data(as_text=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn("/admin/products/import-template", html)
        self.assertIn("下载导入模板", html)
        self.assertIn('accept=".xls,.xlsx,.xlsm"', html)

    def test_product_edit_page_offers_editable_unit(self):
        response = self.client.get(f"/admin/{self.product_id}/edit")
        html = response.get_data(as_text=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn('name="unit"', html)
        self.assertIn('value="PCS"', html)

    def test_product_basic_detail_displays_imported_unit(self):
        response = self.client.get(f"/manual/{self.product_id}")
        html = response.get_data(as_text=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn("<dt>单位</dt><dd>PCS</dd>", html)

    def test_product_create_and_edit_persist_unit(self):
        edited = self.client.post(
            f"/admin/{self.product_id}/edit",
            data={
                "drawing_no": "PART-UNIT",
                "product_name": "单位测试产品",
                "customer": CUSTOMER,
                "sku": "MAT-UNIT",
                "unit": "SET",
                "min_stock": "0",
                "unit_price": "",
                "currency": "CNY",
            },
        )
        created = self.client.post(
            "/admin/upload",
            data={
                "drawing_no": "PART-NEW-UNIT",
                "product_name": "新增单位产品",
                "customer": CUSTOMER,
                "sku": "MAT-NEW-UNIT",
                "unit": "PCS",
                "min_stock": "0",
                "unit_price": "",
                "currency": "CNY",
            },
        )

        self.assertEqual(edited.status_code, 302)
        self.assertEqual(created.status_code, 302)
        with app.get_db() as conn:
            units = {
                row["drawing_no"]: row["unit"]
                for row in conn.execute(
                    "SELECT drawing_no, unit FROM manuals"
                ).fetchall()
            }
        self.assertEqual(units["PART-UNIT"], "SET")
        self.assertEqual(units["PART-NEW-UNIT"], "PCS")

    def test_copying_an_imported_product_keeps_its_unit(self):
        response = self.client.post(f"/admin/{self.product_id}/copy")

        self.assertEqual(response.status_code, 302)
        with app.get_db() as conn:
            copied = conn.execute(
                "SELECT unit FROM manuals WHERE product_name = '单位测试产品 - 副本'"
            ).fetchone()
        self.assertIsNotNone(copied)
        self.assertEqual(copied["unit"], "PCS")


if __name__ == "__main__":
    unittest.main()

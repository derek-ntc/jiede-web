"""Template layout and server-provided supplier display contracts."""
from datetime import datetime
from io import BytesIO
import json
import os
import re
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

import app
from procurement_documents import build_purchase_order_workbook, purchase_document_view
from tests.test_purchase_order_exports import sample_order, sample_items
from tests import test_purchase_orders as order_tests
from tests.test_purchase_orders import payload, form


class PurchaseTemplateLayoutTests(unittest.TestCase):
    @unittest.skipUnless(os.environ.get("PURCHASE_TEST_SOFFICE") and os.environ.get("PURCHASE_TEST_PDF_READER"), "optional Excel print integration requires local renderer and PDF reader")
    def test_three_line_orders_print_complete_footer_on_one_page(self):
        for category in ("raw_material", "carton", "outsourcing"):
            with self.subTest(category=category), tempfile.TemporaryDirectory() as folder:
                order = sample_order(category)
                order["created_by"] = "preview"
                items = sample_items(3)
                items[1]["remark"] = "本行备注较长，用于验证自动换行和打印范围，尺寸需要复核，表面不得划伤。" * 2
                source = Path(folder) / "order.xlsx"
                source.write_bytes(build_purchase_order_workbook(order, items, include_prices=True).getvalue())
                subprocess.run([os.environ["PURCHASE_TEST_SOFFICE"], f"-env:UserInstallation={(Path(folder) / 'profile').as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", folder, str(source)], check=True, capture_output=True)
                output = subprocess.run([os.environ["PURCHASE_TEST_PDF_READER"], str(source.with_suffix(".pdf"))], check=True, capture_output=True, text=True)
                pages = json.loads(output.stdout)
                self.assertEqual(len(pages), 1, "short order must not leave a footer-only second page")
                compact = "".join(pages[0]["text"].split())
                for token in ("确认回传", "preview", "buy@example.com", "精密支架003", "000042", "0013800000000", "2026-09-09", "2026-09-20"):
                    self.assertIn(token, compact)

    def test_template_has_company_first_complete_raw_columns_and_signed_footer(self):
        order = sample_order("raw_material")
        order["created_by"] = "历史经办人"
        items = sample_items()
        items[0].update(item_name="镀锌板", unit="张")
        sheet = load_workbook(build_purchase_order_workbook(order, items, include_prices=False)).active
        self.assertEqual(sheet["A1"].value, order["company_profile"]["company_name"])
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(str(sheet.page_setup.paperSize), str(sheet.PAPERSIZE_A4))
        self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight), (1, 0))
        header = next(row for row in sheet if any(c.value == "物品名称" for c in row))
        self.assertEqual([c.value for c in header], ["物品名称", "材质", "长 mm", "宽 mm", "厚度 mm", "表面", "数量", "单位", "预计到货日期", "备注"])
        self.assertEqual(sheet.print_title_rows.replace("$", ""), f"{header[0].row}:{header[0].row}")
        self.assertIn("$A$1:", sheet.print_area)
        values = [c.value for row in sheet for c in row]
        for value in ("镀锌板", "张", "000042", "0574-01234567", "历史经办人", "确认回传"):
            self.assertIn(value, values)
        self.assertIn(datetime(2026, 9, 9), values)
        self.assertIn(datetime(2026, 9, 20), values)
        self.assertTrue(any(c.value == 1200.5 and c.data_type == "n" for row in sheet for c in row))
        footer_row = next(c.row for row in sheet for c in row if c.value == "确认回传")
        self.assertGreater(footer_row, header[0].row + 1)
        self.assertIn(f"$J${sheet.max_row}", sheet.print_area)

    def test_every_category_keeps_name_unit_and_price_projection(self):
        for category in ("raw_material", "carton", "outsourcing", "other"):
            for prices in (True, False):
                with self.subTest(category=category, prices=prices):
                    stream = build_purchase_order_workbook(sample_order(category), sample_items(), include_prices=prices)
                    sheet = load_workbook(stream).active
                    text = "".join(str(c.value or "") for row in sheet for c in row).replace("\n", "")
                    self.assertIn("精密支架001", text)
                    self.assertIn("单位", text)
                    self.assertIn("件", text)
                    if not prices:
                        with ZipFile(stream) as archive:
                            contents = b"".join(archive.read(n) for n in archive.namelist()).decode()
                        for forbidden in ("unit_price", "单价", "金额", "87654321", "262962963"):
                            self.assertNotIn(forbidden, contents)

    def test_structured_model_uses_only_snapshot_and_long_text_is_not_clipped(self):
        order = sample_order("raw_material")
        order["created_by"] = "旧经办人"
        model = purchase_document_view(order, sample_items(), include_prices=False)
        self.assertEqual(model.get("supplier_fields"), {"name": order["supplier_name"], "code": "000042", "contact": "张工", "phone": "0574-01234567", "email": "supplier@example.com", "address": order["supplier_address"]})
        self.assertEqual(model.get("created_by"), "旧经办人")
        items = sample_items(45)
        items[0]["remark"] = "长备注" * 1300 + "结束标记"
        items[0]["item_name"] = "=1+2"
        sheet = load_workbook(build_purchase_order_workbook(order, items, include_prices=False)).active
        text = "".join(str(c.value or "") for row in sheet for c in row).replace("\n", "")
        self.assertIn(items[0]["remark"], text)
        self.assertIn("精密支架045", text)
        self.assertFalse(any(c.data_type == "f" for row in sheet for c in row))
        self.assertTrue(all(r.height <= 409 for r in sheet.row_dimensions.values()))


class PurchaseTemplateFormTests(unittest.TestCase):
    setUp = order_tests.PurchaseOrderRouteTests.setUp
    tearDown = order_tests.PurchaseOrderRouteTests.tearDown
    login = order_tests.PurchaseOrderRouteTests.login
    create = order_tests.PurchaseOrderRouteTests.create

    def data(self, html):
        match = re.search(r'<script[^>]+data-supplier-data[^>]*>(.*?)</script>', html, re.S)
        self.assertIsNotNone(match, "form must provide whitelisted supplier data")
        return json.loads(match.group(1))

    def test_supplier_snapshot_survives_edit_error_switch_back_and_forged_fields(self):
        oid = self.create()
        with app.get_db() as conn:
            conn.execute("UPDATE suppliers SET name='新名字',contact='新联系人',phone='00999',payment_account_no='BANK-SECRET' WHERE id=1")
            conn.execute("INSERT INTO suppliers(code,name,contact,phone,address,created_at,updated_at) VALUES ('0002','供方乙','乙联系人','00123','乙地址','now','now')")
            item = conn.execute("SELECT id FROM purchase_order_items WHERE purchase_order_id=?", (oid,)).fetchone()[0]
        url = f"/admin/purchases/orders/{oid}/edit"
        data = self.data(self.client.get(url).get_data(as_text=True))
        self.assertEqual(data["saved"]["name"], "供方甲")
        self.assertEqual(data["saved"]["address"], "旧地址")
        self.assertEqual(data["suppliers"][0].keys(), {"id", "code", "name", "contact", "phone", "email", "address"})
        for supplier_id, expected in (("1", "供方甲"), ("2", "供方乙")):
            posted = form(payload(supplier_id=supplier_id, purchased_at="bad", supplier_name="伪造", supplier_phone="伪造电话"))
            response = self.client.post(url, data=posted)
            self.assertEqual(response.status_code, 400)
            html = response.get_data(as_text=True)
            self.assertRegex(html, rf'data-supplier-field="name">{expected}</')
            self.assertNotIn("BANK-SECRET", html)
            self.assertNotIn("伪造电话", html)
            self.assertEqual(self.data(html)["saved"]["name"], "供方甲")
        row = dict(id=item, item_name="螺栓", quantity="12", unit_price="3.25", expected_at="2026-09-20")
        response = self.client.post(url, data=form(payload(supplier_id="2", rows=[row], supplier_phone="伪造电话")))
        self.assertEqual(response.status_code, 302)
        with app.get_db() as conn:
            saved = conn.execute("SELECT * FROM purchase_orders WHERE id=?", (oid,)).fetchone()
            self.assertEqual((saved["supplier_name"], saved["supplier_phone"]), ("供方乙", "00123"))
            self.assertEqual(conn.execute("SELECT id FROM purchase_order_items WHERE purchase_order_id=?", (oid,)).fetchone()[0], item)

    def test_form_zones_fields_company_operator_and_price_permissions(self):
        with app.get_db() as conn:
            conn.execute("UPDATE reconciliation_company_profile SET company_name='测试抬头公司'")
        for slug in ("raw-material", "carton", "outsourcing", "other"):
            html = self.client.get(f"/admin/purchases/{slug}/new").get_data(as_text=True)
            self.assertIn("测试抬头公司", html)
            self.assertIn("经办人：buyer", html)
            self.assertLess(html.index("data-supplier-field"), html.index("data-purchase-rows"))
            self.assertLess(html.index("data-purchase-rows"), html.index('name="delivery_address"'))
            self.assertLess(html.index('name="delivery_address"'), html.index('name="purchased_at"'))
            if slug != "raw-material":
                self.assertIn("data-line-amount", html)
        self.login("blind")
        html = self.client.get("/admin/purchases/carton/new").get_data(as_text=True)
        self.assertNotIn("data-line-amount", html)
        self.assertNotIn("unit_price", html)
        self.assertNotIn("line_total", html)
        self.assertEqual(self.data(html)["saved"], None)

    def test_edit_supplier_selector_displays_snapshot_even_after_deactivation(self):
        oid = self.create()
        for active in (1, 0):
            with app.get_db() as conn:
                conn.execute("UPDATE suppliers SET name='后来改名', active=? WHERE id=1", (active,))
            html = self.client.get(f"/admin/purchases/orders/{oid}/edit").get_data(as_text=True)
            self.assertRegex(html, r'<option value="1"[^>]*selected[^>]*>供方甲')
            self.assertRegex(html, r'data-supplier-field="name">供方甲</')
            self.assertEqual(self.data(html)["saved"]["id"], 1)

    def test_purchase_asset_urls_refresh_when_either_resource_changes(self):
        oid = self.create()
        with tempfile.TemporaryDirectory() as folder:
            static = Path(folder) / "static"
            static.mkdir()
            for name in ("style.css", "editor.js", "purchase-orders.css", "purchase-orders.js"):
                (static / name).touch()
                os.utime(static / name, (100, 100))
            for filename, version in (("purchase-orders.js", 200), ("purchase-orders.css", 300)):
                os.utime(static / filename, (version, version))
                with patch.object(app, "BASE_DIR", Path(folder)):
                    for url in ("/admin/purchases/raw-material/new", "/admin/purchases/raw-material", f"/admin/purchases/orders/{oid}"):
                        html = self.client.get(url).get_data(as_text=True)
                        self.assertIn(f'/static/purchase-orders.css?v={version}', html)
                        if url.endswith("/new"):
                            self.assertIn(f'/static/purchase-orders.js?v={version}', html)

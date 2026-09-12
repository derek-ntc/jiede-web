import unittest

from openpyxl import load_workbook
import procurement_documents as documents
from tests.test_purchase_order_exports import pdf_pages


def sample_document():
    header = dict(outbound_no="POUT-20260912-TEST", outbound_at="2026-09-12", used_by="王师傅",
                  operator="keeper", status="posted", remark="车间领用 <b>纯文本</b>", created_at="2026-09-12T10:00:00")
    rows = [dict(category="raw_material", item_name="=SUM(1,2)", drawing_no="图号-A", material="镀锌板",
                 spec="板材", dimension_text="按图纸", length=1290, width=1250, thickness=1.4, unit="张",
                 supplier_name="供应商甲", order_no="PO-RAW", lot_no="PLOT-20260912-TEST", location_code="A01",
                 location_name="原料区", quantity=12, remark="机架 <b>按图</b>", unit_price_minor=987654)]
    return header, rows


class PurchaseOutboundExportTests(unittest.TestCase):
    def test_workbook_numeric_quantity_safe_literals_and_landscape_a4(self):
        builder = getattr(documents, "build_purchase_outbound_workbook", None)
        self.assertTrue(callable(builder), "outbound workbook builder missing")
        header, rows = sample_document()
        ws = load_workbook(builder(header, rows)).active
        values = "\n".join(str(c.value) for row in ws for c in row if c.value is not None)
        for value in ("采购出库单", "王师傅", "keeper", "PO-RAW", "PLOT-20260912-TEST", "原料区", "1290", "=SUM(1,2)", "<b>"):
            self.assertIn(value, values)
        self.assertNotIn("987654", values)
        self.assertNotIn("单价", values)
        self.assertFalse(any(c.data_type == "f" for row in ws for c in row))
        qty = next(c for row in ws for c in row if c.value == 12)
        self.assertEqual((qty.data_type, qty.number_format), ("n", "General"))
        self.assertEqual((ws.page_setup.orientation, str(ws.page_setup.paperSize)), ("landscape", "9"))
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertTrue(ws.print_title_rows)

    def test_pdf_contains_chinese_snapshots_and_paginated_long_text(self):
        builder = getattr(documents, "build_purchase_outbound_pdf", None)
        self.assertTrue(callable(builder), "outbound PDF builder missing")
        header, rows = sample_document()
        rows[0]["remark"] = "长备注 <tag>& 逐项核对。" * 200
        stream = builder(header, rows)
        self.assertRegex(stream.getvalue(), rb"/MediaBox\s*\[\s*0\s+0\s+841\.\d+\s+595\.\d+")
        pages = pdf_pages(stream)
        self.assertGreater(len(pages), 1)
        text = "".join("".join(pages).split())
        for value in ("采购出库单", "王师傅", "keeper", "供应商甲", "PO-RAW", "<tag>", "领用"):
            self.assertIn(value, text)
        self.assertNotIn("987654", text)
        self.assertNotIn("单价", text)

    def test_voided_documents_include_audit_without_replacing_original_operator(self):
        builder = getattr(documents, "build_purchase_outbound_workbook", None)
        self.assertTrue(callable(builder), "outbound workbook builder missing")
        header, rows = sample_document()
        header.update(status="voided", voided_by="supervisor", voided_at="2026-09-13T11:00:00")
        ws = load_workbook(builder(header, rows)).active
        text = " ".join(str(c.value) for row in ws for c in row)
        for value in ("已作废", "supervisor", "2026-09-13", "keeper"):
            self.assertIn(value, text)

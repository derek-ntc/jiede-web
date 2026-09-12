import json
import sqlite3
import tempfile
import subprocess
import re
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import app
import shipping_workflow
from openpyxl import load_workbook
from tests.test_assembly_shipping import AssemblyTask7TestCase, VALID_PNG_BYTES, png_bytes
from tests.test_product_bom_import import workbook_upload
from tests.test_shipped_pdf_company import ShippedPdfCompanyTests


class DeliveryNoteWorkflowTests(AssemblyTask7TestCase):
    def setUp(self):
        self.files = tempfile.TemporaryDirectory()
        self.paths = patch.multiple(app, **{
            name: Path(self.files.name) / name.lower()
            for name in ('DATA_DIR', 'MANUALS_DIR', 'SIGNATURES_DIR',
                         'RECONCILIATION_SIGNATURES_DIR', 'SHIPMENT_IMAGES_DIR',
                         'INSPECTION_REPORTS_DIR', 'PRODUCTION_DRAWINGS_DIR')
        })
        self.paths.start()
        super().setUp()
        self.product = self.create_product()
        self.order = self.create_order(self.product, 'SO-DN', 20)
        self.stock_product(self.product, 20)
        self.configure_components([(self.product, 1)])
        with app.get_db() as conn:
            conn.execute("UPDATE manuals SET supplier='长规格', unit='件', unit_price_minor=0")
            conn.execute("INSERT INTO customers (name, recipient_name, recipient_phone, address, created_at, updated_at) VALUES ('客户A', '张师傅', '13800000000', '宁波仓库', '', '')")

    def tearDown(self):
        super().tearDown()
        self.paths.stop()
        self.files.cleanup()

    def ordinary_data(self, **extra):
        return dict(shipped_at='2026-09-08', order_id=str(self.order),
                    shipped_quantity='20', operation_token='ordinary-token', **extra)

    def test_import_list_template_snapshot_and_claimed_batch_delete_work_together(self):
        with app.get_db() as conn:
            conn.execute("UPDATE manuals SET unit_price_minor=123, remark='保留旧备注' WHERE id=?", (self.product,))
        preview = self.client.post('/admin/products/import', data={
            'customer': '客户A', 'file': workbook_upload([
                ('物料编码', '物料名称', '规格型号', '单位'),
                ('SKU-UPDATED', '重新导入的产品', 'P1', 'PCS'),
                ('SKU-FREE', '可删除的产品', 'FREE-IMPORT', 'PCS'),
            ])}, content_type='multipart/form-data')
        self.assertEqual(preview.status_code, 200)
        token = re.search(r'name="import_token"\s+value="([^"]+)"', preview.get_data(as_text=True)).group(1)
        confirmed = self.client.post('/admin/products/import/confirm', data={'import_token': token})
        self.assertEqual(confirmed.location, '/admin/products')
        with app.get_db() as conn:
            self.assertEqual(tuple(conn.execute('''SELECT drawing_no,supplier,sku,unit,unit_price_minor,remark
                FROM manuals WHERE id=?''', (self.product,)).fetchone()),
                ('P1', 'P1', 'SKU-UPDATED', 'PCS', 123, '保留旧备注'))
            free = conn.execute("SELECT id FROM manuals WHERE drawing_no='FREE-IMPORT'").fetchone()[0]
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM product_assembly_components WHERE manual_id=?',
                                         (self.product,)).fetchone()[0], 1)
        listed = self.client.get('/admin/products?supplier=P1').get_data(as_text=True)
        self.assertIn('重新导入的产品', listed)
        self.assertNotIn('可删除的产品', listed)
        self.assertIn('规格型号', listed)

        template = f'/admin/{self.product}/process-template'
        self.client.get(f'/manual/{self.product}/technical')
        with self.client.session_transaction() as session:
            process_csrf = session['production_followup_csrf_token']
        self.assertEqual(self.client.post(template, data={
            'process_name': ['下料', '检验 <终检>'],
            'production_followup_csrf_token': process_csrf,
        }).status_code, 302)

        def create_card():
            response = self.client.post('/admin/production-followups', data={
                'ordered_at': '2026-09-09', 'customer': '客户A', 'manual_id': str(self.product),
                'drawing_no': 'P1', 'product_name': '重新导入的产品'})
            self.assertEqual(response.status_code, 302)
            with app.get_db() as conn:
                return conn.execute('SELECT MAX(id) FROM production_followups').fetchone()[0]

        historical = create_card()
        self.client.get('/admin/production-followups')
        with self.client.session_transaction() as session:
            csrf = session['production_followup_csrf_token']
        with app.get_db() as conn:
            step = conn.execute('SELECT id FROM production_followup_process_steps WHERE followup_id=? ORDER BY sort_order',
                                (historical,)).fetchone()[0]
        completed = self.client.post(f'/admin/production-followups/{historical}/processes/{step}/complete',
                                     data={'production_followup_csrf_token': csrf})
        self.assertEqual(completed.status_code, 302)
        with app.get_db() as conn:
            old_steps = [dict(row) for row in conn.execute(
                'SELECT * FROM production_followup_process_steps WHERE followup_id=? ORDER BY sort_order', (historical,))]
            self.assertTrue(old_steps[0]['completed_at'])
            self.assertEqual(old_steps[0]['completed_by'], 'admin')
        self.assertEqual(self.client.post(template, data={
            'process_name': ['包装'],
            'production_followup_csrf_token': process_csrf,
        }).status_code, 302)
        current = create_card()
        app.init_db()
        app.init_db()
        with app.get_db() as conn:
            self.assertEqual([dict(row) for row in conn.execute(
                'SELECT * FROM production_followup_process_steps WHERE followup_id=? ORDER BY sort_order', (historical,))], old_steps)
            self.assertEqual([row[0] for row in conn.execute(
                'SELECT name FROM production_followup_process_steps WHERE followup_id=?', (current,))], ['包装'])
        card = self.client.get(f'/admin/production-followups/{historical}/process-card').get_data(as_text=True)
        self.assertIn('检验 &lt;终检&gt;', card)
        self.assertIn('admin', card)

        data = self.ordinary_data()
        data['shipped_quantity'] = '1'
        self.assertRegex(self.client.post('/admin/shipped-orders/new', data=data).location,
                         r'/admin/delivery-notes/operations/\d+$')
        with app.get_db() as conn:
            source = conn.execute('SELECT id FROM product_order_shipments').fetchone()[0]
            customer = conn.execute("SELECT id FROM customers WHERE name='客户A'").fetchone()[0]
            app.create_finance_invoice(conn, customer, [('ordinary', source)], 'admin')
            app.create_reconciliation_statement(conn, [('ordinary', source)], 'admin')
        blocked = self.client.post('/admin/products/delete-batch', data={
            'production_followup_csrf_token': process_csrf,
            'manual_id': [str(free), str(self.product)], 'return_supplier': 'P1'}, follow_redirects=True)
        self.assertIn('整批未删除', blocked.get_data(as_text=True))
        with app.get_db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM manuals WHERE id IN (?,?)', (free, self.product)).fetchone()[0], 2)
        deleted = self.client.post('/admin/products/delete-batch', data={
            'manual_id': str(free),
            'return_q': 'FREE-IMPORT',
            'production_followup_csrf_token': process_csrf,
        })
        self.assertIn('/admin/products?q=FREE-IMPORT', deleted.location)
        with app.get_db() as conn:
            self.assertIsNone(conn.execute('SELECT id FROM manuals WHERE id=?', (free,)).fetchone())
            self.assertIsNotNone(conn.execute('SELECT id FROM manuals WHERE id=?', (self.product,)).fetchone())

    def test_both_shipment_modes_preserve_zero_added_rows_through_accounting_and_restart(self):
        for mode in ('assembly', 'ordinary'):
            with self.subTest(mode=mode):
                positive = self.create_product(f'{mode}-POS')
                zero = self.create_product(f'{mode}-ZERO')
                extra = self.create_product(f'{mode}-EXTRA')
                self.create_order(positive, f'{mode}-ORDER', 10)
                zero_order = self.create_order(zero, f'{mode}-ZERO-ORDER', 9)
                self.stock_product(positive, 10)
                self.stock_product(zero, 5)
                self.stock_product(extra, 1)
                with app.get_db() as conn:
                    for product, price in [(positive, 123), (zero, 999), (extra, 250)]:
                        conn.execute("UPDATE manuals SET supplier='冻结规格', unit='PCS', unit_price_minor=?, currency='CNY' WHERE id=?",
                                     (price, product))
                quantities = [2, 0, 3]
                remarks = ['先发两件', '本次不发 <待补>', '追加三件 & 缺货两件']
                if mode == 'assembly':
                    self.configure_components([(positive, 1), (zero, 1)])
                    preview = self.client.post('/admin/shipped-orders/assembly-preview', json={
                        'customer': '客户A', 'assembly_drawing_no': 'ASM-100', 'set_quantity': 2,
                        'selected_manual_ids': [positive, zero, extra],
                        'overrides': dict(zip([positive, zero, extra], quantities)),
                        'remarks': dict(zip([positive, zero, extra], remarks))})
                    self.assertEqual(preview.status_code, 200)
                    data = self.save_data(preview.get_json(), extra_data={
                        'selected_manual_ids': [str(positive), str(zero), str(extra)],
                        'line_remark': remarks, 'operation_token': f'e2e-{mode}'})
                    route = '/admin/shipped-orders/assembly/new'
                else:
                    lines = [dict(manual_id=product, customer='客户A', source_kind='extra', quantity=quantity, remark=remark)
                             for product, quantity, remark in zip([positive, zero, extra], quantities, remarks)]
                    preview = self.client.post('/admin/shipped-orders/order-preview', json={'lines': lines})
                    self.assertEqual(preview.status_code, 200)
                    data = dict(shipment_lines=json.dumps(lines), shipped_at='2026-09-09', operation_token=f'e2e-{mode}',
                                preview_token=preview.get_json()['preview_token'])
                    route = '/admin/shipped-orders/new'
                saved = self.client.post(route, data=data)
                self.assertEqual(saved.status_code, 201 if mode == 'assembly' else 302)
                receipt = saved.get_json()['redirect_url'] if mode == 'assembly' else saved.location
                retry = self.client.post(route, data=data)
                self.assertEqual(retry.status_code, saved.status_code)
                self.assertEqual(retry.get_json()['redirect_url'] if mode == 'assembly' else retry.location, receipt)
                with app.get_db() as conn:
                    note_id = conn.execute('SELECT MAX(id) FROM delivery_notes').fetchone()[0]
                    note = shipping_workflow.load_delivery_note(conn, note_id)
                    expected_items = [(positive, 2, remarks[0]), (zero, 0, remarks[1]), (extra, 3, remarks[2])]
                    self.assertEqual([(i['manual_id'], i['quantity'], i['remark']) for i in note['items']], expected_items)
                    self.assertEqual([app.inventory_total_for_manual(conn, product) for product in (positive, zero, extra)], [8, 5, 0])
                    self.assertEqual(conn.execute('SELECT shipped_quantity FROM product_orders WHERE id=?', (zero_order,)).fetchone()[0], 0)
                    self.assertEqual(conn.execute("SELECT COUNT(*) FROM inventory_transactions WHERE manual_id=? AND type='out'", (zero,)).fetchone()[0], 0)
                    sources = [row for row in app.fetch_available_finance_sources(conn, '客户A')
                               if row['drawing_no'] in (f'{mode}-POS', f'{mode}-ZERO', f'{mode}-EXTRA')]
                    self.assertEqual(sorted(row['quantity'] for row in sources), [2, 3])
                    self.assertEqual({row['source_type'] for row in sources},
                                     {'assembly_item'} if mode == 'assembly' else {'ordinary', 'supplemental'})
                    refs = [(row['source_type'], row['source_id']) for row in sources]
                    customer = conn.execute("SELECT id FROM customers WHERE name='客户A'").fetchone()[0]
                    invoice = app.create_finance_invoice(conn, customer, refs, 'admin')
                    statement = app.create_reconciliation_statement(conn, refs, 'admin')['statement_id']
                    self.assertEqual(conn.execute('SELECT total_minor FROM finance_invoices WHERE id=?', (invoice,)).fetchone()[0], 996)
                    self.assertEqual(conn.execute('SELECT amount_incl_tax_minor FROM reconciliation_statements WHERE id=?', (statement,)).fetchone()[0], 996)
                    for table, parent, parent_id in [('finance_invoice_items', 'invoice_id', invoice),
                                                     ('reconciliation_statement_items', 'statement_id', statement)]:
                        self.assertEqual(sorted(row[0] for row in conn.execute(f'SELECT quantity FROM {table} WHERE {parent}=?', (parent_id,))), [2, 3])
                    for check in (app.assert_finance_sources_mutable, app.assert_reconciliation_sources_mutable):
                        with self.assertRaises(ValueError):
                            check(conn, refs)
                    conn.execute("UPDATE manuals SET supplier='后来规格', unit_price_minor=9999 WHERE id IN (?,?,?)", (positive, zero, extra))
                app.init_db()
                app.init_db()
                with app.get_db() as conn:
                    after = shipping_workflow.load_delivery_note(conn, note_id)
                    self.assertEqual([(i['manual_id'], i['quantity'], i['remark']) for i in after['items']], expected_items)
                    self.assertEqual([i['specification'] for i in after['items']], ['冻结规格'] * 3)
                    self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(), [])
                html = self.client.get(receipt).get_data(as_text=True)
                self.assertIn('本次不发 &lt;待补&gt;', html)
                self.assertIn('追加三件 &amp; 缺货两件', html)
                pdf = self.client.get(f'/admin/delivery-notes/{note_id}.pdf')
                self.assertEqual(pdf.status_code, 200)
                self.assertTrue(pdf.data.startswith(b'%PDF'))

    def test_delivery_html_and_real_pdf_show_zero_and_final_remark_column_without_prices(self):
        zero = self.create_product('ZERO-PDF')
        order = self.create_order(zero, 'ZERO-ORDER', 5)
        response = self.client.post('/admin/shipped-orders/new', data={
            'order_id': [str(order), str(self.order)], 'shipped_quantity': ['0', '2'],
            'line_remark': ['ZERO-REMARK <fragile> & hold', 'POSITIVE-REMARK'],
            'shipped_at': '2026-09-09', 'operation_token': 'remark-pdf'})
        self.assertRegex(response.location, r'/admin/delivery-notes/operations/\d+$')
        body = self.client.get(response.location).get_data(as_text=True)
        self.assertRegex(body, r'<th>数量</th><th>备注</th></tr>')
        self.assertIn('ZERO-REMARK &lt;fragile&gt; &amp; hold', body)
        with app.get_db() as conn:
            note_id = conn.execute('SELECT id FROM delivery_notes').fetchone()[0]
        response = self.client.get(f'/admin/delivery-notes/{note_id}.pdf')
        self.assertEqual(response.status_code, 200)
        from base64 import a85decode
        import zlib
        streams = re.findall(rb'stream\r?\n(.*?)endstream', response.data, re.S)
        text = ''.join(zlib.decompress(a85decode(stream.strip()[:-2])).decode('latin-1')
                       for stream in streams if stream.strip().endswith(b'~>'))
        # The A4 remark column may wrap a word across consecutive text draws.
        fragments = re.findall(r'\(((?:\\.|[^\\)])*)\)\s*Tj', text)
        # PDF string literals use octal escapes. The server's CID fallback
        # encodes text in UTF-16BE, whereas an embedded TrueType font emits
        # single-byte strings for these ASCII document-contract markers.
        decoded = []
        escapes = {'n': '\n', 'r': '\r', 't': '\t', 'b': '\b', 'f': '\f'}
        for fragment in fragments:
            literal = re.sub(r'\\([0-7]{1,3}|.)', lambda match:
                chr(int(match[1], 8)) if re.fullmatch('[0-7]{1,3}', match[1])
                else escapes.get(match[1], match[1]), fragment).encode('latin-1')
            decoded.append(literal.decode('utf-16-be') if b'\x00' in literal else literal.decode('latin-1'))
        rendered_text = ''.join(decoded)
        self.assertIn('ZERO-REMARK', rendered_text)
        self.assertIn('POSITIVE-REMARK', rendered_text)
        self.assertIn('ZERO-PDF', rendered_text)
        self.assertIn('0', decoded)
        self.assertNotIn('CNY', rendered_text)

    def test_delivery_pdf_contract_with_cid_fallback_font(self):
        with patch.object(app, 'PDF_FONT_CANDIDATES', []), patch.object(app, 'PDF_FONT_NAME', None):
            self.test_delivery_html_and_real_pdf_show_zero_and_final_remark_column_without_prices()

    def test_delivery_note_exports_template_style_excel_without_word(self):
        zero = self.create_product('ZERO-EXPORT')
        order = self.create_order(zero, 'ZERO-EXPORT-ORDER', 5)
        response = self.client.post('/admin/shipped-orders/new', data={
            'order_id': [str(order), str(self.order)], 'shipped_quantity': ['0', '1'],
            'line_remark': ['暂不发货', '正常发货'], 'shipped_at': '2026-09-09',
            'operation_token': 'zero-export',
        })
        body = self.client.get(response.location).get_data(as_text=True)
        self.assertIn('.xlsx', body)
        self.assertNotIn('.docx', body)
        with app.get_db() as conn:
            note_id = conn.execute('SELECT id FROM delivery_notes').fetchone()[0]
        xlsx = self.client.get(f'/admin/delivery-notes/{note_id}.xlsx')
        self.assertEqual(xlsx.status_code, 200)
        workbook = load_workbook(BytesIO(xlsx.data))
        sheet = workbook.active
        self.assertEqual(sheet.page_setup.orientation, "portrait")
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(str(sheet.print_area), "'送货单'!$A$2:$H$11")
        self.assertEqual(sheet["A2"].value, "宁波市杰德机械科技有限公司")
        self.assertEqual(sheet["A3"].value, "送货单")
        self.assertEqual(
            [sheet.cell(7, column).value for column in range(1, 9)],
            ["序号", "产品图号", "产品名称", "规格型号", "单位", "订单数量", "实发数量", "备注"],
        )
        self.assertIn("暂不发货", " ".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row))
        self.assertEqual(self.client.get(f'/admin/delivery-notes/{note_id}.docx').status_code, 404)
        with app.get_db() as conn:
            note = shipping_workflow.load_delivery_note(conn, note_id)
        story = []
        with patch.object(app.SimpleDocTemplate, 'build', lambda _doc, values: story.extend(values)):
            app.build_delivery_note_pdf(app.delivery_note_export_payload(note))
        text = ShippedPdfCompanyTests.collect_story_text(story)
        self.assertIn('订单数量', text)
        self.assertIn('实发数量', text)

    def test_assembly_delivery_note_excel_uses_assembly_template_columns(self):
        preview = self.post_preview(sets=2).get_json()
        saved = self.client.post('/admin/shipped-orders/assembly/new', data=self.save_data(preview))
        self.assertEqual(saved.status_code, 201, saved.get_json())
        with app.get_db() as conn:
            note_id = conn.execute('SELECT id FROM delivery_notes').fetchone()[0]

        xlsx = self.client.get(f'/admin/delivery-notes/{note_id}.xlsx')
        self.assertEqual(xlsx.status_code, 200)
        sheet = load_workbook(BytesIO(xlsx.data)).active
        self.assertEqual(sheet["A7"].value, "组装件名称")
        self.assertEqual(sheet["B7"].value, "ASM-100")
        self.assertEqual(
            [sheet.cell(8, column).value for column in range(1, 9)],
            ["序号", "产品图号", "产品名称", "规格型号", "单位", "每套数量", "实发数量", "备注"],
        )

    def test_delivery_note_can_send_pdf_and_excel_to_saved_customer_email(self):
        with app.get_db() as conn:
            conn.execute("UPDATE customers SET email='customer@example.com' WHERE name='客户A'")
        response = self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        with app.get_db() as conn:
            note_id = conn.execute('SELECT id FROM delivery_notes').fetchone()[0]

        self.client.get(response.location)
        with self.client.session_transaction() as session:
            csrf_token = session['delivery_note_email_csrf_token']

        sent_messages = []

        class FakeSMTP:
            def __init__(self, *args, **kwargs):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def starttls(self):
                pass

            def login(self, *args):
                pass

            def send_message(self, message):
                sent_messages.append(message)

        with patch.dict(app.os.environ, {
            'SMTP_HOST': 'smtp.example.com', 'SMTP_FROM': 'noreply@example.com',
            'DELIVERY_NOTE_CC': '',
        }, clear=False), patch.object(app.smtplib, 'SMTP', FakeSMTP):
            result = self.client.post(
                f'/admin/delivery-notes/{note_id}/email',
                data={'delivery_note_email_csrf_token': csrf_token}, follow_redirects=True,
            )

        self.assertEqual(result.status_code, 200)
        self.assertIn('已发送至 customer@example.com', result.get_data(as_text=True))
        self.assertEqual(len(sent_messages), 1)
        attachments = list(sent_messages[0].iter_attachments())
        self.assertEqual([item.get_content_type() for item in attachments], [
            'application/pdf',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        ])

    def test_delivery_note_email_requires_customer_email(self):
        response = self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        with app.get_db() as conn:
            note_id = conn.execute('SELECT id FROM delivery_notes').fetchone()[0]
        self.client.get(response.location)
        with self.client.session_transaction() as session:
            csrf_token = session['delivery_note_email_csrf_token']

        result = self.client.post(
            f'/admin/delivery-notes/{note_id}/email',
            data={'delivery_note_email_csrf_token': csrf_token}, follow_redirects=True,
        )
        self.assertIn('未填写客户邮箱', result.get_data(as_text=True))

    def test_assembly_snapshot_keeps_order_zero_remarks_and_replay_is_read_only(self):
        zero_product = self.create_product('ZERO-DN')
        self.configure_components([(zero_product, 3), (self.product, 1)])
        response = self.client.post('/admin/shipped-orders/assembly-preview', json={
            'customer': '客户A', 'assembly_drawing_no': 'ASM-100', 'set_quantity': 2,
            'selected_manual_ids': [zero_product, self.product],
            'overrides': {zero_product: 0, self.product: 2},
            'remarks': {zero_product: '<待补> & 本次不发', self.product: '先发两件'},
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        data = self.save_data(response.get_json(), extra_data={
            'selected_manual_ids': [str(zero_product), str(self.product)],
            'line_remark': ['<待补> & 本次不发', '先发两件'],
            'operation_token': 'zero-assembly-snapshot',
        })
        first = self.client.post('/admin/shipped-orders/assembly/new', data=data)
        self.assertEqual(first.status_code, 201, first.get_json())
        retry = self.client.post('/admin/shipped-orders/assembly/new', data=data)
        self.assertEqual(retry.status_code, 201, retry.get_json())
        self.assertEqual(first.get_json(), retry.get_json())
        with app.get_db() as conn:
            note_id = conn.execute('SELECT id FROM delivery_notes').fetchone()[0]
            note = shipping_workflow.load_delivery_note(conn, note_id)
            self.assertEqual([(i['manual_id'], i['quantity'], i['remark']) for i in note['items']],
                             [(zero_product, 0, '<待补> & 本次不发'), (self.product, 2, '先发两件')])
            self.assertEqual([i['sort_order'] for i in note['items']], [0, 1])
            self.assertIsNone(note['items'][0]['source_id'])
            self.assertEqual(len(note['sources']), 1)
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 18)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0], 1)
            conn.execute("UPDATE manuals SET drawing_no='CHANGED', supplier='CHANGED'")
            note = shipping_workflow.load_delivery_note(conn, note_id)
            self.assertEqual(note['items'][0]['drawing_no'], 'ZERO-DN')
            self.assertEqual(note['items'][1]['specification'], '长规格')
        self.assertEqual(self.client.get(first.get_json()['redirect_url']).status_code, 200)
        self.assertEqual(self.client.get(f'/admin/delivery-notes/{note_id}.pdf').status_code, 200)

    def note_rows(self):
        with app.get_db() as conn:
            return [dict(row) for row in conn.execute('SELECT * FROM delivery_notes ORDER BY id')]

    def create_plan(self, order_ids):
        response = self.client.post('/admin/orders/shipment-plans', data={
            'order_id': [str(order_id) for order_id in order_ids],
            'planned_ship_at': '2026-09-10',
        })
        self.assertEqual(response.status_code, 302)
        with app.get_db() as conn:
            plan = conn.execute('SELECT * FROM shipment_plans ORDER BY id DESC LIMIT 1').fetchone()
            items = conn.execute(
                'SELECT * FROM shipment_plan_items WHERE plan_id=? ORDER BY id',
                (plan['id'],),
            ).fetchall()
        return dict(plan), [dict(item) for item in items]

    @staticmethod
    def plan_approval_data(items, *, token='plan-token', images=None, overrides=None):
        data = {
            'operation_token': token,
            'shipped_at': '2026-09-09',
            'logistics_no': '计划发货',
            'plan_item_id': [str(item['id']) for item in items],
            'shipped_quantity': ['2'] * len(items),
            'recipient_overrides': json.dumps(overrides or {}),
        }
        if images is not None:
            data['images'] = images
        return data

    def test_plan_approval_snapshots_prices_recipients_and_replays_original_result(self):
        product_b = self.create_product('P-PLAN-B', '客户B')
        order_b = self.create_order(product_b, 'SO-PLAN-B', 5, customer='客户B')
        self.stock_product(product_b, 5)
        with app.get_db() as conn:
            conn.execute("UPDATE manuals SET supplier='计划规格-A', unit_price_minor=0, currency='CNY' WHERE id=?", (self.product,))
            conn.execute("UPDATE manuals SET supplier='计划规格-B', unit_price_minor=NULL, currency='USD' WHERE id=?", (product_b,))
            conn.execute("INSERT INTO customers (name,recipient_name,recipient_phone,address,created_at,updated_at) VALUES ('客户B','李师傅','13900000000','上海仓库','','')")
        plan, items = self.create_plan([self.order, order_b])
        request_data = lambda: self.plan_approval_data(
            items,
            images=(BytesIO(VALID_PNG_BYTES), 'plan.png'),
            overrides={'客户A': {'recipient_name': '临时收货人', 'recipient_phone': '123', 'address': '临时仓'}},
        )

        first = self.client.post(f"/admin/shipped-orders/plans/{plan['id']}/approve", data=request_data())
        self.assertEqual(first.status_code, 302)
        self.assertRegex(first.location, r'/admin/delivery-notes/operations/\d+$')
        with app.get_db() as conn:
            shipments = [dict(row) for row in conn.execute(
                '''SELECT s.*, o.customer FROM product_order_shipments s
                   JOIN product_orders o ON o.id=s.order_id ORDER BY s.id''')]
            notes = [dict(row) for row in conn.execute('SELECT * FROM delivery_notes ORDER BY customer')]
            stock_after_first = {
                self.product: app.inventory_total_for_manual(conn, self.product),
                product_b: app.inventory_total_for_manual(conn, product_b),
            }
            plan_after_first = dict(conn.execute('SELECT * FROM shipment_plans WHERE id=?', (plan['id'],)).fetchone())
            plan_items_after_first = [dict(row) for row in conn.execute(
                'SELECT * FROM shipment_plan_items WHERE plan_id=? ORDER BY id', (plan['id'],))]
        self.assertEqual([(row['specification_snapshot'], row['unit_price_minor'], row['currency']) for row in shipments],
                         [('计划规格-A', 0, 'CNY'), ('计划规格-B', None, 'USD')])
        self.assertTrue(shipments[0]['price_recorded_by'])
        self.assertFalse(shipments[1]['price_recorded_by'])
        self.assertEqual([(note['customer'], note['recipient_name'], note['address']) for note in notes],
                         [('客户A', '临时收货人', '临时仓'), ('客户B', '李师傅', '上海仓库')])
        self.assertEqual(plan_after_first['status'], '已完成')
        self.assertEqual([item['shipped_quantity'] for item in plan_items_after_first], [2, 2])
        self.assertEqual(stock_after_first, {self.product: 18, product_b: 3})

        with app.get_db() as conn:
            conn.execute("UPDATE manuals SET supplier='已修改规格', unit_price_minor=99999")
            conn.execute("UPDATE customers SET recipient_name='已修改收货人'")
        retry = self.client.post(f"/admin/shipped-orders/plans/{plan['id']}/approve", data=request_data())
        self.assertEqual(retry.status_code, 302)
        self.assertEqual(retry.location, first.location)
        with app.get_db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM product_order_shipments').fetchone()[0], 2)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0], 1)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_notes').fetchone()[0], 2)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM product_order_shipment_images').fetchone()[0], 2)
            self.assertEqual({
                self.product: app.inventory_total_for_manual(conn, self.product),
                product_b: app.inventory_total_for_manual(conn, product_b),
            }, stock_after_first)
            self.assertEqual(
                [row[0] for row in conn.execute('SELECT specification_snapshot FROM product_order_shipments ORDER BY id')],
                ['计划规格-A', '计划规格-B'],
            )

        conflict_data = self.plan_approval_data(items, token='plan-token')
        conflict_data['shipped_quantity'] = ['1', '2']
        self.assertEqual(
            self.client.post(f"/admin/shipped-orders/plans/{plan['id']}/approve", data=conflict_data).status_code,
            409,
        )

    def test_plan_approval_failure_rolls_back_images_inventory_notes_and_can_retry(self):
        product_b = self.create_product('P-PLAN-ROLLBACK', '客户A')
        order_b = self.create_order(product_b, 'SO-PLAN-ROLLBACK', 5)
        self.stock_product(product_b, 5)
        plan, items = self.create_plan([self.order, order_b])
        invalid = self.plan_approval_data(
            items, token='plan-invalid-image', images=(BytesIO(b'not an image'), 'broken.png'))
        response = self.client.post(f"/admin/shipped-orders/plans/{plan['id']}/approve", data=invalid)
        self.assertEqual(response.status_code, 302)
        with app.get_db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM product_order_shipments').fetchone()[0], 0)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0], 0)
            self.assertEqual(conn.execute('SELECT status FROM shipment_plans WHERE id=?', (plan['id'],)).fetchone()[0], '待发货')

        failing = self.plan_approval_data(
            items, token='plan-note-failure', images=(BytesIO(VALID_PNG_BYTES), 'plan.png'))
        with patch.object(app, 'create_delivery_notes', side_effect=sqlite3.OperationalError('injected plan note failure')):
            response = self.client.post(f"/admin/shipped-orders/plans/{plan['id']}/approve", data=failing)
        self.assertEqual(response.status_code, 500)
        with app.get_db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM product_order_shipments').fetchone()[0], 0)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0], 0)
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 20)
            self.assertEqual(app.inventory_total_for_manual(conn, product_b), 5)
        self.assertEqual(list(app.SHIPMENT_IMAGES_DIR.iterdir()), [])
        self.assertEqual(
            self.client.post(f"/admin/shipped-orders/plans/{plan['id']}/approve", data=self.plan_approval_data(items, token='plan-note-failure')).status_code,
            302,
        )

    def test_plan_approval_commit_exit_returns_receipt_and_retry_is_read_only(self):
        plan, items = self.create_plan([self.order])
        data = self.plan_approval_data(items, token='plan-commit-exit')
        original_get_db = app.get_db

        @contextmanager
        def fail_after_commit():
            with original_get_db() as conn:
                before = conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0]
                yield conn
                after = conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0]
            if after > before:
                raise RuntimeError('injected committed plan context exit')

        with patch.object(app, 'get_db', fail_after_commit):
            first = self.client.post(f"/admin/shipped-orders/plans/{plan['id']}/approve", data=data)
        self.assertEqual(first.status_code, 302)
        self.assertRegex(first.location, r'/admin/delivery-notes/operations/\d+$')
        retry = self.client.post(f"/admin/shipped-orders/plans/{plan['id']}/approve", data=data)
        self.assertEqual(retry.location, first.location)
        with app.get_db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM product_order_shipments').fetchone()[0], 1)
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 18)

    def test_ordinary_retry_returns_same_receipt_after_stock_and_customer_change(self):
        first = self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        self.assertRegex(first.location, r'/admin/delivery-notes/operations/\d+$')
        notes = self.note_rows()
        with app.get_db() as conn:
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 0)
            conn.execute("UPDATE customers SET recipient_name='李师傅'")
            note = shipping_workflow.load_delivery_note(conn, notes[0]['id'])
        self.assertEqual(note['recipient_name'], '张师傅')
        self.assertEqual(note['items'][0]['specification'], '长规格')
        self.assertEqual(note['items'][0]['unit'], '件')
        retry = self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        self.assertEqual(retry.location, first.location)
        self.assertEqual([n['id'] for n in self.note_rows()], [n['id'] for n in notes])
        with app.get_db() as conn:
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 0)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM product_order_shipments').fetchone()[0], 1)
        self.assertIn('张师傅', self.client.get(first.location).get_data(as_text=True))

    def test_token_conflict_checks_content_and_user(self):
        data = self.ordinary_data()
        self.client.post('/admin/shipped-orders/new', data=data)
        data['shipped_quantity'] = '1'
        self.assertEqual(self.client.post('/admin/shipped-orders/new', data=data).status_code, 409)
        with app.get_db() as conn:
            app.seed_user(conn, 'other-admin', 'test-password', 'admin', '2026-09-08')
        with self.client.session_transaction() as session:
            session['admin_username'] = 'other-admin'
        self.assertEqual(self.client.post('/admin/shipped-orders/new', data=self.ordinary_data()).status_code, 409)

    def test_cross_customer_groups_and_one_time_recipient_override(self):
        product = self.create_product('PB', '客户B')
        order = self.create_order(product, 'SO-B', 2, customer='客户B')
        self.stock_product(product, 2)
        data = self.ordinary_data(recipient_overrides=json.dumps({
            '客户A': {'recipient_name': '王师傅', 'recipient_phone': '123', 'address': '临时仓'},
        }))
        data.update(order_id=[str(self.order), str(order)], shipped_quantity=['2', '2'])
        self.client.post('/admin/shipped-orders/new', data=data)
        notes = self.note_rows()
        self.assertEqual([n['customer'] for n in notes], ['客户A', '客户B'])
        self.assertEqual(notes[0]['recipient_name'], '王师傅')
        self.assertEqual(notes[1]['recipient_name'], '')
        with app.get_db() as conn:
            self.assertEqual(conn.execute('SELECT recipient_name FROM customers').fetchone()[0], '张师傅')
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_note_sources').fetchone()[0], 2)

    def test_assembly_retry_uses_batch_source_and_ignores_changed_preview_state(self):
        preview = self.post_preview(sets=2).get_json()
        data = self.save_data(preview, extra_data={'operation_token': 'assembly-token'})
        first = self.client.post('/admin/shipped-orders/assembly/new', data=data)
        self.assertEqual(first.status_code, 201, first.get_json())
        self.assertIn('/admin/delivery-notes/operations/', first.get_json()['redirect_url'])
        retry = self.client.post('/admin/shipped-orders/assembly/new', data=data)
        self.assertEqual(retry.status_code, 201, retry.get_json())
        self.assertEqual(retry.get_json(), first.get_json())
        with app.get_db() as conn:
            source = conn.execute('SELECT source_type, source_id FROM delivery_note_sources').fetchone()
            self.assertEqual(tuple(source), ('assembly', first.get_json()['batch_id']))
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 18)

    def test_attachment_digest_rejects_same_filename_different_content(self):
        data = self.ordinary_data()
        data['shipped_quantity'] = '1'
        data['images'] = (BytesIO(VALID_PNG_BYTES), 'photo.png')
        first = self.client.post('/admin/shipped-orders/new', data=data)
        self.assertEqual(first.status_code, 302)
        data['images'] = (BytesIO(png_bytes(2, 2)), 'photo.png')
        self.assertEqual(self.client.post('/admin/shipped-orders/new', data=data).status_code, 409)
        data['images'] = (BytesIO(VALID_PNG_BYTES), 'photo.png')
        self.assertEqual(self.client.post('/admin/shipped-orders/new', data=data).location, first.location)
        self.assertEqual(len(list(app.SHIPMENT_IMAGES_DIR.iterdir())), 1)

    def test_note_creation_failure_rolls_back_stock_images_and_operation_allows_retry(self):
        for mode in ['ordinary', 'assembly']:
            with self.subTest(mode=mode):
                route = '/admin/shipped-orders/new' if mode == 'ordinary' else '/admin/shipped-orders/assembly/new'
                data = self.ordinary_data() if mode == 'ordinary' else self.save_data(self.post_preview(sets=2).get_json(), extra_data={'operation_token': 'assembly-fail'})
                data['images'] = (BytesIO(VALID_PNG_BYTES), 'photo.png')
                with patch.object(app, 'create_delivery_notes', side_effect=sqlite3.OperationalError('injected note write failure')):
                    response = self.client.post(route, data=data)
                self.assertEqual(response.status_code, 500)
                with app.get_db() as conn:
                    for table in ['delivery_operations', 'delivery_notes', 'delivery_note_sources', 'product_order_shipments', 'assembly_shipment_batches']:
                        self.assertEqual(conn.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0], 0)
                    self.assertEqual(app.inventory_total_for_manual(conn, self.product), 20)
                self.assertEqual(list(app.SHIPMENT_IMAGES_DIR.iterdir()), [])
        data = self.ordinary_data()
        self.assertEqual(self.client.post('/admin/shipped-orders/new', data=data).status_code, 302)
        self.assertEqual(len(self.note_rows()), 1)

    def test_ordinary_edit_keeps_snapshot_and_delete_invalidates_download(self):
        data = self.ordinary_data()
        data['shipped_quantity'] = '2'
        self.client.post('/admin/shipped-orders/new', data=data)
        note_id = self.note_rows()[0]['id']
        with app.get_db() as conn:
            sid = conn.execute('SELECT id FROM product_order_shipments').fetchone()[0]
            conn.execute("UPDATE delivery_notes SET updated_at='2000-01-01'")
        self.client.post(f'/admin/shipped-orders/{sid}/edit', data={'shipped_quantity': '3', 'shipped_at': '2026-09-09'})
        with app.get_db() as conn:
            note = shipping_workflow.load_delivery_note(conn, note_id)
        self.assertEqual(note['items'][0]['shipped_quantity'], 2)
        self.assertNotEqual(note['updated_at'], '2000-01-01')
        self.client.post(f'/admin/shipped-orders/{sid}/delete')
        self.assertEqual(self.client.get(f'/admin/delivery-notes/{note_id}.pdf').status_code, 410)
        self.assertEqual(self.note_rows()[0]['invalidated'], 1)

    def test_pdf_retry_is_read_only_and_permission_protected(self):
        self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        note = self.note_rows()[0]
        route = f"/admin/delivery-notes/{note['id']}.pdf"
        with patch.object(app, 'build_delivery_note_pdf', side_effect=OSError('render failure')):
            self.assertEqual(self.client.get(route).status_code, 503)
        pdf = self.client.get(route)
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.data.startswith(b'%PDF'))
        self.assertEqual(self.note_rows(), [note])
        with self.client.session_transaction() as session:
            session.clear()
        self.assertNotEqual(self.client.get(route).status_code, 200)

    def test_assembly_item_replacement_keeps_note_and_delete_invalidates(self):
        preview = self.post_preview(sets=2).get_json()
        first = self.client.post('/admin/shipped-orders/assembly/new', data=self.save_data(preview))
        batch_id = first.get_json()['batch_id']
        note_id = self.note_rows()[0]['id']
        with app.get_db() as conn:
            item_id = conn.execute('SELECT id FROM assembly_shipment_items').fetchone()[0]
        edit_route = f'/admin/shipped-orders/assembly/{batch_id}/edit'
        edited_preview = self.client.post(edit_route, json={
            'customer': '客户A', 'assembly_drawing_no': 'ASM-100', 'set_quantity': 3,
            'selected_manual_ids': [self.product], 'overrides': {self.product: 3}}).get_json()
        edited = self.client.post(edit_route, data=self.save_data(edited_preview))
        self.assertEqual(edited.status_code, 201, edited.get_json())
        with app.get_db() as conn:
            self.assertNotEqual(conn.execute('SELECT id FROM assembly_shipment_items').fetchone()[0], item_id)
            note = shipping_workflow.load_delivery_note(conn, note_id)
        self.assertEqual(note['invalidated'], 0)
        self.assertEqual(note['items'][0]['shipped_quantity'], 2)
        self.assertEqual(note['recipient_name'], '张师傅')
        self.assertEqual(self.client.get(f'/admin/delivery-notes/{note_id}.pdf').status_code, 200)
        self.client.post(f'/admin/shipped-orders/assembly/{batch_id}/delete')
        self.assertEqual(self.client.get(f'/admin/delivery-notes/{note_id}.pdf').status_code, 410)

    def test_saved_note_link_is_reachable_from_history(self):
        self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        note = self.note_rows()[0]
        html = self.client.get('/admin/shipped-orders').get_data(as_text=True)
        self.assertIn(f'/admin/delivery-notes/operations/{note["operation_id"]}', html)
        self.assertIn(note['document_no'], html)

    def test_legacy_pdf_keeps_specification_unit_and_warns_missing_recipient(self):
        self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        with app.get_db() as conn:
            shipments = app.fetch_shipped_orders(conn)
        self.assertIn('unit', shipments[0].keys())
        self.assertEqual(shipments[0]['unit'], '件')
        from tests.test_shipped_pdf_company import ShippedPdfCompanyTests
        story = []
        with patch.object(app.SimpleDocTemplate, 'build', lambda _doc, values: story.extend(values)):
            app.build_shipped_orders_pdf(shipments, '', '客户A', '')
        text = ShippedPdfCompanyTests.collect_story_text(story)
        for content in ['长规格', '件', '收货资料未完整填写']:
            self.assertIn(content, text)
        self.assertNotIn('单价', text)

    def test_committed_context_exit_error_preserves_receipt_files_and_retry(self):
        for mode in ['ordinary', 'assembly']:
            with self.subTest(mode=mode):
                route = '/admin/shipped-orders/new' if mode == 'ordinary' else '/admin/shipped-orders/assembly/new'
                data = self.ordinary_data() if mode == 'ordinary' else self.save_data(self.post_preview(sets=2).get_json(), extra_data={'operation_token': 'assembly-exit'})
                if mode == 'ordinary':
                    data['shipped_quantity'] = '1'
                original_get_db = app.get_db

                @contextmanager
                def fail_after_commit():
                    with original_get_db() as conn:
                        before = conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0]
                        yield conn
                        after = conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0]
                    if after > before:
                        raise RuntimeError('injected committed context exit')

                data['images'] = (BytesIO(VALID_PNG_BYTES), 'exit.png')
                with patch.object(app, 'get_db', fail_after_commit):
                    first = self.client.post(route, data=data)
                self.assertEqual(first.status_code, 302 if mode == 'ordinary' else 201)
                data['images'] = (BytesIO(VALID_PNG_BYTES), 'exit.png')
                retry = self.client.post(route, data=data)
                self.assertEqual(retry.location, first.location)
                self.assertEqual(retry.get_json(silent=True), first.get_json(silent=True))
        with app.get_db() as conn:
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 17)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0], 2)
        self.assertEqual(len(list(app.SHIPMENT_IMAGES_DIR.iterdir())), 2)

    def test_no_shipment_permission_cannot_read_result_or_pdf(self):
        first = self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        note = self.note_rows()[0]
        with app.get_db() as conn:
            conn.execute("""INSERT INTO users (username, password_hash, role, active,
                can_view_shipped, can_manage_shipped, created_at, updated_at)
                VALUES ('no-shipped', 'unused', 'operator', 1, 0, 0, '', '')""")
        with self.client.session_transaction() as session:
            session['admin_username'] = 'no-shipped'
            session['admin_role'] = 'operator'
        for route in [first.location, f'/admin/delivery-notes/{note["id"]}.pdf']:
            self.assertNotEqual(self.client.get(route).status_code, 200)

    def test_concurrent_same_intent_creates_only_one_operation_per_mode(self):
        for mode in ['ordinary', 'assembly']:
            with self.subTest(mode=mode):
                data = self.ordinary_data() if mode == 'ordinary' else self.save_data(self.post_preview(sets=2).get_json(), extra_data={'operation_token': 'concurrent-assembly'})
                if mode == 'ordinary':
                    data['shipped_quantity'] = '2'
                route = '/admin/shipped-orders/new' if mode == 'ordinary' else '/admin/shipped-orders/assembly/new'
                def submit(_):
                    client = app.app.test_client()
                    with client.session_transaction() as session:
                        session.update(admin_logged_in=True, admin_username='admin', admin_role='admin')
                    response = client.post(route, data=data)
                    return response.status_code, response.location, response.get_json(silent=True)
                with ThreadPoolExecutor(max_workers=2) as pool:
                    responses = list(pool.map(submit, [1, 2]))
                self.assertEqual(responses[0], responses[1])
                self.assertEqual(responses[0][0], 302 if mode == 'ordinary' else 201)
        with app.get_db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_operations').fetchone()[0], 2)
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 16)

    def test_completed_assembly_retry_survives_connection_exit_error(self):
        data = self.save_data(self.post_preview(sets=2).get_json(), extra_data={'operation_token': 'retry-exit'})
        first = self.client.post('/admin/shipped-orders/assembly/new', data=data)
        original_get_db = app.get_db

        @contextmanager
        def fail_after_receipt_lookup():
            with original_get_db() as conn:
                yield conn
                was_transaction = conn.in_transaction
            if was_transaction:
                raise OSError('receipt lookup connection exit')

        with patch.object(app, 'get_db', fail_after_receipt_lookup):
            retry = self.client.post('/admin/shipped-orders/assembly/new', data=data)
        self.assertEqual(retry.status_code, 201)
        self.assertEqual(retry.get_json(), first.get_json())

    def test_result_contains_only_current_operation_and_no_price_fields(self):
        first = self.client.post('/admin/shipped-orders/new', data={**self.ordinary_data(), 'shipped_quantity': '2'})
        second = self.client.post('/admin/shipped-orders/new', data={**self.ordinary_data(), 'operation_token': 'second', 'shipped_quantity': '3'})
        notes = self.note_rows()
        html = self.client.get(second.location).get_data(as_text=True)
        self.assertIn(notes[1]['document_no'], html)
        self.assertNotIn(notes[0]['document_no'], html)
        self.assertNotIn('单价', html)
        self.assertNotEqual(first.location, second.location)

    def test_invalid_image_and_inventory_failure_leave_no_receipt_and_retry(self):
        data = self.ordinary_data()
        data['images'] = (BytesIO(b'not an image'), 'broken.png')
        self.client.post('/admin/shipped-orders/new', data=data)
        with patch.object(app, 'deduct_inventory_for_shipment', side_effect=sqlite3.OperationalError('inventory write failed')):
            self.assertEqual(self.client.post('/admin/shipped-orders/new', data=self.ordinary_data()).status_code, 500)
        with app.get_db() as conn:
            for table in ['delivery_operations', 'delivery_notes', 'delivery_note_sources', 'product_order_shipments']:
                self.assertEqual(conn.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0], 0)
            self.assertEqual(app.inventory_total_for_manual(conn, self.product), 20)
        self.assertEqual(list(app.SHIPMENT_IMAGES_DIR.iterdir()), [])
        self.assertEqual(self.client.post('/admin/shipped-orders/new', data=self.ordinary_data()).status_code, 302)

    def test_ordinary_recipient_browser_defaults_grouping_and_edits(self):
        html = self.client.get('/admin/shipped-orders/create').get_data(as_text=True)
        defaults = re.search(r'data-recipient-defaults>(.*?)</script>', html, re.S)[1]
        script = next(text for text in re.findall(r'<script>(.*?)</script>', html, re.S) if 'data-shipment-create-form' in text)
        harness = r'''
        const vm=require('vm'), assert=require('assert');
        const payload=JSON.parse(require('fs').readFileSync(0,'utf8'));
        class Element {
          constructor(){this.value='';this.children=[];this.listeners={};this.dataset={};}
          addEventListener(type, fn){(this.listeners[type] ||= []).push(fn);}
          emit(type){for(const fn of this.listeners[type] || []) fn();}
          append(...items){this.children.push(...items);}
          replaceChildren(...items){this.children=items;}
          removeAttribute(name){delete this[name];}
          set innerHTML(value){throw Error('unsafe HTML');}
        }
        const keys=['lines','add-shipment-line','order-search','customer-search','product-search','recipient-defaults','recipient-groups','recipient-overrides'];
        const nodes=Object.fromEntries(keys.map(key=>[key,new Element()]));
        nodes['recipient-defaults'].textContent=payload.defaults;
        let rows=[];
        function line(){
          const select=new Element(),quantity=new Element(),remove=new Element();
          select.options=[{value:'',dataset:{}},{value:'1',dataset:{customer:'客户A',unshipped:'20'}},{value:'2',dataset:{customer:'客户B',unshipped:'20'}}];
          Object.defineProperty(select,'selectedOptions',{get:()=>[select.options.find(o=>o.value===select.value)]});
          const row={select,querySelector:s=>s.includes('order-select')?select:s.includes('quantity')?quantity:remove,remove:()=>{rows=rows.filter(r=>r!==row)}};
          return row;
        }
        rows=[line(),line()];
        const form={querySelectorAll:()=>rows,querySelector:s=>{
          const name=s.slice(1,-1).replace(/^data-shipment-/,'').replace(/^data-/,'');
          return nodes[name];
        }};
        global.document={querySelector:()=>form,createElement:()=>new Element()};
        const recipientController=require('fs').readFileSync('static/order_shipping.js','utf8')
          .split('function initializeOrderShipmentLines',1)[0];
        vm.runInThisContext(recipientController);
        vm.runInThisContext(payload.script);
        rows[0].select.value='1'; rows[0].select.emit('change');
        assert.equal(nodes['recipient-groups'].children.length,1);
        let active=JSON.parse(nodes['recipient-overrides'].value);
        assert.equal(active['客户A'].recipient_name,'张师傅');
        const all=el=>[el,...el.children.flatMap(all)];
        const nameInput=all(nodes['recipient-groups']).find(el=>el.value==='张师傅');
        nameInput.value='临时收货人';nameInput.emit('input');
        rows[1].select.value='2'; rows[1].select.emit('change');
        active=JSON.parse(nodes['recipient-overrides'].value);
        assert.equal(nodes['recipient-groups'].children.length,2);
        assert.equal(active['客户A'].recipient_name,'临时收货人');
        assert.equal(active['客户B'].recipient_name,'');
        rows[1].querySelector('[data-remove-shipment-line]').emit('click');
        active=JSON.parse(nodes['recipient-overrides'].value);
        assert.deepEqual(Object.keys(active),['客户A']);
        assert.equal(active['客户A'].recipient_name,'临时收货人');
        '''
        result = subprocess.run(['node', '-e', harness], input=json.dumps({'defaults': defaults, 'script': script}), capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_assembly_note_uses_actual_allocated_orders_not_assembly_drawing_number(self):
        self.create_order(self.product, 'SO-SECOND', 10)
        preview = self.post_preview(sets=25).get_json()
        self.client.post('/admin/shipped-orders/assembly/new', data=self.save_data(preview))
        note_id = self.note_rows()[0]['id']
        with app.get_db() as conn:
            note = shipping_workflow.load_delivery_note(conn, note_id)
        self.assertIn('SO-DN', note['items'][0]['order_no'])
        self.assertIn('SO-SECOND', note['items'][0]['order_no'])
        self.assertNotIn('ASM-100', note['items'][0]['order_no'])
        extra = self.create_product('NO-ORDER')
        preview = self.client.post('/admin/shipped-orders/assembly-preview', json={
            'customer': '客户A', 'assembly_drawing_no': 'ASM-100', 'set_quantity': 1,
            'selected_manual_ids': [extra], 'overrides': {extra: 1}}).get_json()
        response = self.client.post('/admin/shipped-orders/assembly/new', data=self.save_data(
            preview, extra_data={'selected_manual_ids': [str(extra)]}))
        self.assertEqual(response.status_code, 201, response.get_json())
        last_note_id = self.note_rows()[-1]['id']
        with app.get_db() as conn:
            note = shipping_workflow.load_delivery_note(conn, last_note_id)
        self.assertEqual(note['items'][0]['order_no'], '未关联订单')

    def test_assembly_note_retains_mixed_allocated_and_unallocated_order_quantities(self):
        preview = self.post_preview(sets=10).get_json()
        response = self.client.post('/admin/shipped-orders/assembly/new', data=self.save_data(preview))
        self.assertEqual(response.status_code, 201, response.get_json())
        note_id = self.note_rows()[0]['id']
        with app.get_db() as conn:
            # Historical notes without line snapshots still read live allocations.
            conn.execute('DELETE FROM delivery_note_items WHERE note_id=?', (note_id,))
            item = conn.execute('SELECT id FROM assembly_shipment_items').fetchone()
            conn.execute(
                'UPDATE assembly_shipment_allocations SET quantity=6 WHERE item_id=?',
                (item['id'],),
            )
            conn.execute(
                '''INSERT INTO assembly_shipment_allocations
                   (item_id, order_id, quantity, created_at) VALUES (?, NULL, 4, '')''',
                (item['id'],),
            )
            note = shipping_workflow.load_delivery_note(conn, note_id)
        self.assertEqual(note['items'][0]['order_no'], 'SO-DN / 未关联订单（4）')
        story = []
        from tests.test_shipped_pdf_company import ShippedPdfCompanyTests
        with patch.object(app.SimpleDocTemplate, 'build', lambda _doc, values: story.extend(values)):
            app.build_shipped_orders_pdf(note['items'], '', note['customer'], '', recipient_metadata=note)
        text = ShippedPdfCompanyTests.collect_story_text(story)
        self.assertIn('SO-DN', text)
        self.assertIn('未关联订单（4）', text)

    def test_assembly_legacy_specification_fallback_is_marked_but_saved_empty_is_not(self):
        preview = self.post_preview(sets=2).get_json()
        response = self.client.post('/admin/shipped-orders/assembly/new', data=self.save_data(preview))
        self.assertEqual(response.status_code, 201, response.get_json())
        note_id = self.note_rows()[0]['id']
        with app.get_db() as conn:
            conn.execute('DELETE FROM delivery_note_items WHERE note_id=?', (note_id,))
            item_id = conn.execute('SELECT id FROM assembly_shipment_items').fetchone()[0]
            conn.execute('UPDATE assembly_shipment_items SET specification_snapshot=NULL WHERE id=?', (item_id,))
            conn.execute("UPDATE manuals SET supplier='当前组装规格' WHERE id=?", (self.product,))
            note = shipping_workflow.load_delivery_note(conn, note_id)
        self.assertEqual(note['items'][0]['specification'], '当前组装规格')
        self.assertEqual(note['items'][0]['specification_is_fallback'], 1)
        self.assertIn('当前规格，历史未保存', self.client.get(response.get_json()['redirect_url']).get_data(as_text=True))
        with app.get_db() as conn:
            conn.execute("UPDATE assembly_shipment_items SET specification_snapshot='' WHERE id=?", (item_id,))
            note = shipping_workflow.load_delivery_note(conn, note_id)
        self.assertEqual(note['items'][0]['specification'], '')
        self.assertEqual(note['items'][0]['specification_is_fallback'], 0)

    def test_delivery_timestamp_display_converts_utc_without_changing_snapshot(self):
        self.client.post('/admin/shipped-orders/new', data=self.ordinary_data())
        note = self.note_rows()[0]
        with app.get_db() as conn:
            conn.execute("UPDATE delivery_notes SET updated_at='2026-09-08T15:03:25.123456+00:00' WHERE id=?", (note['id'],))
            note = shipping_workflow.load_delivery_note(conn, note['id'])
        from tests.test_shipped_pdf_company import ShippedPdfCompanyTests
        story = []
        with patch.object(app.SimpleDocTemplate, 'build', lambda _doc, values: story.extend(values)):
            app.build_shipped_orders_pdf(note['items'], '', note['customer'], '', recipient_metadata=note)
        text = ShippedPdfCompanyTests.collect_story_text(story)
        self.assertIn('2026-09-08 23:03:25（北京时间）', text)
        html = self.client.get(f'/admin/delivery-notes/operations/{note["operation_id"]}').get_data(as_text=True)
        self.assertIn('2026-09-08 23:03:25（北京时间）', html)
        self.assertEqual(self.note_rows()[0]['updated_at'], '2026-09-08T15:03:25.123456+00:00')

    def save_both_note_modes(self):
        self.client.post('/admin/shipped-orders/new', data={**self.ordinary_data(), 'shipped_quantity': '2'})
        response = self.client.post('/admin/shipped-orders/assembly/new', data=self.save_data(self.post_preview(sets=2).get_json()))
        self.assertEqual(response.status_code, 201)
        return self.note_rows()

    def rename_customer(self, name):
        with app.get_db() as conn:
            customer_id = conn.execute("SELECT id FROM customers WHERE name='客户A'").fetchone()[0]
        return self.client.post(f'/admin/customers/{customer_id}/edit', data={
            'name': name, 'recipient_name': '新联系人', 'recipient_phone': '456', 'address': '新地址'})

    def test_customer_rename_preserves_identity_and_frozen_display_for_both_modes(self):
        notes = self.save_both_note_modes()
        with app.get_db() as conn:
            conn.execute("UPDATE delivery_notes SET updated_at='2000-01-01'")
        self.assertEqual(self.rename_customer('客户新名称').status_code, 302)
        for saved in notes:
            with app.get_db() as conn:
                note = shipping_workflow.load_delivery_note(conn, saved['id'])
            self.assertEqual(note['invalidated'], 0)
            self.assertEqual(note['customer'], '客户A')
            self.assertEqual(note['recipient_name'], '张师傅')
            self.assertEqual(note['recipient_phone'], '13800000000')
            self.assertEqual(note['address'], '宁波仓库')
            self.assertNotEqual(note['updated_at'], '2000-01-01')
            self.assertEqual(self.client.get(f'/admin/delivery-notes/{saved["id"]}.pdf').status_code, 200)

    def test_customer_rename_keeps_supplemental_note_valid_and_moves_finance_source(self):
        from tests.test_supplemental_shipments import insert_supplemental
        with app.get_db() as conn:
            operation_id, sid = insert_supplemental(conn, self.product)
            note_id = shipping_workflow.create_delivery_notes(conn, operation_id,
                [('supplemental', sid)], {}, 'shipper')[0]
            original = shipping_workflow.load_delivery_note(conn, note_id)
        self.assertEqual(self.rename_customer('客户新名称').status_code, 302)
        with app.get_db() as conn:
            source = conn.execute('SELECT customer FROM supplemental_shipments WHERE id=?', (sid,)).fetchone()
            self.assertEqual(source['customer'], '客户新名称')
            note = shipping_workflow.load_delivery_note(conn, note_id)
            self.assertEqual(note['invalidated'], 0)
            self.assertEqual(note['items'], original['items'])
            for field in ('customer_id', 'customer', 'recipient_name', 'recipient_phone', 'address'):
                self.assertEqual(note[field], original[field])
            self.assertIn(('supplemental', sid), {(s['source_type'], s['source_id'])
                for s in app.fetch_available_finance_sources(conn, '客户新名称')})
            self.assertNotIn(('supplemental', sid), {(s['source_type'], s['source_id'])
                for s in app.fetch_available_finance_sources(conn, '客户A')})
        pdf = self.client.get(f'/admin/delivery-notes/{note_id}.pdf')
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.data.startswith(b'%PDF'))

    def test_supplemental_rename_failure_rolls_back_customer_and_related_sources(self):
        from tests.test_supplemental_shipments import insert_supplemental
        with app.get_db() as conn:
            _, sid = insert_supplemental(conn, self.product)
            conn.execute("""CREATE TRIGGER reject_supplemental_rename BEFORE UPDATE OF customer
                ON supplemental_shipments BEGIN SELECT RAISE(ABORT, 'rename failure'); END""")
        self.assertEqual(self.rename_customer('客户新名称').status_code, 302)
        with app.get_db() as conn:
            self.assertIsNotNone(conn.execute("SELECT id FROM customers WHERE name='客户A'").fetchone())
            self.assertIsNone(conn.execute("SELECT id FROM customers WHERE name='客户新名称'").fetchone())
            self.assertEqual(conn.execute('SELECT customer FROM manuals WHERE id=?', (self.product,)).fetchone()[0], '客户A')
            self.assertEqual(conn.execute('SELECT customer FROM product_orders WHERE id=?', (self.order,)).fetchone()[0], '客户A')
            self.assertEqual(conn.execute('SELECT customer FROM supplemental_shipments WHERE id=?', (sid,)).fetchone()[0], '客户A')

    def test_true_customer_reassignment_invalidates_both_note_modes(self):
        notes = self.save_both_note_modes()
        with app.get_db() as conn:
            conn.execute("INSERT INTO customers (name,created_at,updated_at) VALUES ('客户B','','')")
            conn.execute("UPDATE product_orders SET customer='客户B'")
            conn.execute("UPDATE assembly_shipment_batches SET customer='客户B'")
        for note in notes:
            self.assertEqual(self.client.get(f'/admin/delivery-notes/{note["id"]}.pdf').status_code, 410)

    def test_customer_id_does_not_fall_back_to_same_name_after_identity_replacement(self):
        notes = self.save_both_note_modes()
        with app.get_db() as conn:
            conn.execute("DELETE FROM customers WHERE name='客户A'")
            conn.execute("INSERT INTO customers (name,created_at,updated_at) VALUES ('客户A','','')")
        for note in notes:
            self.assertEqual(self.client.get(f'/admin/delivery-notes/{note["id"]}.pdf').status_code, 410)

    def test_null_legacy_notes_bind_only_valid_ownership_during_controlled_rename(self):
        notes = self.save_both_note_modes()
        with app.get_db() as conn:
            conn.execute('UPDATE delivery_notes SET customer_id=NULL')
        self.rename_customer('客户新名称')
        for saved in notes:
            with app.get_db() as conn:
                note = shipping_workflow.load_delivery_note(conn, saved['id'])
            self.assertEqual(note['invalidated'], 0)
            self.assertIsNotNone(note['customer_id'])
            self.assertEqual(note['customer'], '客户A')
            self.assertEqual(self.client.get(f'/admin/delivery-notes/{saved["id"]}.pdf').status_code, 200)

    def test_unproven_or_invalidated_legacy_note_is_not_rebound_on_rename(self):
        notes = self.save_both_note_modes()
        with app.get_db() as conn:
            conn.execute('UPDATE delivery_notes SET customer_id=NULL')
            conn.execute("UPDATE product_orders SET customer='另一客户'")
            conn.execute('UPDATE delivery_notes SET invalidated=1 WHERE id=?', (notes[1]['id'],))
        self.rename_customer('客户新名称')
        for saved in notes:
            with app.get_db() as conn:
                note = shipping_workflow.load_delivery_note(conn, saved['id'])
            self.assertEqual(note['invalidated'], 1)
            self.assertIsNone(note['customer_id'])

    def test_legacy_identity_binding_rolls_back_with_failed_customer_rename(self):
        self.save_both_note_modes()
        with app.get_db() as conn:
            conn.execute('UPDATE delivery_notes SET customer_id=NULL')
            conn.execute("INSERT INTO customers (name,created_at,updated_at) VALUES ('重名客户','','')")
        self.rename_customer('重名客户')
        with app.get_db() as conn:
            self.assertIsNotNone(conn.execute("SELECT id FROM customers WHERE name='客户A'").fetchone())
            for row in conn.execute('SELECT customer_id FROM delivery_notes'):
                self.assertIsNone(row['customer_id'])

    def test_unregistered_and_nameless_customer_notes_use_strict_name_fallback(self):
        # Complete the app's first-request customer backfill before simulating
        # a name-only source with no customer master record.
        self.client.get('/admin/shipped-orders')
        with app.get_db() as conn:
            conn.execute("DELETE FROM customers WHERE name='客户A'")
        notes = self.save_both_note_modes()
        for saved in notes:
            with app.get_db() as conn:
                note = shipping_workflow.load_delivery_note(conn, saved['id'])
            self.assertIsNone(note['customer_id'])
            self.assertEqual(note['invalidated'], 0)
        with app.get_db() as conn:
            conn.execute("UPDATE product_orders SET customer=''")
            conn.execute("UPDATE manuals SET customer=''")
        response = self.client.post('/admin/shipped-orders/new', data={
            **self.ordinary_data(), 'shipped_quantity': '1', 'operation_token': 'nameless'})
        self.assertEqual(response.status_code, 302)
        nameless_id = self.note_rows()[-1]['id']
        with app.get_db() as conn:
            note = shipping_workflow.load_delivery_note(conn, nameless_id)
            self.assertEqual(note['customer'], '')
            self.assertEqual(note['invalidated'], 0)
            conn.execute("UPDATE product_orders SET customer='新归属'")
        self.assertEqual(self.client.get(f'/admin/delivery-notes/{nameless_id}.pdf').status_code, 410)

    def test_customer_identity_migration_is_additive_and_idempotent(self):
        notes = self.save_both_note_modes()
        with app.get_db() as conn:
            conn.execute('ALTER TABLE delivery_notes DROP COLUMN customer_id')
            before = [dict(row) for row in conn.execute('SELECT * FROM delivery_notes')]
            shipping_workflow.ensure_shipping_workflow_tables(conn)
            shipping_workflow.ensure_shipping_workflow_tables(conn)
            after = [dict(row) for row in conn.execute('SELECT * FROM delivery_notes')]
        for old, migrated in zip(before, after):
            self.assertIsNone(migrated.pop('customer_id'))
            self.assertEqual(old, migrated)
        for saved in notes:
            self.assertEqual(self.client.get(f'/admin/delivery-notes/{saved["id"]}.pdf').status_code, 200)

import base64
import re
import sqlite3
import subprocess
import sys
import tempfile
import textwrap
import unittest
from decimal import Decimal
from io import BytesIO
from datetime import datetime, timedelta
from pathlib import Path

import app
from openpyxl import load_workbook
from reconciliation import (
    build_reconciliation_workbook,
    calculate_reconciliation_amounts,
    format_reconciliation_amount_minor,
    format_unit_price_ex_tax_scaled,
)


class ReconciliationCalculationTests(unittest.TestCase):
    def test_13_percent_tax_is_back_calculated_from_tax_inclusive_price(self):
        amounts = calculate_reconciliation_amounts(11300, 2, "CNY")
        self.assertEqual(
            amounts,
            {
                "unit_price_incl_tax_minor": 11300,
                "unit_price_ex_tax_scaled": 100_000_000,
                "amount_incl_tax_minor": 22600,
                "amount_ex_tax_minor": 20000,
                "tax_amount_minor": 2600,
            },
        )

    def test_zero_price_keeps_every_reconciliation_amount_at_zero(self):
        self.assertEqual(
            calculate_reconciliation_amounts(0, 3, "CNY"),
            {
                "unit_price_incl_tax_minor": 0,
                "unit_price_ex_tax_scaled": 0,
                "amount_incl_tax_minor": 0,
                "amount_ex_tax_minor": 0,
                "tax_amount_minor": 0,
            },
        )

    def test_jpy_uses_zero_decimal_currency_precision_for_unit_amount(self):
        self.assertEqual(
            calculate_reconciliation_amounts(113, 2, "JPY"),
            {
                "unit_price_incl_tax_minor": 113,
                "unit_price_ex_tax_scaled": 100_000_000,
                "amount_incl_tax_minor": 226,
                "amount_ex_tax_minor": 200,
                "tax_amount_minor": 26,
            },
        )

    def test_non_exact_jpy_tax_division_rounds_minor_and_scaled_amounts_half_up(self):
        self.assertEqual(
            calculate_reconciliation_amounts(13, 1, "JPY"),
            {
                "unit_price_incl_tax_minor": 13,
                "unit_price_ex_tax_scaled": 11_504_425,
                "amount_incl_tax_minor": 13,
                "amount_ex_tax_minor": 12,
                "tax_amount_minor": 1,
            },
        )

    def test_non_exact_tax_division_rounds_minor_and_scaled_amounts_half_up(self):
        self.assertEqual(
            calculate_reconciliation_amounts(10000, 1, "CNY"),
            {
                "unit_price_incl_tax_minor": 10000,
                "unit_price_ex_tax_scaled": 88_495_575,
                "amount_incl_tax_minor": 10000,
                "amount_ex_tax_minor": 8850,
                "tax_amount_minor": 1150,
            },
        )

    def test_invalid_currency_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "请选择有效币种"):
            calculate_reconciliation_amounts(100, 1, "XYZ")

    def test_non_positive_quantity_is_rejected(self):
        for quantity in (0, -1):
            with self.subTest(quantity=quantity):
                with self.assertRaisesRegex(ValueError, "对账价格和数量必须有效"):
                    calculate_reconciliation_amounts(100, quantity, "CNY")

    def test_negative_price_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "对账价格和数量必须有效"):
            calculate_reconciliation_amounts(-1, 1, "CNY")

    def test_total_larger_than_sqlite_integer_range_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "对账金额过大"):
            calculate_reconciliation_amounts(2**63 - 1, 2, "CNY")

    def test_scaled_unit_price_larger_than_sqlite_integer_range_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "对账金额过大"):
            calculate_reconciliation_amounts(2**63 - 1, 1, "CNY")

    def test_display_helpers_use_the_shared_stored_amount_scales(self):
        self.assertEqual(format_unit_price_ex_tax_scaled(100_000_000), "100.000000")
        self.assertEqual(format_reconciliation_amount_minor(20000, "CNY"), "200.00")
        self.assertEqual(format_reconciliation_amount_minor(200, "JPY"), "200")


class ReconciliationMigrationTests(unittest.TestCase):
    def setUp(self):
        self.original_db_path = app.DB_PATH
        self.original_database_ready = app.DATABASE_READY
        self.tmpdir = tempfile.TemporaryDirectory()
        app.DB_PATH = Path(self.tmpdir.name) / "manuals.db"
        app.DATABASE_READY = False

    def tearDown(self):
        app.DB_PATH = self.original_db_path
        app.DATABASE_READY = self.original_database_ready
        self.tmpdir.cleanup()

    def test_reconciliation_schema_is_idempotent_and_claim_is_unique(self):
        app.init_db()
        app.init_db()
        with app.get_db() as conn:
            names = {
                row["name"]
                for row in conn.execute(
                    "SELECT name FROM sqlite_master WHERE type = 'table'"
                )
            }
            self.assertTrue(
                {
                    "reconciliation_company_profile",
                    "reconciliation_statements",
                    "reconciliation_statement_items",
                }
                <= names
            )
            profile = conn.execute(
                "SELECT company_name, address, contact, phone, email FROM reconciliation_company_profile WHERE id = 1"
            ).fetchone()
            self.assertEqual(
                dict(profile),
                {
                    "company_name": "宁波市杰德机械科技有限公司",
                    "address": "",
                    "contact": "",
                    "phone": "",
                    "email": "",
                },
            )

            customer_id = conn.execute(
                "INSERT INTO customers (name, created_at, updated_at) VALUES ('客户A', '2026-09-08', '2026-09-08')"
            ).lastrowid
            statement_id = self.insert_statement(conn, customer_id)
            self.insert_item(conn, statement_id, "ordinary", "ordinary:1")
            with self.assertRaises(sqlite3.IntegrityError):
                self.insert_item(conn, statement_id, "ordinary", "ordinary:1")
            with self.assertRaises(sqlite3.IntegrityError):
                self.insert_item(conn, statement_id, "invalid", "invalid:1")

    @staticmethod
    def insert_statement(conn, customer_id):
        return conn.execute(
            """
            INSERT INTO reconciliation_statements (
                statement_no, period_start, period_end, customer_id, customer_name,
                supplier_company_name, currency, tax_rate_ppm, amount_ex_tax_minor,
                amount_incl_tax_minor, tax_amount_minor, token_digest, password_hash,
                access_expires_at, created_by, created_at, updated_at
            ) VALUES (
                'REC-001', '2026-09-01', '2026-09-30', ?, '客户A',
                '宁波市杰德机械科技有限公司', 'CNY', 130000, 100, 113, 13,
                'token-001', 'hash', '2026-10-01', 'admin', '2026-09-08', '2026-09-08'
            )
            """,
            (customer_id,),
        ).lastrowid

    @staticmethod
    def insert_item(conn, statement_id, source_type, claim_key):
        return conn.execute(
            """
            INSERT INTO reconciliation_statement_items (
                statement_id, sort_order, source_type, source_id, active_claim_key,
                shipped_at, quantity, currency, unit_price_incl_tax_minor,
                unit_price_ex_tax_scaled, amount_incl_tax_minor, amount_ex_tax_minor,
                tax_amount_minor, created_at
            ) VALUES (?, 1, ?, 1, ?, '2026-09-08', 1, 'CNY', 113, 1000000, 113, 100, 13, '2026-09-08')
            """,
            (statement_id, source_type, claim_key),
        ).lastrowid


class ReconciliationDomainTestCase(unittest.TestCase):
    """Real SQLite fixtures covering reconciliation's source boundary."""

    def setUp(self):
        self.original_db_path = app.DB_PATH
        self.original_database_ready = app.DATABASE_READY
        self.tmpdir = tempfile.TemporaryDirectory()
        app.DB_PATH = Path(self.tmpdir.name) / "manuals.db"
        app.DATABASE_READY = False
        app.init_db()
        self.now = "2026-09-08T10:00:00"
        with app.get_db() as conn:
            self.customer_a = self.create_customer(conn, "客户A")
            self.customer_b = self.create_customer(conn, "客户B")
            self.manual_a = self.create_manual(conn, "RA-100", "对账产品A", "客户A")
            self.manual_b = self.create_manual(conn, "RB-100", "对账产品B", "客户B")
            self.order_a = self.create_order(conn, self.manual_a, "SO-A", "客户A")
            self.order_b = self.create_order(conn, self.manual_b, "SO-B", "客户B")
            self.ordinary_a_cny = self.create_shipment(
                conn, self.order_a, 2, 11300, "CNY", "2026-09-01", "物流-A"
            )
            self.ordinary_a_usd = self.create_shipment(
                conn, self.order_a, 1, 113, "USD", "2026-09-02", ""
            )
            self.ordinary_b_cny = self.create_shipment(
                conn, self.order_b, 3, 22600, "CNY", "2026-09-03", ""
            )
            self.ordinary_a_unpriced = self.create_shipment(
                conn, self.order_a, 4, None, "CNY", "2026-09-04", ""
            )
            self.batch_a = conn.execute(
                """
                INSERT INTO assembly_shipment_batches (
                    customer, assembly_drawing_no, set_quantity, shipped_at,
                    logistics_no, created_by, created_at, updated_at
                ) VALUES ('客户A', 'ASM-A', 4, '2026-09-05', '组装物流-A', 'shipper', ?, ?)
                """,
                (self.now, self.now),
            ).lastrowid
            self.assembly_a_cny = self.create_assembly_item(
                conn, self.batch_a, self.manual_a, 4, 5650, "CNY"
            )
            conn.execute(
                """INSERT INTO assembly_shipment_allocations
                   (item_id, order_id, quantity, created_at) VALUES (?, ?, 4, ?)""",
                (self.assembly_a_cny, self.order_a, self.now),
            )

    def tearDown(self):
        app.DB_PATH = self.original_db_path
        app.DATABASE_READY = self.original_database_ready
        self.tmpdir.cleanup()

    def create_customer(self, conn, name):
        return conn.execute(
            """
            INSERT INTO customers (
                name, contact, address, phone, email, remark,
                invoice_title, tax_id, registered_address, registered_phone,
                bank_name, bank_account, invoice_email, created_at, updated_at
            ) VALUES (?, '采购', '客户地址', '0574-1', 'customer@example.com', '',
                      '客户抬头', 'TAX-1', '开票地址', '0574-2', '银行', '账号',
                      'invoice@example.com', ?, ?)
            """,
            (name, self.now, self.now),
        ).lastrowid

    def create_manual(self, conn, drawing_no, product_name, customer):
        return conn.execute(
            """
            INSERT INTO manuals (
                drawing_no, product_name, customer, sku, unit, model, category,
                version, filename, original_filename, created_at, updated_at
            ) VALUES (?, ?, ?, 'SKU-100', '件', 'M-100', '', '', '', '', ?, ?)
            """,
            (drawing_no, product_name, customer, self.now, self.now),
        ).lastrowid

    def create_order(self, conn, manual_id, order_no, customer):
        return conn.execute(
            """
            INSERT INTO product_orders (
                manual_id, order_no, ordered_at, quantity, customer,
                planned_ship_at, created_at, updated_at
            ) VALUES (?, ?, '2026-09-01', 100, ?, '2026-09-30', ?, ?)
            """,
            (manual_id, order_no, customer, self.now, self.now),
        ).lastrowid

    def create_shipment(self, conn, order_id, quantity, price, currency, shipped_at, logistics_no):
        return conn.execute(
            """
            INSERT INTO product_order_shipments (
                order_id, shipped_quantity, shipped_at, logistics_no, remark, created_at,
                unit_price_minor, currency, price_recorded_by, price_recorded_at
            ) VALUES (?, ?, ?, ?, '普通备注', ?, ?, ?, 'shipper', ?)
            """,
            (order_id, quantity, shipped_at, logistics_no, self.now, price, currency,
             self.now if price is not None else ""),
        ).lastrowid

    def create_assembly_item(self, conn, batch_id, manual_id, quantity, price, currency):
        return conn.execute(
            """
            INSERT INTO assembly_shipment_items (
                batch_id, manual_id, drawing_no, product_name,
                quantity_per_set, calculated_quantity, shipped_quantity,
                inventory_deducted_quantity, inventory_shortage_quantity,
                unit_price_minor, currency, price_recorded_by, price_recorded_at,
                created_at, updated_at
            ) VALUES (?, ?, 'RA-100', '组装配件A', 1, ?, ?, 0, ?, ?, ?, 'shipper', ?, ?, ?)
            """,
            (batch_id, manual_id, quantity, quantity, quantity, price, currency,
             self.now if price is not None else "", self.now, self.now),
        ).lastrowid


class ReconciliationCreationTests(ReconciliationDomainTestCase):
    def test_create_statement_snapshots_mixed_sources_and_exact_totals(self):
        with app.get_db() as conn:
            result = app.create_reconciliation_statement(
                conn,
                [("ordinary", self.ordinary_a_cny), ("assembly_item", self.assembly_a_cny)],
                "finance-user",
            )
            statement = app.fetch_reconciliation_statement(conn, result["statement_id"])
            items = app.fetch_reconciliation_statement_items(conn, result["statement_id"])
        self.assertEqual(statement["customer_name"], "客户A")
        self.assertEqual(statement["currency"], "CNY")
        self.assertEqual(statement["statement_no"], "DZ-202609-0001")
        self.assertEqual(len(items), 2)
        self.assertEqual(statement["amount_incl_tax_minor"], 45200)
        self.assertEqual(statement["amount_incl_tax_minor"], sum(row["amount_incl_tax_minor"] for row in items))
        self.assertEqual(statement["amount_ex_tax_minor"], 40000)
        self.assertEqual(statement["tax_amount_minor"], 5200)
        self.assertRegex(result["token"], r"^[A-Za-z0-9_-]{40,}$")
        self.assertRegex(result["password"], r"^[0-9]{6}$")
        self.assertEqual(items[0]["delivery_no"], "物流-A")
        self.assertEqual(items[1]["delivery_no"], f"ZP-{self.batch_a}")
        self.assertEqual(items[1]["order_no"], "SO-A")

    def test_rejects_invalid_or_incompatible_source_selection(self):
        cases = (
            ([], "请至少选择一条发货记录"),
            ([("ordinary", self.ordinary_a_cny), ("ordinary", self.ordinary_a_cny)], "不能重复选择同一发货记录"),
            ([("invalid", 1)], "发货记录来源无效"),
            ([("ordinary", 99999)], "发货记录不存在"),
            ([("ordinary", self.ordinary_a_cny), ("ordinary", self.ordinary_b_cny)], "只能合并同一客户、同一币种的发货记录"),
            ([("ordinary", self.ordinary_a_cny), ("ordinary", self.ordinary_a_usd)], "只能合并同一客户、同一币种的发货记录"),
            ([("ordinary", self.ordinary_a_unpriced)], "该发货记录尚未记录价格，不能生成对账单"),
        )
        for refs, message in cases:
            with self.subTest(message=message), app.get_db() as conn:
                with self.assertRaisesRegex(ValueError, message):
                    app.create_reconciliation_statement(conn, refs, "finance-user")

    def test_rejects_source_without_matching_customer_record(self):
        with app.get_db() as conn:
            conn.execute("DELETE FROM customers WHERE id = ?", (self.customer_a,))
            with self.assertRaisesRegex(ValueError, "客户不存在"):
                app.create_reconciliation_statement(
                    conn, [("ordinary", self.ordinary_a_cny)], "finance-user"
                )

    def test_snapshots_remain_unchanged_after_source_customer_and_profile_edits(self):
        with app.get_db() as conn:
            result = app.create_reconciliation_statement(
                conn, [("ordinary", self.ordinary_a_cny)], "finance-user"
            )
            before_statement = dict(app.fetch_reconciliation_statement(conn, result["statement_id"]))
            before_item = dict(app.fetch_reconciliation_statement_items(conn, result["statement_id"])[0])
            conn.execute("UPDATE customers SET address = '新客户地址' WHERE id = ?", (self.customer_a,))
            conn.execute("UPDATE manuals SET product_name = '新产品', sku = 'NEW' WHERE id = ?", (self.manual_a,))
            conn.execute("UPDATE reconciliation_company_profile SET company_name = '新供应商' WHERE id = 1")
            after_statement = app.fetch_reconciliation_statement(conn, result["statement_id"])
            after_item = app.fetch_reconciliation_statement_items(conn, result["statement_id"])[0]
        self.assertEqual(after_statement["customer_address"], before_statement["customer_address"])
        self.assertEqual(after_statement["supplier_company_name"], before_statement["supplier_company_name"])
        self.assertEqual(after_item["product_name"], before_item["product_name"])
        self.assertEqual(after_item["sku"], before_item["sku"])

    def test_rejects_a_source_with_an_active_reconciliation_claim(self):
        refs = [("ordinary", self.ordinary_a_cny)]
        with app.get_db() as conn:
            app.create_reconciliation_statement(conn, refs, "first")
            with self.assertRaisesRegex(ValueError, "该发货记录已加入其他对账单"):
                app.create_reconciliation_statement(conn, refs, "second")

    def test_assembly_item_uses_shipment_snapshot_identity_and_batch_logistics_remark(self):
        with app.get_db() as conn:
            conn.execute(
                """
                UPDATE assembly_shipment_items
                SET drawing_no = 'SHIP-RA-100', product_name = '发货时配件名'
                WHERE id = ?
                """,
                (self.assembly_a_cny,),
            )
            conn.execute(
                """
                UPDATE manuals
                SET drawing_no = 'LIVE-RA-100', product_name = '当前产品名'
                WHERE id = ?
                """,
                (self.manual_a,),
            )
            statement_id = app.create_reconciliation_statement(
                conn, [("assembly_item", self.assembly_a_cny)], "finance-user"
            )["statement_id"]
            item = app.fetch_reconciliation_statement_items(conn, statement_id)[0]
        self.assertEqual(item["drawing_no"], "SHIP-RA-100")
        self.assertEqual(item["product_name"], "发货时配件名")
        self.assertEqual(item["remark"], "组装物流-A")


class ReconciliationLifecycleTests(ReconciliationDomainTestCase):
    def test_void_releases_source_and_requires_reason(self):
        with app.get_db() as conn:
            statement_id = app.create_reconciliation_statement(
                conn, [("ordinary", self.ordinary_a_cny)], "finance-user"
            )["statement_id"]
            with self.assertRaisesRegex(ValueError, "请填写作废原因"):
                app.void_reconciliation_statement(conn, statement_id, "", "finance-user")
            app.void_reconciliation_statement(conn, statement_id, "金额有误", "finance-user")
            self.assertEqual(app.fetch_reconciliation_statement(conn, statement_id)["status"], "void")
            self.assertIsNone(app.fetch_reconciliation_statement_items(conn, statement_id)[0]["active_claim_key"])
            replacement = app.create_reconciliation_statement(
                conn, [("ordinary", self.ordinary_a_cny)], "finance-user"
            )
            self.assertNotEqual(replacement["statement_id"], statement_id)

    def test_active_claim_blocks_source_mutation_and_void_releases_it(self):
        refs = [("ordinary", self.ordinary_a_cny)]
        with app.get_db() as conn:
            statement_id = app.create_reconciliation_statement(conn, refs, "finance-user")["statement_id"]
            with self.assertRaisesRegex(ValueError, "已加入对账单"):
                app.assert_reconciliation_sources_mutable(conn, refs)
            app.void_reconciliation_statement(conn, statement_id, "金额有误", "finance-user")
            app.assert_reconciliation_sources_mutable(conn, refs)


class ReconciliationWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.statement = {
            "statement_no": "DZ-202609-0001",
            "period_start": "2026-09-01",
            "period_end": "2026-09-30",
            "customer_name": "客户A",
            "customer_address": "客户地址",
            "customer_phone": "0574-1",
            "customer_email": "customer@example.com",
            "customer_purchase_contact": "采购联系人",
            "customer_reconciliation_contact": "对账联系人",
            "supplier_company_name": "宁波市杰德机械科技有限公司",
            "supplier_address": "供应商地址",
            "supplier_contact": "供应商联系人",
            "supplier_phone": "0574-2",
            "supplier_email": "finance@example.com",
            "currency": "CNY",
            "tax_rate_ppm": 130000,
            "amount_ex_tax_minor": 40000,
            "amount_incl_tax_minor": 45200,
            "tax_amount_minor": 5200,
            "status": "pending",
            "remark": "对账备注",
            "created_by": "finance-user",
            "created_at": "2026-09-08T10:00:00",
            "confirmed_name": "",
            "confirmed_at": "",
        }
        self.items = [
            {
                "sort_order": 1,
                "source_type": "ordinary",
                "shipped_at": "2026-09-01",
                "delivery_no": "FH-1",
                "order_no": "SO-A",
                "sku": "SKU-100",
                "drawing_no": "RA-100",
                "product_name": "对账产品A",
                "model": "M-100",
                "unit": "件",
                "quantity": 2,
                "currency": "CNY",
                "unit_price_ex_tax_scaled": 100_000_000,
                "unit_price_incl_tax_minor": 11_300,
                "amount_ex_tax_minor": 20_000,
                "amount_incl_tax_minor": 22_600,
                "tax_amount_minor": 2_600,
                "assembly_drawing_no": "",
                "remark": "普通备注",
            },
            {
                "sort_order": 2,
                "source_type": "assembly_item",
                "shipped_at": "2026-09-05",
                "delivery_no": "ZP-2",
                "order_no": "SO-A",
                "sku": "SKU-101",
                "drawing_no": "RA-101",
                "product_name": "组装配件A",
                "model": "M-101",
                "unit": "件",
                "quantity": 4,
                "currency": "CNY",
                "unit_price_ex_tax_scaled": 50_000_000,
                "unit_price_incl_tax_minor": 5_650,
                "amount_ex_tax_minor": 20_000,
                "amount_incl_tax_minor": 22_600,
                "tax_amount_minor": 2_600,
                "assembly_drawing_no": "ASM-A",
                "remark": "组装备注",
            },
        ]

    def single_item_snapshot(self, currency, unit_price_minor, quantity):
        amounts = calculate_reconciliation_amounts(
            unit_price_minor,
            quantity,
            currency,
        )
        item = dict(
            self.items[0],
            currency=currency,
            quantity=quantity,
            **amounts,
        )
        statement = dict(
            self.statement,
            currency=currency,
            amount_ex_tax_minor=amounts["amount_ex_tax_minor"],
            amount_incl_tax_minor=amounts["amount_incl_tax_minor"],
            tax_amount_minor=amounts["tax_amount_minor"],
        )
        return statement, [item]

    def test_cny_database_value_beyond_excel_precision_is_rejected(self):
        statement, items = self.single_item_snapshot(
            "CNY",
            1_000_000_000_000_001,
            1,
        )
        with self.assertRaisesRegex(ValueError, "超出 Excel 可精确表示范围"):
            build_reconciliation_workbook(statement, items)

    def test_jpy_database_value_beyond_excel_precision_is_rejected(self):
        statement, items = self.single_item_snapshot(
            "JPY",
            10_000,
            100_000_000_001,
        )
        with self.assertRaisesRegex(ValueError, "超出 Excel 可精确表示范围"):
            build_reconciliation_workbook(statement, items)

    def test_formula_inputs_that_would_lose_a_cent_are_rejected(self):
        large_quantity = (999_999_999_999_999 - 2_000) // 11_300
        large_amounts = calculate_reconciliation_amounts(
            11_300,
            large_quantity,
            "CNY",
        )
        small_amounts = calculate_reconciliation_amounts(1, 1, "CNY")
        items = [
            dict(
                self.items[0],
                quantity=large_quantity,
                **large_amounts,
            )
        ]
        items.extend(
            dict(
                self.items[0],
                sort_order=index + 2,
                quantity=1,
                **small_amounts,
            )
            for index in range(50)
        )
        statement = dict(
            self.statement,
            amount_ex_tax_minor=sum(item["amount_ex_tax_minor"] for item in items),
            amount_incl_tax_minor=sum(
                item["amount_incl_tax_minor"] for item in items
            ),
            tax_amount_minor=sum(item["tax_amount_minor"] for item in items),
        )

        with self.assertRaisesRegex(ValueError, "汇总公式"):
            build_reconciliation_workbook(statement, items)

    def test_normal_values_stay_numeric_and_six_decimal_unit_price_round_trips(self):
        statement, items = self.single_item_snapshot("CNY", 10_000, 1)
        output = BytesIO()
        build_reconciliation_workbook(statement, items).save(output)
        output.seek(0)
        sheet = load_workbook(output, data_only=False).active

        for coordinate in ("I9", "J9", "K9", "L9", "M9"):
            self.assertEqual(sheet[coordinate].data_type, "n")
            self.assertIsInstance(sheet[coordinate].value, (int, float))
        self.assertEqual(
            Decimal(str(sheet["J9"].value)).quantize(Decimal("0.000001")),
            Decimal("88.495575"),
        )

    def test_workbook_has_dynamic_a4_landscape_structure_and_numeric_detail_cells(self):
        workbook = build_reconciliation_workbook(self.statement, self.items)
        sheet = workbook.active
        self.assertEqual(sheet.title, "对账单")
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(sheet.page_setup.paperSize, int(sheet.PAPERSIZE_A4))
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.page_setup.fitToHeight, 0)
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.autoPageBreaks)
        self.assertEqual(sheet.print_title_rows, "$8:$8")
        self.assertEqual(sheet.freeze_panes, "A9")
        self.assertEqual(sheet["A1"].value, "2026年9月对账单")
        self.assertEqual(
            [sheet.cell(8, column).value for column in range(1, 15)],
            [
                "序号", "送货日期", "送货单号", "采购订单号", "物料编码", "物料名称",
                "规格型号", "单位", "送货数量", "不含税单价", "含税单价",
                "不含税金额", "含税金额", "备注",
            ],
        )
        for column in range(9, 14):
            self.assertIsInstance(sheet.cell(9, column).value, (int, float))
            self.assertEqual(sheet.cell(9, column).data_type, "n")
        self.assertEqual(sheet.cell(9, 10).number_format, "#,##0.000000")
        self.assertEqual(sheet.cell(9, 11).number_format, "#,##0.00")
        self.assertEqual(sheet["J9"].value, 100.0)
        self.assertEqual(sheet["K9"].value, 113.0)
        self.assertEqual(sheet["L9"].value, 200.0)
        self.assertEqual(sheet["M9"].value, 226.0)
        self.assertIn("客户A", sheet["A3"].value)
        self.assertEqual(sheet["F9"].value, "对账产品A")

        total_row = 8 + len(self.items) + 2
        self.assertEqual(sheet.cell(total_row, 5).value, "=SUM(L9:L10)")
        self.assertEqual(sheet.cell(total_row, 10).value, "=SUM(M9:M10)")
        self.assertEqual(
            sheet.cell(total_row, 14).value,
            "=SUM(M9:M10)-SUM(L9:L10)",
        )
        self.assertEqual(
            sheet.print_area,
            f"'对账单'!$A$1:$N${sheet.max_row}",
        )

    def test_untrusted_formula_prefixes_are_written_as_literal_text(self):
        statement = dict(
            self.statement,
            customer_name="=1+1",
            supplier_company_name="+CMD",
            remark="@SUM(A1:A2)",
        )
        items = [
            dict(
                self.items[0],
                delivery_no="+CMD",
                order_no="\t=1+1",
                sku="=1+1",
                product_name="@SUM(A1:A2)",
                model="-1+2",
                remark="=HYPERLINK(\"https://invalid.example\")",
            )
        ]
        statement.update(
            amount_ex_tax_minor=20_000,
            amount_incl_tax_minor=22_600,
            tax_amount_minor=2_600,
        )
        workbook = build_reconciliation_workbook(statement, items)
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        sheet = load_workbook(buffer, data_only=False).active
        self.assertEqual(sheet["C9"].value, "'+CMD")
        self.assertEqual(sheet["D9"].value, "'\t=1+1")
        self.assertEqual(sheet["E9"].value, "'=1+1")
        self.assertEqual(sheet["F9"].value, "'@SUM(A1:A2)")
        self.assertEqual(sheet["G9"].value, "'-1+2")
        self.assertIn("'=HYPERLINK", sheet["N9"].value)
        expected_formula_cells = {
            sheet.cell(11, 5).coordinate,
            sheet.cell(11, 10).coordinate,
            sheet.cell(11, 14).coordinate,
        }
        actual_formula_cells = {
            cell.coordinate
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        }
        self.assertEqual(actual_formula_cells, expected_formula_cells)

    def test_snapshot_totals_must_match_item_totals(self):
        statement = dict(self.statement, amount_incl_tax_minor=45_201)
        with self.assertRaisesRegex(ValueError, "汇总金额与明细不一致"):
            build_reconciliation_workbook(statement, self.items)

    def test_confirmed_signature_is_bounded_and_signer_metadata_is_printed(self):
        signature_bytes = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            signature_path = Path(tmpdir) / "signature.png"
            signature_path.write_bytes(signature_bytes)
            statement = dict(
                self.statement,
                status="confirmed",
                confirmed_name="客户确认人",
                confirmed_at="2026-09-09T09:30:00",
            )
            sheet = build_reconciliation_workbook(
                statement,
                self.items,
                signature_path=signature_path,
            ).active
        self.assertEqual(len(sheet._images), 1)
        self.assertLessEqual(sheet._images[0].width, 180)
        self.assertLessEqual(sheet._images[0].height, 65)
        self.assertIn("客户确认人", sheet.cell(sheet.max_row, 1).value)
        self.assertIn("2026-09-09", sheet.cell(sheet.max_row, 1).value)

    def test_long_product_and_assembly_remark_expand_only_the_affected_detail_row(self):
        long_product_name = "超长产品名称" * 12
        long_assembly_remark = "超长组装备注" * 12
        items = [
            dict(self.items[0]),
            dict(
                self.items[1],
                product_name=long_product_name,
                remark=long_assembly_remark,
            ),
        ]
        sheet = build_reconciliation_workbook(self.statement, items).active
        expected_product_lines = (len(long_product_name) * 2 + 18) // 19
        self.assertEqual(sheet.row_dimensions[9].height, 25)
        self.assertGreaterEqual(
            sheet.row_dimensions[10].height,
            expected_product_lines * 15,
        )
        self.assertTrue(sheet["F10"].alignment.wrap_text)
        self.assertTrue(sheet["N10"].alignment.wrap_text)

    def test_standalone_builder_can_save_a_confirmed_signature_without_importing_app(self):
        script = textwrap.dedent(
            r'''
            import base64
            import sys
            import tempfile
            from io import BytesIO
            from pathlib import Path

            from reconciliation import build_reconciliation_workbook

            assert "app" not in sys.modules
            statement = {
                "currency": "CNY",
                "period_start": "2026-09-01",
                "period_end": "2026-09-30",
                "amount_ex_tax_minor": 10000,
                "amount_incl_tax_minor": 11300,
                "tax_amount_minor": 1300,
                "status": "confirmed",
                "confirmed_name": "客户确认人",
                "confirmed_at": "2026-09-09T09:30:00",
            }
            items = [{
                "sort_order": 1,
                "currency": "CNY",
                "quantity": 1,
                "unit_price_ex_tax_scaled": 100_000_000,
                "unit_price_incl_tax_minor": 11_300,
                "amount_ex_tax_minor": 10_000,
                "amount_incl_tax_minor": 11_300,
                "tax_amount_minor": 1_300,
            }]
            signature_bytes = base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
            )
            with tempfile.TemporaryDirectory() as tmpdir:
                signature_path = Path(tmpdir) / "signature.png"
                signature_path.write_bytes(signature_bytes)
                workbook = build_reconciliation_workbook(
                    statement,
                    items,
                    signature_path=signature_path,
                )
                output = BytesIO()
                workbook.save(output)
                assert output.getvalue().startswith(b"PK")
            '''
        )
        result = subprocess.run(
            [sys.executable, "-c", script],
            cwd=app.BASE_DIR,
            text=True,
            capture_output=True,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


class ReconciliationBackendRouteTests(ReconciliationDomainTestCase):
    SIGNATURE_BYTES = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )

    def setUp(self):
        super().setUp()
        self.original_testing = app.app.config["TESTING"]
        self.original_secret_key = app.app.config["SECRET_KEY"]
        self.original_signatures_dir = app.RECONCILIATION_SIGNATURES_DIR
        self.signature_dir = Path(self.tmpdir.name) / "reconciliation-signatures"
        app.RECONCILIATION_SIGNATURES_DIR = self.signature_dir
        app.app.config.update(TESTING=True, SECRET_KEY="reconciliation-test-secret")
        with app.get_db() as conn:
            self.create_user(conn, "finance-manager", "operator", can_manage_finance=1)
            self.create_user(conn, "plain-operator", "operator")
            self.create_user(conn, "system-admin", "admin")
        self.client = app.app.test_client()

    def tearDown(self):
        app.RECONCILIATION_SIGNATURES_DIR = self.original_signatures_dir
        app.app.config.update(
            TESTING=self.original_testing,
            SECRET_KEY=self.original_secret_key,
        )
        super().tearDown()

    def create_user(self, conn, username, role, can_manage_finance=0):
        conn.execute(
            """
            INSERT INTO users (
                username, password_hash, role, active,
                can_manage_products, can_manage_orders, can_view_orders,
                can_manage_shipped, can_view_shipped, can_manage_customers,
                can_manage_common_info, can_manage_purchase_followups,
                can_manage_powder_coating, can_manage_carton_purchases,
                can_manage_warehouse_inventory, can_manage_production_followups,
                can_create_products, can_edit_products, can_view_prices,
                can_manage_finance,
                created_at, updated_at
            ) VALUES (?, 'hash', ?, 1,
                      0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0,
                      ?, ?, ?)
            """,
            (username, role, can_manage_finance, self.now, self.now),
        )

    def login_as(self, username, role):
        with self.client.session_transaction() as flask_session:
            flask_session.clear()
            flask_session.update(
                admin_logged_in=True,
                admin_username=username,
                admin_role=role,
            )

    def create_statement_directly(self):
        with app.get_db() as conn:
            return app.create_reconciliation_statement(
                conn,
                [("ordinary", self.ordinary_a_cny)],
                "fixture",
            )["statement_id"]

    def admin_csrf_token_from(self, path):
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)
        match = re.search(
            r'name="reconciliation_csrf_token" value="([^"]+)"',
            response.get_data(as_text=True),
        )
        self.assertIsNotNone(match)
        return match.group(1)

    def mark_statement_confirmed(self, statement_id):
        filename = f"confirmed-{statement_id}.png"
        self.signature_dir.mkdir(parents=True, exist_ok=True)
        (self.signature_dir / filename).write_bytes(self.SIGNATURE_BYTES)
        with app.get_db() as conn:
            conn.execute(
                """
                UPDATE reconciliation_statements
                SET status = 'confirmed', confirmed_name = '张确认',
                    confirmed_at = '2026-09-08T12:30:00', signature_image = ?,
                    confirmed_ip = '203.0.113.8',
                    confirmed_user_agent = 'HistoryTest/1.0'
                WHERE id = ?
                """,
                (filename, statement_id),
            )
        return filename

    def mark_statement_disputed(self, statement_id):
        with app.get_db() as conn:
            conn.execute(
                """
                UPDATE reconciliation_statements
                SET status = 'disputed', dispute_name = '李异议',
                    dispute_content = '历史数量有异议',
                    disputed_at = '2026-09-08T13:45:00',
                    disputed_ip = '203.0.113.9',
                    disputed_user_agent = 'HistoryTest/2.0'
                WHERE id = ?
                """,
                (statement_id,),
            )

    def test_finance_manager_can_create_and_view_statement_with_one_time_credentials(self):
        self.login_as("finance-manager", "operator")
        csrf_token = self.admin_csrf_token_from("/admin/shipped-orders")
        response = self.client.post(
            "/admin/reconciliation-statements",
            data={
                "source_ref": [
                    f"ordinary:{self.ordinary_a_cny}",
                    f"assembly_item:{self.assembly_a_cny}",
                ],
                "customer": "伪造客户",
                "currency": "USD",
                "amount_incl_tax_minor": "1",
                "reconciliation_csrf_token": csrf_token,
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertRegex(
            response.location,
            r"/admin/reconciliation-statements/[0-9]+$",
        )

        first_detail = self.client.get(response.location)
        first_html = first_detail.get_data(as_text=True)
        self.assertEqual(first_detail.status_code, 200)
        self.assertIn("客户A", first_html)
        self.assertNotIn("伪造客户", first_html)
        self.assertIn("凭证仅显示一次", first_html)
        self.assertRegex(first_html, r"[0-9]{6}")

        second_html = self.client.get(response.location).get_data(as_text=True)
        self.assertNotIn("凭证仅显示一次", second_html)
        self.assertIn("重置访问凭证", second_html)

    def test_cross_customer_create_error_inserts_no_statement_or_items(self):
        self.login_as("finance-manager", "operator")
        csrf_token = self.admin_csrf_token_from("/admin/shipped-orders")
        response = self.client.post(
            "/admin/reconciliation-statements",
            data={
                "source_ref": [
                    f"ordinary:{self.ordinary_a_cny}",
                    f"ordinary:{self.ordinary_b_cny}",
                ],
                "reconciliation_csrf_token": csrf_token,
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "只能合并同一客户、同一币种的发货记录",
            response.get_data(as_text=True),
        )
        with app.get_db() as conn:
            self.assertEqual(
                conn.execute("SELECT COUNT(*) FROM reconciliation_statements").fetchone()[0],
                0,
            )
            self.assertEqual(
                conn.execute("SELECT COUNT(*) FROM reconciliation_statement_items").fetchone()[0],
                0,
            )

    def test_reset_replaces_credentials_and_increments_password_version(self):
        statement_id = self.create_statement_directly()
        with app.get_db() as conn:
            before = dict(app.fetch_reconciliation_statement(conn, statement_id))
        self.login_as("finance-manager", "operator")
        csrf_token = self.admin_csrf_token_from(
            f"/admin/reconciliation-statements/{statement_id}"
        )

        response = self.client.post(
            f"/admin/reconciliation-statements/{statement_id}/reset-access",
            data={"reconciliation_csrf_token": csrf_token},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            response.location,
            f"/admin/reconciliation-statements/{statement_id}",
        )
        with app.get_db() as conn:
            after = dict(app.fetch_reconciliation_statement(conn, statement_id))
        self.assertNotEqual(after["token_digest"], before["token_digest"])
        self.assertNotEqual(after["password_hash"], before["password_hash"])
        self.assertEqual(after["password_version"], before["password_version"] + 1)

        detail = self.client.get(response.location).get_data(as_text=True)
        self.assertIn("凭证仅显示一次", detail)
        self.assertRegex(detail, r"[0-9]{6}")

    def test_finance_manager_can_void_and_list_filter_statements(self):
        statement_id = self.create_statement_directly()
        self.login_as("finance-manager", "operator")
        csrf_token = self.admin_csrf_token_from(
            f"/admin/reconciliation-statements/{statement_id}"
        )
        response = self.client.post(
            f"/admin/reconciliation-statements/{statement_id}/void",
            data={
                "reason": "客户要求重开",
                "reconciliation_csrf_token": csrf_token,
            },
        )
        self.assertEqual(response.status_code, 302)
        response = self.client.get(
            "/admin/reconciliation-statements",
            query_string={"customer": "客户A", "status": "void", "date": "2026-09-01"},
        )
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("DZ-202609-0001", html)
        self.assertIn("已作废", html)

    def test_plain_operator_cannot_use_any_statement_management_route(self):
        statement_id = self.create_statement_directly()
        self.login_as("plain-operator", "operator")
        routes = (
            ("get", "/admin/reconciliation-statements", None),
            ("post", "/admin/reconciliation-statements", {"source_ref": "ordinary:1"}),
            ("get", f"/admin/reconciliation-statements/{statement_id}", None),
            ("get", f"/admin/reconciliation-statements/{statement_id}/export.xlsx", None),
            ("post", f"/admin/reconciliation-statements/{statement_id}/reset-access", None),
            ("post", f"/admin/reconciliation-statements/{statement_id}/void", {"reason": "x"}),
        )
        for method, path, data in routes:
            with self.subTest(path=path):
                response = getattr(self.client, method)(path, data=data)
                self.assertEqual(response.status_code, 302)
                self.assertEqual(response.location, "/admin")

    def test_finance_user_can_export_xlsx(self):
        statement_id = self.create_statement_directly()
        with app.get_db() as conn:
            conn.execute(
                "UPDATE customers SET name = '新客户', address = '新地址' WHERE id = ?",
                (self.customer_a,),
            )
            conn.execute(
                "UPDATE manuals SET product_name = '新产品', sku = 'NEW-SKU' WHERE id = ?",
                (self.manual_a,),
            )
            conn.execute(
                "UPDATE reconciliation_company_profile SET company_name = '新供应商' WHERE id = 1"
            )
        self.login_as("finance-manager", "operator")
        response = self.client.get(
            f"/admin/reconciliation-statements/{statement_id}/export.xlsx"
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data.startswith(b"PK"))
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", response.headers["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.data), data_only=False)
        sheet = workbook.active
        self.assertIn("客户A", sheet["A3"].value)
        self.assertIn("宁波市杰德机械科技有限公司", sheet["H3"].value)
        self.assertEqual(sheet["E9"].value, "SKU-100")
        self.assertEqual(sheet["F9"].value, "对账产品A")
        self.assertNotIn("新客户", sheet["A3"].value)
        self.assertNotEqual(sheet["E9"].value, "NEW-SKU")

    def test_export_precision_error_redirects_with_user_message(self):
        statement_id = self.create_statement_directly()
        amounts = calculate_reconciliation_amounts(
            1_000_000_000_000_001,
            1,
            "CNY",
        )
        with app.get_db() as conn:
            conn.execute(
                """
                UPDATE reconciliation_statement_items
                SET quantity = ?, unit_price_incl_tax_minor = ?,
                    unit_price_ex_tax_scaled = ?, amount_incl_tax_minor = ?,
                    amount_ex_tax_minor = ?, tax_amount_minor = ?
                WHERE statement_id = ?
                """,
                (
                    1,
                    amounts["unit_price_incl_tax_minor"],
                    amounts["unit_price_ex_tax_scaled"],
                    amounts["amount_incl_tax_minor"],
                    amounts["amount_ex_tax_minor"],
                    amounts["tax_amount_minor"],
                    statement_id,
                ),
            )
            conn.execute(
                """
                UPDATE reconciliation_statements
                SET amount_incl_tax_minor = ?, amount_ex_tax_minor = ?,
                    tax_amount_minor = ?
                WHERE id = ?
                """,
                (
                    amounts["amount_incl_tax_minor"],
                    amounts["amount_ex_tax_minor"],
                    amounts["tax_amount_minor"],
                    statement_id,
                ),
            )

        self.login_as("finance-manager", "operator")
        response = self.client.get(
            f"/admin/reconciliation-statements/{statement_id}/export.xlsx",
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.request.path,
            f"/admin/reconciliation-statements/{statement_id}",
        )
        self.assertIn(
            "超出 Excel 可精确表示范围",
            response.get_data(as_text=True),
        )

    def test_finance_user_gets_404_for_missing_statement_export(self):
        self.login_as("finance-manager", "operator")
        response = self.client.get(
            "/admin/reconciliation-statements/999999/export.xlsx"
        )
        self.assertEqual(response.status_code, 404)

    def test_only_admin_can_update_supplier_profile(self):
        self.login_as("finance-manager", "operator")
        denied = self.client.post(
            "/admin/reconciliation-company-profile",
            data={"company_name": "无权修改"},
        )
        self.assertEqual(denied.status_code, 302)
        self.assertEqual(denied.location, "/admin")

        self.login_as("system-admin", "admin")
        csrf_token = self.admin_csrf_token_from(
            "/admin/reconciliation-company-profile"
        )
        response = self.client.post(
            "/admin/reconciliation-company-profile",
            data={
                "company_name": "杰德供应商快照公司",
                "address": "宁波地址",
                "contact": "财务联系人",
                "phone": "0574-10086",
                "email": "finance@example.com",
                "reconciliation_csrf_token": csrf_token,
            },
        )
        self.assertEqual(response.status_code, 302)
        with app.get_db() as conn:
            profile = conn.execute(
                "SELECT * FROM reconciliation_company_profile WHERE id = 1"
            ).fetchone()
        self.assertEqual(profile["company_name"], "杰德供应商快照公司")
        self.assertEqual(profile["updated_by"], "system-admin")

    def test_finance_only_operator_can_reach_read_only_shipped_selection_ui(self):
        self.login_as("finance-manager", "operator")
        response = self.client.get("/admin/shipped-orders")
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn('id="reconciliation-create-form"', html)
        self.assertIn(
            f'value="ordinary:{self.ordinary_a_cny}"',
            html,
        )
        self.assertIn(
            f'value="assembly_item:{self.assembly_a_cny}"',
            html,
        )
        self.assertNotIn("修改</a>", html)
        self.assertNotIn("补录</button>", html)

    def test_backend_reconciliation_forms_render_server_csrf_tokens(self):
        self.login_as("finance-manager", "operator")
        shipped_html = self.client.get("/admin/shipped-orders").get_data(as_text=True)
        self.assertIn(
            'name="reconciliation_csrf_token"',
            shipped_html,
        )

        statement_id = self.create_statement_directly()
        detail_html = self.client.get(
            f"/admin/reconciliation-statements/{statement_id}"
        ).get_data(as_text=True)
        self.assertEqual(detail_html.count('name="reconciliation_csrf_token"'), 2)

        self.login_as("system-admin", "admin")
        profile_html = self.client.get(
            "/admin/reconciliation-company-profile"
        ).get_data(as_text=True)
        self.assertIn('name="reconciliation_csrf_token"', profile_html)

    def test_backend_reconciliation_posts_reject_bad_csrf_without_side_effects(self):
        self.login_as("finance-manager", "operator")
        valid_token = self.admin_csrf_token_from("/admin/shipped-orders")
        for submitted in ("wrong-token", "错误令牌", None):
            data = {"source_ref": f"ordinary:{self.ordinary_a_cny}"}
            if submitted is not None:
                data["reconciliation_csrf_token"] = submitted
            with self.subTest(operation="create", submitted=submitted):
                response = self.client.post(
                    "/admin/reconciliation-statements",
                    data=data,
                )
                self.assertEqual(response.status_code, 403)
                with app.get_db() as conn:
                    self.assertEqual(
                        conn.execute(
                            "SELECT COUNT(*) FROM reconciliation_statements"
                        ).fetchone()[0],
                        0,
                    )

        statement_id = self.create_statement_directly()
        with app.get_db() as conn:
            before = dict(app.fetch_reconciliation_statement(conn, statement_id))
        reset_response = self.client.post(
            f"/admin/reconciliation-statements/{statement_id}/reset-access",
            data={"reconciliation_csrf_token": "wrong-token"},
        )
        self.assertEqual(reset_response.status_code, 403)
        void_response = self.client.post(
            f"/admin/reconciliation-statements/{statement_id}/void",
            data={"reason": "不应作废"},
        )
        self.assertEqual(void_response.status_code, 403)
        with app.get_db() as conn:
            after = dict(app.fetch_reconciliation_statement(conn, statement_id))
            claim = app.fetch_reconciliation_statement_items(conn, statement_id)[0][
                "active_claim_key"
            ]
        self.assertEqual(after, before)
        self.assertEqual(claim, f"ordinary:{self.ordinary_a_cny}")

        self.login_as("system-admin", "admin")
        self.admin_csrf_token_from("/admin/reconciliation-company-profile")
        for submitted in ("wrong-token", "错误令牌", None):
            data = {"company_name": "不应保存"}
            if submitted is not None:
                data["reconciliation_csrf_token"] = submitted
            with self.subTest(operation="profile", submitted=submitted):
                response = self.client.post(
                    "/admin/reconciliation-company-profile",
                    data=data,
                )
                self.assertEqual(response.status_code, 403)
        with app.get_db() as conn:
            company_name = conn.execute(
                "SELECT company_name FROM reconciliation_company_profile WHERE id = 1"
            ).fetchone()[0]
        self.assertEqual(company_name, "宁波市杰德机械科技有限公司")
        self.assertTrue(valid_token)

    def test_confirmed_detail_shows_retained_signature_and_confirmation_evidence(self):
        statement_id = self.create_statement_directly()
        filename = self.mark_statement_confirmed(statement_id)
        self.login_as("finance-manager", "operator")

        response = self.client.get(
            f"/admin/reconciliation-statements/{statement_id}"
        )
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("张确认", html)
        self.assertIn("2026-09-08", html)
        self.assertIn("客户确认签名", html)
        signature_url = f"/reconciliation-signature/{filename}"
        self.assertIn(signature_url, html)
        signature_response = self.client.get(signature_url)
        self.assertEqual(signature_response.status_code, 200)
        signature_response.close()

    def test_confirmed_then_void_detail_retains_confirmation_and_signature_access(self):
        statement_id = self.create_statement_directly()
        filename = self.mark_statement_confirmed(statement_id)
        self.login_as("finance-manager", "operator")
        csrf_token = self.admin_csrf_token_from(
            f"/admin/reconciliation-statements/{statement_id}"
        )
        with app.get_db() as conn:
            before = dict(app.fetch_reconciliation_statement(conn, statement_id))

        response = self.client.post(
            f"/admin/reconciliation-statements/{statement_id}/void",
            data={
                "reason": "确认后发现单据错误",
                "reconciliation_csrf_token": csrf_token,
            },
        )
        self.assertEqual(response.status_code, 302)
        with app.get_db() as conn:
            after = dict(app.fetch_reconciliation_statement(conn, statement_id))
        self.assertEqual(after["status"], "void")
        for field in ("confirmed_name", "confirmed_at", "signature_image"):
            self.assertEqual(after[field], before[field])

        html = self.client.get(response.location).get_data(as_text=True)
        self.assertIn("张确认", html)
        self.assertIn("客户确认签名", html)
        self.assertIn("确认后发现单据错误", html)
        signature_response = self.client.get(
            f"/reconciliation-signature/{filename}"
        )
        self.assertEqual(signature_response.status_code, 200)
        signature_response.close()

    def test_disputed_then_void_detail_retains_dispute_evidence(self):
        statement_id = self.create_statement_directly()
        self.mark_statement_disputed(statement_id)
        self.login_as("finance-manager", "operator")
        csrf_token = self.admin_csrf_token_from(
            f"/admin/reconciliation-statements/{statement_id}"
        )
        with app.get_db() as conn:
            before = dict(app.fetch_reconciliation_statement(conn, statement_id))

        response = self.client.post(
            f"/admin/reconciliation-statements/{statement_id}/void",
            data={
                "reason": "异议后重开",
                "reconciliation_csrf_token": csrf_token,
            },
        )
        self.assertEqual(response.status_code, 302)
        with app.get_db() as conn:
            after = dict(app.fetch_reconciliation_statement(conn, statement_id))
        self.assertEqual(after["status"], "void")
        for field in ("dispute_name", "dispute_content", "disputed_at"):
            self.assertEqual(after[field], before[field])

        html = self.client.get(response.location).get_data(as_text=True)
        self.assertIn("李异议", html)
        self.assertIn("历史数量有异议", html)
        self.assertIn("异议后重开", html)

    def test_backend_reconciliation_tables_keep_their_own_mobile_scroll_container(self):
        statement_id = self.create_statement_directly()
        self.login_as("finance-manager", "operator")
        response = self.client.get(
            f"/admin/reconciliation-statements/{statement_id}"
        )
        html = response.get_data(as_text=True)
        self.assertIn(
            'class="table-wrap reconciliation-table-wrap"',
            html,
        )

        stylesheet = (app.BASE_DIR / "static" / "reconciliation.css").read_text(
            encoding="utf-8"
        )
        self.assertIn("@media (max-width: 760px)", stylesheet)
        phone_rules = stylesheet.rsplit("@media (max-width: 760px)", 1)[1]
        self.assertIn(
            ".reconciliation-page .reconciliation-table-wrap",
            phone_rules,
        )
        self.assertIn("overflow-x: auto", phone_rules)
        self.assertIn(
            ".reconciliation-page .reconciliation-table {",
            phone_rules,
        )
        self.assertIn("display: table", phone_rules)

    def test_viewing_another_statement_does_not_consume_one_time_credentials(self):
        with app.get_db() as conn:
            first = app.create_reconciliation_statement(
                conn,
                [("ordinary", self.ordinary_a_cny)],
                "fixture",
            )
            second = app.create_reconciliation_statement(
                conn,
                [("ordinary", self.ordinary_b_cny)],
                "fixture",
            )
        self.login_as("finance-manager", "operator")
        with self.client.session_transaction() as flask_session:
            flask_session["reconciliation_credentials"] = first

        unrelated = self.client.get(
            f"/admin/reconciliation-statements/{second['statement_id']}"
        )
        self.assertNotIn("凭证仅显示一次", unrelated.get_data(as_text=True))
        intended = self.client.get(
            f"/admin/reconciliation-statements/{first['statement_id']}"
        )
        intended_html = intended.get_data(as_text=True)
        self.assertIn("凭证仅显示一次", intended_html)
        self.assertIn(first["password"], intended_html)

    def test_voiding_another_statement_does_not_clear_one_time_credentials(self):
        with app.get_db() as conn:
            first = app.create_reconciliation_statement(
                conn,
                [("ordinary", self.ordinary_a_cny)],
                "fixture",
            )
            second = app.create_reconciliation_statement(
                conn,
                [("ordinary", self.ordinary_b_cny)],
                "fixture",
            )
        self.login_as("finance-manager", "operator")
        with self.client.session_transaction() as flask_session:
            flask_session["reconciliation_credentials"] = first

        self.client.post(
            f"/admin/reconciliation-statements/{second['statement_id']}/void",
            data={
                "reason": "另一张单据作废",
                "reconciliation_csrf_token": self.admin_csrf_token_from(
                    "/admin/shipped-orders"
                ),
            },
        )
        intended = self.client.get(
            f"/admin/reconciliation-statements/{first['statement_id']}"
        )
        self.assertIn("凭证仅显示一次", intended.get_data(as_text=True))

    def test_detail_renders_complete_persisted_customer_snapshot(self):
        statement_id = self.create_statement_directly()
        self.login_as("finance-manager", "operator")
        response = self.client.get(
            f"/admin/reconciliation-statements/{statement_id}"
        )
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("注册地址", html)
        self.assertIn("开票地址", html)
        self.assertIn("注册电话", html)
        self.assertIn("0574-2", html)
        self.assertIn("收票邮箱", html)
        self.assertIn("invoice@example.com", html)


class ReconciliationPublicTestCase(ReconciliationDomainTestCase):
    PNG_DATA_URL = "data:image/png;base64," + base64.b64encode(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
    ).decode("ascii")

    def setUp(self):
        super().setUp()
        self.original_testing = app.app.config["TESTING"]
        self.original_secret_key = app.app.config["SECRET_KEY"]
        self.original_signatures_dir = app.RECONCILIATION_SIGNATURES_DIR
        self.signature_dir = Path(self.tmpdir.name) / "reconciliation-signatures"
        app.RECONCILIATION_SIGNATURES_DIR = self.signature_dir
        app.app.config.update(TESTING=True, SECRET_KEY="reconciliation-public-test")
        with app.get_db() as conn:
            self.credentials = app.create_reconciliation_statement(
                conn, [("ordinary", self.ordinary_a_cny)], "fixture"
            )
        self.statement_id = self.credentials["statement_id"]
        self.password = self.credentials["password"]
        self.public_path = f"/reconciliation/{self.credentials['token']}"
        self.client = app.app.test_client()

    def tearDown(self):
        app.RECONCILIATION_SIGNATURES_DIR = self.original_signatures_dir
        app.app.config.update(
            TESTING=self.original_testing,
            SECRET_KEY=self.original_secret_key,
        )
        super().tearDown()

    def authorize(self):
        response = self.client.post(
            self.public_path, data={"password": self.password}
        )
        self.assertEqual(response.status_code, 302)
        return self.client.get(self.public_path)

    def current_nonce(self):
        with self.client.session_transaction() as flask_session:
            return flask_session["reconciliation_action_nonce"]


class ReconciliationPublicAccessTests(ReconciliationPublicTestCase):
    def test_public_statement_requires_password_then_grants_versioned_session(self):
        response = self.client.get(self.public_path)
        self.assertIn("请输入6位临时密码", response.get_data(as_text=True))

        response = self.client.post(
            self.public_path, data={"password": self.password}
        )
        self.assertEqual(response.status_code, 302)
        response = self.client.get(self.public_path)
        self.assertIn("客户A", response.get_data(as_text=True))
        with self.client.session_transaction() as flask_session:
            access = flask_session["reconciliation_access"]
            self.assertEqual(
                set(access), {"statement_id", "password_version", "verified_until"}
            )
            self.assertEqual(access["statement_id"], self.statement_id)
            self.assertEqual(access["password_version"], 1)

    def test_invalid_tokens_are_non_enumerating_404_and_expired_or_void_are_410(self):
        invalid = self.client.get("/reconciliation/not-a-real-token")
        unknown = self.client.get("/reconciliation/" + "x" * 43)
        self.assertEqual(invalid.status_code, 404)
        self.assertEqual(unknown.status_code, 404)
        self.assertEqual(invalid.get_data(), unknown.get_data())

        with app.get_db() as conn:
            conn.execute(
                "UPDATE reconciliation_statements SET access_expires_at = ? WHERE id = ?",
                ("2000-01-01T00:00:00", self.statement_id),
            )
        self.assertEqual(self.client.get(self.public_path).status_code, 410)

        with app.get_db() as conn:
            conn.execute(
                "UPDATE reconciliation_statements SET access_expires_at = ?, status = 'void' WHERE id = ?",
                ("2999-01-01T00:00:00", self.statement_id),
            )
        self.assertEqual(self.client.get(self.public_path).status_code, 410)

    def test_five_wrong_passwords_lock_for_fifteen_minutes(self):
        wrong_password = "000000" if self.password != "000000" else "999999"
        started_at = datetime.utcnow()
        for _ in range(4):
            response = self.client.post(
                self.public_path,
                data={"action": "verify", "password": wrong_password},
            )
            self.assertEqual(response.status_code, 401)
        fifth = self.client.post(
            self.public_path,
            data={"action": "verify", "password": wrong_password},
        )
        self.assertEqual(fifth.status_code, 429)
        with app.get_db() as conn:
            statement = app.fetch_reconciliation_statement(conn, self.statement_id)
        self.assertEqual(statement["failed_attempts"], 5)
        locked_until = datetime.fromisoformat(statement["locked_until"])
        self.assertGreaterEqual(locked_until, started_at + timedelta(minutes=14, seconds=59))
        self.assertLessEqual(locked_until, datetime.utcnow() + timedelta(minutes=15, seconds=1))
        blocked = self.client.post(
            self.public_path,
            data={"action": "verify", "password": self.password},
        )
        self.assertEqual(blocked.status_code, 429)

    def test_successful_password_verification_clears_failures(self):
        wrong_password = "000000" if self.password != "000000" else "999999"
        self.client.post(
            self.public_path,
            data={"action": "verify", "password": wrong_password},
        )
        response = self.client.post(self.public_path, data={"password": self.password})
        self.assertEqual(response.status_code, 302)
        with app.get_db() as conn:
            statement = app.fetch_reconciliation_statement(conn, self.statement_id)
        self.assertEqual(statement["failed_attempts"], 0)
        self.assertEqual(statement["locked_until"], "")

    def test_reset_invalidates_old_token_and_old_versioned_session(self):
        self.authorize()
        with app.get_db() as conn:
            replacement = app.reset_reconciliation_statement_access(
                conn, self.statement_id
            )
        self.assertEqual(self.client.get(self.public_path).status_code, 404)
        new_path = f"/reconciliation/{replacement['token']}"
        response = self.client.get(new_path)
        self.assertIn("请输入6位临时密码", response.get_data(as_text=True))

    def test_database_stores_no_raw_token_or_password(self):
        with app.get_db() as conn:
            statement = dict(app.fetch_reconciliation_statement(conn, self.statement_id))
        self.assertNotIn(self.credentials["token"], statement.values())
        self.assertNotIn(self.password, statement.values())
        self.assertEqual(
            statement["token_digest"],
            app.reconciliation_token_digest(self.credentials["token"]),
        )

    def test_public_item_identifier_falls_back_to_snapshotted_drawing_number(self):
        with app.get_db() as conn:
            conn.execute(
                """
                UPDATE reconciliation_statement_items
                SET sku = '', drawing_no = 'SNAPSHOT-DRAWING-100'
                WHERE statement_id = ?
                """,
                (self.statement_id,),
            )
        self.authorize()

        html = self.client.get(self.public_path).get_data(as_text=True)
        self.assertIn(">SNAPSHOT-DRAWING-100</td>", html)


class ReconciliationCustomerActionTests(ReconciliationPublicTestCase):
    def post_confirm(self, **overrides):
        data = {
            "action": "confirm",
            "action_nonce": self.current_nonce(),
            "confirmed_name": "客户确认人",
            "signature_data": self.PNG_DATA_URL,
        }
        data.update(overrides)
        return self.client.post(
            self.public_path,
            data=data,
            headers={"User-Agent": "ReconciliationTest/1.0"},
        )

    def post_dispute(self, **overrides):
        data = {
            "action": "dispute",
            "action_nonce": self.current_nonce(),
            "dispute_name": "客户异议人",
            "dispute_content": "数量与我方记录不一致",
        }
        data.update(overrides)
        return self.client.post(self.public_path, data=data)

    def test_confirmation_requires_name_nonce_and_valid_signature(self):
        self.authorize()
        cases = (
            ({"confirmed_name": ""}, "请填写确认人姓名"),
            ({"action_nonce": "wrong"}, "页面已失效"),
            ({"signature_data": ""}, "请手写签名"),
            ({"signature_data": "data:image/png;base64,bad"}, "签名图片"),
        )
        for overrides, expected in cases:
            with self.subTest(expected=expected):
                response = self.post_confirm(**overrides)
                self.assertEqual(response.status_code, 400)
                self.assertIn(expected, response.get_data(as_text=True))
                with app.get_db() as conn:
                    self.assertEqual(
                        app.fetch_reconciliation_statement(conn, self.statement_id)["status"],
                        "pending",
                    )

    def assert_confirmation_image_rejected(self, image_bytes, expected):
        self.authorize()
        response = self.post_confirm(
            signature_data="data:image/png;base64,"
            + base64.b64encode(image_bytes).decode("ascii")
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn(expected, response.get_data(as_text=True))
        with app.get_db() as conn:
            statement = app.fetch_reconciliation_statement(conn, self.statement_id)
        self.assertEqual(statement["status"], "pending")
        self.assertEqual(statement["signature_image"], "")
        self.assertFalse(
            self.signature_dir.exists() and any(self.signature_dir.iterdir())
        )

    def test_confirmation_rejects_png_header_without_side_effects(self):
        self.assert_confirmation_image_rejected(
            b"\x89PNG\r\n\x1a\n", "签名图片无法解析"
        )

    def test_confirmation_rejects_truncated_png_without_side_effects(self):
        valid_png = base64.b64decode(self.PNG_DATA_URL.split(",", 1)[1])
        self.assert_confirmation_image_rejected(
            valid_png[:-12], "签名图片无法解析"
        )

    def test_confirmation_rejects_png_missing_last_iend_byte_without_side_effects(self):
        valid_png = base64.b64decode(self.PNG_DATA_URL.split(",", 1)[1])
        self.assert_confirmation_image_rejected(
            valid_png[:-1], "签名图片无法解析"
        )

    def test_confirmation_rejects_png_missing_iend_crc_without_side_effects(self):
        valid_png = base64.b64decode(self.PNG_DATA_URL.split(",", 1)[1])
        self.assert_confirmation_image_rejected(
            valid_png[:-4], "签名图片无法解析"
        )

    def test_confirmation_rejects_png_with_invalid_chunk_crc_without_side_effects(self):
        valid_png = bytearray(base64.b64decode(self.PNG_DATA_URL.split(",", 1)[1]))
        valid_png[-1] ^= 0x01
        self.assert_confirmation_image_rejected(
            bytes(valid_png), "签名图片无法解析"
        )

    def test_confirmation_rejects_png_bytes_after_iend_without_side_effects(self):
        valid_png = base64.b64decode(self.PNG_DATA_URL.split(",", 1)[1])
        self.assert_confirmation_image_rejected(
            valid_png + b"trailing-data", "签名图片无法解析"
        )

    def test_confirmation_rejects_oversized_png_without_side_effects(self):
        oversized_buffer = BytesIO()
        app.PillowImage.new("1", (2100, 2100)).save(oversized_buffer, format="PNG")
        self.assert_confirmation_image_rejected(
            oversized_buffer.getvalue(), "签名图片像素尺寸过大"
        )

    def test_signature_script_does_not_treat_a_no_motion_tap_as_ink(self):
        script = textwrap.dedent(
            r'''
            const assert = require("node:assert/strict");
            const fs = require("node:fs");
            const vm = require("node:vm");

            class Element {
              constructor() {
                this.listeners = {};
                this.dataset = {};
              }
              addEventListener(type, listener) {
                (this.listeners[type] ||= []).push(listener);
              }
              emit(type, event = {}) {
                let prevented = false;
                event.preventDefault ||= () => { prevented = true; };
                for (const listener of this.listeners[type] || []) listener(event);
                return prevented;
              }
            }

            const context = {
              setTransform() {}, beginPath() {}, moveTo() {}, lineTo() {},
              stroke() {}, clearRect() {},
            };
            const canvas = new Element();
            canvas.getContext = () => context;
            canvas.getBoundingClientRect = () => ({left: 0, top: 0, width: 320, height: 110});
            canvas.setPointerCapture = () => {};
            canvas.toDataURL = () => "data:image/png;base64,drawn";
            const clearButton = new Element();
            const signatureData = {value: ""};
            const confirmationForm = new Element();
            confirmationForm.dataset.confirmMessage = "confirm";
            confirmationForm.elements = {signature_data: signatureData};

            global.document = {
              querySelector: (selector) => ({
                ".reconciliation-signature-canvas": canvas,
                ".reconciliation-confirm-form": confirmationForm,
                "[data-clear-signature]": clearButton,
                ".reconciliation-dispute-form": null,
              })[selector] ?? null,
            };
            let alerts = 0;
            let confirms = 0;
            global.window = {
              devicePixelRatio: 2,
              addEventListener() {},
              alert: () => { alerts += 1; },
              confirm: () => { confirms += 1; return true; },
            };

            vm.runInThisContext(fs.readFileSync("static/reconciliation.js", "utf8"));
            canvas.emit("pointerdown", {pointerId: 1, clientX: 12, clientY: 12});
            canvas.emit("pointerup", {pointerId: 1, clientX: 12, clientY: 12});
            assert.equal(confirmationForm.emit("submit"), true);
            assert.equal(alerts, 1);
            assert.equal(confirms, 0);
            assert.equal(signatureData.value, "");

            canvas.emit("pointerdown", {pointerId: 2, clientX: 12, clientY: 12});
            canvas.emit("pointermove", {pointerId: 2, clientX: 12, clientY: 12});
            canvas.emit("pointerup", {pointerId: 2, clientX: 12, clientY: 12});
            assert.equal(confirmationForm.emit("submit"), true);
            assert.equal(alerts, 2);
            assert.equal(confirms, 0);
            assert.equal(signatureData.value, "");

            canvas.emit("pointerdown", {pointerId: 3, clientX: 12, clientY: 12});
            canvas.emit("pointermove", {pointerId: 3, clientX: 24, clientY: 18});
            canvas.emit("pointerup", {pointerId: 3, clientX: 24, clientY: 18});
            assert.equal(confirmationForm.emit("submit"), false);
            assert.equal(confirms, 1);
            assert.equal(signatureData.value, "data:image/png;base64,drawn");
            '''
        )
        result = subprocess.run(
            ["node", "-e", script],
            cwd=app.BASE_DIR,
            text=True,
            capture_output=True,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_confirmation_records_file_metadata_and_second_post_cannot_overwrite(self):
        self.authorize()
        first_nonce = self.current_nonce()
        response = self.post_confirm()
        self.assertEqual(response.status_code, 200)
        with app.get_db() as conn:
            first = dict(app.fetch_reconciliation_statement(conn, self.statement_id))
        signature_path = self.signature_dir / first["signature_image"]
        self.assertTrue(signature_path.is_file())
        self.assertEqual(signature_path.read_bytes(), app.decode_signature_image(self.PNG_DATA_URL))
        self.assertEqual(first["status"], "confirmed")
        self.assertEqual(first["confirmed_name"], "客户确认人")
        self.assertTrue(first["confirmed_at"])
        self.assertEqual(first["confirmed_ip"], "127.0.0.1")
        self.assertEqual(first["confirmed_user_agent"], "ReconciliationTest/1.0")

        second = self.client.post(
            self.public_path,
            data={
                "action": "confirm",
                "action_nonce": first_nonce,
                "confirmed_name": "覆盖者",
                "signature_data": self.PNG_DATA_URL,
            },
        )
        self.assertEqual(second.status_code, 409)
        with app.get_db() as conn:
            after = dict(app.fetch_reconciliation_statement(conn, self.statement_id))
        self.assertEqual(after, first)
        self.assertEqual(list(self.signature_dir.iterdir()), [signature_path])

    def test_dispute_requires_name_and_nonempty_content(self):
        self.authorize()
        for overrides, expected in (
            ({"dispute_name": ""}, "请填写异议人姓名"),
            ({"dispute_content": "  "}, "请填写异议内容"),
        ):
            with self.subTest(expected=expected):
                response = self.post_dispute(**overrides)
                self.assertEqual(response.status_code, 400)
                self.assertIn(expected, response.get_data(as_text=True))

    def test_first_terminal_action_wins_and_terminal_page_is_read_only(self):
        self.authorize()
        self.assertEqual(self.post_dispute().status_code, 200)
        html = self.client.get(self.public_path).get_data(as_text=True)
        self.assertIn("已提交异议", html)
        self.assertNotIn('name="signature_data"', html)
        with self.client.session_transaction() as flask_session:
            stale_nonce = flask_session.get("reconciliation_action_nonce", "stale")
        rejected = self.client.post(
            self.public_path,
            data={
                "action": "confirm",
                "action_nonce": stale_nonce,
                "confirmed_name": "迟到确认",
                "signature_data": self.PNG_DATA_URL,
            },
        )
        self.assertEqual(rejected.status_code, 409)
        with app.get_db() as conn:
            statement = app.fetch_reconciliation_statement(conn, self.statement_id)
        self.assertEqual(statement["status"], "disputed")

    def test_confirmed_statement_rejects_dispute_and_remains_readable_until_expiry(self):
        self.authorize()
        self.assertEqual(self.post_confirm().status_code, 200)
        html = self.client.get(self.public_path).get_data(as_text=True)
        self.assertIn("已确认", html)
        self.assertNotIn('name="dispute_content"', html)
        rejected = self.client.post(
            self.public_path,
            data={
                "action": "dispute",
                "action_nonce": "stale",
                "dispute_name": "迟到异议",
                "dispute_content": "不能覆盖",
            },
        )
        self.assertEqual(rejected.status_code, 409)
        with app.get_db() as conn:
            conn.execute(
                "UPDATE reconciliation_statements SET access_expires_at = ? WHERE id = ?",
                ("2000-01-01T00:00:00", self.statement_id),
            )
        self.assertEqual(self.client.get(self.public_path).status_code, 410)

    def test_reset_and_void_invalidate_session_and_void_prevents_reads(self):
        self.authorize()
        with app.get_db() as conn:
            replacement = app.reset_reconciliation_statement_access(
                conn, self.statement_id
            )
        replacement_path = f"/reconciliation/{replacement['token']}"
        self.assertIn(
            "请输入6位临时密码",
            self.client.get(replacement_path).get_data(as_text=True),
        )
        with app.get_db() as conn:
            app.void_reconciliation_statement(conn, self.statement_id, "作废", "admin")
        self.assertEqual(self.client.get(replacement_path).status_code, 410)

    def test_signature_image_requires_exact_public_session_or_finance_backend_access(self):
        self.authorize()
        self.post_confirm()
        with app.get_db() as conn:
            statement = dict(app.fetch_reconciliation_statement(conn, self.statement_id))
        image_path = f"/reconciliation-signature/{statement['signature_image']}"
        allowed = self.client.get(image_path)
        self.assertEqual(allowed.status_code, 200)
        self.assertEqual(allowed.headers["X-Content-Type-Options"], "nosniff")
        self.assertIn("no-store", allowed.headers["Cache-Control"])
        allowed.close()

        with app.get_db() as conn:
            conn.execute(
                "UPDATE reconciliation_statements SET status = 'void' WHERE id = ?",
                (self.statement_id,),
            )
        self.assertEqual(self.client.get(image_path).status_code, 404)
        with app.get_db() as conn:
            conn.execute(
                """
                UPDATE reconciliation_statements
                SET status = 'confirmed', access_expires_at = '2000-01-01T00:00:00'
                WHERE id = ?
                """,
                (self.statement_id,),
            )
        self.assertEqual(self.client.get(image_path).status_code, 404)
        with app.get_db() as conn:
            conn.execute(
                "UPDATE reconciliation_statements SET access_expires_at = ? WHERE id = ?",
                (statement["access_expires_at"], self.statement_id),
            )

        with self.client.session_transaction() as flask_session:
            access = dict(flask_session["reconciliation_access"])
            access["password_version"] += 1
            flask_session["reconciliation_access"] = access
        self.assertEqual(self.client.get(image_path).status_code, 404)

        self.client = app.app.test_client()
        self.assertEqual(self.client.get(image_path).status_code, 404)
        self.assertEqual(
            self.client.get(
                f"/reconciliation-signature/../{statement['signature_image']}"
            ).status_code,
            404,
        )

        with app.get_db() as conn:
            conn.execute(
                """
                INSERT INTO users (
                    username, password_hash, role, active, can_manage_finance,
                    created_at, updated_at
                ) VALUES ('signature-finance', 'hash', 'operator', 1, 1, ?, ?)
                """,
                (self.now, self.now),
            )
        with self.client.session_transaction() as flask_session:
            flask_session.update(
                admin_logged_in=True,
                admin_username="signature-finance",
                admin_role="operator",
            )
        backend_allowed = self.client.get(image_path)
        self.assertEqual(backend_allowed.status_code, 200)
        backend_allowed.close()

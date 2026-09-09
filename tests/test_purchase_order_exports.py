"""Purchase document contracts: saved data, category layout, and permission boundary."""
from datetime import datetime
from io import BytesIO
import importlib
import shutil
import subprocess
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import app
from tests import test_purchase_orders as order_tests
from tests.test_purchase_orders import payload


def sample_order(category="other"):
    return dict(order_no="PO-20260909-0042", category=category, status="ordered",
                purchased_at="2026-09-09", supplier_name="宁波精密制造有限公司",
                supplier_code="000042", supplier_contact="张工", supplier_phone="0574-01234567",
                supplier_email="supplier@example.com", supplier_address="浙江省宁波市供应商旧地址",
                delivery_address="宁波市余姚市低塘街道厂区一号仓库", recipient="李师傅",
                recipient_phone="0013800000000", remark="请工作日送货，来货前电话联系。",
                currency="CNY", company_profile=dict(company_name="宁波市杰德机械科技有限公司",
                contact="采购部王工", phone="0574-86543210", email="buy@example.com", address="公司办公地址"))


def sample_items(count=1):
    return [dict(item_name=f"精密支架{i + 1:03}", drawing_no=f"000{i + 1:03}", spec="M12 非标件",
                 material="不锈钢304", length=1200.5, width=600, height=400, thickness=2.5,
                 dimension_text="按图加工 120×60", surface="拉丝", unit="件", ordered_quantity=3,
                 unit_price_minor=87654321, line_total_minor=262962963,
                 expected_at="2026-09-20", remark="去除毛刺，表面不得有划伤。") for i in range(count)]


def pdf_pages(stream):
    executable = shutil.which("pdftotext")
    if not executable:
        raise unittest.SkipTest("PDF text QA requires Poppler pdftotext")
    result = subprocess.run([executable, "-layout", "-", "-"], input=stream.getvalue(),
                            capture_output=True, check=True)
    return [page for page in result.stdout.decode().split("\f") if page.strip()]


def workbook_text(stream):
    return " ".join(str(cell.value or "") for row in load_workbook(stream).active for cell in row)


class PurchaseOrderBuilderTests(unittest.TestCase):
    def setUp(self):
        self.docs = importlib.import_module("procurement_documents")
        self.order, self.items = sample_order(), sample_items()

    def test_category_columns_do_not_export_irrelevant_storage_fields(self):
        contracts = {
            "raw_material": (["材质", "长", "宽", "厚度", "表面", "数量", "预计到货日期", "其他要求"], ["高", "图号", "规格", "单价"]),
            "carton": (["材质", "长", "宽", "高", "数量", "单价", "金额", "预计到货日期", "其他要求"], ["厚度", "表面", "图号"]),
            "outsourcing": (["物品／图号", "材质", "尺寸／厚度／表面", "单位", "数量", "单价", "金额"], ["规格", "高"]),
            "other": (["物品／规格", "材质／尺寸／厚度／表面", "单位", "数量", "单价", "金额"], ["图号", "高"]),
        }
        for category, (required, forbidden) in contracts.items():
            with self.subTest(category=category):
                sheet = load_workbook(self.docs.build_purchase_order_workbook(sample_order(category), self.items, include_prices=True)).active
                rows = [[cell.value for cell in row] for row in sheet]
                header = next(row for row in rows if "数量" in row)
                for value in required:
                    self.assertIn(value, header)
                for value in forbidden:
                    self.assertNotIn(value, header)
                text = "".join(pdf_pages(self.docs.build_purchase_order_pdf(sample_order(category), self.items, include_prices=True)))
                for value in required:
                    self.assertIn(value, "".join(text.split()))

    def test_workbook_typed_cells_snapshots_and_print_contract(self):
        sheet = load_workbook(self.docs.build_purchase_order_workbook(self.order, self.items, include_prices=True))["采购订单"]
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(str(sheet.page_setup.paperSize), str(sheet.PAPERSIZE_A4))
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight), (1, 0))
        self.assertTrue(sheet.print_area)
        self.assertTrue(sheet.print_title_rows)
        cells = [cell for row in sheet for cell in row]
        self.assertTrue(any(cell.value == datetime(2026, 9, 20) and cell.is_date for cell in cells))
        self.assertTrue(any(cell.value == 3 and cell.data_type == "n" for cell in cells))
        self.assertTrue(any(cell.value == 876543.21 and "¥" in cell.number_format for cell in cells))
        self.assertTrue(any(cell.value == "PO-20260909-0042" and cell.data_type == "s" for cell in cells))
        text = " ".join(str(cell.value or "") for cell in cells)
        for value in (self.order["delivery_address"], self.order["supplier_address"], "000042", "0013800000000", "buy@example.com", "RMB／人民币"):
            self.assertIn(value, text)

    def test_price_hidden_has_no_price_data_in_any_xlsx_part_or_pdf(self):
        for category in ("raw_material", "carton", "outsourcing", "other"):
            order = sample_order(category)
            stream = self.docs.build_purchase_order_workbook(order, self.items, include_prices=False)
            with ZipFile(stream) as archive:
                serialized = b"".join(archive.read(name) for name in archive.namelist()).decode("utf-8")
            text = " ".join(pdf_pages(self.docs.build_purchase_order_pdf(order, self.items, include_prices=False)))
            for token in ("单价", "金额", "合计", "876543.21", "876,543.21", "2629629.63", "2,629,629.63", "87654321", "262962963"):
                self.assertNotIn(token, serialized)
                self.assertNotIn(token, text)

    def test_missing_price_is_not_treated_as_free_and_saved_totals_are_used(self):
        self.items[0]["line_total_minor"] = 12345
        text = workbook_text(self.docs.build_purchase_order_workbook(self.order, self.items, include_prices=True))
        self.assertIn("123.45", text)
        self.items[0].update(unit_price_minor=None, line_total_minor=None)
        for builder, extract in ((self.docs.build_purchase_order_workbook, workbook_text), (self.docs.build_purchase_order_pdf, lambda s: " ".join(pdf_pages(s)))):
            text = extract(builder(self.order, self.items, include_prices=True))
            self.assertIn("未录价", text)
            self.assertIn("未完整录价", text)

    def test_literal_formula_like_and_markup_text_is_not_executed(self):
        self.items[0].update(item_name='=HYPERLINK("https://example.com","危险")', spec="<b>原样</b>&", remark="=1+2")
        sheet = load_workbook(self.docs.build_purchase_order_workbook(self.order, self.items, include_prices=True)).active
        self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))
        text = " ".join(pdf_pages(self.docs.build_purchase_order_pdf(self.order, self.items, include_prices=True)))
        self.assertIn("<b>原样</b>&", text)
        self.assertIn("=1+2", text)

    def test_pdf_landscape_multiple_pages_repeat_headers_and_extract_chinese(self):
        stream = self.docs.build_purchase_order_pdf(self.order, sample_items(70), include_prices=True)
        pages = pdf_pages(stream)
        self.assertGreater(len(pages), 1)
        self.assertRegex(stream.getvalue(), rb"/MediaBox\s*\[\s*0\s+0\s+841\.\d+\s+595\.\d+")
        for page in pages:
            if "精密支架" in page:
                self.assertIn("数量", page)
                self.assertIn("预计到货日期", "".join(page.split()))
        text = "".join("".join(pages).split())
        for value in ("采购订单", "精密支架070", "宁波精密制造有限公司", self.order["delivery_address"], "采购部王工"):
            self.assertIn(value, text)

    def test_default_company_and_blank_optional_columns(self):
        self.order.pop("company_profile")
        for field in ("material", "dimension_text", "thickness", "surface", "length", "width", "height"):
            self.items[0][field] = None
        sheet = load_workbook(self.docs.build_purchase_order_workbook(self.order, self.items, include_prices=False)).active
        values = [cell.value for row in sheet for cell in row]
        self.assertIn("宁波市杰德机械科技有限公司", values)
        self.assertNotIn("材质／尺寸／厚度／表面", values)

    def test_large_quantity_is_not_rounded_in_pdf_and_title_has_room(self):
        self.items[0]["ordered_quantity"] = 2147483647
        text = "".join(pdf_pages(self.docs.build_purchase_order_pdf(self.order, self.items, include_prices=False)))
        self.assertIn("2147483647", "".join(text.split()).replace(",", ""))
        sheet = load_workbook(self.docs.build_purchase_order_workbook(self.order, self.items, include_prices=False)).active
        self.assertGreaterEqual(sheet.row_dimensions[1].height, sheet["A1"].font.sz + 10)

    def test_max_length_remark_survives_page_splits_without_excel_height_clipping(self):
        self.items[0]["remark"] = "到货检查" * 999 + "结束标记"
        pages = pdf_pages(self.docs.build_purchase_order_pdf(self.order, self.items, include_prices=False))
        self.assertGreater(len(pages), 1)
        self.assertIn("结束标记", "".join("".join(pages).split()))
        sheet = load_workbook(self.docs.build_purchase_order_workbook(self.order, self.items, include_prices=False)).active
        values = "".join(str(cell.value or "") for row in sheet for cell in row).replace("\n", "")
        self.assertIn(self.items[0]["remark"], values)
        self.assertTrue(all(row.height <= 409 for row in sheet.row_dimensions.values()))

    def test_carton_header_date_print_width_and_pdf_delivery_block_stay_readable(self):
        sheet = load_workbook(self.docs.build_purchase_order_workbook(sample_order("carton"), self.items, include_prices=True)).active
        date_cell = next(cell for cell in sheet[4] if cell.is_date)
        span = next((area for area in sheet.merged_cells.ranges if date_cell.coordinate in area), None)
        width = sum(sheet.column_dimensions[get_column_letter(col)].width for col in range(span.min_col, span.max_col + 1)) if span else sheet.column_dimensions[date_cell.column_letter].width
        self.assertGreaterEqual(width, 14)
        items = sample_items(32)
        for item in items:
            item["remark"] = "请按图纸检验孔位和外形尺寸，装箱前确认无磕碰、无锈蚀。"
        pages = pdf_pages(self.docs.build_purchase_order_pdf(self.order, items, include_prices=True))
        delivery_page = next(page for page in pages if "收货地址" in page)
        self.assertIn("收件人", delivery_page)
        self.assertIn("buy@example.com", delivery_page)


class PurchaseOrderExportRouteTests(unittest.TestCase):
    setUp = order_tests.PurchaseOrderRouteTests.setUp
    tearDown = order_tests.PurchaseOrderRouteTests.tearDown
    login = order_tests.PurchaseOrderRouteTests.login
    create = order_tests.PurchaseOrderRouteTests.create

    def test_export_permission_and_missing_invalid_cancelled_orders(self):
        oid = self.create()
        for extension in ("xlsx", "pdf"):
            path = f"/admin/purchases/orders/{oid}/export.{extension}"
            self.login("outsider")
            self.assertEqual(self.client.get(path).status_code, 302)
            self.client = app.app.test_client()
            self.assertEqual(self.client.get(path).status_code, 302)
            self.login("buyer")
            for bad_id in ("0", "999", "-1", "abc", str(2**64)):
                self.assertEqual(self.client.get(f"/admin/purchases/orders/{bad_id}/export.{extension}").status_code, 404)
        self.client.post(f"/admin/purchases/orders/{oid}/cancel")
        for extension in ("xlsx", "pdf"):
            response = self.client.get(f"/admin/purchases/orders/{oid}/export.{extension}")
            self.assertEqual(response.status_code, 409)
            self.assertNotIn("Content-Disposition", response.headers)

    def test_disposition_mime_and_immutable_filename(self):
        oid = self.create()
        for extension, mime in (("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), ("pdf", "application/pdf")):
            for query, disposition in (("", "attachment" if extension == "xlsx" else "inline"), ("?download=1", "attachment"), ("?download=true", "attachment" if extension == "xlsx" else "inline")):
                response = self.client.get(f"/admin/purchases/orders/{oid}/export.{extension}{query}")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.mimetype, mime)
                self.assertTrue(response.headers["Content-Disposition"].startswith(disposition))
                self.assertIn(f"PO-20260909-0001-purchase-order.{extension}", response.headers["Content-Disposition"])

    def test_user_derived_prices_resist_queries_and_snapshots_survive_master_changes(self):
        oid = self.create(payload(rows=[dict(item_name="测试件", quantity="3", unit_price="876543.21", expected_at="2026-09-20")]))
        with app.get_db() as conn:
            conn.execute("UPDATE suppliers SET name='改名不可出现',address='供应商新地址'")
            conn.execute("UPDATE purchase_delivery_profiles SET delivery_address='收货新地址'")
            conn.execute("UPDATE reconciliation_company_profile SET contact='公司联系甲',email='office@example.com'")
        for name, allowed in (("buyer", True), ("reader", False), ("blind", False), ("admin", True)):
            self.login(name)
            for extension in ("xlsx", "pdf"):
                response = self.client.get(f"/admin/purchases/orders/{oid}/export.{extension}?include_prices=1&can_view_purchase_prices=1&prices=true")
                self.assertEqual(response.status_code, 200)
                stream = BytesIO(response.data)
                text = workbook_text(stream) if extension == "xlsx" else " ".join(pdf_pages(stream))
                self.assertEqual("单价" in text, allowed)
                self.assertEqual("876543.21" in text.replace(",", ""), allowed)
                for snapshot in ("供方甲", "旧地址", "厂区地址", "公司联系甲", "office@example.com"):
                    self.assertIn(snapshot, text)
                for current in ("改名不可出现", "供应商新地址", "收货新地址"):
                    self.assertNotIn(current, text)

    def test_detail_export_links_only_for_non_cancelled_orders(self):
        oid = self.create()
        html = self.client.get(f"/admin/purchases/orders/{oid}").get_data(as_text=True)
        for extension in ("xlsx", "pdf"):
            self.assertIn(f"/admin/purchases/orders/{oid}/export.{extension}", html)
        self.client.post(f"/admin/purchases/orders/{oid}/cancel")
        html = self.client.get(f"/admin/purchases/orders/{oid}").get_data(as_text=True)
        self.assertNotIn("/export.", html)


if __name__ == "__main__":
    unittest.main()

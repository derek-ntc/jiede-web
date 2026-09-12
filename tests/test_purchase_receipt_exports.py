"""Receipt and inventory documents exercise actual builders and HTTP permissions."""
from datetime import datetime
from io import BytesIO
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import app
import procurement_inventory as inventory
import procurement_documents as docs
from tests.test_purchase_order_exports import pdf_pages, workbook_text
from tests import test_purchase_inventory_routes as fixtures


def sample_receipt():
    return dict(receipt_no="RC-20260912-0001", order_no="PO-ORIGINAL", category="raw_material",
                supplier_name="原供应商", received_at="2026-09-12", status="posted",
                posted_by="收货员甲", posted_at="2026-09-12T10:00:00", remark="首批到货")


def receipt_items(count=1):
    return [dict(item_name=f"实际镀锌板{i+1:03}", material="DC51D", length=1290, width=1250,
                 thickness=1.4, unit="张", actual_quantity=100, qualified_quantity=98,
                 location_code="RAW-A", location_name="原库位", invoice_status="pending", lot_no=f"LOT-{i+1:03}",
                 remark="=1+2 <b>原样</b>", ordered=dict(item_name="订购镀锌板", material="AB", length=1300,
                 width=1250, thickness=1.5, ordered_quantity=120, unit_price_minor=87654321)) for i in range(count)]


class PurchaseReceiptBuilderTests(unittest.TestCase):
    def builder(self, name):
        result = getattr(docs, name, None)
        self.assertTrue(callable(result), f"missing {name}")
        return result

    def assert_print_contract(self, sheet):
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(str(sheet.page_setup.paperSize), str(sheet.PAPERSIZE_A4))
        self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight), (1, 0))
        self.assertTrue(sheet.print_area)
        self.assertTrue(sheet.print_title_rows)
        self.assertFalse(any(c.data_type == "f" for row in sheet for c in row))
        # A 10pt font must not shrink below roughly 8pt on landscape A4.
        self.assertLessEqual(sum(sheet.column_dimensions[get_column_letter(i)].width
                                 for i in range(1, sheet.max_column + 1)), 180)

    def test_receipt_workbook_has_comparison_quantities_typed_date_and_audit(self):
        stream = self.builder("build_purchase_receipt_workbook")(sample_receipt(), receipt_items(), include_prices=True)
        sheet = load_workbook(stream).active
        self.assert_print_contract(sheet)
        text = "".join(workbook_text(stream).split())
        for value in ("RC-20260912-0001", "PO-ORIGINAL", "原供应商", "收货员甲", "实际镀锌板001", "订购镀锌板", "1300", "1290", "LOT-001", "RAW-A", "原库位", "待开票", "订单资料", "实际资料"):
            self.assertIn(value, text)
        for value in (100, 98, 120):
            self.assertTrue(any(c.value == value and c.data_type == "n" for row in sheet for c in row))
        self.assertTrue(any(c.value == datetime(2026, 9, 12) and c.is_date for row in sheet for c in row))

    def test_receipt_pdf_repeats_headers_and_preserves_chinese_and_long_text(self):
        items = receipt_items(30)
        items[0]["remark"] = "到货检查" * 150 + "结束标记"
        stream = self.builder("build_purchase_receipt_pdf")(sample_receipt(), items, include_prices=False)
        self.assertRegex(stream.getvalue(), rb"/MediaBox\s*\[\s*0\s+0\s+841\.\d+\s+595\.\d+")
        pages = pdf_pages(stream)
        self.assertGreater(len(pages), 1)
        for page in pages:
            compact = "".join(page.split())
            if "实际镀锌板" in compact or "到货检查" in compact:
                self.assertIn("实际到货", compact)
        text = "".join("".join(pages).split())
        for value in ("实际镀锌板030", "收货员甲", "订购镀锌板", "结束标记", "=1+2<b>原样</b>"):
            self.assertIn(value, text)
        self.assertEqual(text.count("到货检查"), 150)

    def test_inventory_export_includes_active_filters_current_stock_and_location(self):
        row = dict(receipt_items()[0], **sample_receipt(), opening_quantity=98, available_quantity=75,
                   amount_minor=6574074075, unit_price_minor=87654321)
        stream = self.builder("build_purchase_inventory_workbook")(
            dict(category="raw_material", q="DC51D", received_from="2026-09-01", include_zero="1"), [row], include_prices=True)
        sheet = load_workbook(stream).active
        self.assert_print_contract(sheet)
        text = "".join(workbook_text(stream).split())
        for value in ("物品关键词：DC51D", "入库日期从：2026-09-01", "包含零库存：是", "当前可用数量", "原库位", "PO-ORIGINAL", "RC-20260912-0001", "LOT-001"):
            self.assertIn(value, text)
        self.assertTrue(any(c.value == 75 and c.data_type == "n" for row in sheet for c in row))
        date_header = next(c for row in sheet for c in row if c.value == "入库日期")
        self.assertGreaterEqual(sheet.column_dimensions[date_header.column_letter].width, 13)

    def test_long_receipt_price_only_appears_on_first_physical_row(self):
        for price, expected in ((87654321, 876543.21), (0, 0), (None, "未录价")):
            with self.subTest(price=price):
                items = receipt_items()
                items[0]["remark"] = "到货检查" * 200 + "结束标记"
                items[0]["ordered"]["unit_price_minor"] = price
                sheet = load_workbook(docs.build_purchase_receipt_workbook(sample_receipt(), items, include_prices=True)).active
                header = next(c for row in sheet for c in row if c.value == "订单单价")
                self.assertGreater(sheet.max_row, header.row + 2, "fixture must produce continuation rows")
                self.assertEqual(sheet.cell(header.row + 1, header.column).value, expected)
                self.assertTrue(all(sheet.cell(row, header.column).value is None
                                    for row in range(header.row + 2, sheet.max_row + 1)))
                self.assertIn("结束标记", "".join(str(c.value or "") for row in sheet for c in row))

    def test_long_inventory_price_and_amount_only_appear_on_first_physical_row(self):
        for price, amount, expected_price, expected_amount in (
                (87654321, 6574074075, 876543.21, 65740740.75),
                (0, 0, 0, 0), (None, None, "未录价", "未录价")):
            with self.subTest(price=price):
                item = dict(receipt_items()[0], **sample_receipt(), opening_quantity=98, available_quantity=75,
                            material="长材质说明" * 200 + "结束标记", unit_price_minor=price, amount_minor=amount)
                sheet = load_workbook(docs.build_purchase_inventory_workbook(
                    {"category": "raw_material"}, [item], include_prices=True)).active
                for label, expected in (("单价", expected_price), ("当前金额", expected_amount)):
                    header = next(c for row in sheet for c in row if c.value == label)
                    self.assertGreater(sheet.max_row, header.row + 2, "fixture must produce continuation rows")
                    self.assertEqual(sheet.cell(header.row + 1, header.column).value, expected)
                    self.assertTrue(all(sheet.cell(row, header.column).value is None
                                        for row in range(header.row + 2, sheet.max_row + 1)))
                self.assertIn("结束标记", "".join(str(c.value or "") for row in sheet for c in row).replace("\n", ""))

    def test_price_hidden_has_no_secrets_or_raw_order_terms_in_files(self):
        receipt, items = sample_receipt(), receipt_items()
        inventory = dict(items[0], **receipt, opening_quantity=98, available_quantity=75,
                         unit_price_minor=87654321, amount_minor=6574074075)
        cases = (("build_purchase_receipt_workbook", (receipt, items)),
                 ("build_purchase_receipt_pdf", (receipt, items)),
                 ("build_purchase_inventory_workbook", ({"category": "raw_material"}, [inventory])))
        for name, args in cases:
            with self.subTest(builder=name):
                stream = self.builder(name)(*args, include_prices=False)
                if name.endswith("pdf"):
                    text = "".join(pdf_pages(stream))
                else:
                    with ZipFile(stream) as archive:
                        text = b"".join(archive.read(n) for n in archive.namelist()).decode()
                for forbidden in ("单价", "金额", "87654321", "876543.21", "876,543.21", "6574074075", "产品交付时必须标识明确", "订单要求纳入供应商考核"):
                    self.assertNotIn(forbidden, text)


class PurchaseReceiptExportRouteTests(unittest.TestCase):
    fixture_setup = fixtures.PurchaseInventoryRouteTests.fixture_setup
    setUp = fixtures.PurchaseInventoryRouteTests.setUp
    tearDown = fixtures.PurchaseInventoryRouteTests.tearDown
    create_order = fixtures.PurchaseInventoryRouteTests.create_order
    login = fixtures.PurchaseInventoryRouteTests.login
    scalar = fixtures.PurchaseInventoryRouteTests.scalar
    page_data = fixtures.PurchaseInventoryRouteTests.page_data
    payload = fixtures.PurchaseInventoryRouteTests.payload

    def test_saved_document_context_survives_master_and_order_edits(self):
        rid = self.scalar("SELECT id FROM purchase_receipts")
        loader = getattr(inventory, "load_purchase_receipt_document", None)
        self.assertTrue(callable(loader), "missing snapshot document projection")
        with app.get_db() as conn:
            before = loader(conn, rid, include_prices=True)
            conn.execute("UPDATE purchase_orders SET supplier_name='后改供方',order_no='PO-CHANGED' WHERE id=?", (self.order_id,))
            conn.execute("UPDATE purchase_order_items SET item_name='后改物品',ordered_quantity=200,unit_price_minor=1 WHERE id=?", (self.item_id,))
            conn.execute("UPDATE warehouse_locations SET code='NEW-A',name='后改库位' WHERE id=?", (self.location_id,))
            inventory.ensure_purchase_inventory_tables(conn)
            self.assertEqual(loader(conn, rid, include_prices=True), before)
            safe_header, safe_rows = loader(conn, rid, include_prices=False)
            self.assertNotIn("unit_price_minor", str(safe_rows))
            self.assertNotIn("ordered_snapshot_json", str(safe_rows))
        for fmt in ("xlsx", "pdf"):
            response = self.client.get(f"/admin/purchase-receipts/records/{rid}/export.{fmt}")
            self.assertEqual(response.status_code, 200)
            text = workbook_text(BytesIO(response.data)) if fmt == "xlsx" else "".join(pdf_pages(BytesIO(response.data)))
            text = "".join(text.split())
            for expected in ("PO-RAW", "供应商甲", "RAW-A", "原料区", "镀锌板"):
                self.assertIn(expected, text)
            for forbidden in ("后改", "PO-CHANGED", "NEW-A"):
                self.assertNotIn(forbidden, text)

    def test_legacy_document_snapshot_backfill_is_deterministic_and_idempotent(self):
        required = {"order_no", "category", "supplier_name"}
        with app.get_db() as conn:
            self.assertTrue(required <= {r[1] for r in conn.execute("PRAGMA table_info(purchase_receipts)")})
            self.assertIn("ordered_snapshot_json", {r[1] for r in conn.execute("PRAGMA table_info(purchase_receipt_items)")})
            conn.execute("UPDATE purchase_receipts SET order_no='',category='',supplier_name=''")
            conn.execute("UPDATE purchase_receipt_items SET ordered_snapshot_json='',location_code='',location_name=''")
            inventory.ensure_purchase_inventory_tables(conn)
            rid = conn.execute("SELECT id FROM purchase_receipts").fetchone()[0]
            before = inventory.load_purchase_receipt_document(conn, rid, include_prices=True)
            self.assertEqual(before[0]["order_no"], "PO-RAW")
            self.assertEqual(before[1][0]["ordered"]["ordered_quantity"], 100)
            self.assertEqual(before[1][0]["location_code"], "RAW-A")
            conn.execute("UPDATE purchase_orders SET order_no='LATER'")
            conn.execute("UPDATE warehouse_locations SET name='LATER' WHERE id=?", (self.location_id,))
            inventory.ensure_purchase_inventory_tables(conn)
            self.assertEqual(inventory.load_purchase_receipt_document(conn, rid, include_prices=True), before)

    def test_lot_locations_snapshot_at_receipt_and_transfer_and_survive_rename(self):
        with app.get_db() as conn:
            conn.execute("UPDATE warehouse_locations SET code='TARGET-B',name='转库时名称' WHERE id=?", (self.target_id,))
            result = inventory.transfer_purchase_inventory(conn, self.lot_id, 20, self.target_id, "移库", "receiver",
                                                           "2026-09-12T12:00:00", 1, "snapshot-transfer")
            conn.execute("UPDATE warehouse_locations SET code='LATER-'||id,name='改名后库位'")
            conn.execute("UPDATE purchase_orders SET order_no='LATER-ORDER'")
            inventory.ensure_purchase_inventory_tables(conn)
            rows = inventory.fetch_purchase_inventory(conn, {"category": "raw_material"}, False)
            by_id = {row["id"]: row for row in rows}
            self.assertEqual((by_id[self.lot_id]["location_code"], by_id[self.lot_id]["location_name"]), ("RAW-A", "原料区"))
            self.assertEqual((by_id[result["target_lot_id"]]["location_code"], by_id[result["target_lot_id"]]["location_name"]), ("TARGET-B", "转库时名称"))
            self.assertEqual(by_id[self.lot_id]["order_no"], "PO-RAW")
            self.assertEqual(len(inventory.fetch_purchase_inventory(conn, {"category": "raw_material", "order_no": "PO-RAW"}, False)), 2)

    def test_legacy_lot_location_backfill_is_not_overwritten(self):
        with app.get_db() as conn:
            self.assertTrue({"location_code", "location_name"} <= {r[1] for r in conn.execute("PRAGMA table_info(purchase_inventory_lots)")})
            conn.execute("UPDATE purchase_inventory_lots SET location_code='',location_name=''")
            inventory.ensure_purchase_inventory_tables(conn)
            conn.execute("UPDATE warehouse_locations SET name='修改后'")
            inventory.ensure_purchase_inventory_tables(conn)
            row = inventory.fetch_purchase_inventory(conn, {"category": "raw_material"}, False)[0]
            self.assertEqual((row["location_code"], row["location_name"]), ("RAW-A", "原料区"))

    def test_receipt_export_permissions_links_and_price_redaction(self):
        rid = self.scalar("SELECT id FROM purchase_receipts")
        root = f"/admin/purchase-receipts/records/{rid}"
        for fmt in ("xlsx", "pdf"):
            url = root + "/export." + fmt
            self.assertIn(url, self.client.get(root).text)
            self.login("blind")
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["Cache-Control"], "private, no-store")
            text = workbook_text(BytesIO(response.data)) if fmt == "xlsx" else "".join(pdf_pages(BytesIO(response.data)))
            self.assertIn("实际镀锌板", "".join(text.split()))
            for secret in ("单价", "9876.54", "987654"):
                self.assertNotIn(secret, text)
            self.login("reader")
            self.assertEqual(self.client.get(url).status_code, 302)
            self.login("receiver")
        self.assertEqual(self.client.get(root + "/export.csv").status_code, 404)
        self.assertEqual(self.client.get("/admin/purchase-receipts/records/999999/export.xlsx").status_code, 404)

    def test_inventory_export_uses_filters_and_separate_permission(self):
        url = self.inventory_url + "/export.xlsx"
        self.assertIn(url, self.client.get(self.inventory_url, query_string={"q": self.lot_no}).text)
        self.login("blind")
        for q, present in ((self.lot_no, True), ("not-found", False)):
            response = self.client.get(url, query_string={"q": q})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["Cache-Control"], "private, no-store")
            text = workbook_text(BytesIO(response.data))
            self.assertEqual("实际镀锌板" in "".join(text.split()), present)
            self.assertNotIn("9876.54", text)
        self.login("reader")
        self.assertEqual(self.client.get(url).status_code, 302)
        self.login("receiver")
        self.assertEqual(self.client.get("/admin/purchase-inventory/invalid/export.xlsx").status_code, 404)
        self.assertEqual(self.client.get(url, query_string={"received_from": "bad"}).status_code, 400)

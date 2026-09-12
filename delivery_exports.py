"""Price-free delivery-note exports with a shared portrait-A4 layout."""

from io import BytesIO
from math import ceil
from unicodedata import east_asian_width
from xml.sax.saxutils import escape as xml_escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


COLUMN_WIDTHS = (14, 20, 20, 22, 12, 12, 12, 16)
PRINT_WIDTH = 194 * mm
HEADER_COLOR = "5C7280"
COMPANY_NAME = "宁波市杰德机械科技有限公司"


def delivery_note_export_columns(payload):
    quantity_label = "每套数量" if payload["is_assembly"] else "订单数量"
    quantity_key = "quantity_per_set" if payload["is_assembly"] else "order_quantity"
    return (
        ("序号", "index"), ("产品图号", "drawing_no"), ("产品名称", "product_name"),
        ("规格型号", "specification"), ("单位", "unit"), (quantity_label, quantity_key),
        ("实发数量", "shipped_quantity"), ("备注", "remark"),
    )


def _delivery_note_rows(payload):
    return [[item[key] for _, key in delivery_note_export_columns(payload)] for item in payload["items"]]


def _delivery_note_metadata_rows(payload):
    rows = [
        ("客户名称", payload["customer"] or "-", "送货单号", payload["document_no"]),
        ("收货人", payload["recipient_name"], "收货电话", payload["recipient_phone"]),
        ("收货地址", payload["address"], "订单号", " / ".join(dict.fromkeys(
            item["order_no"] for item in payload["items"] if item["order_no"] and item["order_no"] != "-"
        )) or "-"),
    ]
    if payload["is_assembly"]:
        rows.append(("组装件名称", payload["assembly_drawing_no"], "整套数量", f"{payload['assembly_set_quantity']} 套"))
    return rows


def _text(value):
    return "" if value is None else str(value)


def _excel_row_height(values_and_widths):
    """Allow wrapped address, specification and remarks to remain visible."""
    lines = 1
    for value, width in values_and_widths:
        # Excel widths use the default font's digit width (approximately 7 px).
        available = max(1, width * 7 - 6)
        line_count = sum(max(1, ceil(sum(
            14.7 if east_asian_width(character) in "WF" else 7.4
            for character in line
        ) / available)) for line in _text(value).split("\n"))
        lines = max(lines, line_count)
    return max(22, lines * 14 + 8)


def build_delivery_note_xlsx(payload):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "送货单"
    sheet.sheet_view.showGridLines = False
    border = Border(*(Side(style="thin", color="000000") for _ in range(4)))
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    header_fill = PatternFill("solid", fgColor=HEADER_COLOR)
    sheet.merge_cells("A2:H2")
    sheet["A2"] = COMPANY_NAME
    sheet["A2"].font = Font(name="宋体", size=18)
    sheet["A2"].fill = white_fill
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 28
    sheet.merge_cells("A3:H3")
    sheet["A3"] = "送货单"
    sheet["A3"].font = Font(name="宋体", size=18, bold=True)
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[3].height = 28

    metadata_rows = _delivery_note_metadata_rows(payload)
    for row_no, values in enumerate(metadata_rows, 4):
        for column, value in ((1, values[0]), (2, values[1]), (5, values[2]), (6, values[3])):
            cell = sheet.cell(row_no, column, value)
            is_label = column in (1, 5)
            cell.font = Font(name="宋体", size=11, color="000000")
            cell.alignment = Alignment(
                horizontal="center" if is_label else "left",
                vertical="center", wrap_text=not is_label,
            )
            cell.border = border
            cell.fill = white_fill
        sheet.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=4)
        sheet.merge_cells(start_row=row_no, start_column=6, end_row=row_no, end_column=8)
        for column in range(1, 9):
            sheet.cell(row_no, column).border = border
        sheet.row_dimensions[row_no].height = _excel_row_height((
            (values[1], sum(COLUMN_WIDTHS[1:4])),
            (values[3], sum(COLUMN_WIDTHS[5:8])),
        ))

    header_row = 4 + len(metadata_rows)
    for column, (label, _) in enumerate(delivery_note_export_columns(payload), 1):
        cell = sheet.cell(header_row, column, label)
        is_quantity = column in (6, 7)
        cell.font = Font(name="宋体", size=11, bold=True, color="000000" if is_quantity else "FFFFFF")
        cell.fill = white_fill if is_quantity else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    sheet.row_dimensions[header_row].height = 24
    for row_no, values in enumerate(_delivery_note_rows(payload), header_row + 1):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_no, column, value)
            cell.font = Font(name="宋体", size=11)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if column in (1, 5, 6, 7) else "left",
                vertical="center", wrap_text=True,
            )
        sheet.row_dimensions[row_no].height = _excel_row_height(zip(values, COLUMN_WIDTHS))

    footer_row = sheet.max_row + 2
    for start, end in ((1, 2), (3, 5), (6, 8)):
        sheet.merge_cells(start_row=footer_row, start_column=start, end_row=footer_row, end_column=end)
    for column, value in (
        (1, f"发货日期：{payload['items'][0]['shipped_at'] if payload['items'] else '-'}"),
        (3, "发货签名："), (6, "收货签名："),
    ):
        cell = sheet.cell(footer_row, column, value)
        cell.font = Font(name="宋体", size=11)
        cell.fill = white_fill
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[footer_row].height = 24
    for column, width in enumerate(COLUMN_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.page_setup.orientation = sheet.ORIENTATION_PORTRAIT
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A2:H{footer_row}"
    sheet.print_title_rows = f"2:{header_row}"
    sheet.page_margins.left = sheet.page_margins.right = 8 / 25.4
    sheet.page_margins.top = sheet.page_margins.bottom = 8 / 25.4
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_delivery_note_pdf(payload, font_name):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=8 * mm, rightMargin=8 * mm,
        topMargin=8 * mm, bottomMargin=8 * mm,
    )
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle(
        "DeliveryCell", parent=styles["BodyText"], fontName=font_name,
        fontSize=8.5, leading=11, textColor=colors.black, wordWrap="CJK",
    )
    center_style = ParagraphStyle("DeliveryCenter", parent=cell_style, alignment=1)
    header_style = ParagraphStyle("DeliveryHeader", parent=center_style, textColor=colors.white)
    title_style = ParagraphStyle(
        "DeliveryTitle", parent=center_style, fontSize=18, leading=24, spaceAfter=2 * mm,
    )
    company_style = ParagraphStyle("DeliveryCompany", parent=center_style, fontSize=16, leading=23)

    def paragraph(value, style=cell_style):
        return Paragraph(xml_escape(_text(value)).replace("\n", "<br/>"), style)

    # Excel's character widths include a small fixed cell padding. Using these
    # same proportions keeps metadata boundaries aligned with the item grid.
    pixel_widths = [width * 7 + 5 for width in COLUMN_WIDTHS]
    widths = [PRINT_WIDTH * width / sum(pixel_widths) for width in pixel_widths]
    story = [
        Table([[paragraph(COMPANY_NAME, company_style)]], colWidths=[PRINT_WIDTH]),
        paragraph("送货单", title_style),
    ]
    metadata = [[paragraph(value, center_style if index in (0, 2) else cell_style)
                 for index, value in enumerate(values)]
                for values in _delivery_note_metadata_rows(payload)]
    meta_table = Table(metadata, colWidths=[widths[0], sum(widths[1:4]), widths[4], sum(widths[5:8])])
    common_style = [
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    meta_table.setStyle(TableStyle(common_style))
    story.append(meta_table)
    columns = delivery_note_export_columns(payload)
    data = [[paragraph(label, center_style if index in (5, 6) else header_style)
             for index, (label, _) in enumerate(columns)]]
    for row in _delivery_note_rows(payload):
        data.append([paragraph(value, center_style if index in (0, 4, 5, 6) else cell_style)
                     for index, value in enumerate(row)])
    if len(data) == 1:
        data.append([paragraph("暂无发货记录")] + [""] * 7)
    item_table = Table(data, colWidths=widths, repeatRows=1)
    item_table.setStyle(TableStyle(common_style + [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + HEADER_COLOR)),
        ("BACKGROUND", (5, 0), (6, 0), colors.white),
    ]))
    story.extend([item_table, Spacer(1, 3 * mm)])
    footer = Table([[
        paragraph(f"发货日期：{payload['items'][0]['shipped_at'] if payload['items'] else '-'}"),
        paragraph("发货签名："), paragraph("收货签名："),
    ]], colWidths=[sum(widths[:2]), sum(widths[2:5]), sum(widths[5:])])
    footer.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(footer)
    doc.build(story)
    buffer.seek(0)
    return buffer

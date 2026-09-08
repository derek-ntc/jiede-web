from copy import copy
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import os
from pathlib import Path
import sys
from typing import Mapping, Sequence
from unicodedata import east_asian_width

# The repository contains a legacy PIL compatibility stub.  Load the Pillow
# distribution declared in requirements.txt before OpenPyXL imports PIL.Image.
_original_import_path = list(sys.path)
_application_root = Path(__file__).resolve().parent
try:
    sys.path = [
        entry
        for entry in sys.path
        if Path(entry or os.getcwd()).resolve() != _application_root
    ]
    import PIL as PillowPackage
finally:
    sys.path = _original_import_path

if not getattr(PillowPackage, "__version__", ""):
    raise ImportError("对账单签名导出需要 requirements.txt 中声明的 Pillow")

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from pricing import SQLITE_INTEGER_MAX, SUPPORTED_CURRENCIES, format_money_minor, normalize_currency


TAX_RATE_PPM = 130_000
TAX_DIVISOR = Decimal("1.13")
EX_TAX_UNIT_SCALE = 1_000_000
EXCEL_MAX_SIGNIFICANT_SCALED_INTEGER = 999_999_999_999_999
EXCEL_PRECISION_ERROR = "对账单数值超出 Excel 可精确表示范围，不能导出"


def calculate_reconciliation_amounts(unit_price_minor: int, quantity: int, currency: str) -> dict[str, int]:
    currency = normalize_currency(currency)
    unit_price_minor = int(unit_price_minor)
    quantity = int(quantity)
    if unit_price_minor < 0 or quantity <= 0:
        raise ValueError("对账价格和数量必须有效")

    amount_incl_tax_minor = unit_price_minor * quantity
    if amount_incl_tax_minor > SQLITE_INTEGER_MAX:
        raise ValueError("对账金额过大")

    digits = SUPPORTED_CURRENCIES[currency]
    unit_major = Decimal(unit_price_minor).scaleb(-digits)
    unit_price_ex_tax_scaled = int(
        (unit_major / TAX_DIVISOR * EX_TAX_UNIT_SCALE).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )
    )
    amount_ex_tax_minor = int(
        (Decimal(amount_incl_tax_minor) / TAX_DIVISOR).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )
    )
    amounts = {
        "unit_price_incl_tax_minor": unit_price_minor,
        "unit_price_ex_tax_scaled": unit_price_ex_tax_scaled,
        "amount_incl_tax_minor": amount_incl_tax_minor,
        "amount_ex_tax_minor": amount_ex_tax_minor,
        "tax_amount_minor": amount_incl_tax_minor - amount_ex_tax_minor,
    }
    if any(amount > SQLITE_INTEGER_MAX for amount in amounts.values()):
        raise ValueError("对账金额过大")
    return amounts


def format_reconciliation_amount_minor(value: int | None, currency: str) -> str:
    return format_money_minor(value, currency)


def format_unit_price_ex_tax_scaled(value: int | None) -> str:
    if value is None:
        return ""
    return f"{Decimal(value).scaleb(-6):.6f}"


RECONCILIATION_HEADERS = (
    "序号",
    "送货日期",
    "送货单号",
    "采购订单号",
    "物料编码",
    "物料名称",
    "规格型号",
    "单位",
    "送货数量",
    "不含税单价",
    "含税单价",
    "不含税金额",
    "含税金额",
    "备注",
)
_EXCEL_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _mapping_value(mapping: Mapping, key: str, default=""):
    try:
        value = mapping[key]
    except (KeyError, IndexError):
        return default
    return default if value is None else value


def set_excel_literal_value(cell, value) -> None:
    """Store untrusted text without allowing Excel formula evaluation."""
    text = "" if value is None else str(value)
    if text.lstrip(" \t\r\n").startswith(_EXCEL_FORMULA_PREFIXES):
        text = "'" + text
    cell.value = text


def _statement_title(statement: Mapping) -> str:
    period_start = str(_mapping_value(statement, "period_start"))
    period_end = str(_mapping_value(statement, "period_end"))
    try:
        start = datetime.strptime(period_start[:10], "%Y-%m-%d")
        end = datetime.strptime(period_end[:10], "%Y-%m-%d")
    except ValueError:
        return f"{period_start} 至 {period_end} 对账单"
    if start.year == end.year and start.month == end.month:
        return f"{start.year}年{start.month}月对账单"
    return f"{start:%Y-%m-%d} 至 {end:%Y-%m-%d} 对账单"


def _money_number_format(currency: str) -> str:
    digits = SUPPORTED_CURRENCIES[currency]
    return "#,##0" if digits == 0 else "#,##0." + ("0" * digits)


def _excel_scaled_numeric_value(value, decimal_places: int, label: str) -> float:
    """Convert a scaled integer only when Excel can preserve its business precision."""
    scaled_integer = int(value)
    if abs(scaled_integer) > EXCEL_MAX_SIGNIFICANT_SCALED_INTEGER:
        raise ValueError(f"{EXCEL_PRECISION_ERROR}（{label}）")

    numeric_value = float(Decimal(scaled_integer).scaleb(-decimal_places))
    restored_integer = int(
        (Decimal(str(numeric_value)).scaleb(decimal_places)).quantize(
            Decimal("1"),
            rounding=ROUND_HALF_UP,
        )
    )
    if restored_integer != scaled_integer:
        raise ValueError(f"{EXCEL_PRECISION_ERROR}（{label}）")
    return numeric_value


def _excel_integer_value(value, label: str) -> int:
    integer_value = int(value)
    _excel_scaled_numeric_value(integer_value, 0, label)
    return integer_value


def _scaled_integer_from_excel_result(value: float, decimal_places: int) -> int:
    return int(
        (Decimal(str(value)).scaleb(decimal_places)).quantize(
            Decimal("1"),
            rounding=ROUND_HALF_UP,
        )
    )


def _prepare_excel_numeric_rows(
    statement: Mapping,
    items: Sequence[Mapping],
    currency: str,
) -> list[dict[str, int | float]]:
    currency_digits = SUPPORTED_CURRENCIES[currency]
    numeric_rows = []
    for offset, item in enumerate(items):
        row_number = offset + 1
        numeric_rows.append(
            {
                "sort_order": _excel_integer_value(
                    _mapping_value(item, "sort_order", row_number),
                    f"第{row_number}行序号",
                ),
                "quantity": _excel_integer_value(
                    _mapping_value(item, "quantity"),
                    f"第{row_number}行数量",
                ),
                "unit_price_ex_tax": _excel_scaled_numeric_value(
                    _mapping_value(item, "unit_price_ex_tax_scaled"),
                    6,
                    f"第{row_number}行不含税单价",
                ),
                "unit_price_incl_tax": _excel_scaled_numeric_value(
                    _mapping_value(item, "unit_price_incl_tax_minor"),
                    currency_digits,
                    f"第{row_number}行含税单价",
                ),
                "amount_ex_tax": _excel_scaled_numeric_value(
                    _mapping_value(item, "amount_ex_tax_minor"),
                    currency_digits,
                    f"第{row_number}行不含税金额",
                ),
                "amount_incl_tax": _excel_scaled_numeric_value(
                    _mapping_value(item, "amount_incl_tax_minor"),
                    currency_digits,
                    f"第{row_number}行含税金额",
                ),
                "tax_amount": _excel_scaled_numeric_value(
                    _mapping_value(item, "tax_amount_minor"),
                    currency_digits,
                    f"第{row_number}行税额",
                ),
            }
        )

    statement_totals = {
        "amount_ex_tax": _excel_scaled_numeric_value(
            _mapping_value(statement, "amount_ex_tax_minor"),
            currency_digits,
            "不含税合计",
        ),
        "amount_incl_tax": _excel_scaled_numeric_value(
            _mapping_value(statement, "amount_incl_tax_minor"),
            currency_digits,
            "含税合计",
        ),
        "tax_amount": _excel_scaled_numeric_value(
            _mapping_value(statement, "tax_amount_minor"),
            currency_digits,
            "税额合计",
        ),
    }
    # Reproduce ordinary IEEE-754 accumulation conservatively.  A workbook is
    # rejected if its SUM formulas could lose a payable currency unit even
    # though every individual cell passes the precision check.
    formula_results = {"amount_ex_tax": 0.0, "amount_incl_tax": 0.0}
    for numeric_row in numeric_rows:
        formula_results["amount_ex_tax"] += numeric_row["amount_ex_tax"]
        formula_results["amount_incl_tax"] += numeric_row["amount_incl_tax"]
    formula_results["tax_amount"] = (
        formula_results["amount_incl_tax"] - formula_results["amount_ex_tax"]
    )
    for field, formula_result in formula_results.items():
        if (
            _scaled_integer_from_excel_result(formula_result, currency_digits)
            != _scaled_integer_from_excel_result(
                statement_totals[field],
                currency_digits,
            )
        ):
            raise ValueError(f"{EXCEL_PRECISION_ERROR}（汇总公式）")
    return numeric_rows


def _wrapped_line_count(value, column_width: float) -> int:
    text = "" if value is None else str(value)
    capacity = max(1, int(column_width))
    line_count = 0
    for logical_line in text.split("\n"):
        display_units = sum(
            2 if east_asian_width(character) in {"W", "F", "A"} else 1
            for character in logical_line
        )
        line_count += max(1, (display_units + capacity - 1) // capacity)
    return max(1, line_count)


def _validate_workbook_snapshot_totals(statement: Mapping, items: Sequence[Mapping]) -> None:
    total_fields = (
        "amount_ex_tax_minor",
        "amount_incl_tax_minor",
        "tax_amount_minor",
    )
    for field in total_fields:
        item_total = sum(int(_mapping_value(item, field, 0)) for item in items)
        if item_total != int(_mapping_value(statement, field, 0)):
            raise ValueError("对账单汇总金额与明细不一致")


def build_reconciliation_workbook(
    statement: Mapping,
    items: Sequence[Mapping],
    signature_path: Path | None = None,
) -> Workbook:
    items = list(items)
    if not items:
        raise ValueError("对账单没有明细")
    currency = normalize_currency(_mapping_value(statement, "currency", "CNY"))
    if any(normalize_currency(_mapping_value(item, "currency", currency)) != currency for item in items):
        raise ValueError("对账单明细币种不一致")
    _validate_workbook_snapshot_totals(statement, items)
    numeric_rows = _prepare_excel_numeric_rows(statement, items, currency)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "对账单"
    sheet.sheet_view.showGridLines = False

    black = "000000"
    dark_border = Side(style="thin", color="555555")
    border = Border(
        left=dark_border,
        right=dark_border,
        top=dark_border,
        bottom=dark_border,
    )
    body_font = Font(name="宋体", size=10, color=black)
    header_font = Font(name="宋体", size=10, bold=True, color=black)
    title_font = Font(name="宋体", size=16, bold=True, color=black)
    section_fill = PatternFill("solid", fgColor="D9E2F3")
    alternate_fill = PatternFill("solid", fgColor="F5F7FA")

    widths = (5, 11, 13, 14, 14, 19, 14, 7, 10, 13, 12, 13, 13, 18)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    sheet.merge_cells("A1:N1")
    set_excel_literal_value(sheet["A1"], _statement_title(statement))
    sheet["A1"].font = title_font
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells("A2:G2")
    sheet.merge_cells("H2:N2")
    set_excel_literal_value(
        sheet["A2"],
        f"对账单号：{_mapping_value(statement, 'statement_no')}",
    )
    set_excel_literal_value(
        sheet["H2"],
        f"账期：{_mapping_value(statement, 'period_start')} 至 {_mapping_value(statement, 'period_end')}    币种：{'RMB' if currency == 'CNY' else currency}",
    )

    customer_lines = (
        f"客户全称：{_mapping_value(statement, 'customer_name')}",
        f"电话：{_mapping_value(statement, 'customer_phone')}    邮箱：{_mapping_value(statement, 'customer_email')}",
        f"采购联系人：{_mapping_value(statement, 'customer_purchase_contact')}",
        f"对账联系人：{_mapping_value(statement, 'customer_reconciliation_contact')}\n地址：{_mapping_value(statement, 'customer_address')}",
    )
    supplier_lines = (
        f"供应商全称：{_mapping_value(statement, 'supplier_company_name')}",
        f"电话：{_mapping_value(statement, 'supplier_phone')}",
        f"联系人：{_mapping_value(statement, 'supplier_contact')}",
        f"邮箱：{_mapping_value(statement, 'supplier_email')}\n地址：{_mapping_value(statement, 'supplier_address')}",
    )
    for row, (customer_line, supplier_line) in enumerate(
        zip(customer_lines, supplier_lines), start=3
    ):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        sheet.merge_cells(start_row=row, start_column=8, end_row=row, end_column=14)
        set_excel_literal_value(sheet.cell(row, 1), customer_line)
        set_excel_literal_value(sheet.cell(row, 8), supplier_line)
        sheet.row_dimensions[row].height = 27 if row == 6 else 21

    for column, header in enumerate(RECONCILIATION_HEADERS, start=1):
        cell = sheet.cell(8, column, header)
        cell.font = header_font
        cell.fill = section_fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    sheet.row_dimensions[8].height = 30

    first_item_row = 9
    last_item_row = first_item_row + len(items) - 1
    money_format = _money_number_format(currency)
    for offset, item in enumerate(items):
        row = first_item_row + offset
        numeric_values = numeric_rows[offset]
        literal_values = {
            1: numeric_values["sort_order"],
            2: _mapping_value(item, "shipped_at"),
            3: _mapping_value(item, "delivery_no"),
            4: _mapping_value(item, "order_no"),
            5: _mapping_value(item, "sku") or _mapping_value(item, "drawing_no"),
            6: _mapping_value(item, "product_name"),
            7: _mapping_value(item, "model"),
            8: _mapping_value(item, "unit"),
        }
        for column, value in literal_values.items():
            if column == 1:
                sheet.cell(row, column).value = value
            else:
                set_excel_literal_value(sheet.cell(row, column), value)

        sheet.cell(row, 9).value = numeric_values["quantity"]
        sheet.cell(row, 10).value = numeric_values["unit_price_ex_tax"]
        sheet.cell(row, 11).value = numeric_values["unit_price_incl_tax"]
        sheet.cell(row, 12).value = numeric_values["amount_ex_tax"]
        sheet.cell(row, 13).value = numeric_values["amount_incl_tax"]
        assembly_drawing_no = str(_mapping_value(item, "assembly_drawing_no"))
        remark = str(_mapping_value(item, "remark"))
        if assembly_drawing_no:
            remark = f"组装件：{assembly_drawing_no}" + (f"；{remark}" if remark else "")
        set_excel_literal_value(sheet.cell(row, 14), remark)

        for column in range(1, 15):
            cell = sheet.cell(row, column)
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal="right" if column in range(9, 14) else "left",
                vertical="center",
                wrap_text=True,
            )
            if offset % 2:
                cell.fill = alternate_fill
        for column in (1, 2, 8):
            sheet.cell(row, column).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        sheet.cell(row, 9).number_format = "#,##0"
        sheet.cell(row, 10).number_format = "#,##0.000000"
        for column in (11, 12, 13):
            sheet.cell(row, column).number_format = money_format
        wrapped_line_count = max(
            _wrapped_line_count(literal_values[column], widths[column - 1])
            for column in range(2, 9)
        )
        wrapped_line_count = max(
            wrapped_line_count,
            _wrapped_line_count(remark, widths[13]),
        )
        sheet.row_dimensions[row].height = min(
            409.5,
            max(25, wrapped_line_count * 15),
        )

    summary_header_row = last_item_row + 1
    totals_row = summary_header_row + 1
    placeholders_row = totals_row + 1
    notes_row = placeholders_row + 1
    preparer_row = notes_row + 1
    signature_header_row = preparer_row + 1
    signature_start_row = signature_header_row + 1
    signature_end_row = signature_start_row + 3

    sheet.merge_cells(
        start_row=summary_header_row, start_column=1, end_row=summary_header_row, end_column=14
    )
    set_excel_literal_value(
        sheet.cell(summary_header_row, 1),
        f"本期销售合计（{'RMB' if currency == 'CNY' else currency} / 税率 {int(_mapping_value(statement, 'tax_rate_ppm', TAX_RATE_PPM)) / 10000:.2f}%）",
    )
    sheet.cell(summary_header_row, 1).fill = section_fill
    sheet.cell(summary_header_row, 1).font = header_font
    sheet.cell(summary_header_row, 1).alignment = Alignment(horizontal="center")

    for range_ref in (
        f"A{totals_row}:D{totals_row}",
        f"E{totals_row}:F{totals_row}",
        f"G{totals_row}:I{totals_row}",
        f"J{totals_row}:K{totals_row}",
        f"L{totals_row}:M{totals_row}",
        f"A{placeholders_row}:D{placeholders_row}",
        f"E{placeholders_row}:F{placeholders_row}",
        f"G{placeholders_row}:I{placeholders_row}",
        f"J{placeholders_row}:K{placeholders_row}",
        f"L{placeholders_row}:M{placeholders_row}",
    ):
        sheet.merge_cells(range_ref)

    sheet.cell(totals_row, 1).value = "不含税金额"
    sheet.cell(totals_row, 5).value = f"=SUM(L{first_item_row}:L{last_item_row})"
    sheet.cell(totals_row, 7).value = "含税金额"
    sheet.cell(totals_row, 10).value = f"=SUM(M{first_item_row}:M{last_item_row})"
    sheet.cell(totals_row, 12).value = "税额"
    sheet.cell(totals_row, 14).value = (
        f"=SUM(M{first_item_row}:M{last_item_row})-SUM(L{first_item_row}:L{last_item_row})"
    )
    for column in (5, 10, 14):
        sheet.cell(totals_row, column).number_format = money_format

    sheet.cell(placeholders_row, 1).value = "开票金额"
    sheet.cell(placeholders_row, 5).value = "未登记"
    sheet.cell(placeholders_row, 7).value = "结算金额"
    sheet.cell(placeholders_row, 10).value = "未登记"
    sheet.cell(placeholders_row, 12).value = "未结算金额"
    sheet.cell(placeholders_row, 14).value = "未登记"

    sheet.merge_cells(start_row=notes_row, start_column=1, end_row=notes_row, end_column=2)
    sheet.merge_cells(start_row=notes_row, start_column=3, end_row=notes_row, end_column=14)
    sheet.cell(notes_row, 1).value = "备注"
    set_excel_literal_value(sheet.cell(notes_row, 3), _mapping_value(statement, "remark"))
    sheet.row_dimensions[notes_row].height = 32

    sheet.merge_cells(start_row=preparer_row, start_column=1, end_row=preparer_row, end_column=4)
    sheet.merge_cells(start_row=preparer_row, start_column=5, end_row=preparer_row, end_column=8)
    sheet.merge_cells(start_row=preparer_row, start_column=9, end_row=preparer_row, end_column=14)
    set_excel_literal_value(
        sheet.cell(preparer_row, 1),
        f"制单人：{_mapping_value(statement, 'created_by')}",
    )
    sheet.cell(preparer_row, 5).value = "复核人："
    set_excel_literal_value(
        sheet.cell(preparer_row, 9),
        f"制单日期：{str(_mapping_value(statement, 'created_at'))[:10]}",
    )

    sheet.merge_cells(
        start_row=signature_header_row,
        start_column=1,
        end_row=signature_header_row,
        end_column=7,
    )
    sheet.merge_cells(
        start_row=signature_header_row,
        start_column=8,
        end_row=signature_header_row,
        end_column=14,
    )
    sheet.cell(signature_header_row, 1).value = "客户确认 / 财务签字"
    sheet.cell(signature_header_row, 8).value = "供应商确认 / 财务签字（盖章）"
    sheet.cell(signature_header_row, 1).fill = section_fill
    sheet.cell(signature_header_row, 8).fill = section_fill

    sheet.merge_cells(
        start_row=signature_start_row,
        start_column=1,
        end_row=signature_end_row - 1,
        end_column=7,
    )
    sheet.merge_cells(
        start_row=signature_start_row,
        start_column=8,
        end_row=signature_end_row - 1,
        end_column=14,
    )
    sheet.merge_cells(
        start_row=signature_end_row, start_column=1, end_row=signature_end_row, end_column=7
    )
    sheet.merge_cells(
        start_row=signature_end_row, start_column=8, end_row=signature_end_row, end_column=14
    )
    confirmed_name = str(_mapping_value(statement, "confirmed_name"))
    confirmed_at = str(_mapping_value(statement, "confirmed_at"))
    confirmed_text = "客户签字："
    if _mapping_value(statement, "status") == "confirmed":
        confirmed_text += f"{confirmed_name}    确认日期：{confirmed_at[:10]}"
    set_excel_literal_value(sheet.cell(signature_end_row, 1), confirmed_text)
    sheet.cell(signature_end_row, 8).value = "供应商签章：                    确认日期："
    for row in range(signature_start_row, signature_end_row + 1):
        sheet.row_dimensions[row].height = 22

    if (
        _mapping_value(statement, "status") == "confirmed"
        and signature_path is not None
        and Path(signature_path).is_file()
    ):
        signature = ExcelImage(str(signature_path))
        ratio = min(180 / signature.width, 65 / signature.height, 1)
        signature.width = max(1, round(signature.width * ratio))
        signature.height = max(1, round(signature.height * ratio))
        sheet.add_image(signature, f"B{signature_start_row}")

    for row in sheet.iter_rows(min_row=1, max_row=signature_end_row, min_col=1, max_col=14):
        for cell in row:
            font = copy(cell.font)
            font.name = "宋体"
            font.color = black
            if font.sz is None:
                font.sz = 10
            cell.font = font
            if cell.alignment == Alignment():
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in range(summary_header_row, signature_end_row + 1):
        for column in range(1, 15):
            sheet.cell(row, column).border = border
            if sheet.cell(row, column).alignment == Alignment():
                sheet.cell(row, column).alignment = Alignment(
                    vertical="center", wrap_text=True
                )

    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = True
    sheet.print_title_rows = "8:8"
    sheet.freeze_panes = "A9"
    sheet.print_area = f"A1:N{signature_end_row}"
    sheet.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.35,
        bottom=0.35,
        header=0.15,
        footer=0.15,
    )
    sheet.oddFooter.center.text = "第 &[Page] 页 / 共 &[Pages] 页"
    return workbook

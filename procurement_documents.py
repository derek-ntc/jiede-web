"""Snapshot-only purchase documents. No Flask or application imports are needed."""
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from functools import lru_cache
from io import BytesIO
from pathlib import Path
import unicodedata
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from procurement import PURCHASE_CATEGORY_LABELS, PURCHASE_STATUS_LABELS


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    width: float
    kind: str = "text"
    optional: bool = False


_MATERIAL = Column("material", "材质", 15)
_LENGTH = Column("length", "长 mm", 10, "number")
_WIDTH = Column("width", "宽 mm", 10, "number")
_QUANTITY = Column("ordered_quantity", "数量", 15, "number")
_UNIT = Column("unit", "单位", 7)
_TAIL = (_QUANTITY, _UNIT)
_PRICES = (Column("unit_price", "单价", 16, "money"), Column("line_total", "金额", 17, "money"))
_END = (Column("expected_at", "预计到货日期", 19, "date"), Column("remark", "备注", 29))
CATEGORY_DOCUMENT_COLUMNS = {
    "raw_material": (Column("item_name", "物品名称", 20), _MATERIAL, _LENGTH, _WIDTH, Column("thickness", "厚度 mm", 12, "number"), Column("surface", "表面", 14), *_TAIL, *_END),
    "carton": (Column("item_name", "物品名称", 20), _MATERIAL, _LENGTH, _WIDTH, Column("height", "高 mm", 10, "number"), Column("dimension_text", "尺寸说明/历史尺寸", 24), *_TAIL, *_PRICES, *_END),
    "outsourcing": (Column("identity", "物品／图号", 27), _MATERIAL, Column("details", "尺寸／厚度／表面", 30, optional=True), *_TAIL, *_PRICES, *_END),
    "other": (Column("identity", "物品／规格", 30), Column("details", "材质／尺寸／厚度／表面", 34, optional=True), *_TAIL, *_PRICES, *_END),
}


def sanitize_document_text(value):
    """Make saved text safe for XML and PDF, preserving readable separators."""
    text = "" if value is None else str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    ")
    return "".join(
        " " if (unicodedata.category(char) in {"Cc", "Cs"} and char != "\n")
        or ord(char) & 0xFFFF in {0xFFFE, 0xFFFF} else char
        for char in text)


def _text(value):
    return sanitize_document_text(value)


def _join(*values):
    return "／".join(_text(value) for value in values if value not in (None, ""))


def _money(value):
    return None if value is None else Decimal(value) / 100


def _excel_money(value):
    """Excel guarantees 15 significant digits; use exact RMB text beyond that."""
    if isinstance(value, Decimal) and Decimal(format(float(value), ".15g")) != value:
        return f"¥{value:,.2f}"
    return value


def purchase_document_view(order, items, *, include_prices):
    """Whitelist presentation fields; hidden-price models never contain amounts."""
    order = dict(order)
    category = order["category"]
    columns = [column for column in CATEGORY_DOCUMENT_COLUMNS[category]
               if include_prices or column.kind != "money"]
    rows = []
    for source in items:
        item = dict(source)
        dimensions = item.get("dimension_text") or "×".join(
            _text(item.get(key)) for key in ("length", "width", "height") if item.get(key) is not None)
        details = _join(item.get("material") if category == "other" else None, dimensions,
                        f"厚{item['thickness']}" if item.get("thickness") is not None else None, item.get("surface"))
        row = {}
        for column in columns:
            if column.key == "identity":
                value = _join(item.get("item_name"), item.get("drawing_no" if category == "outsourcing" else "spec"))
            elif column.key == "details":
                value = details
            elif column.kind == "money":
                value = _money(item.get(column.key + "_minor"))
            elif column.kind == "date":
                value = date.fromisoformat(item[column.key]) if item.get(column.key) else None
            else:
                value = item.get(column.key)
            row[column.key] = _text(value) if column.kind == "text" else value
        rows.append(row)
    columns = [column for column in columns if not column.optional or any(row[column.key] not in (None, "") for row in rows)]
    rows = [{column.key: row[column.key] for column in columns} for row in rows]
    profile = dict(order.get("company_profile") or {})
    supplier = [
        "供应商：" + _text(order.get("supplier_name")) + ("    代码：" + _text(order["supplier_code"]) if order.get("supplier_code") else ""),
        "    ".join(label + _text(order[key]) for key, label in (("supplier_contact", "联系人："), ("supplier_phone", "电话："), ("supplier_email", "邮箱：")) if order.get(key)),
        "供应商地址：" + _text(order.get("supplier_address")),
    ]
    delivery = ["收货地址：" + _text(order.get("delivery_address")),
                "收件人：" + _text(order.get("recipient")) + "    电话：" + _text(order.get("recipient_phone")),
                "订单备注：" + _text(order.get("remark"))]
    company = [_text(profile.get("company_name") or "宁波市杰德机械科技有限公司")]
    company.extend(label + _text(profile[key]) for key, label in (("address", "地址："), ("contact", "联系人："), ("phone", "电话："), ("email", "邮箱：")) if profile.get(key))
    result = dict(order_no=_text(order["order_no"]), category=PURCHASE_CATEGORY_LABELS[category],
                  status=PURCHASE_STATUS_LABELS[order["status"]], purchased_at=date.fromisoformat(order["purchased_at"]),
                  supplier=[line for line in supplier if line], delivery=delivery, company=company, columns=columns, rows=rows,
                  supplier_fields={key: _text(order.get("supplier_" + key)) for key in ("name", "code", "contact", "phone", "email", "address")},
                  delivery_fields={key: _text(order.get(key)) for key in ("delivery_address", "recipient", "recipient_phone", "remark")},
                  created_by=_text(order.get("created_by")))
    if any(column.kind == "money" for column in columns):
        totals = [row["line_total"] for row in rows]
        result["total"] = sum(totals, Decimal(0)) if all(value is not None for value in totals) else None
    return result


def _wrapped_lines(value, width):
    """Conservative CJK-aware line wrapping, also used to size printed Excel rows."""
    lines, line, used = [], "", 0
    for char in _text(value):
        size = 2 if unicodedata.east_asian_width(char) in "WF" else 1
        if char == "\n" or (line and used + size > width):
            lines.append(line)
            line, used = "", 0
        if char != "\n":
            line += char
            used += size
    lines.append(line)
    return lines


def build_purchase_order_workbook(order, items, *, include_prices: bool) -> BytesIO:
    model = purchase_document_view(order, items, include_prices=include_prices)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "采购订单"
    sheet.sheet_view.showGridLines = False
    columns = model["columns"]
    count = len(columns)
    # Keep the physical page width stable across category/permission projections.
    widths = [column.width * 145 / sum(c.width for c in columns) for column in columns]
    # A4 fitting may scale a wide order down slightly; never replace a valid amount with ###.
    for index, column in enumerate(columns):
        if column.kind == "money":
            labels = [f"¥{row[column.key]:,.2f}" for row in model["rows"] if row[column.key] is not None]
            widths[index] = max([widths[index]] + [len(label) + 3 for label in labels])
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    font = Font(name="Arial Unicode MS", size=10, color="000000")
    rule = Side(style="thin", color="000000")

    def write(row, col, value, *, kind="text", bold=False, center=False):
        if kind == "money":
            value = _excel_money(value)
        cell = sheet.cell(row, col, value)
        if isinstance(value, str):
            cell.data_type = "s"  # Never interpret supplier/item text as Excel formulas.
        cell.font = Font(name=font.name, size=font.sz, color="000000", bold=bold)
        cell.alignment = Alignment(horizontal="center" if center else "right" if kind in ("number", "money") else "left", vertical="center", wrap_text=True)
        cell.number_format = {"money": '"¥"#,##0.00', "date": "yyyy-mm-dd", "number": "#,##0.########"}.get(kind, "General")
        return cell

    def full_line(value, *, bold=False, center=False, size=10, bordered=False):
        for start in range(0, len(lines := _wrapped_lines(value, 138)), 12):
            row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=count)
            cell = write(row, 1, "\n".join(lines[start:start + 12]), bold=bold, center=center)
            cell.font = Font(name=font.name, size=size, color="000000", bold=bold)
            sheet.row_dimensions[row].height = max(size + 10, len(lines[start:start + 12]) * 16 + 7)
            if bordered:
                for col in range(1, count + 1):
                    sheet.cell(row, col).border = Border(top=rule, bottom=rule, left=rule, right=rule)

    def paired_line(left_label, left_value, right_label, right_value, *, left_kind="text"):
        # The middle table column may be a narrow unit column. Pick a nearby
        # label column that can hold Chinese header labels without extra lines.
        middle = min((index for index in range(3, count) if widths[index - 1] >= 10),
                     key=lambda index: abs(index - (count // 2 + 1)), default=count // 2 + 1)
        sections = ((1, left_label, "text"), (2, left_value, left_kind),
                    (middle, right_label, "text"), (middle + 1, right_value, "text"))
        spans = {2: middle - 1, middle + 1: count}
        text_lines = {col: _wrapped_lines(value, max(4, sum(widths[col - 1:spans.get(col, col)]) - 2))
                      for col, value, kind in sections if kind == "text"}
        for chunk in range(max([1] + [(len(lines) + 11) // 12 for lines in text_lines.values()])):
            row = sheet.max_row + 1
            height = 1
            for col, value, kind in sections:
                if col in spans:
                    sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=spans[col])
                if kind == "text":
                    lines = text_lines[col][chunk * 12:(chunk + 1) * 12]
                    value = "\n".join(lines)
                    height = max(height, len(lines))
                elif chunk:
                    value = None
                write(row, col, value, kind=kind)
            for col in range(1, count + 1):
                sheet.cell(row, col).border = Border(top=rule, bottom=rule, left=rule, right=rule)
            sheet.row_dimensions[row].height = max(22, height * 14 + 7)

    full_line(model["company"][0], bold=True, center=True, size=16)
    full_line("采购订单", bold=True, center=True, size=20)
    supplier = model["supplier_fields"]
    paired_line("供应商：", supplier["name"], "代码：", supplier["code"])
    paired_line("联系人：", supplier["contact"], "电话：", supplier["phone"])
    paired_line("供应商地址：", supplier["address"], "订单号：", model["order_no"])
    full_line("邮箱：" + supplier["email"], bordered=True)
    full_line(f"{model['category']}明细    状态：{model['status']}" + ("    币种：RMB／人民币" if "total" in model else ""), center=True)
    header_row = sheet.max_row + 1
    for index, column in enumerate(columns, 1):
        cell = write(header_row, index, column.label, bold=True, center=True)
        cell.fill = PatternFill("solid", fgColor="F0F0F0")
        cell.border = Border(top=rule, bottom=rule, left=rule, right=rule)
    sheet.row_dimensions[header_row].height = 34
    for source in model["rows"]:
        # Split exceptionally long text into continuation rows, never clip at Excel's height limit.
        texts = {c.key: _wrapped_lines(source[c.key], max(4, widths[i] - 2)) for i, c in enumerate(columns) if c.kind == "text"}
        chunks = max([1] + [(len(lines) + 11) // 12 for lines in texts.values()])
        for chunk in range(chunks):
            row = sheet.max_row + 1
            height = 1
            for index, column in enumerate(columns, 1):
                value = source[column.key]
                if column.kind == "text":
                    lines = texts[column.key][chunk * 12:(chunk + 1) * 12]
                    value = "\n".join(lines)
                    height = max(height, len(lines))
                elif chunk:
                    value = None
                elif column.kind == "money" and value is None:
                    value = "未录价"
                cell = write(row, index, value, kind=column.kind)
                if column.kind == "money" and isinstance(cell.value, str):
                    height = max(height, len(_wrapped_lines(cell.value, max(4, widths[index - 1] - 2))))
                if column.kind == "date":
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = Border(top=rule, bottom=rule, left=rule, right=rule)
            sheet.row_dimensions[row].height = max(26, height * 14 + 8)
    if "total" in model:
        row = sheet.max_row + 1
        total_column = count // 2 + 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_column - 1)
        sheet.merge_cells(start_row=row, start_column=total_column, end_row=row, end_column=count)
        write(row, 1, "合计（RMB／人民币）", bold=True)
        total = write(row, total_column, model["total"] if model["total"] is not None else "未完整录价", kind="money", bold=True)
        total_width = sum(widths[total_column - 1:])
        sheet.row_dimensions[row].height = max(29, len(_wrapped_lines(total.value, max(4, total_width - 2))) * 16 + 8)
    full_line("\n".join(model["delivery"][:2]), bordered=True)
    for line in model["delivery"][2:]:
        full_line(line, bordered=True)
    paired_line("下单日期：", model["purchased_at"], "经办人：", model["created_by"], left_kind="date")
    paired_line("确认回传", "", "签字：", "")
    if model["company"][1:]:
        full_line("    ".join(model["company"][1:]))
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.print_area = f"A1:{get_column_letter(count)}{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=.3, right=.3, top=.35, bottom=.4, header=.15, footer=.18)
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "&P / &N"
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


@lru_cache(maxsize=1)
def _pdf_font():
    # Match the app's OS font discovery without importing app (which initializes Flask).
    candidates = (
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
    )
    for index, candidate in enumerate(candidates):
        if Path(candidate).exists():
            name = f"PurchaseCJK{index}"
            try:
                pdfmetrics.registerFont(TTFont(name, candidate))
                return name
            except Exception:
                continue
    # Never silently fall back to Helvetica, which cannot display Chinese.
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def _pdf_cell_parts(paragraph, width, max_height=156):
    """Split a cell before Table sees it; Table itself only breaks between rows."""
    parts = []
    while paragraph.wrap(width, max_height)[1] > max_height:
        split = paragraph.split(width, max_height)
        if len(split) != 2:
            raise ValueError("采购订单文本无法安全分页")
        parts.append(split[0])
        paragraph = split[1]
    parts.append(paragraph)
    return parts


def build_purchase_order_pdf(order, items, *, include_prices: bool) -> BytesIO:
    model = purchase_document_view(order, items, include_prices=include_prices)
    stream = BytesIO()
    font = _pdf_font()
    document = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                                 topMargin=25, bottomMargin=34, title="采购订单", author="")
    styles = {name: ParagraphStyle(name, fontName=font, fontSize=size, leading=size + 4,
                                   alignment=align, wordWrap="CJK", textColor=colors.HexColor("#202833"))
              for name, size, align in (("body", 10, TA_LEFT), ("title", 20, TA_CENTER), ("center", 10, TA_CENTER), ("number", 9, TA_RIGHT), ("cell", 9, TA_LEFT))}

    def paragraph(value, style="body"):
        return Paragraph(escape(_text(value)).replace("\n", "<br/>"), styles[style])

    story = [paragraph("采购订单", "title"), Spacer(1, 5), paragraph(model["order_no"], "center"), Spacer(1, 9),
             paragraph(f"{model['category']}    采购日期：{model['purchased_at']}    状态：{model['status']}")]
    story.extend(paragraph(line) for line in model["supplier"])
    if "total" in model:
        story.append(paragraph("币种：RMB／人民币"))
    story.append(Spacer(1, 8))
    columns = model["columns"]
    column_widths = [document.width * c.width / sum(c.width for c in columns) for c in columns]
    table_rows = [[paragraph(column.label, "center") for column in columns]]
    for row in model["rows"]:
        values = []
        for column, width in zip(columns, column_widths):
            value = row[column.key]
            if column.kind == "money":
                value = "未录价" if value is None else f"¥{value:,.2f}"
            elif column.kind == "number" and value is not None:
                value = str(value)
            cell = paragraph(value, "number" if column.kind in ("number", "money") else "cell")
            values.append(_pdf_cell_parts(cell, width - 10))
        for continuation in range(max(len(parts) for parts in values)):
            table_rows.append([parts[continuation] if continuation < len(parts) else "" for parts in values])
    table = Table(table_rows, colWidths=column_widths,
                  repeatRows=1, hAlign="LEFT", splitByRow=1, splitInRow=0)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF2")),
        ("LINEABOVE", (0, 0), (-1, 0), .5, colors.HexColor("#A9B6C3")),
        ("LINEBELOW", (0, 0), (-1, -1), .3, colors.HexColor("#BEC7CF")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.append(table)
    if "total" in model:
        total = "未完整录价" if model["total"] is None else f"¥{model['total']:,.2f}"
        story.extend((Spacer(1, 7), paragraph("合计（RMB／人民币）：" + total)))
    footer = [Spacer(1, 9)]
    footer.extend(paragraph(line) for line in model["delivery"])
    footer.append(Spacer(1, 7))
    # Company footer flows with the document, so long contact data cannot cover details.
    footer.extend(paragraph(line) for line in model["company"])
    story.append(KeepTogether(footer))

    def page_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#59636F"))
        canvas.drawCentredString(landscape(A4)[0] / 2, 17, f"第 {doc.page} 页")
        canvas.restoreState()

    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    stream.seek(0)
    return stream

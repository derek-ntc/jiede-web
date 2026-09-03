import os
import re
import time
import base64
import binascii
import math
import secrets
import smtplib
import shutil
import sqlite3
import struct
import sys
import uuid
import warnings
import zlib
import json
import fcntl
import threading
from contextlib import contextmanager
from email.message import EmailMessage
from io import BytesIO
from textwrap import wrap
from datetime import datetime
from datetime import timedelta
from functools import wraps
from itertools import zip_longest
from mimetypes import guess_type
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from xml.sax.saxutils import escape as xml_escape

# The repository contains a legacy PIL compatibility stub.  Shipment image
# validation must use the Pillow distribution declared in requirements.txt,
# and ReportLab should see that same implementation when it imports PIL below.
_original_import_path = list(sys.path)
_application_root = Path(__file__).resolve().parent
try:
    sys.path = [
        entry
        for entry in sys.path
        if Path(entry or os.getcwd()).resolve() != _application_root
    ]
    import PIL as PillowPackage
    from PIL import Image as PillowImage
    from PIL import ImageFile as PillowImageFile
    from PIL import ImageSequence as PillowImageSequence
    from PIL import UnidentifiedImageError as PillowUnidentifiedImageError
finally:
    sys.path = _original_import_path

if not getattr(PillowPackage, "__version__", ""):
    raise ImportError("发货图片验证需要 requirements.txt 中声明的 Pillow")
PillowImageFile.LOAD_TRUNCATED_IMAGES = False

from dotenv import load_dotenv
from flask import (
    Flask,
    abort,
    flash,
    g,
    jsonify,
    Response,
    redirect,
    render_template,
    request,
    send_file,
    send_from_directory,
    session,
    url_for,
)
from flask.wrappers import Request as FlaskRequest
from itsdangerous import BadSignature
from itsdangerous import URLSafeSerializer
from werkzeug.utils import secure_filename
from werkzeug.exceptions import RequestEntityTooLarge
from werkzeug.security import check_password_hash, generate_password_hash
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Flowable, Image
from reportlab.lib.utils import ImageReader

from assembly_shipping import (
    allocate_quantity,
    expand_components,
    inventory_result,
    normalize_component_rows,
    parse_positive_int,
    preview_token,
)
from pricing import (
    SUPPORTED_CURRENCIES,
    format_money_minor,
    normalize_currency,
    parse_money_minor,
)


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
MANUALS_DIR = BASE_DIR / "manuals"
UPLOADS_DIR = BASE_DIR / "uploads"
SIGNATURES_DIR = UPLOADS_DIR / "signatures"
SHIPMENT_IMAGES_DIR = UPLOADS_DIR / "shipment-images"
INSPECTION_REPORTS_DIR = UPLOADS_DIR / "inspection-reports"
PRODUCTION_DRAWINGS_DIR = UPLOADS_DIR / "production-drawings"
DB_PATH = DATA_DIR / "manuals.db"
DATABASE_READY = False
MAX_ORDER_QUANTITY = 2_147_483_647
SIGNATURE_LINK_TTL = timedelta(days=7)
PRODUCTION_STAGE_LABELS = {
    "laser": "激光",
    "bending": "折弯",
    "welding": "焊接",
}
CUSTOMER_INVOICE_FIELDS = (
    "invoice_title",
    "tax_id",
    "registered_address",
    "registered_phone",
    "bank_name",
    "bank_account",
    "invoice_email",
)
DISABLED_ENDPOINT_PREFIXES = (
    "admin_powder_coating",
    "export_powder_coating",
    "selected_powder_coating",
    "create_powder_coating",
    "edit_powder_coating",
    "delete_powder_coating",
    "copy_powder_coating",
    "toggle_powder_coating",
)


load_dotenv(BASE_DIR / ".env", override=True)


SHIPMENT_IMAGE_UPLOAD_ENDPOINTS = {
    "create_shipment_from_shipped_page",
    "create_assembly_shipment_from_shipped_page",
    "edit_assembly_shipment",
    "approve_shipment_plan",
    "upload_shipment_images_admin",
    "upload_shipment_photos",
}


class FactoryRequest(FlaskRequest):
    @property
    def max_content_length(self):
        global_limit = super().max_content_length
        if getattr(self, "endpoint", None) not in SHIPMENT_IMAGE_UPLOAD_ENDPOINTS:
            return global_limit
        shipment_limit = app.config.get("SHIPMENT_IMAGE_MAX_REQUEST_BYTES")
        if shipment_limit is None:
            return global_limit
        shipment_limit = int(shipment_limit)
        if global_limit is None:
            return shipment_limit
        return min(int(global_limit), shipment_limit)


app = Flask(__name__)
app.request_class = FactoryRequest
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "change-me-in-env")
DEFAULT_WRITE_LOCK_PATH = "/tmp/jiede-web-write.lock"
app.config["WRITE_LOCK_PATH"] = os.getenv(
    "JIEDE_WRITE_LOCK_PATH", DEFAULT_WRITE_LOCK_PATH
)
_write_lock_state = threading.local()
_HTTP_SAFE_METHODS = {"GET", "HEAD", "OPTIONS", "TRACE"}


@contextmanager
def application_write_lock():
    depth = getattr(_write_lock_state, "depth", 0)
    if depth == 0:
        lock_path = Path(app.config["WRITE_LOCK_PATH"])
        lock_path.parent.mkdir(parents=True, exist_ok=True)
        lock_file = lock_path.open("a+b")
        try:
            fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX)
        except Exception:
            lock_file.close()
            raise
        _write_lock_state.lock_file = lock_file
    _write_lock_state.depth = depth + 1
    try:
        yield
    finally:
        _write_lock_state.depth -= 1
        if _write_lock_state.depth == 0:
            lock_file = _write_lock_state.lock_file
            try:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)
            finally:
                lock_file.close()
                del _write_lock_state.lock_file
                del _write_lock_state.depth


@app.before_request
def serialize_http_write_requests():
    if request.method in _HTTP_SAFE_METHODS:
        return
    request_limit = request.max_content_length
    if (
        request.endpoint in SHIPMENT_IMAGE_UPLOAD_ENDPOINTS
        and request_limit is not None
        and request.content_length is not None
        and request.content_length > request_limit
    ):
        raise RequestEntityTooLarge()
    lock_context = application_write_lock()
    lock_context.__enter__()
    g._application_write_lock_context = lock_context


@app.teardown_request
def release_http_write_lock(error):
    lock_context = getattr(g, "_application_write_lock_context", None)
    if lock_context is None:
        return
    del g._application_write_lock_context
    lock_context.__exit__(
        type(error) if error is not None else None,
        error,
        error.__traceback__ if error is not None else None,
    )
app.config["MAX_CONTENT_LENGTH"] = int(os.getenv("MAX_UPLOAD_MB", "1024")) * 1024 * 1024
app.config["MAX_FORM_MEMORY_SIZE"] = int(os.getenv("MAX_FORM_MB", "100")) * 1024 * 1024
app.config["SHIPMENT_IMAGE_MAX_FILES"] = int(
    os.getenv("SHIPMENT_IMAGE_MAX_FILES", "20")
)
app.config["SHIPMENT_IMAGE_MAX_FILE_BYTES"] = int(
    os.getenv("SHIPMENT_IMAGE_MAX_FILE_MB", "20")
) * 1024 * 1024
app.config["SHIPMENT_IMAGE_MAX_TOTAL_BYTES"] = int(
    os.getenv("SHIPMENT_IMAGE_MAX_TOTAL_MB", "100")
) * 1024 * 1024
app.config["SHIPMENT_IMAGE_MAX_REQUEST_BYTES"] = int(
    os.getenv("SHIPMENT_IMAGE_MAX_REQUEST_MB", "110")
) * 1024 * 1024
app.config["SHIPMENT_IMAGE_MAX_PIXELS"] = int(
    os.getenv("SHIPMENT_IMAGE_MAX_PIXELS", "40000000")
)
app.config["SHIPMENT_IMAGE_MAX_FRAMES"] = int(
    os.getenv("SHIPMENT_IMAGE_MAX_FRAMES", "100")
)
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=90)
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 3600

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".dwg",
    ".dxf",
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
    ".svg",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".csv",
    ".ppt",
    ".pptx",
    ".txt",
    ".zip",
    ".rar",
    ".7z",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".svg", ".heic", ".heif"}
SHIPMENT_RASTER_SUFFIX_FORMATS = {
    ".png": "png",
    ".jpg": "jpeg",
    ".jpeg": "jpeg",
    ".gif": "gif",
}
SHIPMENT_RASTER_MIME_SUFFIXES = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/gif": ".gif",
}
ORDER_IMPORT_EXTENSIONS = {".xlsx", ".xlsm"}
ORDER_IMPORT_DRAWING_HEADERS = {"图号", "产品图号", "drawing_no", "drawingno", "drawing", "part_no", "partno"}
ORDER_IMPORT_QUANTITY_HEADERS = {"数量", "订单数量", "quantity", "qty", "order_quantity", "orderquantity"}
ORDER_IMPORT_PLANNED_SHIP_HEADERS = {"计划发货时间", "计划交货时间", "交货时间", "planned_ship_at", "plannedshipat"}
CARTON_BOARD_TYPE_OPTIONS = ["高", "中", "低"]
CARTON_BOARD_PRICE_COLUMNS = {
    "高": "board_price_high",
    "中": "board_price_middle",
    "低": "board_price_low",
}
SORT_COLUMNS = {
    "drawing_no": "manuals.drawing_no COLLATE NOCASE",
    "product_name": "manuals.product_name COLLATE NOCASE",
    "supplier": "manuals.supplier COLLATE NOCASE",
    "customer": "manuals.customer COLLATE NOCASE",
    "updated_at": "manuals.updated_at",
    "remark": "manuals.remark COLLATE NOCASE",
    "created_by": "manuals.created_by COLLATE NOCASE",
}

PDF_FONT_CANDIDATES = [
    Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    Path("/System/Library/Fonts/STHeiti Medium.ttc"),
    Path("/System/Library/Fonts/Hiragino Sans GB.ttc"),
    Path("/System/Library/Fonts/Supplemental/Songti.ttc"),
]
PDF_FONT_NAME = None


def format_datetime(value):
    if not value:
        return "-"
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError:
        return value
    return parsed.strftime("%Y-%m-%d")


app.jinja_env.filters["datetime"] = format_datetime
app.jinja_env.filters["money_minor"] = format_money_minor


def format_bytes(size):
    if size is None:
        return "未知大小"
    if size >= 1024 * 1024:
        return f"{size / 1024 / 1024:.1f}MB"
    if size >= 1024:
        return f"{size / 1024:.1f}KB"
    return f"{size}B"


@app.context_processor
def upload_limits():
    max_upload_bytes = app.config["MAX_CONTENT_LENGTH"]
    max_form_bytes = app.config["MAX_FORM_MEMORY_SIZE"]
    return {
        "max_upload_mb": max_upload_bytes // 1024 // 1024,
        "max_form_mb": max_form_bytes // 1024 // 1024,
        "max_upload_bytes": max_upload_bytes,
        "max_form_bytes": max_form_bytes,
        "static_version": int(max(
            (BASE_DIR / "static" / "style.css").stat().st_mtime,
            (BASE_DIR / "static" / "editor.js").stat().st_mtime,
            (BASE_DIR / "static" / "inventory.js").stat().st_mtime if (BASE_DIR / "static" / "inventory.js").exists() else 0,
            (BASE_DIR / "static" / "order_entry.js").stat().st_mtime if (BASE_DIR / "static" / "order_entry.js").exists() else 0,
            (BASE_DIR / "static" / "product_list.js").stat().st_mtime if (BASE_DIR / "static" / "product_list.js").exists() else 0,
            (BASE_DIR / "static" / "finance.js").stat().st_mtime if (BASE_DIR / "static" / "finance.js").exists() else 0,
        )),
        "current_user_role": current_user_role(),
        "current_admin_username": current_admin_username(),
        "can_manage_products": user_has_permission("products"),
        "can_manage_customers": user_has_permission("customers"),
        "can_manage_common_info": user_has_permission("common_info"),
        "can_manage_purchase_followups": user_has_permission("purchase_followups"),
        "can_manage_powder_coating": user_has_permission("powder_coating"),
        "can_manage_carton_purchases": user_has_permission("carton_purchases"),
        "can_manage_warehouse_inventory": user_has_permission("warehouse_inventory"),
        "can_manage_production_followups": user_has_permission("production_followups_manage"),
        "can_create_products": user_has_permission("product_create"),
        "can_edit_products": user_has_permission("product_edit"),
        "can_view_orders": user_has_permission("orders_view"),
        "can_manage_orders": user_has_permission("orders_manage"),
        "can_view_shipped": user_has_permission("shipped_view"),
        "can_manage_shipped": user_has_permission("shipped_manage"),
        "can_view_prices": user_can_view_prices(),
        "can_manage_finance": user_has_permission("finance_manage"),
        "can_access_admin_modules": user_can_access_admin_modules(),
        "shipment_signature_url": shipment_signature_url,
        "shipment_photo_upload_url": shipment_photo_upload_url,
        "signature_link_expired": signature_link_expired,
        "purchase_followup_pending_purchase_at_sse_url": purchase_followup_pending_purchase_at_sse_url,
        "purchase_followup_pending_purchase_at_plain_sse_url": purchase_followup_pending_purchase_at_plain_sse_url,
    }


@app.errorhandler(RequestEntityTooLarge)
def handle_request_entity_too_large(error):
    max_upload_bytes = request.max_content_length or app.config["MAX_CONTENT_LENGTH"]
    max_form_bytes = app.config["MAX_FORM_MEMORY_SIZE"]
    actual_bytes = request.content_length
    custom_description = str(getattr(error, "description", "") or "").strip()
    if custom_description and custom_description != RequestEntityTooLarge.description:
        reason = custom_description
        limit_text = ""
    elif actual_bytes and actual_bytes > max_upload_bytes:
        reason = "触发了单次请求体上传限制"
        limit_text = format_bytes(max_upload_bytes)
    else:
        reason = "触发了编辑内容表单字段限制"
        limit_text = format_bytes(max_form_bytes)
    if limit_text:
        message = (
            f"上传内容过大：{reason}。"
            f"本次请求大小 {format_bytes(actual_bytes)}，"
            f"对应限制 {limit_text}。"
        )
    else:
        message = f"上传内容过大：{reason}。"
    if request.endpoint in {
        "create_assembly_shipment_from_shipped_page",
        "edit_assembly_shipment",
    }:
        return jsonify(
            {
                "error": message,
                "reason": reason,
                "actual_bytes": actual_bytes,
                "max_upload_bytes": max_upload_bytes,
            }
        ), 413
    if request.path.startswith("/admin/editor-images"):
        return jsonify(
            {
                "error": message,
                "reason": reason,
                "actual_bytes": actual_bytes,
                "max_upload_bytes": max_upload_bytes,
                "max_form_bytes": max_form_bytes,
            }
        ), 413
    if request.endpoint in SHIPMENT_IMAGE_UPLOAD_ENDPOINTS:
        return Response(message, status=413, mimetype="text/plain")
    flash(message, "error")
    return redirect(request.referrer or url_for("admin_index"))


def sort_state(default_sort="updated_at"):
    sort = request.args.get("sort", default_sort).strip()
    direction = request.args.get("direction", "desc").strip().lower()
    if sort not in SORT_COLUMNS:
        sort = default_sort
    if direction not in {"asc", "desc"}:
        direction = "desc"
    return sort, direction


def order_clause(sort, direction):
    fallback_direction = "DESC" if sort == "updated_at" else "ASC"
    sql_direction = "ASC" if direction == "asc" else "DESC"
    return (
        f" ORDER BY {SORT_COLUMNS[sort]} {sql_direction}, "
        f"manuals.updated_at {fallback_direction}, manuals.id DESC"
    )


@contextmanager
def get_db():
    # The project often lives on a shared LAN volume. SQLite file locking can be
    # rejected by SMB mounts, so this single-container app uses nolock mode. Every
    # production connection is therefore serialized by the container-local lock
    # for its complete transaction, including commit/rollback and close.
    with application_write_lock():
        conn = sqlite3.connect(f"file:{DB_PATH}?mode=rwc&nolock=1", uri=True)
        conn.row_factory = sqlite3.Row
        try:
            with conn:
                yield conn
        finally:
            conn.close()


def init_db():
    DATA_DIR.mkdir(exist_ok=True)
    MANUALS_DIR.mkdir(exist_ok=True)
    SIGNATURES_DIR.mkdir(parents=True, exist_ok=True)
    SHIPMENT_IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    INSPECTION_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    PRODUCTION_DRAWINGS_DIR.mkdir(parents=True, exist_ok=True)
    with get_db() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS manuals (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                product_name TEXT NOT NULL,
                customer TEXT NOT NULL DEFAULT '',
                model TEXT NOT NULL,
                category TEXT NOT NULL,
                version TEXT NOT NULL,
                remark TEXT,
                filename TEXT NOT NULL,
                original_filename TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        ensure_columns(conn)
        ensure_file_table(conn)
        ensure_product_material_table(conn)
        ensure_order_table(conn)
        ensure_inspection_tables(conn)
        ensure_user_table(conn)
        ensure_customer_table(conn)
        ensure_common_info_table(conn)
        ensure_purchase_followup_table(conn)
        ensure_powder_coating_tables(conn)
        ensure_carton_purchase_table(conn)
        ensure_arrival_record_tables(conn)
        ensure_inventory_tables(conn)
        ensure_assembly_shipping_tables(conn)
        ensure_finance_tables(conn)
        ensure_feedback_table(conn)
        ensure_production_followup_tables(conn)
        ensure_indexes(conn)


def ensure_columns(conn):
    existing = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(manuals)").fetchall()
    }
    migrations = {
        "drawing_no": "ALTER TABLE manuals ADD COLUMN drawing_no TEXT NOT NULL DEFAULT ''",
        "supplier": "ALTER TABLE manuals ADD COLUMN supplier TEXT NOT NULL DEFAULT ''",
        "customer": "ALTER TABLE manuals ADD COLUMN customer TEXT NOT NULL DEFAULT ''",
        "pack_quantity": "ALTER TABLE manuals ADD COLUMN pack_quantity TEXT NOT NULL DEFAULT ''",
        "pack_carton_size": "ALTER TABLE manuals ADD COLUMN pack_carton_size TEXT NOT NULL DEFAULT ''",
        "pack_weight": "ALTER TABLE manuals ADD COLUMN pack_weight TEXT NOT NULL DEFAULT ''",
        "description_html": "ALTER TABLE manuals ADD COLUMN description_html TEXT NOT NULL DEFAULT ''",
        "file_type": "ALTER TABLE manuals ADD COLUMN file_type TEXT NOT NULL DEFAULT 'file'",
        "created_by": "ALTER TABLE manuals ADD COLUMN created_by TEXT NOT NULL DEFAULT ''",
        "updated_by": "ALTER TABLE manuals ADD COLUMN updated_by TEXT NOT NULL DEFAULT ''",
        "sku": "ALTER TABLE manuals ADD COLUMN sku TEXT NOT NULL DEFAULT ''",
        "barcode": "ALTER TABLE manuals ADD COLUMN barcode TEXT NOT NULL DEFAULT ''",
        "qr_code": "ALTER TABLE manuals ADD COLUMN qr_code TEXT NOT NULL DEFAULT ''",
        "default_location_id": "ALTER TABLE manuals ADD COLUMN default_location_id INTEGER",
        "min_stock": "ALTER TABLE manuals ADD COLUMN min_stock INTEGER NOT NULL DEFAULT 0",
        "unit_price_minor": "ALTER TABLE manuals ADD COLUMN unit_price_minor INTEGER",
        "currency": "ALTER TABLE manuals ADD COLUMN currency TEXT NOT NULL DEFAULT 'CNY'",
    }
    for column, statement in migrations.items():
        if column not in existing:
            conn.execute(statement)
    ensure_file_table(conn)


def ensure_file_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS manual_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            manual_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            original_filename TEXT NOT NULL,
            file_type TEXT NOT NULL DEFAULT 'file',
            created_at TEXT NOT NULL,
            FOREIGN KEY (manual_id) REFERENCES manuals(id) ON DELETE CASCADE
        )
        """
    )
    rows = conn.execute(
        """
        SELECT id, filename, original_filename, file_type, created_at
        FROM manuals
        WHERE filename IS NOT NULL AND filename != ''
        """
    ).fetchall()
    for row in rows:
        exists = conn.execute(
            """
            SELECT 1
            FROM manual_files
            WHERE manual_id = ? AND filename = ?
            LIMIT 1
            """,
            (row["id"], row["filename"]),
        ).fetchone()
        if not exists:
            conn.execute(
                """
                INSERT INTO manual_files (
                    manual_id, filename, original_filename, file_type, created_at
                )
                VALUES (?, ?, ?, ?, ?)
                """,
                (
                    row["id"],
                    row["filename"],
                    row["original_filename"],
                    row["file_type"],
                    row["created_at"],
                ),
            )


def ensure_product_material_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS product_materials (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            manual_id INTEGER NOT NULL,
            material TEXT NOT NULL DEFAULT '',
            thickness TEXT NOT NULL DEFAULT '',
            surface_type TEXT NOT NULL DEFAULT '',
            supplier TEXT NOT NULL DEFAULT '',
            sort_order INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (manual_id) REFERENCES manuals(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_product_materials_manual_id
        ON product_materials (manual_id)
        """
    )


def ensure_assembly_shipping_tables(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS product_assembly_components (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            manual_id INTEGER NOT NULL,
            assembly_drawing_no TEXT NOT NULL,
            quantity_per_set INTEGER NOT NULL,
            sort_order INTEGER NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            UNIQUE(manual_id, assembly_drawing_no),
            FOREIGN KEY (manual_id) REFERENCES manuals(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS assembly_shipment_batches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer TEXT NOT NULL,
            assembly_drawing_no TEXT NOT NULL,
            set_quantity INTEGER NOT NULL,
            shipped_at TEXT NOT NULL,
            logistics_no TEXT NOT NULL,
            created_by TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS assembly_shipment_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_id INTEGER NOT NULL,
            manual_id INTEGER NOT NULL,
            drawing_no TEXT NOT NULL,
            product_name TEXT NOT NULL,
            quantity_per_set INTEGER NOT NULL,
            calculated_quantity INTEGER NOT NULL,
            shipped_quantity INTEGER NOT NULL,
            inventory_deducted_quantity INTEGER NOT NULL,
            inventory_shortage_quantity INTEGER NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (batch_id) REFERENCES assembly_shipment_batches(id) ON DELETE CASCADE,
            FOREIGN KEY (manual_id) REFERENCES manuals(id) ON DELETE CASCADE
        )
        """
    )
    existing_item_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(assembly_shipment_items)").fetchall()
    }
    item_migrations = {
        "unit_price_minor": "ALTER TABLE assembly_shipment_items ADD COLUMN unit_price_minor INTEGER",
        "currency": "ALTER TABLE assembly_shipment_items ADD COLUMN currency TEXT NOT NULL DEFAULT 'CNY'",
        "price_recorded_by": "ALTER TABLE assembly_shipment_items ADD COLUMN price_recorded_by TEXT NOT NULL DEFAULT ''",
        "price_recorded_at": "ALTER TABLE assembly_shipment_items ADD COLUMN price_recorded_at TEXT NOT NULL DEFAULT ''",
    }
    for column, statement in item_migrations.items():
        if column not in existing_item_columns:
            conn.execute(statement)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS assembly_shipment_allocations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            item_id INTEGER NOT NULL,
            order_id INTEGER,
            quantity INTEGER NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (item_id) REFERENCES assembly_shipment_items(id) ON DELETE CASCADE,
            FOREIGN KEY (order_id) REFERENCES product_orders(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS assembly_shipment_images (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            original_filename TEXT NOT NULL,
            content_type TEXT NOT NULL DEFAULT '',
            file_size INTEGER NOT NULL DEFAULT 0,
            uploaded_at TEXT NOT NULL,
            uploaded_ip TEXT NOT NULL DEFAULT '',
            uploaded_user_agent TEXT NOT NULL DEFAULT '',
            FOREIGN KEY (batch_id) REFERENCES assembly_shipment_batches(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_product_assembly_components_manual_id
        ON product_assembly_components (manual_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_product_assembly_components_drawing_manual
        ON product_assembly_components (assembly_drawing_no, manual_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_assembly_shipment_batches_shipped_customer
        ON assembly_shipment_batches (shipped_at, customer)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_assembly_shipment_items_batch_id
        ON assembly_shipment_items (batch_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_assembly_shipment_allocations_item_id
        ON assembly_shipment_allocations (item_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_assembly_shipment_allocations_order_id
        ON assembly_shipment_allocations (order_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_assembly_shipment_images_batch_id
        ON assembly_shipment_images (batch_id)
        """
    )


def ensure_order_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS product_orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            manual_id INTEGER NOT NULL,
            order_no TEXT NOT NULL,
            ordered_at TEXT NOT NULL,
            quantity INTEGER NOT NULL,
            customer TEXT NOT NULL DEFAULT '',
            assembly_drawing_no TEXT NOT NULL DEFAULT '',
            assembly_set_quantity INTEGER NOT NULL DEFAULT 0,
            planned_ship_at TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (manual_id) REFERENCES manuals(id) ON DELETE CASCADE
        )
        """
    )
    existing = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(product_orders)").fetchall()
    }
    migrations = {
        "shipped_quantity": "ALTER TABLE product_orders ADD COLUMN shipped_quantity INTEGER NOT NULL DEFAULT 0",
        "shipped_at": "ALTER TABLE product_orders ADD COLUMN shipped_at TEXT NOT NULL DEFAULT ''",
        "customer_email": "ALTER TABLE product_orders ADD COLUMN customer_email TEXT NOT NULL DEFAULT ''",
        "material_stock_status": "ALTER TABLE product_orders ADD COLUMN material_stock_status TEXT NOT NULL DEFAULT ''",
        "material_stock_assignee": "ALTER TABLE product_orders ADD COLUMN material_stock_assignee TEXT NOT NULL DEFAULT ''",
        "remark": "ALTER TABLE product_orders ADD COLUMN remark TEXT NOT NULL DEFAULT ''",
        "carton_status": "ALTER TABLE product_orders ADD COLUMN carton_status TEXT NOT NULL DEFAULT ''",
        "carton_assignee": "ALTER TABLE product_orders ADD COLUMN carton_assignee TEXT NOT NULL DEFAULT ''",
        "recent_ship_status": "ALTER TABLE product_orders ADD COLUMN recent_ship_status TEXT NOT NULL DEFAULT ''",
        "inventory_received_quantity": "ALTER TABLE product_orders ADD COLUMN inventory_received_quantity INTEGER NOT NULL DEFAULT 0",
        "inventory_status": "ALTER TABLE product_orders ADD COLUMN inventory_status TEXT NOT NULL DEFAULT '未入库'",
        "assembly_drawing_no": "ALTER TABLE product_orders ADD COLUMN assembly_drawing_no TEXT NOT NULL DEFAULT ''",
        "assembly_set_quantity": "ALTER TABLE product_orders ADD COLUMN assembly_set_quantity INTEGER NOT NULL DEFAULT 0",
    }
    for column, statement in migrations.items():
        if column not in existing:
            conn.execute(statement)

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS product_order_shipments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            order_id INTEGER NOT NULL,
            shipped_quantity INTEGER NOT NULL,
            shipped_at TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (order_id) REFERENCES product_orders(id) ON DELETE CASCADE
        )
        """
    )
    existing_shipment_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(product_order_shipments)").fetchall()
    }
    shipment_migrations = {
        "email_sent_at": "ALTER TABLE product_order_shipments ADD COLUMN email_sent_at TEXT NOT NULL DEFAULT ''",
        "logistics_no": "ALTER TABLE product_order_shipments ADD COLUMN logistics_no TEXT NOT NULL DEFAULT ''",
        "signature_token": "ALTER TABLE product_order_shipments ADD COLUMN signature_token TEXT NOT NULL DEFAULT ''",
        "signature_status": "ALTER TABLE product_order_shipments ADD COLUMN signature_status TEXT NOT NULL DEFAULT '未签收'",
        "signature_image": "ALTER TABLE product_order_shipments ADD COLUMN signature_image TEXT NOT NULL DEFAULT ''",
        "signed_at": "ALTER TABLE product_order_shipments ADD COLUMN signed_at TEXT NOT NULL DEFAULT ''",
        "signed_ip": "ALTER TABLE product_order_shipments ADD COLUMN signed_ip TEXT NOT NULL DEFAULT ''",
        "signed_user_agent": "ALTER TABLE product_order_shipments ADD COLUMN signed_user_agent TEXT NOT NULL DEFAULT ''",
        "signature_expires_at": "ALTER TABLE product_order_shipments ADD COLUMN signature_expires_at TEXT NOT NULL DEFAULT ''",
        "photo_upload_token": "ALTER TABLE product_order_shipments ADD COLUMN photo_upload_token TEXT NOT NULL DEFAULT ''",
        "remark": "ALTER TABLE product_order_shipments ADD COLUMN remark TEXT NOT NULL DEFAULT ''",
        "unit_price_minor": "ALTER TABLE product_order_shipments ADD COLUMN unit_price_minor INTEGER",
        "currency": "ALTER TABLE product_order_shipments ADD COLUMN currency TEXT NOT NULL DEFAULT 'CNY'",
        "price_recorded_by": "ALTER TABLE product_order_shipments ADD COLUMN price_recorded_by TEXT NOT NULL DEFAULT ''",
        "price_recorded_at": "ALTER TABLE product_order_shipments ADD COLUMN price_recorded_at TEXT NOT NULL DEFAULT ''",
    }
    for column, statement in shipment_migrations.items():
        if column not in existing_shipment_columns:
            conn.execute(statement)

    conn.execute(
        """
        UPDATE product_order_shipments
        SET signature_expires_at = datetime(COALESCE(NULLIF(created_at, ''), 'now'), '+7 days')
        WHERE signature_token != ''
          AND signature_expires_at = ''
        """
    )

    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_product_order_shipments_signature_token
        ON product_order_shipments (signature_token)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_product_order_shipments_photo_upload_token
        ON product_order_shipments (photo_upload_token)
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS product_order_shipment_images (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            original_filename TEXT NOT NULL,
            content_type TEXT NOT NULL DEFAULT '',
            file_size INTEGER NOT NULL DEFAULT 0,
            uploaded_at TEXT NOT NULL,
            uploaded_ip TEXT NOT NULL DEFAULT '',
            uploaded_user_agent TEXT NOT NULL DEFAULT '',
            FOREIGN KEY (shipment_id) REFERENCES product_order_shipments(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_product_order_shipment_images_shipment_id
        ON product_order_shipment_images (shipment_id)
        """
    )

    shipment_rows = conn.execute(
        """
        SELECT id, shipped_quantity, shipped_at
        FROM product_orders
        WHERE shipped_quantity > 0 OR shipped_at != ''
        """
    ).fetchall()
    for row in shipment_rows:
        exists = conn.execute(
            """
            SELECT 1
            FROM product_order_shipments
            WHERE order_id = ?
            LIMIT 1
            """,
            (row["id"],),
        ).fetchone()
        if not exists and row["shipped_quantity"] > 0:
            created_at = datetime.utcnow().isoformat(timespec="seconds")
            conn.execute(
                """
                INSERT INTO product_order_shipments (
                    order_id, shipped_quantity, shipped_at, created_at
                )
                VALUES (?, ?, ?, ?)
                """,
                (
                    row["id"],
                    row["shipped_quantity"],
                    row["shipped_at"] or created_at,
                    created_at,
                ),
            )

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS shipment_plans (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            plan_no TEXT NOT NULL UNIQUE,
            planned_ship_at TEXT NOT NULL DEFAULT '',
            customer TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL DEFAULT '待发货',
            remark TEXT NOT NULL DEFAULT '',
            created_by TEXT NOT NULL DEFAULT '',
            reviewed_by TEXT NOT NULL DEFAULT '',
            reviewed_at TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS shipment_plan_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            plan_id INTEGER NOT NULL,
            order_id INTEGER NOT NULL,
            planned_quantity INTEGER NOT NULL DEFAULT 0,
            shipped_quantity INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (plan_id) REFERENCES shipment_plans(id) ON DELETE CASCADE,
            FOREIGN KEY (order_id) REFERENCES product_orders(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_shipment_plan_items_plan_id
        ON shipment_plan_items (plan_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_shipment_plan_items_order_id
        ON shipment_plan_items (order_id)
        """
    )


def ensure_production_followup_tables(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS production_followups (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_no TEXT NOT NULL DEFAULT '',
            ordered_at TEXT NOT NULL,
            drawing_no TEXT NOT NULL,
            product_name TEXT NOT NULL,
            laser_completed_at TEXT NOT NULL DEFAULT '',
            bending_completed_at TEXT NOT NULL DEFAULT '',
            welding_completed_at TEXT NOT NULL DEFAULT '',
            created_by TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    existing = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(production_followups)").fetchall()
    }
    if "batch_no" not in existing:
        conn.execute("ALTER TABLE production_followups ADD COLUMN batch_no TEXT NOT NULL DEFAULT ''")
    conn.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_production_followups_batch_no
        ON production_followups (batch_no)
        WHERE batch_no != ''
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS production_followup_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            followup_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            original_filename TEXT NOT NULL,
            file_type TEXT NOT NULL DEFAULT 'file',
            content_type TEXT NOT NULL DEFAULT '',
            file_size INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            FOREIGN KEY (followup_id) REFERENCES production_followups(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_production_followup_files_followup_id
        ON production_followup_files (followup_id)
        """
    )
    rows = conn.execute(
        """
        SELECT id, ordered_at
        FROM production_followups
        WHERE batch_no = ''
        ORDER BY ordered_at ASC, id ASC
        """
    ).fetchall()
    for row in rows:
        conn.execute(
            """
            UPDATE production_followups
            SET batch_no = ?
            WHERE id = ?
            """,
            (next_production_batch_no(conn, row["ordered_at"]), row["id"]),
        )


def ensure_inspection_tables(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS manual_inspection_requirements (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            manual_id INTEGER NOT NULL,
            item_name TEXT NOT NULL DEFAULT '',
            standard TEXT NOT NULL DEFAULT '',
            method TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            sort_order INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (manual_id) REFERENCES manuals(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_manual_inspection_requirements_manual_id
        ON manual_inspection_requirements (manual_id, sort_order, id)
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS shipment_plan_inspection_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            plan_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            original_filename TEXT NOT NULL,
            content_type TEXT NOT NULL DEFAULT '',
            file_size INTEGER NOT NULL DEFAULT 0,
            uploaded_at TEXT NOT NULL,
            uploaded_by TEXT NOT NULL DEFAULT '',
            FOREIGN KEY (plan_id) REFERENCES shipment_plans(id) ON DELETE CASCADE
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_shipment_plan_inspection_reports_plan_id
        ON shipment_plan_inspection_reports (plan_id, uploaded_at DESC, id DESC)
        """
    )


def ensure_user_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'operator',
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    existing = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(users)").fetchall()
    }
    migrations = {
        "role": "ALTER TABLE users ADD COLUMN role TEXT NOT NULL DEFAULT 'operator'",
        "active": "ALTER TABLE users ADD COLUMN active INTEGER NOT NULL DEFAULT 1",
        "can_manage_products": "ALTER TABLE users ADD COLUMN can_manage_products INTEGER NOT NULL DEFAULT 1",
        "can_manage_orders": "ALTER TABLE users ADD COLUMN can_manage_orders INTEGER NOT NULL DEFAULT 1",
        "can_view_orders": "ALTER TABLE users ADD COLUMN can_view_orders INTEGER NOT NULL DEFAULT 1",
        "can_manage_shipped": "ALTER TABLE users ADD COLUMN can_manage_shipped INTEGER NOT NULL DEFAULT 1",
        "can_view_shipped": "ALTER TABLE users ADD COLUMN can_view_shipped INTEGER NOT NULL DEFAULT 1",
        "can_manage_customers": "ALTER TABLE users ADD COLUMN can_manage_customers INTEGER NOT NULL DEFAULT 1",
        "can_manage_common_info": "ALTER TABLE users ADD COLUMN can_manage_common_info INTEGER NOT NULL DEFAULT 1",
        "can_manage_purchase_followups": "ALTER TABLE users ADD COLUMN can_manage_purchase_followups INTEGER NOT NULL DEFAULT 1",
        "can_manage_powder_coating": "ALTER TABLE users ADD COLUMN can_manage_powder_coating INTEGER NOT NULL DEFAULT 1",
        "can_manage_carton_purchases": "ALTER TABLE users ADD COLUMN can_manage_carton_purchases INTEGER NOT NULL DEFAULT 1",
        "can_manage_warehouse_inventory": "ALTER TABLE users ADD COLUMN can_manage_warehouse_inventory INTEGER NOT NULL DEFAULT 1",
        "can_manage_production_followups": "ALTER TABLE users ADD COLUMN can_manage_production_followups INTEGER NOT NULL DEFAULT 0",
        "can_create_products": "ALTER TABLE users ADD COLUMN can_create_products INTEGER NOT NULL DEFAULT 1",
        "can_edit_products": "ALTER TABLE users ADD COLUMN can_edit_products INTEGER NOT NULL DEFAULT 1",
        "can_view_prices": "ALTER TABLE users ADD COLUMN can_view_prices INTEGER NOT NULL DEFAULT 0",
        "can_manage_finance": "ALTER TABLE users ADD COLUMN can_manage_finance INTEGER NOT NULL DEFAULT 0",
    }
    for column, statement in migrations.items():
        if column not in existing:
            conn.execute(statement)

    now = datetime.utcnow().isoformat(timespec="seconds")
    legacy_username, legacy_password = admin_credentials()
    if legacy_username and legacy_password:
        seed_user(conn, legacy_username, legacy_password, "admin", now)

    configured = os.getenv("ADMIN_ACCOUNTS", "").strip()
    if configured:
        for item in configured.replace(";", ",").split(","):
            if ":" not in item:
                continue
            username, password = item.split(":", 1)
            seed_user(conn, username.strip(), password.strip(), "admin", now)


def ensure_customer_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS customers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            contact TEXT NOT NULL DEFAULT '',
            address TEXT NOT NULL DEFAULT '',
            email TEXT NOT NULL DEFAULT '',
            phone TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    existing = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(customers)").fetchall()
    }
    migrations = {
        "invoice_title": "ALTER TABLE customers ADD COLUMN invoice_title TEXT NOT NULL DEFAULT ''",
        "tax_id": "ALTER TABLE customers ADD COLUMN tax_id TEXT NOT NULL DEFAULT ''",
        "registered_address": "ALTER TABLE customers ADD COLUMN registered_address TEXT NOT NULL DEFAULT ''",
        "registered_phone": "ALTER TABLE customers ADD COLUMN registered_phone TEXT NOT NULL DEFAULT ''",
        "bank_name": "ALTER TABLE customers ADD COLUMN bank_name TEXT NOT NULL DEFAULT ''",
        "bank_account": "ALTER TABLE customers ADD COLUMN bank_account TEXT NOT NULL DEFAULT ''",
        "invoice_email": "ALTER TABLE customers ADD COLUMN invoice_email TEXT NOT NULL DEFAULT ''",
    }
    for column, statement in migrations.items():
        if column not in existing:
            conn.execute(statement)
    sync_product_customers(conn)


def ensure_finance_tables(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS finance_invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer_id INTEGER NOT NULL,
            customer_name TEXT NOT NULL,
            invoice_title TEXT NOT NULL DEFAULT '',
            tax_id TEXT NOT NULL DEFAULT '',
            registered_address TEXT NOT NULL DEFAULT '',
            registered_phone TEXT NOT NULL DEFAULT '',
            bank_name TEXT NOT NULL DEFAULT '',
            bank_account TEXT NOT NULL DEFAULT '',
            invoice_email TEXT NOT NULL DEFAULT '',
            currency TEXT NOT NULL,
            total_minor INTEGER NOT NULL,
            status TEXT NOT NULL DEFAULT 'pending',
            invoice_no TEXT NOT NULL DEFAULT '',
            invoice_date TEXT NOT NULL DEFAULT '',
            payment_date TEXT NOT NULL DEFAULT '',
            finance_remark TEXT NOT NULL DEFAULT '',
            voided_by TEXT NOT NULL DEFAULT '',
            voided_at TEXT NOT NULL DEFAULT '',
            created_by TEXT NOT NULL,
            updated_by TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (customer_id) REFERENCES customers(id)
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS finance_invoice_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            invoice_id INTEGER NOT NULL,
            source_type TEXT NOT NULL,
            source_id INTEGER NOT NULL,
            active_claim_key TEXT UNIQUE,
            order_no TEXT NOT NULL DEFAULT '',
            assembly_batch_id INTEGER,
            assembly_drawing_no TEXT NOT NULL DEFAULT '',
            shipped_at TEXT NOT NULL,
            drawing_no TEXT NOT NULL DEFAULT '',
            product_name TEXT NOT NULL DEFAULT '',
            quantity INTEGER NOT NULL,
            unit_price_minor INTEGER NOT NULL,
            currency TEXT NOT NULL,
            line_total_minor INTEGER NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (invoice_id) REFERENCES finance_invoices(id) ON DELETE CASCADE,
            CHECK (source_type IN ('ordinary', 'assembly_item'))
        )
        """
    )


def ensure_customer_exists(conn, name, now=None):
    name = (name or "").strip()
    if not name:
        return
    timestamp = now or datetime.utcnow().isoformat(timespec="seconds")
    conn.execute(
        """
        INSERT OR IGNORE INTO customers (
            name, contact, address, email, phone, remark, created_at, updated_at
        )
        VALUES (?, '', '', '', '', '', ?, ?)
        """,
        (name, timestamp, timestamp),
    )


def sync_product_customers(conn):
    now = datetime.utcnow().isoformat(timespec="seconds")
    rows = conn.execute(
        """
        SELECT DISTINCT customer
        FROM manuals
        WHERE customer IS NOT NULL AND TRIM(customer) != ''
        """
    ).fetchall()
    for row in rows:
        ensure_customer_exists(conn, row["customer"], now)


def ensure_common_info_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS common_infos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            category TEXT NOT NULL DEFAULT '',
            title TEXT NOT NULL,
            content TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )


def ensure_purchase_followup_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS purchase_followups (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            recorded_at TEXT NOT NULL,
            item_name TEXT NOT NULL,
            remark TEXT NOT NULL DEFAULT '',
            recorded_by TEXT NOT NULL DEFAULT '',
            quantity INTEGER NOT NULL DEFAULT 0,
            purchased INTEGER NOT NULL DEFAULT 0,
            purchased_at TEXT NOT NULL DEFAULT '',
            image_filename TEXT NOT NULL DEFAULT '',
            image_original_filename TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    existing = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(purchase_followups)").fetchall()
    }
    if "purchased_at" not in existing:
        conn.execute("ALTER TABLE purchase_followups ADD COLUMN purchased_at TEXT NOT NULL DEFAULT ''")
    if "image_filename" not in existing:
        conn.execute("ALTER TABLE purchase_followups ADD COLUMN image_filename TEXT NOT NULL DEFAULT ''")
    if "image_original_filename" not in existing:
        conn.execute("ALTER TABLE purchase_followups ADD COLUMN image_original_filename TEXT NOT NULL DEFAULT ''")


def ensure_powder_coating_tables(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS powder_coating_suppliers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            contact TEXT NOT NULL DEFAULT '',
            phone TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS powder_coating_products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_name TEXT NOT NULL,
            drawing_no TEXT NOT NULL DEFAULT '',
            supplier_name TEXT NOT NULL DEFAULT '',
            unit_price REAL NOT NULL DEFAULT 0,
            image_filename TEXT NOT NULL DEFAULT '',
            image_original_filename TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS powder_coating_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_id INTEGER,
            product_name TEXT NOT NULL,
            drawing_no TEXT NOT NULL DEFAULT '',
            supplier_name TEXT NOT NULL DEFAULT '',
            unit_price REAL NOT NULL DEFAULT 0,
            quantity INTEGER NOT NULL DEFAULT 0,
            delivered_at TEXT NOT NULL,
            received_at TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            recorded_by TEXT NOT NULL DEFAULT '',
            image_filename TEXT NOT NULL DEFAULT '',
            image_original_filename TEXT NOT NULL DEFAULT '',
            completed INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (product_id) REFERENCES powder_coating_products(id) ON DELETE SET NULL
        )
        """
    )
    record_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(powder_coating_records)").fetchall()
    }
    if "supplier_name" not in record_columns:
        conn.execute("ALTER TABLE powder_coating_records ADD COLUMN supplier_name TEXT NOT NULL DEFAULT ''")
    if "completed" not in record_columns:
        conn.execute("ALTER TABLE powder_coating_records ADD COLUMN completed INTEGER NOT NULL DEFAULT 0")
    product_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(powder_coating_products)").fetchall()
    }
    if "supplier_name" not in product_columns:
        conn.execute("ALTER TABLE powder_coating_products ADD COLUMN supplier_name TEXT NOT NULL DEFAULT ''")


def ensure_carton_purchase_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS carton_suppliers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            contact TEXT NOT NULL DEFAULT '',
            phone TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS carton_products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            print_mark TEXT NOT NULL,
            supplier_name TEXT NOT NULL DEFAULT '',
            carton_size TEXT NOT NULL DEFAULT '',
            carton_length REAL NOT NULL DEFAULT 0,
            carton_width REAL NOT NULL DEFAULT 0,
            carton_height REAL NOT NULL DEFAULT 0,
            board_type TEXT NOT NULL DEFAULT '',
            board_price_high REAL NOT NULL DEFAULT 0,
            board_price_middle REAL NOT NULL DEFAULT 0,
            board_price_low REAL NOT NULL DEFAULT 0,
            remark TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS carton_purchases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_id INTEGER,
            ordered_at TEXT NOT NULL,
            print_mark TEXT NOT NULL DEFAULT '',
            supplier_name TEXT NOT NULL DEFAULT '',
            board_type TEXT NOT NULL DEFAULT '',
            quantity INTEGER NOT NULL DEFAULT 0,
            unit_price REAL NOT NULL DEFAULT 0,
            carton_size TEXT NOT NULL DEFAULT '',
            carton_length REAL NOT NULL DEFAULT 0,
            carton_width REAL NOT NULL DEFAULT 0,
            carton_height REAL NOT NULL DEFAULT 0,
            received_at TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            recorded_by TEXT NOT NULL DEFAULT '',
            completed INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (product_id) REFERENCES carton_products(id) ON DELETE SET NULL
        )
        """
    )
    purchase_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(carton_purchases)").fetchall()
    }
    if "product_id" not in purchase_columns:
        conn.execute("ALTER TABLE carton_purchases ADD COLUMN product_id INTEGER")
    if "supplier_name" not in purchase_columns:
        conn.execute("ALTER TABLE carton_purchases ADD COLUMN supplier_name TEXT NOT NULL DEFAULT ''")
    if "unit_price" not in purchase_columns:
        conn.execute("ALTER TABLE carton_purchases ADD COLUMN unit_price REAL NOT NULL DEFAULT 0")
    if "board_type" not in purchase_columns:
        conn.execute("ALTER TABLE carton_purchases ADD COLUMN board_type TEXT NOT NULL DEFAULT ''")
    if "completed" not in purchase_columns:
        conn.execute("ALTER TABLE carton_purchases ADD COLUMN completed INTEGER NOT NULL DEFAULT 0")
    for column in ("carton_length", "carton_width", "carton_height"):
        if column not in purchase_columns:
            conn.execute(f"ALTER TABLE carton_purchases ADD COLUMN {column} REAL NOT NULL DEFAULT 0")
    product_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(carton_products)").fetchall()
    }
    if "board_type" not in product_columns:
        conn.execute("ALTER TABLE carton_products ADD COLUMN board_type TEXT NOT NULL DEFAULT ''")
    if "supplier_name" not in product_columns:
        conn.execute("ALTER TABLE carton_products ADD COLUMN supplier_name TEXT NOT NULL DEFAULT ''")
    for column in ("carton_length", "carton_width", "carton_height"):
        if column not in product_columns:
            conn.execute(f"ALTER TABLE carton_products ADD COLUMN {column} REAL NOT NULL DEFAULT 0")
    for column in ("board_price_high", "board_price_middle", "board_price_low"):
        if column not in product_columns:
            conn.execute(f"ALTER TABLE carton_products ADD COLUMN {column} REAL NOT NULL DEFAULT 0")


def ensure_arrival_record_tables(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS arrival_suppliers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            contact TEXT NOT NULL DEFAULT '',
            phone TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS arrival_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            arrived_at TEXT NOT NULL,
            item_name TEXT NOT NULL,
            spec TEXT NOT NULL DEFAULT '',
            quantity INTEGER NOT NULL DEFAULT 0,
            unit_price REAL NOT NULL DEFAULT 0,
            supplier_name TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            recorded_by TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )


def seed_user(conn, username, password, role, now):
    if not username or not password:
        return
    exists = conn.execute(
        "SELECT 1 FROM users WHERE username = ?",
        (username,),
    ).fetchone()
    if exists:
        return
    conn.execute(
        """
        INSERT INTO users (
            username, password_hash, role, active,
            can_manage_products, can_manage_orders,
            can_manage_customers, can_manage_common_info, can_manage_purchase_followups,
            can_create_products, can_edit_products,
            created_at, updated_at
        )
        VALUES (?, ?, ?, 1, 1, 1, 1, 1, 1, 1, 1, ?, ?)
        """,
        (username, generate_password_hash(password), role, now, now),
    )


def ensure_indexes(conn):
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_manuals_updated_id
        ON manuals (updated_at DESC, id DESC)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_manuals_supplier
        ON manuals (supplier)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_manuals_customer
        ON manuals (customer)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_manual_files_manual_id
        ON manual_files (manual_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_product_orders_manual_id
        ON product_orders (manual_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_product_orders_order_no
        ON product_orders (order_no)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_finance_invoices_customer_id
        ON finance_invoices (customer_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_finance_invoices_status
        ON finance_invoices (status)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_finance_invoices_invoice_date
        ON finance_invoices (invoice_date)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_finance_invoice_items_invoice_id
        ON finance_invoice_items (invoice_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_finance_invoice_items_source
        ON finance_invoice_items (source_type, source_id)
        """
    )
    conn.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_users_username
        ON users (username)
        """
    )
    conn.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_customers_name
        ON customers (name)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_common_infos_category
        ON common_infos (category)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_purchase_followups_purchased
        ON purchase_followups (purchased, recorded_at)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_powder_coating_products_name
        ON powder_coating_products (product_name, drawing_no)
        """
    )
    conn.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_powder_coating_suppliers_name
        ON powder_coating_suppliers (name)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_powder_coating_records_dates
        ON powder_coating_records (delivered_at, received_at)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_carton_purchases_dates
        ON carton_purchases (ordered_at, received_at)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_carton_products_mark_size
        ON carton_products (print_mark, carton_size)
        """
    )
    conn.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_carton_suppliers_name
        ON carton_suppliers (name)
        """
    )
    conn.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_arrival_suppliers_name
        ON arrival_suppliers (name)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_arrival_records_arrived_at
        ON arrival_records (arrived_at)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_manuals_sku
        ON manuals (sku)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_manuals_barcode
        ON manuals (barcode)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_manuals_qr_code
        ON manuals (qr_code)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_inventory_balances_manual
        ON inventory_balances (manual_id, location_id)
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_inventory_transactions_manual
        ON inventory_transactions (manual_id, created_at DESC)
        """
    )


def ensure_inventory_tables(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS warehouse_locations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            code TEXT NOT NULL UNIQUE,
            remark TEXT NOT NULL DEFAULT '',
            enabled INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS inventory_balances (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            manual_id INTEGER NOT NULL,
            location_id INTEGER NOT NULL,
            quantity INTEGER NOT NULL DEFAULT 0,
            updated_at TEXT NOT NULL,
            UNIQUE(manual_id, location_id),
            FOREIGN KEY (manual_id) REFERENCES manuals(id) ON DELETE CASCADE,
            FOREIGN KEY (location_id) REFERENCES warehouse_locations(id) ON DELETE RESTRICT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS inventory_transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            transaction_no TEXT NOT NULL UNIQUE,
            type TEXT NOT NULL,
            manual_id INTEGER NOT NULL,
            quantity INTEGER NOT NULL,
            from_location_id INTEGER,
            to_location_id INTEGER,
            related_order_type TEXT NOT NULL DEFAULT '',
            related_order_id TEXT NOT NULL DEFAULT '',
            related_order_no TEXT NOT NULL DEFAULT '',
            customer TEXT NOT NULL DEFAULT '',
            operator TEXT NOT NULL DEFAULT '',
            remark TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY (manual_id) REFERENCES manuals(id) ON DELETE CASCADE,
            FOREIGN KEY (from_location_id) REFERENCES warehouse_locations(id) ON DELETE SET NULL,
            FOREIGN KEY (to_location_id) REFERENCES warehouse_locations(id) ON DELETE SET NULL
        )
        """
    )


def ensure_feedback_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS feedbacks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            content TEXT NOT NULL,
            created_by TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            handled INTEGER NOT NULL DEFAULT 0,
            handled_by TEXT NOT NULL DEFAULT '',
            handled_at TEXT NOT NULL DEFAULT ''
        )
        """
    )
    existing = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(feedbacks)").fetchall()
    }
    feedback_migrations = {
        "handled": "ALTER TABLE feedbacks ADD COLUMN handled INTEGER NOT NULL DEFAULT 0",
        "handled_by": "ALTER TABLE feedbacks ADD COLUMN handled_by TEXT NOT NULL DEFAULT ''",
        "handled_at": "ALTER TABLE feedbacks ADD COLUMN handled_at TEXT NOT NULL DEFAULT ''",
    }
    for column, statement in feedback_migrations.items():
        if column not in existing:
            conn.execute(statement)
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_feedbacks_created_at
        ON feedbacks (created_at)
        """
    )


def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if not session.get("admin_logged_in") or not current_user():
            return redirect(url_for("admin_login", next=request.full_path))
        return view(*args, **kwargs)

    return wrapped_view


def permission_required(permission):
    def decorator(view):
        @wraps(view)
        def wrapped_view(*args, **kwargs):
            if not session.get("admin_logged_in") or not current_user():
                return redirect(url_for("admin_login", next=request.full_path))
            if not user_has_permission(permission):
                flash("当前账号没有权限访问该模块", "error")
                return redirect(url_for("admin_index"))
            return view(*args, **kwargs)

        return wrapped_view

    return decorator


def admin_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if not session.get("admin_logged_in") or not current_user():
            return redirect(url_for("admin_login", next=request.full_path))
        if current_user_role() != "admin":
            flash("当前账号没有权限管理用户", "error")
            return redirect(url_for("admin_index"))
        return view(*args, **kwargs)

    return wrapped_view


def user_can_access_admin_modules():
    return any([
        user_has_permission("customers"),
        user_has_permission("common_info"),
        user_has_permission("purchase_followups"),
        user_has_permission("carton_purchases"),
        user_has_permission("warehouse_inventory"),
        user_has_permission("product_create"),
        user_has_permission("product_edit"),
        user_has_permission("finance_manage"),
    ])


def admin_credentials():
    return (
        os.getenv("ADMIN_USERNAME", "admin"),
        os.getenv("ADMIN_PASSWORD", "admin123"),
    )


def admin_accounts():
    accounts = {}
    configured = os.getenv("ADMIN_ACCOUNTS", "").strip()
    if configured:
        for item in configured.replace(";", ",").split(","):
            if ":" not in item:
                continue
            username, password = item.split(":", 1)
            username = username.strip()
            password = password.strip()
            if username and password:
                accounts[username] = password

    legacy_username, legacy_password = admin_credentials()
    if legacy_username and legacy_password:
        accounts.setdefault(legacy_username, legacy_password)
    return accounts


def valid_admin_login(username, password):
    with get_db() as conn:
        user = conn.execute(
            """
            SELECT *
            FROM users
            WHERE username = ? AND active = 1
            """,
            (username,),
        ).fetchone()
    if not user:
        return None
    if not check_password_hash(user["password_hash"], password):
        return None
    return user


def current_admin_username():
    return session.get("admin_username") or session.get("admin_logged_in") and "admin" or ""


def current_user():
    username = session.get("admin_username")
    if not username:
        return None
    with get_db() as conn:
        return conn.execute(
            """
            SELECT *
            FROM users
            WHERE username = ? AND active = 1
            """,
            (username,),
        ).fetchone()


def current_user_role():
    user = current_user()
    return user["role"] if user else ""


def user_has_permission(permission):
    user = current_user()
    if not user:
        return False
    if user["role"] == "admin":
        return True
    if permission == "products":
        return bool(user["can_manage_products"])
    if permission == "customers":
        return bool(user["can_manage_customers"])
    if permission == "common_info":
        return bool(user["can_manage_common_info"])
    if permission == "purchase_followups":
        return bool(user["can_manage_purchase_followups"])
    if permission == "powder_coating":
        return bool(user["can_manage_powder_coating"])
    if permission == "carton_purchases":
        return bool(user["can_manage_carton_purchases"])
    if permission == "warehouse_inventory":
        return bool(user["can_manage_warehouse_inventory"])
    if permission == "production_followups_manage":
        return bool(user["can_manage_production_followups"])
    if permission == "product_create":
        return bool(user["can_create_products"])
    if permission == "product_edit":
        return bool(user["can_edit_products"])
    if permission == "orders":
        return bool(user["can_manage_orders"])
    if permission == "orders_view":
        return bool(user["can_view_orders"] or user["can_manage_orders"])
    if permission == "orders_manage":
        return bool(user["can_manage_orders"])
    if permission == "shipped_view":
        return bool(user["can_view_shipped"] or user["can_manage_shipped"])
    if permission == "shipped_manage":
        return bool(user["can_manage_shipped"])
    if permission == "price_view":
        return bool(user["can_view_prices"])
    if permission == "finance_manage":
        return bool(user["can_manage_finance"])
    return False


def user_can_view_prices():
    return user_has_permission("price_view") or user_has_permission("finance_manage")


def get_suppliers():
    with get_db() as conn:
        rows = get_suppliers_with_conn(conn)
    return [row["supplier"] for row in rows]


def get_customers():
    with get_db() as conn:
        rows = get_customers_with_conn(conn)
    return [row["customer"] for row in rows]


def get_customer_name_options(conn):
    rows = conn.execute(
        """
        SELECT name AS customer
        FROM customers
        WHERE name IS NOT NULL AND name != ''
        UNION
        SELECT customer
        FROM manuals
        WHERE customer IS NOT NULL AND customer != ''
        ORDER BY customer COLLATE NOCASE
        """
    ).fetchall()
    return [row["customer"] for row in rows]


def get_customers_with_conn(conn):
    return conn.execute(
        """
        SELECT DISTINCT customer
        FROM manuals
        WHERE customer IS NOT NULL AND customer != ''
        ORDER BY customer
        """
    ).fetchall()


def customer_invoice_snapshot(customer):
    snapshot = {"customer_name": str(customer["name"] or "")}
    snapshot.update(
        {
            field: str(customer[field] or "")
            for field in CUSTOMER_INVOICE_FIELDS
        }
    )
    return snapshot


def get_suppliers_with_conn(conn):
    return conn.execute(
        """
        SELECT DISTINCT supplier
        FROM manuals
        WHERE supplier IS NOT NULL AND supplier != ''
        ORDER BY supplier
        """
    ).fetchall()


def is_allowed_file(file_storage):
    filename = file_storage.filename or ""
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def image_suffix_for(file_storage):
    suffix = Path(file_storage.filename or "").suffix.lower()
    if suffix in IMAGE_EXTENSIONS:
        return suffix
    mimetype = (file_storage.mimetype or "").lower()
    fallback_suffixes = {
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/gif": ".gif",
        "image/webp": ".webp",
        "image/bmp": ".bmp",
        "image/tiff": ".tiff",
        "image/svg+xml": ".svg",
        "image/heic": ".heic",
        "image/heif": ".heif",
    }
    return fallback_suffixes.get(mimetype, "")


def shipment_image_suffix_for(file_storage):
    suffix = Path(file_storage.filename or "").suffix.lower()
    if suffix:
        return suffix if suffix in SHIPMENT_RASTER_SUFFIX_FORMATS else ""
    return SHIPMENT_RASTER_MIME_SUFFIXES.get(
        (file_storage.mimetype or "").lower(), ""
    )


def _pillow_image_dimensions(image):
    width, height = image.size
    width = int(width)
    height = int(height)
    return width, height, width * height


def _pillow_has_required_palette(image_format, image):
    if image_format == "PNG":
        return image.mode != "P" or image.palette is not None
    if image_format == "GIF":
        return not (
            image.mode in {"L", "P"}
            and image.palette is None
            and "background" not in image.info
        )
    return True


def detected_shipment_raster_format(path):
    path = Path(path)
    if path.stat().st_size > int(app.config["SHIPMENT_IMAGE_MAX_FILE_BYTES"]):
        return ""
    allowed_formats = {"PNG": "png", "JPEG": "jpeg", "GIF": "gif"}
    max_pixels = int(app.config["SHIPMENT_IMAGE_MAX_PIXELS"])
    max_frames = int(app.config["SHIPMENT_IMAGE_MAX_FRAMES"])
    if max_pixels <= 0 or max_frames <= 0:
        raise RuntimeError("发货图片像素和帧数限制必须大于 0")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", PillowImage.DecompressionBombWarning)
            with PillowImage.open(path) as image:
                image_format = image.format or ""
                detected_format = allowed_formats.get(image_format, "")
                if not detected_format:
                    return ""
                width, height, frame_pixels = _pillow_image_dimensions(image)
                frame_count = int(getattr(image, "n_frames", 1) or 0)
                if (
                    width <= 0
                    or height <= 0
                    or frame_pixels > max_pixels
                    or frame_count <= 0
                    or frame_count > max_frames
                    or (image_format != "GIF" and frame_count != 1)
                    or not _pillow_has_required_palette(image_format, image)
                ):
                    return ""
                image.verify()

            with PillowImage.open(path) as image:
                if allowed_formats.get(image.format or "", "") != detected_format:
                    return ""
                decoded_pixels = 0
                decoded_frames = 0
                frames = (
                    PillowImageSequence.Iterator(image)
                    if detected_format == "gif"
                    else (image,)
                )
                for frame in frames:
                    width, height, frame_pixels = _pillow_image_dimensions(frame)
                    decoded_frames += 1
                    decoded_pixels += frame_pixels
                    if (
                        width <= 0
                        or height <= 0
                        or decoded_frames > max_frames
                        or decoded_pixels > max_pixels
                        or not _pillow_has_required_palette(image.format or "", frame)
                    ):
                        return ""
                    frame.load()
                if decoded_frames != frame_count:
                    return ""
            return detected_format
    except (
        EOFError,
        MemoryError,
        OSError,
        OverflowError,
        PillowImage.DecompressionBombError,
        PillowImage.DecompressionBombWarning,
        PillowUnidentifiedImageError,
        SyntaxError,
        ValueError,
    ):
        return ""


def is_allowed_editor_image(file_storage):
    return bool(image_suffix_for(file_storage))


def file_type_for(filename):
    return "image" if Path(filename).suffix.lower() in IMAGE_EXTENSIONS else "file"


def stored_filename_for(original):
    safe_name = secure_filename(original) or "manual-file"
    suffix = Path(original).suffix.lower()
    stem = Path(safe_name).stem or "manual"
    return f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}-{stem}{suffix}"


def save_upload(file_storage):
    original = file_storage.filename or "manual-file"
    filename = stored_filename_for(original)
    file_storage.save(MANUALS_DIR / filename)
    return filename, original, file_type_for(filename)


def save_editor_image(file_storage):
    suffix = image_suffix_for(file_storage)
    if not suffix:
        raise ValueError(file_storage.filename or "")
    original = file_storage.filename or f"editor-image{suffix}"
    if not Path(original).suffix:
        original = f"{original}{suffix}"
    filename = stored_filename_for(original)
    file_storage.save(MANUALS_DIR / filename)
    return filename, original


def save_purchase_followup_image(file_storage):
    if not file_storage or not (file_storage.filename or file_storage.mimetype):
        return "", ""
    if not is_allowed_editor_image(file_storage):
        raise ValueError(file_storage.filename or "")
    return save_editor_image(file_storage)


def save_powder_coating_image(file_storage):
    return save_purchase_followup_image(file_storage)


def parse_positive_int(value, field_name):
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"{field_name}必须为整数")
    if parsed <= 0:
        raise ValueError(f"{field_name}必须大于 0")
    return parsed


def parse_nonnegative_price(value):
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        raise ValueError("单价必须为数字")
    if parsed < 0:
        raise ValueError("单价不能小于 0")
    return round(parsed, 2)


def parse_nonnegative_number(value, field_name):
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{field_name}必须为数字")
    if parsed < 0:
        raise ValueError(f"{field_name}不能小于 0")
    return round(parsed, 2)


def carton_size_text(length, width, height):
    values = [length, width, height]
    if not all(float(value or 0) > 0 for value in values):
        return ""
    return "x".join(f"{float(value):g}" for value in values)


def parse_carton_size_dimensions(size_text):
    values = re.findall(r"\d+(?:\.\d+)?", size_text or "")
    if len(values) < 3:
        return 0, 0, 0
    try:
        length, width, height = (round(float(value), 2) for value in values[:3])
    except ValueError:
        return 0, 0, 0
    if not all(value > 0 for value in (length, width, height)):
        return 0, 0, 0
    return length, width, height


def calculate_carton_unit_price(length, width, height, board_price):
    if not all(float(value or 0) > 0 for value in (length, width, height, board_price)):
        return 0
    return round((float(length) + float(width) + 8) * (float(width) + float(height) + 4) * float(board_price) / 10000, 2)


def carton_board_price(product, board_type):
    column = CARTON_BOARD_PRICE_COLUMNS.get(board_type)
    if not column or product is None:
        return 0
    return float(product[column] or 0)


def pdf_wrapped_paragraph(text, style, chunk_size=12):
    raw_text = str(text or "").strip()
    if not raw_text:
        return Paragraph("", style)
    lines = []
    for raw_line in raw_text.splitlines():
        words = re.split(r"([\\s/，,;；、]+)", raw_line)
        current = ""
        for word in words:
            if not word:
                continue
            pieces = wrap(word, chunk_size, break_long_words=True, break_on_hyphens=False) or [word]
            for piece in pieces:
                if len(current) + len(piece) > chunk_size and current.strip():
                    lines.append(current.strip())
                    current = piece
                else:
                    current += piece
        if current.strip():
            lines.append(current.strip())
    return Paragraph("<br/>".join(xml_escape(line) for line in lines), style)


def pdf_single_line_paragraph(text, style):
    raw_text = str(text or "").strip()
    if not raw_text:
        return Paragraph("", style)
    return Paragraph("<br/>".join(xml_escape(line.strip()) for line in raw_text.splitlines()), style)


def stored_shipment_image_filename(shipment_id, original, suffix):
    safe_name = secure_filename(original) or f"shipment-image{suffix}"
    stem = Path(safe_name).stem or "shipment-image"
    return f"{shipment_id}-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}-{stem}{suffix}"


def _shipment_image_limit(name):
    value = int(app.config[name])
    if value <= 0:
        raise RuntimeError(f"{name} 必须大于 0")
    return value


def _copy_shipment_image_stream(
    file_storage, destination, max_file_bytes, remaining_total_bytes
):
    original = file_storage.filename or "未命名图片"
    written = 0
    try:
        file_storage.stream.seek(0)
    except (AttributeError, OSError):
        pass
    try:
        with Path(destination).open("xb") as output:
            while True:
                chunk = file_storage.stream.read(64 * 1024)
                if not chunk:
                    break
                next_size = written + len(chunk)
                if next_size > max_file_bytes:
                    raise RequestEntityTooLarge(
                        description=(
                            f"发货图片 {original} 超过单文件大小限制 "
                            f"{format_bytes(max_file_bytes)}"
                        )
                    )
                if next_size > remaining_total_bytes:
                    raise RequestEntityTooLarge(
                        description=(
                            "本次发货图片总大小超过限制 "
                            f"{format_bytes(_shipment_image_limit('SHIPMENT_IMAGE_MAX_TOTAL_BYTES'))}"
                        )
                    )
                output.write(chunk)
                written = next_size
    finally:
        try:
            file_storage.stream.seek(0)
        except (AttributeError, OSError):
            pass
    return written


def _stage_shipment_images(files, assembly=False):
    files = list(files)
    max_files = _shipment_image_limit("SHIPMENT_IMAGE_MAX_FILES")
    if len(files) > max_files:
        raise RequestEntityTooLarge(
            description=f"发货图片数量不能超过 {max_files} 张"
        )
    validated = []
    for file_storage in files:
        submitted_suffix = Path(file_storage.filename or "").suffix.lower()
        submitted_mimetype = (file_storage.mimetype or "").lower()
        if submitted_suffix == ".webp" or submitted_mimetype == "image/webp":
            raise ValueError(
                f"{file_storage.filename or '未命名文件'}："
                "新发货图片不支持 WebP，请转换为 PNG、JPEG 或 GIF"
            )
        suffix = shipment_image_suffix_for(file_storage)
        if not suffix:
            raise ValueError(
                f"{file_storage.filename or '未命名文件'} 不是安全的位图文件"
            )
        original = file_storage.filename or f"shipment-image{suffix}"
        if assembly:
            original = _assembly_original_image_basename(original, suffix)
        elif not Path(original).suffix:
            original = f"{original}{suffix}"
        validated.append((file_storage, suffix, original))

    SHIPMENT_IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    max_file_bytes = _shipment_image_limit("SHIPMENT_IMAGE_MAX_FILE_BYTES")
    max_total_bytes = _shipment_image_limit("SHIPMENT_IMAGE_MAX_TOTAL_BYTES")
    total_bytes = 0
    staged_files = []
    try:
        for file_storage, suffix, original in validated:
            prefix = ".assembly-stage" if assembly else ".shipment-stage"
            stage_path = SHIPMENT_IMAGES_DIR / f"{prefix}-{uuid.uuid4().hex}.tmp"
            staged = {
                "stage_path": stage_path,
                "suffix": suffix,
                "original_filename": original,
                "content_type": "",
                "file_size": 0,
                "created_final": False,
            }
            staged_files.append(staged)
            file_size = _copy_shipment_image_stream(
                file_storage,
                stage_path,
                max_file_bytes,
                max_total_bytes - total_bytes,
            )
            image_format = detected_shipment_raster_format(stage_path)
            expected_format = SHIPMENT_RASTER_SUFFIX_FORMATS[suffix]
            if image_format != expected_format:
                raise ValueError(f"{original} 不是安全的位图文件")
            total_bytes += file_size
            staged["file_size"] = file_size
            staged["content_type"] = {
                "png": "image/png",
                "jpeg": "image/jpeg",
                "gif": "image/gif",
            }[image_format]
    except Exception:
        cleanup_assembly_shipment_image_stages(staged_files)
        raise
    return staged_files


def selected_shipment_images():
    return [file for file in request.files.getlist("images") if file and (file.filename or file.mimetype)]


def save_shipment_images(conn, shipment_id, files):
    staged_files = _stage_shipment_images(files)
    saved = []
    uploaded_at = datetime.utcnow().isoformat(timespec="seconds")
    complete = False
    try:
        for staged in staged_files:
            filename = stored_shipment_image_filename(
                shipment_id,
                staged["original_filename"],
                staged["suffix"],
            )
            final_path = SHIPMENT_IMAGES_DIR / filename
            staged["final_path"] = final_path
            if final_path.exists():
                raise FileExistsError("发货图片文件名冲突")
            os.replace(staged["stage_path"], final_path)
            staged["created_final"] = True
            conn.execute(
                """
                INSERT INTO product_order_shipment_images (
                    shipment_id, filename, original_filename, content_type,
                    file_size, uploaded_at, uploaded_ip, uploaded_user_agent
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    shipment_id,
                    filename,
                    staged["original_filename"],
                    staged["content_type"],
                    staged["file_size"],
                    uploaded_at,
                    client_ip(),
                    request.headers.get("User-Agent", "")[:500],
                ),
            )
            saved.append(filename)
        complete = True
        return saved
    finally:
        cleanup_assembly_shipment_image_stages(
            staged_files, remove_final=not complete
        )


def stored_assembly_shipment_image_filename(batch_id, original, suffix):
    batch_id = parse_positive_int(batch_id, "组装发货批次 ID")
    return stored_shipment_image_filename(
        f"assembly-{batch_id}", original, suffix
    )


def _assembly_original_image_basename(original, suffix):
    basename = Path(str(original or "").replace("\\", "/")).name
    if not basename:
        basename = f"assembly-image{suffix}"
    if not Path(basename).suffix:
        basename = f"{basename}{suffix}"
    return basename


def cleanup_assembly_shipment_image_stages(staged_files, remove_final=False):
    for staged in staged_files or []:
        stage_path = staged.get("stage_path")
        if stage_path:
            try:
                Path(stage_path).unlink(missing_ok=True)
            except OSError:
                pass
        if remove_final and staged.get("created_final"):
            final_path = staged.get("final_path")
            if final_path:
                try:
                    Path(final_path).unlink(missing_ok=True)
                except OSError:
                    pass


def stage_assembly_shipment_images(files):
    return _stage_shipment_images(files, assembly=True)


def _finalize_assembly_shipment_image(stage_path, final_path):
    stage_path = Path(stage_path)
    final_path = Path(final_path)
    if final_path.exists():
        raise FileExistsError("组装发货图片文件名冲突")
    os.replace(stage_path, final_path)


def save_assembly_shipment_images(conn, batch_id, files):
    batch_id = parse_positive_int(batch_id, "组装发货批次 ID")
    uploaded_at = datetime.utcnow().isoformat(timespec="seconds")
    saved = []
    for staged in files:
        filename = stored_assembly_shipment_image_filename(
            batch_id,
            staged["original_filename"],
            staged["suffix"],
        )
        if (
            secure_filename(filename) != filename
            or Path(filename).name != filename
            or Path(filename).suffix.lower() not in SHIPMENT_RASTER_SUFFIX_FORMATS
        ):
            raise ValueError("组装发货图片文件名无效")
        final_path = SHIPMENT_IMAGES_DIR / filename
        staged["final_path"] = final_path
        if final_path.exists():
            raise FileExistsError("组装发货图片文件名冲突")
        try:
            _finalize_assembly_shipment_image(staged["stage_path"], final_path)
        except Exception:
            if not Path(staged["stage_path"]).exists() and final_path.exists():
                staged["created_final"] = True
            raise
        staged["created_final"] = True
        conn.execute(
            """
            INSERT INTO assembly_shipment_images (
                batch_id, filename, original_filename, content_type,
                file_size, uploaded_at, uploaded_ip, uploaded_user_agent
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                batch_id,
                filename,
                staged["original_filename"],
                staged["content_type"],
                staged["file_size"],
                uploaded_at,
                client_ip(),
                request.headers.get("User-Agent", "")[:500],
            ),
        )
        saved.append(
            {
                "filename": filename,
                "original_filename": staged["original_filename"],
                "content_type": staged["content_type"],
                "file_size": staged["file_size"],
            }
        )
    return saved


def safe_assembly_shipment_image_filename(batch_id, filename):
    batch_id = parse_positive_int(batch_id, "组装发货批次 ID")
    filename = str(filename or "")
    if not (
        filename
        and Path(filename).name == filename
        and secure_filename(filename) == filename
        and Path(filename).suffix.lower() in IMAGE_EXTENSIONS
    ):
        return False
    prefix = f"assembly-{batch_id}-"
    if not filename.startswith(prefix):
        return False
    generated_parts = filename[len(prefix) :].split("-", 2)
    return bool(
        len(generated_parts) == 3
        and re.fullmatch(r"\d{14}", generated_parts[0])
        and re.fullmatch(r"[0-9a-f]{8}", generated_parts[1])
        and generated_parts[2]
    )


def delete_assembly_shipment_image_files(batch_id, filenames):
    candidates = sorted(
        {
            str(filename or "")
            for filename in filenames
            if safe_assembly_shipment_image_filename(batch_id, filename)
        }
    )
    if not candidates:
        return
    placeholders = ",".join("?" for _ in candidates)
    with get_db() as conn:
        referenced = {
            row["filename"]
            for row in conn.execute(
                f"""
                SELECT filename
                FROM assembly_shipment_images
                WHERE filename IN ({placeholders})
                UNION
                SELECT filename
                FROM product_order_shipment_images
                WHERE filename IN ({placeholders})
                """,
                candidates + candidates,
            ).fetchall()
        }
    for filename in candidates:
        if filename in referenced:
            continue
        try:
            (SHIPMENT_IMAGES_DIR / filename).unlink(missing_ok=True)
        except OSError:
            pass


def stored_inspection_report_filename(plan_id, original):
    safe_name = secure_filename(original) or "inspection-report"
    suffix = Path(safe_name).suffix.lower()
    stem = Path(safe_name).stem or "inspection-report"
    return f"{plan_id}-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}-{stem}{suffix}"


def selected_inspection_report_files():
    return [
        file
        for file in request.files.getlist("inspection_reports")
        if file and file.filename
    ]


def save_inspection_report_files(conn, plan_id, files):
    files = list(files)
    for file_storage in files:
        if not is_allowed_file(file_storage):
            raise ValueError(file_storage.filename or "")
    INSPECTION_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    uploaded_at = datetime.utcnow().isoformat(timespec="seconds")
    saved = []
    for file_storage in files:
        original = file_storage.filename or "inspection-report"
        filename = stored_inspection_report_filename(plan_id, original)
        report_path = INSPECTION_REPORTS_DIR / filename
        file_storage.save(report_path)
        conn.execute(
            """
            INSERT INTO shipment_plan_inspection_reports (
                plan_id, filename, original_filename, content_type,
                file_size, uploaded_at, uploaded_by
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                plan_id,
                filename,
                original,
                file_storage.mimetype or "",
                report_path.stat().st_size,
                uploaded_at,
                current_admin_username(),
            ),
        )
        saved.append(filename)
    return saved


def stored_production_drawing_filename(followup_id, original):
    safe_name = secure_filename(original) or "production-drawing"
    suffix = Path(safe_name).suffix.lower()
    stem = Path(safe_name).stem or "production-drawing"
    return f"{followup_id}-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}-{stem}{suffix}"


def selected_production_drawings():
    return [
        file
        for file in request.files.getlist("drawings")
        if file and file.filename
    ]


def save_production_drawing_files(conn, followup_id, files):
    files = list(files)
    for file_storage in files:
        if not is_allowed_file(file_storage):
            raise ValueError(file_storage.filename or "")
    PRODUCTION_DRAWINGS_DIR.mkdir(parents=True, exist_ok=True)
    created_at = datetime.utcnow().isoformat(timespec="seconds")
    saved = []
    for file_storage in files:
        original = file_storage.filename or "production-drawing"
        filename = stored_production_drawing_filename(followup_id, original)
        drawing_path = PRODUCTION_DRAWINGS_DIR / filename
        file_storage.save(drawing_path)
        conn.execute(
            """
            INSERT INTO production_followup_files (
                followup_id, filename, original_filename, file_type,
                content_type, file_size, created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                followup_id,
                filename,
                original,
                file_type_for(filename),
                file_storage.mimetype or "",
                drawing_path.stat().st_size,
                created_at,
            ),
        )
        saved.append(filename)
    return saved


def next_production_batch_no(conn, ordered_at):
    compact_date = re.sub(r"[^0-9]", "", ordered_at or "")[:8]
    if len(compact_date) != 8:
        compact_date = datetime.now().strftime("%Y%m%d")
    prefix = f"PF-{compact_date}-"
    row = conn.execute(
        """
        SELECT batch_no
        FROM production_followups
        WHERE batch_no LIKE ?
        ORDER BY batch_no DESC
        LIMIT 1
        """,
        (f"{prefix}%",),
    ).fetchone()
    if row is None:
        next_number = 1
    else:
        try:
            next_number = int(str(row["batch_no"]).rsplit("-", 1)[-1]) + 1
        except ValueError:
            next_number = 1
    return f"{prefix}{next_number:03d}"


def production_stage_column(stage):
    columns = {
        "laser": "laser_completed_at",
        "bending": "bending_completed_at",
        "welding": "welding_completed_at",
    }
    return columns.get(stage)


def production_stage_can_complete(row, stage):
    if stage == "laser":
        return not row["laser_completed_at"]
    if stage == "bending":
        return bool(row["laser_completed_at"]) and not row["bending_completed_at"]
    if stage == "welding":
        return bool(row["bending_completed_at"]) and not row["welding_completed_at"]
    return False


def production_stage_can_revert(row, stage):
    if stage == "laser":
        return bool(row["laser_completed_at"]) and not row["bending_completed_at"] and not row["welding_completed_at"]
    if stage == "bending":
        return bool(row["bending_completed_at"]) and not row["welding_completed_at"]
    if stage == "welding":
        return bool(row["welding_completed_at"])
    return False


def production_stage_waiting_text(row, stage):
    if stage == "bending" and not row["laser_completed_at"]:
        return "待激光"
    if stage == "welding" and not row["bending_completed_at"]:
        return "待折弯"
    return PRODUCTION_STAGE_LABELS.get(stage, "")


def fetch_production_followups(conn, query=""):
    query = (query or "").strip()
    params = []
    sql = """
        SELECT production_followups.*
        FROM production_followups
    """
    if query:
        like = f"%{query}%"
        sql += """
            WHERE production_followups.batch_no LIKE ?
               OR production_followups.ordered_at LIKE ?
               OR production_followups.drawing_no LIKE ?
               OR production_followups.product_name LIKE ?
               OR EXISTS (
                    SELECT 1
                    FROM production_followup_files
                    WHERE production_followup_files.followup_id = production_followups.id
                      AND production_followup_files.original_filename LIKE ?
               )
        """
        params.extend([like, like, like, like, like])
    sql += """
        ORDER BY production_followups.ordered_at DESC,
                 production_followups.batch_no DESC,
                 production_followups.id DESC
    """
    rows = conn.execute(sql, params).fetchall()
    if not rows:
        return []
    followup_ids = [row["id"] for row in rows]
    placeholders = ",".join("?" for _ in followup_ids)
    file_rows = conn.execute(
        f"""
        SELECT *
        FROM production_followup_files
        WHERE followup_id IN ({placeholders})
        ORDER BY id ASC
        """,
        followup_ids,
    ).fetchall()
    files_by_followup = {followup_id: [] for followup_id in followup_ids}
    for file_row in file_rows:
        files_by_followup.setdefault(file_row["followup_id"], []).append(file_row)
    return [
        {
            "row": row,
            "files": files_by_followup.get(row["id"], []),
        }
        for row in rows
    ]


def fetch_production_followup_detail(conn, followup_id):
    row = conn.execute(
        "SELECT * FROM production_followups WHERE id = ?",
        (followup_id,),
    ).fetchone()
    if row is None:
        return None
    files = conn.execute(
        """
        SELECT *
        FROM production_followup_files
        WHERE followup_id = ?
        ORDER BY id ASC
        """,
        (followup_id,),
    ).fetchall()
    return {
        "row": row,
        "files": files,
    }


def copy_upload_file(filename, original_filename):
    new_filename = stored_filename_for(original_filename or filename)
    shutil.copy2(MANUALS_DIR / filename, MANUALS_DIR / new_filename)
    return new_filename, original_filename, file_type_for(new_filename)


def selected_uploads():
    uploads = [file for file in request.files.getlist("uploads") if file and file.filename]
    legacy_upload = request.files.get("upload")
    if legacy_upload and legacy_upload.filename:
        uploads.append(legacy_upload)
    return uploads


def save_uploads(files):
    saved = []
    for file_storage in files:
        if not is_allowed_file(file_storage):
            raise ValueError(file_storage.filename or "")
        saved.append(save_upload(file_storage))
    return saved


def posted_inspection_requirements():
    item_names = request.form.getlist("inspection_item_name")
    standards = request.form.getlist("inspection_standard")
    methods = request.form.getlist("inspection_method")
    remarks = request.form.getlist("inspection_remark")
    rows = []
    for item_name, standard, method, remark in zip_longest(
        item_names,
        standards,
        methods,
        remarks,
        fillvalue="",
    ):
        rows.append(
            {
                "item_name": str(item_name or "").strip(),
                "standard": str(standard or "").strip(),
                "method": str(method or "").strip(),
                "remark": str(remark or "").strip(),
            }
        )
    return rows


def save_manual_inspection_requirements(conn, manual_id, requirements):
    now = datetime.utcnow().isoformat(timespec="seconds")
    cleaned = []
    for requirement in requirements:
        item_name = str(requirement.get("item_name") or "").strip()
        standard = str(requirement.get("standard") or "").strip()
        method = str(requirement.get("method") or "").strip()
        remark = str(requirement.get("remark") or "").strip()
        if not any([item_name, standard, method, remark]):
            continue
        cleaned.append((manual_id, item_name, standard, method, remark, len(cleaned), now, now))

    conn.execute(
        "DELETE FROM manual_inspection_requirements WHERE manual_id = ?",
        (manual_id,),
    )
    if cleaned:
        conn.executemany(
            """
            INSERT INTO manual_inspection_requirements (
                manual_id, item_name, standard, method, remark,
                sort_order, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            cleaned,
        )


def get_manual_inspection_requirements(conn, manual_id):
    return conn.execute(
        """
        SELECT *
        FROM manual_inspection_requirements
        WHERE manual_id = ?
        ORDER BY sort_order ASC, id ASC
        """,
        (manual_id,),
    ).fetchall()


def get_manual_inspection_requirements_map(conn, manual_ids):
    manual_ids = [manual_id for manual_id in dict.fromkeys(manual_ids) if manual_id]
    if not manual_ids:
        return {}
    placeholders = ",".join("?" for _ in manual_ids)
    rows = conn.execute(
        f"""
        SELECT *
        FROM manual_inspection_requirements
        WHERE manual_id IN ({placeholders})
        ORDER BY manual_id ASC, sort_order ASC, id ASC
        """,
        manual_ids,
    ).fetchall()
    result = {manual_id: [] for manual_id in manual_ids}
    for row in rows:
        result.setdefault(row["manual_id"], []).append(row)
    return result


def posted_product_materials():
    materials = request.form.getlist("material")
    thicknesses = request.form.getlist("material_thickness")
    surface_types = request.form.getlist("surface_type")
    suppliers = request.form.getlist("material_supplier")
    rows = []
    for material, thickness, surface_type, supplier in zip_longest(
        materials,
        thicknesses,
        surface_types,
        suppliers,
        fillvalue="",
    ):
        rows.append(
            {
                "material": str(material or "").strip(),
                "thickness": str(thickness or "").strip(),
                "surface_type": str(surface_type or "").strip(),
                "supplier": str(supplier or "").strip(),
            }
        )
    return rows


def save_product_materials(conn, manual_id, materials):
    now = datetime.utcnow().isoformat(timespec="seconds")
    cleaned = []
    for material in materials:
        material_name = str(material.get("material") or "").strip()
        thickness = str(material.get("thickness") or "").strip()
        surface_type = str(material.get("surface_type") or "").strip()
        supplier = str(material.get("supplier") or "").strip()
        if not any([material_name, thickness, surface_type, supplier]):
            continue
        cleaned.append((manual_id, material_name, thickness, surface_type, supplier, len(cleaned), now, now))

    conn.execute("DELETE FROM product_materials WHERE manual_id = ?", (manual_id,))
    if cleaned:
        conn.executemany(
            """
            INSERT INTO product_materials (
                manual_id, material, thickness, surface_type, supplier,
                sort_order, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            cleaned,
        )


def save_product_assembly_components(conn, manual_id, components, now):
    conn.execute(
        "DELETE FROM product_assembly_components WHERE manual_id = ?",
        (manual_id,),
    )
    if not components:
        return
    conn.executemany(
        """
        INSERT INTO product_assembly_components (
            manual_id, assembly_drawing_no, quantity_per_set, sort_order,
            created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        [
            (
                manual_id,
                component["assembly_drawing_no"],
                component["quantity_per_set"],
                component["sort_order"],
                now,
                now,
            )
            for component in components
        ],
    )


def get_product_assembly_components(conn, manual_id):
    return conn.execute(
        """
        SELECT *
        FROM product_assembly_components
        WHERE manual_id = ?
        ORDER BY sort_order ASC, id ASC
        """,
        (manual_id,),
    ).fetchall()


def get_product_assembly_components_map(conn, manual_ids):
    manual_ids = [manual_id for manual_id in dict.fromkeys(manual_ids) if manual_id]
    if not manual_ids:
        return {}
    placeholders = ",".join("?" for _ in manual_ids)
    rows = conn.execute(
        f"""
        SELECT *
        FROM product_assembly_components
        WHERE manual_id IN ({placeholders})
        ORDER BY manual_id ASC, sort_order ASC, id ASC
        """,
        manual_ids,
    ).fetchall()
    result = {manual_id: [] for manual_id in manual_ids}
    for row in rows:
        result.setdefault(row["manual_id"], []).append(row)
    return result


def get_product_materials(conn, manual_id):
    return conn.execute(
        """
        SELECT *
        FROM product_materials
        WHERE manual_id = ?
        ORDER BY sort_order ASC, id ASC
        """,
        (manual_id,),
    ).fetchall()


def get_product_materials_map(conn, manual_ids):
    manual_ids = [manual_id for manual_id in dict.fromkeys(manual_ids) if manual_id]
    if not manual_ids:
        return {}
    placeholders = ",".join("?" for _ in manual_ids)
    rows = conn.execute(
        f"""
        SELECT *
        FROM product_materials
        WHERE manual_id IN ({placeholders})
        ORDER BY manual_id ASC, sort_order ASC, id ASC
        """,
        manual_ids,
    ).fetchall()
    result = {manual_id: [] for manual_id in manual_ids}
    for row in rows:
        result.setdefault(row["manual_id"], []).append(row)
    return result


def get_product_material_options(conn):
    fields = {
        "materials": "material",
        "thicknesses": "thickness",
        "surface_types": "surface_type",
        "suppliers": "supplier",
    }
    options = {}
    for key, column in fields.items():
        rows = conn.execute(
            f"""
            SELECT DISTINCT TRIM({column}) AS value
            FROM product_materials
            WHERE TRIM(COALESCE({column}, '')) != ''
            ORDER BY value COLLATE NOCASE ASC
            """
        ).fetchall()
        options[key] = [row["value"] for row in rows]
    return options


def build_shipment_plan_inspection_sections(conn, plan):
    if not plan:
        return []
    items = [dict(item) for item in plan.get("plan_items", [])]
    requirements_by_manual = get_manual_inspection_requirements_map(
        conn,
        [item.get("manual_id") for item in items],
    )
    sections = []
    for item in items:
        requirements = requirements_by_manual.get(item.get("manual_id"), [])
        if not requirements:
            requirements = [
                {
                    "item_name": "",
                    "standard": "该产品暂无检验要求，可现场填写",
                    "method": "",
                    "remark": "",
                }
            ]
        sections.append(
            {
                "manual_id": item.get("manual_id"),
                "order_no": item.get("order_no") or "",
                "drawing_no": item.get("drawing_no") or "",
                "product_name": item.get("product_name") or "",
                "planned_quantity": item.get("planned_quantity") or 0,
                "pack_quantity": item.get("pack_quantity") or "",
                "box_count": item.get("box_count") or "",
                "requirements": requirements,
            }
        )
    return sections


def get_manual_files(conn, manual_id):
    return conn.execute(
        """
        SELECT *
        FROM manual_files
        WHERE manual_id = ?
        ORDER BY id ASC
        """,
        (manual_id,),
    ).fetchall()


def get_product_options(conn):
    return conn.execute(
        """
        SELECT id, drawing_no, product_name, customer
        FROM manuals
        ORDER BY drawing_no COLLATE NOCASE ASC, product_name COLLATE NOCASE ASC
        """
    ).fetchall()


def get_active_user_options(conn):
    return conn.execute(
        """
        SELECT username
        FROM users
        WHERE active = 1
        ORDER BY username COLLATE NOCASE ASC
        """
    ).fetchall()


def normalize_import_header(value):
    return re.sub(r"[\s_\-]+", "", str(value or "").strip().lower())


def normalize_import_drawing(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_import_quantity(value):
    if value is None or str(value).strip() == "":
        raise ValueError
    if isinstance(value, int):
        quantity = value
    elif isinstance(value, float):
        if not value.is_integer():
            raise ValueError
        quantity = int(value)
    else:
        text = str(value).strip()
        if text.endswith(".0"):
            text = text[:-2]
        quantity = int(text)
    if quantity <= 0:
        raise ValueError
    if quantity > MAX_ORDER_QUANTITY:
        raise OverflowError
    return quantity


def parse_import_date(value):
    if value is None or str(value).strip() == "":
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def parse_order_import_workbook(upload):
    filename = upload.filename or ""
    extension = Path(filename).suffix.lower()
    if extension not in ORDER_IMPORT_EXTENSIONS:
        return [], [f"请上传 .xlsx 或 .xlsm 格式的 Excel 文件，当前文件格式为 {extension or '未知'}"]

    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError):
        return [], ["Excel 文件无法读取，请确认文件未损坏且格式为 .xlsx 或 .xlsm"]

    sheet = workbook.active
    first_row = None
    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if any(value is not None and str(value).strip() for value in row):
            first_row = (row_index, row)
            break

    if first_row is None:
        return [], ["Excel 文件为空，请至少填写图号和数量两列"]

    first_row_index, first_values = first_row
    drawing_col = None
    quantity_col = None
    planned_ship_col = None
    for index, value in enumerate(first_values):
        header = normalize_import_header(value)
        if header in {normalize_import_header(item) for item in ORDER_IMPORT_DRAWING_HEADERS}:
            drawing_col = index
        if header in {normalize_import_header(item) for item in ORDER_IMPORT_QUANTITY_HEADERS}:
            quantity_col = index
        if header in {normalize_import_header(item) for item in ORDER_IMPORT_PLANNED_SHIP_HEADERS}:
            planned_ship_col = index

    has_header = drawing_col is not None and quantity_col is not None
    if not has_header:
        drawing_col = 0
        quantity_col = 1

    rows = []
    errors = []
    iterator = sheet.iter_rows(
        min_row=first_row_index + 1 if has_header else first_row_index,
        values_only=True,
    )
    for row_number, row in enumerate(iterator, start=first_row_index + 1 if has_header else first_row_index):
        drawing_no = normalize_import_drawing(row[drawing_col] if drawing_col < len(row) else None)
        raw_quantity = row[quantity_col] if quantity_col < len(row) else None
        if drawing_no.startswith("填写说明") or re.match(r"^\d+[.．、]", drawing_no):
            continue
        if not drawing_no and (raw_quantity is None or str(raw_quantity).strip() == ""):
            continue
        if not drawing_no:
            errors.append(f"第 {row_number} 行缺少图号")
            continue
        try:
            quantity = parse_import_quantity(raw_quantity)
        except OverflowError:
            errors.append(
                f"第 {row_number} 行数量不能超过 {MAX_ORDER_QUANTITY}"
            )
            continue
        except (TypeError, ValueError):
            errors.append(f"第 {row_number} 行数量必须是大于 0 的整数")
            continue
        planned_ship_at = ""
        if planned_ship_col is not None and planned_ship_col < len(row):
            planned_ship_at = parse_import_date(row[planned_ship_col])
        rows.append({
            "row_number": row_number,
            "drawing_no": drawing_no,
            "quantity": quantity,
            "planned_ship_at": planned_ship_at,
        })

    if not rows and not errors:
        errors.append("没有找到可导入的订单明细")
    return rows, errors


def match_order_import_items(conn, import_rows):
    manual_rows = conn.execute(
        """
        SELECT id, drawing_no, product_name, customer, supplier
        FROM manuals
        WHERE drawing_no IS NOT NULL AND drawing_no != ''
        """
    ).fetchall()
    manuals_by_drawing = {}
    for row in manual_rows:
        key = row["drawing_no"].strip().lower()
        manuals_by_drawing.setdefault(key, []).append(row)

    errors = []
    items_by_manual_id = {}
    for item in import_rows:
        key = item["drawing_no"].strip().lower()
        matches = manuals_by_drawing.get(key, [])
        if not matches:
            errors.append(f"第 {item['row_number']} 行图号 {item['drawing_no']} 未在产品资料中找到")
            continue
        if len(matches) > 1:
            errors.append(f"第 {item['row_number']} 行图号 {item['drawing_no']} 匹配到多个产品，请先在产品资料中确认图号唯一")
            continue
        manual = matches[0]
        existing = items_by_manual_id.setdefault(
            manual["id"],
            {"manual": manual, "quantity": 0, "rows": []},
        )
        existing["quantity"] += item["quantity"]
        existing["rows"].append(item["row_number"])

    items = []
    for item in items_by_manual_id.values():
        if item["quantity"] > MAX_ORDER_QUANTITY:
            rows = "、".join(str(row_number) for row_number in item["rows"])
            errors.append(
                f"第 {rows} 行图号 {item['manual']['drawing_no']} "
                f"合计数量不能超过 {MAX_ORDER_QUANTITY}"
            )
            continue
        items.append(item)

    return items, errors


def get_order_shipments(conn, order_id):
    return conn.execute(
        """
        SELECT *
        FROM product_order_shipments
        WHERE order_id = ?
        ORDER BY shipped_at ASC, id ASC
        """,
        (order_id,),
    ).fetchall()


def _assembly_shipment_summary_tables_exist(conn):
    rows = conn.execute(
        """
        SELECT name
        FROM sqlite_master
        WHERE type = 'table'
          AND name IN (
              'assembly_shipment_batches',
              'assembly_shipment_items',
              'assembly_shipment_allocations'
          )
        """
    ).fetchall()
    return len(rows) == 3


def order_shipment_summary_subquery(include_assembly=False, conn=None):
    if include_assembly and conn is not None:
        include_assembly = _assembly_shipment_summary_tables_exist(conn)
    source = "product_order_shipments"
    if include_assembly:
        source = """
        (
            SELECT order_id, shipped_quantity, shipped_at
            FROM product_order_shipments

            UNION ALL

            SELECT assembly_shipment_allocations.order_id,
                   assembly_shipment_allocations.quantity AS shipped_quantity,
                   assembly_shipment_batches.shipped_at
            FROM assembly_shipment_allocations
            JOIN assembly_shipment_items
              ON assembly_shipment_items.id = assembly_shipment_allocations.item_id
            JOIN assembly_shipment_batches
              ON assembly_shipment_batches.id = assembly_shipment_items.batch_id
            WHERE assembly_shipment_allocations.order_id IS NOT NULL
        ) AS all_shipments
        """
    return f"""
        SELECT order_id,
               SUM(shipped_quantity) AS shipped_total,
               MAX(shipped_at) AS last_shipped_at
        FROM {source}
        GROUP BY order_id
    """


def get_shipped_orders_query(include_prices=False):
    price_columns = ""
    if include_prices:
        price_columns = """,
            product_order_shipments.unit_price_minor,
            product_order_shipments.currency,
            product_order_shipments.price_recorded_by,
            product_order_shipments.price_recorded_at,
            product_order_shipments.unit_price_minor
                * product_order_shipments.shipped_quantity AS line_total_minor
        """
    return f"""
        SELECT
            product_order_shipments.id,
            product_order_shipments.order_id,
            product_order_shipments.shipped_quantity,
            product_order_shipments.shipped_at,
            product_order_shipments.created_at,
            product_order_shipments.logistics_no,
            product_order_shipments.signature_token,
            product_order_shipments.signature_status,
            product_order_shipments.signature_image,
            product_order_shipments.signature_expires_at,
            product_order_shipments.signed_at,
            product_order_shipments.signed_ip,
            product_order_shipments.signed_user_agent,
            product_order_shipments.photo_upload_token,
            product_order_shipments.remark,
            product_orders.manual_id,
            product_orders.order_no,
            product_orders.ordered_at,
            product_orders.planned_ship_at,
            product_orders.remark AS order_remark,
            product_orders.customer AS order_customer,
            product_orders.customer_email,
            manuals.drawing_no,
            manuals.product_name,
            manuals.remark AS product_remark,
            manuals.customer AS product_customer,
            manuals.supplier AS supplier
            {price_columns}
        FROM product_order_shipments
        LEFT JOIN product_orders ON product_orders.id = product_order_shipments.order_id
        LEFT JOIN manuals ON manuals.id = product_orders.manual_id
    """


def get_order_customer_options(conn):
    return [
        row["customer"]
        for row in conn.execute(
            """
            SELECT DISTINCT COALESCE(NULLIF(product_orders.customer, ''), manuals.customer) AS customer
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            WHERE COALESCE(NULLIF(product_orders.customer, ''), manuals.customer) IS NOT NULL
              AND COALESCE(NULLIF(product_orders.customer, ''), manuals.customer) != ''
            ORDER BY customer COLLATE NOCASE
            """
        ).fetchall()
    ]


def get_shipment_customer_options(conn):
    """Return stable customer labels available to either shipment workflow."""
    rows = conn.execute(
        """
        SELECT customer
        FROM (
            SELECT TRIM(customer) AS customer
            FROM manuals
            WHERE TRIM(customer) != ''
            UNION ALL
            SELECT COALESCE(NULLIF(TRIM(product_orders.customer), ''), TRIM(manuals.customer))
                AS customer
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            WHERE COALESCE(NULLIF(TRIM(product_orders.customer), ''), TRIM(manuals.customer)) != ''
            UNION ALL
            SELECT TRIM(customer) AS customer
            FROM assembly_shipment_batches
            WHERE TRIM(customer) != ''
        )
        ORDER BY customer COLLATE NOCASE ASC, customer ASC
        """
    ).fetchall()
    customers = []
    seen = set()
    for row in rows:
        customer = str(row["customer"] or "").strip()
        key = customer.casefold()
        if not customer or key in seen:
            continue
        seen.add(key)
        customers.append(customer)
    return customers


def get_unshipped_order_options(conn):
    return conn.execute(
        f"""
        SELECT product_orders.id,
               product_orders.order_no,
               product_orders.quantity,
               product_orders.planned_ship_at,
               product_orders.customer AS order_customer,
               manuals.drawing_no,
               manuals.product_name,
               manuals.customer AS product_customer,
               manuals.supplier,
               product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS unshipped_quantity
        FROM product_orders
        JOIN manuals ON manuals.id = product_orders.manual_id
        LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments ON shipments.order_id = product_orders.id
        WHERE product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) > 0
        ORDER BY product_orders.planned_ship_at ASC, product_orders.ordered_at ASC, product_orders.id ASC
        """
    ).fetchall()


class AssemblyDefinitionNotFound(LookupError):
    pass


def get_assembly_options_for_customer(conn, customer):
    customer = str(customer or "").strip()
    if not customer:
        raise ValueError("请选择客户")
    rows = conn.execute(
        """
        SELECT product_assembly_components.assembly_drawing_no
        FROM product_assembly_components
        JOIN manuals ON manuals.id = product_assembly_components.manual_id
        WHERE manuals.customer = ?
          AND TRIM(product_assembly_components.assembly_drawing_no) != ''
        ORDER BY product_assembly_components.assembly_drawing_no COLLATE NOCASE ASC,
                 product_assembly_components.assembly_drawing_no ASC,
                 product_assembly_components.manual_id ASC
        """,
        (customer,),
    ).fetchall()
    options = []
    seen = set()
    for row in rows:
        drawing_no = str(row["assembly_drawing_no"] or "").strip()
        key = drawing_no.casefold()
        if key in seen:
            continue
        seen.add(key)
        options.append(drawing_no)
    return options


def get_assembly_definition(conn, customer, assembly_drawing_no):
    customer = str(customer or "").strip()
    assembly_drawing_no = str(assembly_drawing_no or "").strip()
    if not customer:
        raise ValueError("请选择客户")
    if not assembly_drawing_no:
        raise ValueError("请选择组装件图号")
    return conn.execute(
        """
        SELECT manuals.id AS manual_id,
               manuals.drawing_no,
               manuals.product_name,
               product_assembly_components.quantity_per_set,
               product_assembly_components.sort_order
        FROM product_assembly_components
        JOIN manuals ON manuals.id = product_assembly_components.manual_id
        WHERE manuals.customer = ?
          AND product_assembly_components.assembly_drawing_no = ? COLLATE NOCASE
        ORDER BY product_assembly_components.sort_order ASC,
                 manuals.drawing_no COLLATE NOCASE ASC,
                 manuals.id ASC
        """,
        (customer, assembly_drawing_no),
    ).fetchall()


def _validated_excluded_assembly_batch(
    conn, exclude_batch_id, customer, assembly_drawing_no
):
    if exclude_batch_id is None:
        return None
    batch_id = parse_positive_int(exclude_batch_id, "组装发货批次 ID")
    row = conn.execute(
        """
        SELECT id
        FROM assembly_shipment_batches
        WHERE id = ?
          AND customer = ?
          AND assembly_drawing_no = ? COLLATE NOCASE
        """,
        (batch_id, customer, assembly_drawing_no),
    ).fetchone()
    return int(row["id"]) if row else None


def get_component_open_orders(conn, manual_id, customer, exclude_batch_id=None):
    manual_id = parse_positive_int(manual_id, "配件 ID")
    customer = str(customer or "").strip()
    if not customer:
        raise ValueError("请选择客户")
    excluded_batch_id = (
        parse_positive_int(exclude_batch_id, "组装发货批次 ID")
        if exclude_batch_id is not None
        else -1
    )
    rows = conn.execute(
        f"""
        SELECT ordered.id,
               ordered.order_no,
               ordered.unshipped_quantity
        FROM (
            SELECT product_orders.id,
                   product_orders.order_no,
                   product_orders.planned_ship_at,
                   product_orders.ordered_at,
                   product_orders.created_at,
                   product_orders.quantity
                     - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0)
                     + COALESCE(excluded_allocations.quantity, 0)
                     AS unshipped_quantity
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments
              ON shipments.order_id = product_orders.id
            LEFT JOIN (
                SELECT assembly_shipment_allocations.order_id,
                       SUM(assembly_shipment_allocations.quantity) AS quantity
                FROM assembly_shipment_allocations
                JOIN assembly_shipment_items
                  ON assembly_shipment_items.id = assembly_shipment_allocations.item_id
                JOIN assembly_shipment_batches
                  ON assembly_shipment_batches.id = assembly_shipment_items.batch_id
                WHERE assembly_shipment_batches.id = ?
                  AND assembly_shipment_batches.customer = ?
                  AND assembly_shipment_items.manual_id = ?
                  AND assembly_shipment_allocations.order_id IS NOT NULL
                GROUP BY assembly_shipment_allocations.order_id
            ) AS excluded_allocations
              ON excluded_allocations.order_id = product_orders.id
            WHERE product_orders.manual_id = ?
              AND manuals.customer = ?
              AND COALESCE(NULLIF(TRIM(product_orders.customer), ''), manuals.customer) = ?
        ) AS ordered
        WHERE ordered.unshipped_quantity > 0
        ORDER BY CASE WHEN TRIM(ordered.planned_ship_at) = '' THEN 1 ELSE 0 END ASC,
                 ordered.planned_ship_at ASC,
                 ordered.ordered_at ASC,
                 ordered.created_at ASC,
                 ordered.id ASC
        """,
        (
            excluded_batch_id,
            customer,
            manual_id,
            manual_id,
            customer,
            customer,
        ),
    ).fetchall()
    return [
        {
            "id": int(row["id"]),
            "order_no": row["order_no"] or "",
            "unshipped_quantity": int(row["unshipped_quantity"] or 0),
        }
        for row in rows
    ]


def _normalize_assembly_overrides(overrides):
    if overrides is None:
        return {}
    if not isinstance(overrides, dict):
        raise ValueError("配件实际发货数量格式无效")
    normalized = {}
    for manual_id, quantity in overrides.items():
        normalized_id = parse_positive_int(manual_id, "配件 ID")
        normalized[normalized_id] = parse_positive_int(
            quantity, "配件实际发货数量"
        )
    return normalized


def _excluded_inventory_quantity(conn, batch_id, manual_id):
    if batch_id is None:
        return 0
    row = conn.execute(
        """
        SELECT COALESCE(SUM(inventory_deducted_quantity), 0) AS quantity
        FROM assembly_shipment_items
        WHERE batch_id = ? AND manual_id = ?
        """,
        (batch_id, manual_id),
    ).fetchone()
    return max(0, int(row["quantity"] or 0))


def build_assembly_shipment_preview(
    conn,
    customer,
    assembly_drawing_no,
    set_quantity,
    overrides=None,
    exclude_batch_id=None,
):
    customer = str(customer or "").strip()
    assembly_drawing_no = str(assembly_drawing_no or "").strip()
    if not customer:
        raise ValueError("请选择客户")
    if not assembly_drawing_no:
        raise ValueError("请选择组装件图号")
    set_quantity = parse_positive_int(set_quantity, "组装件套数")
    normalized_overrides = _normalize_assembly_overrides(overrides)
    canonical_assembly_drawing_no = next(
        (
            option
            for option in get_assembly_options_for_customer(conn, customer)
            if option.casefold() == assembly_drawing_no.casefold()
        ),
        None,
    )
    if canonical_assembly_drawing_no is None:
        raise AssemblyDefinitionNotFound("该组装件没有有效配件")
    assembly_drawing_no = canonical_assembly_drawing_no
    definition = get_assembly_definition(conn, customer, assembly_drawing_no)
    if not definition:
        raise AssemblyDefinitionNotFound("该组装件没有有效配件")
    definition_ids = {int(row["manual_id"]) for row in definition}
    if not set(normalized_overrides).issubset(definition_ids):
        raise ValueError("实际发货配件不属于所选组装件")
    valid_exclude_batch_id = _validated_excluded_assembly_batch(
        conn, exclude_batch_id, customer, assembly_drawing_no
    )
    expanded = expand_components(definition, set_quantity, normalized_overrides)
    preview = {
        "customer": customer,
        "assembly_drawing_no": assembly_drawing_no,
        "set_quantity": set_quantity,
        "items": [],
        "warnings": [],
    }
    for component in expanded:
        manual_id = int(component["manual_id"])
        orders = get_component_open_orders(
            conn, manual_id, customer, valid_exclude_batch_id
        )
        order_numbers = {order["id"]: order["order_no"] for order in orders}
        allocations = allocate_quantity(component["shipped_quantity"], orders)
        allocations = [
            {
                "order_id": allocation["order_id"],
                "order_no": order_numbers.get(allocation["order_id"], ""),
                "quantity": allocation["quantity"],
            }
            for allocation in allocations
        ]
        no_order_quantity = sum(
            allocation["quantity"]
            for allocation in allocations
            if allocation["order_id"] is None
        )
        available_inventory = max(0, inventory_total_for_manual(conn, manual_id))
        available_inventory += _excluded_inventory_quantity(
            conn, valid_exclude_batch_id, manual_id
        )
        inventory = inventory_result(
            component["shipped_quantity"], available_inventory
        )
        item = {
            "manual_id": manual_id,
            "drawing_no": component["drawing_no"] or "",
            "product_name": component["product_name"] or "",
            "quantity_per_set": int(component["quantity_per_set"]),
            "calculated_quantity": int(component["calculated_quantity"]),
            "shipped_quantity": int(component["shipped_quantity"]),
            "available_inventory": available_inventory,
            "inventory_deducted_quantity": inventory["deducted_quantity"],
            "inventory_shortage_quantity": inventory["shortage_quantity"],
            "no_order_quantity": no_order_quantity,
            "allocations": allocations,
        }
        preview["items"].append(item)
        if no_order_quantity:
            preview["warnings"].append(
                {
                    "code": "no_order",
                    "manual_id": manual_id,
                    "drawing_no": item["drawing_no"],
                    "quantity": no_order_quantity,
                    "message": (
                        f"配件 {item['drawing_no']}：{no_order_quantity} 个没有可匹配订单，"
                        "将作为直接发货保存"
                    ),
                }
            )
        if inventory["shortage_quantity"]:
            shortage = inventory["shortage_quantity"]
            preview["warnings"].append(
                {
                    "code": "inventory_shortage",
                    "manual_id": manual_id,
                    "drawing_no": item["drawing_no"],
                    "quantity": shortage,
                    "message": (
                        f"配件 {item['drawing_no']}：库存不足 {shortage} 个，"
                        "确认后库存将扣到 0"
                    ),
                }
            )
    preview["preview_token"] = preview_token(preview)
    return preview


def unique_shipment_plan_no(conn):
    prefix = f"SP-{datetime.now().strftime('%Y%m%d')}"
    for _ in range(20):
        candidate = f"{prefix}-{secrets.token_hex(2).upper()}"
        exists = conn.execute(
            "SELECT 1 FROM shipment_plans WHERE plan_no = ?",
            (candidate,),
        ).fetchone()
        if not exists:
            return candidate
    return f"{prefix}-{secrets.token_hex(4).upper()}"


def shipment_plan_summary_subquery():
    return """
        SELECT plan_id,
               COUNT(*) AS item_count,
               COALESCE(SUM(planned_quantity), 0) AS total_planned_quantity,
               COALESCE(SUM(shipped_quantity), 0) AS total_shipped_quantity
        FROM shipment_plan_items
        GROUP BY plan_id
    """


def sync_shipment_plan_status(conn, plan_id, updated_at=""):
    summary = conn.execute(
        """
        SELECT
            COALESCE(SUM(planned_quantity), 0) AS planned_total,
            COALESCE(SUM(shipped_quantity), 0) AS shipped_total
        FROM shipment_plan_items
        WHERE plan_id = ?
        """,
        (plan_id,),
    ).fetchone()
    planned_total = int(summary["planned_total"] or 0)
    shipped_total = int(summary["shipped_total"] or 0)
    if shipped_total <= 0:
        status = "待发货"
    elif planned_total and shipped_total >= planned_total:
        status = "已完成"
    else:
        status = "部分发货"
    conn.execute(
        """
        UPDATE shipment_plans
        SET status = ?, updated_at = ?
        WHERE id = ?
        """,
        (status, updated_at or datetime.utcnow().isoformat(timespec="seconds"), plan_id),
    )
    return status


def fetch_open_shipment_plans(conn):
    ensure_inspection_tables(conn)
    plans = [
        dict(row)
        for row in conn.execute(
            f"""
            SELECT shipment_plans.*,
                   COALESCE(summary.item_count, 0) AS item_count,
                   COALESCE(summary.total_planned_quantity, 0) AS total_planned_quantity,
                   COALESCE(summary.total_shipped_quantity, 0) AS total_shipped_quantity
            FROM shipment_plans
            LEFT JOIN ({shipment_plan_summary_subquery()}) AS summary
                ON summary.plan_id = shipment_plans.id
            WHERE shipment_plans.status IN ('待发货', '部分发货')
            ORDER BY shipment_plans.created_at DESC, shipment_plans.id DESC
            """
        ).fetchall()
    ]
    if not plans:
        return []
    plan_ids = [plan["id"] for plan in plans]
    placeholders = ",".join("?" for _ in plan_ids)
    rows = conn.execute(
        f"""
        SELECT shipment_plan_items.*,
               product_orders.order_no,
               product_orders.manual_id,
               product_orders.quantity AS order_quantity,
               product_orders.planned_ship_at AS order_planned_ship_at,
               product_orders.customer AS order_customer,
               manuals.drawing_no,
               manuals.product_name,
               manuals.customer AS product_customer,
               manuals.pack_quantity,
               manuals.pack_carton_size,
               manuals.pack_weight,
               product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS unshipped_quantity
        FROM shipment_plan_items
        JOIN product_orders ON product_orders.id = shipment_plan_items.order_id
        JOIN manuals ON manuals.id = product_orders.manual_id
        LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments ON shipments.order_id = product_orders.id
        WHERE shipment_plan_items.plan_id IN ({placeholders})
        ORDER BY shipment_plan_items.id ASC
        """,
        plan_ids,
    ).fetchall()
    items_by_plan = {plan_id: [] for plan_id in plan_ids}
    for row in rows:
        item = dict(row)
        item["remaining_plan_quantity"] = max(
            0,
            int(item["planned_quantity"] or 0) - int(item["shipped_quantity"] or 0),
        )
        item["max_ship_quantity"] = min(
            int(item["remaining_plan_quantity"] or 0),
            int(item["unshipped_quantity"] or 0),
        )
        items_by_plan.setdefault(item["plan_id"], []).append(item)
    for plan in plans:
        plan["plan_items"] = items_by_plan.get(plan["id"], [])
        plan["inspection_reports"] = []
    report_rows = conn.execute(
        f"""
        SELECT *
        FROM shipment_plan_inspection_reports
        WHERE plan_id IN ({placeholders})
        ORDER BY uploaded_at DESC, id DESC
        """,
        plan_ids,
    ).fetchall()
    plans_by_id = {plan["id"]: plan for plan in plans}
    for report in report_rows:
        plans_by_id.get(report["plan_id"], {}).setdefault("inspection_reports", []).append(report)
    return plans


def parse_positive_number(value):
    if value is None:
        return 0
    match = re.search(r"\d+(?:\.\d+)?", str(value))
    if not match:
        return 0
    try:
        number = float(match.group(0))
    except ValueError:
        return 0
    return number if number > 0 else 0


def pack_quantity_pcs_text(pack_quantity):
    pack_count = parse_positive_number(pack_quantity)
    if pack_count <= 0:
        return ""
    return f"{pack_count:g}pcs"


def product_inventory_code(product):
    if product is None:
        return ""
    sku = (product["sku"] or "").strip() if "sku" in product.keys() else ""
    qr_code = (product["qr_code"] or "").strip() if "qr_code" in product.keys() else ""
    return sku or qr_code


def generated_inventory_code(manual_id):
    return f"P{int(manual_id):06d}"


def ensure_product_inventory_code(conn, manual_id):
    product = conn.execute("SELECT id, sku, qr_code FROM manuals WHERE id = ?", (manual_id,)).fetchone()
    if product is None:
        return ""
    sku = (product["sku"] or "").strip()
    qr_code = (product["qr_code"] or "").strip()
    if sku:
        return sku
    if qr_code:
        return qr_code
    qr_code = generated_inventory_code(manual_id)
    conn.execute("UPDATE manuals SET qr_code = ? WHERE id = ?", (qr_code, manual_id))
    return qr_code


def ensure_all_product_inventory_codes(conn):
    rows = conn.execute("SELECT id FROM manuals WHERE TRIM(COALESCE(sku, '')) = '' AND TRIM(COALESCE(qr_code, '')) = ''").fetchall()
    for row in rows:
        ensure_product_inventory_code(conn, row["id"])


def parse_optional_int(value):
    value = str(value or "").strip()
    if not value:
        return None
    try:
        return int(value)
    except ValueError:
        return None


def unique_inventory_transaction_no(conn):
    prefix = f"INV{datetime.now().strftime('%Y%m%d')}"
    row = conn.execute(
        """
        SELECT transaction_no
        FROM inventory_transactions
        WHERE transaction_no LIKE ?
        ORDER BY transaction_no DESC
        LIMIT 1
        """,
        (f"{prefix}%",),
    ).fetchone()
    if not row:
        return f"{prefix}001"
    suffix = row["transaction_no"].replace(prefix, "", 1)
    try:
        number = int(suffix) + 1
    except ValueError:
        number = 1
    return f"{prefix}{number:03d}"


def get_or_create_default_location(conn):
    row = conn.execute(
        """
        SELECT *
        FROM warehouse_locations
        WHERE enabled = 1
        ORDER BY id ASC
        LIMIT 1
        """
    ).fetchone()
    if row:
        return row
    now = datetime.utcnow().isoformat(timespec="seconds")
    cursor = conn.execute(
        """
        INSERT INTO warehouse_locations (name, code, remark, enabled, created_at, updated_at)
        VALUES ('默认库位', 'DEFAULT', '', 1, ?, ?)
        """,
        (now, now),
    )
    return conn.execute("SELECT * FROM warehouse_locations WHERE id = ?", (cursor.lastrowid,)).fetchone()


def inventory_total_for_manual(conn, manual_id):
    row = conn.execute(
        """
        SELECT COALESCE(SUM(quantity), 0) AS total
        FROM inventory_balances
        WHERE manual_id = ?
        """,
        (manual_id,),
    ).fetchone()
    return int(row["total"] or 0)


def inventory_balance_for_location(conn, manual_id, location_id):
    row = conn.execute(
        """
        SELECT quantity
        FROM inventory_balances
        WHERE manual_id = ? AND location_id = ?
        """,
        (manual_id, location_id),
    ).fetchone()
    return int(row["quantity"] or 0) if row else 0


def update_inventory_balance(conn, manual_id, location_id, delta):
    now = datetime.utcnow().isoformat(timespec="seconds")
    current = inventory_balance_for_location(conn, manual_id, location_id)
    next_quantity = current + int(delta)
    if next_quantity < 0:
        raise ValueError("库存不足，不能直接出库")
    conn.execute(
        """
        INSERT INTO inventory_balances (manual_id, location_id, quantity, updated_at)
        VALUES (?, ?, ?, ?)
        ON CONFLICT(manual_id, location_id)
        DO UPDATE SET quantity = excluded.quantity, updated_at = excluded.updated_at
        """,
        (manual_id, location_id, next_quantity, now),
    )


def update_order_inventory_status(conn, order_id, added_quantity, now):
    if not order_id:
        return
    order = conn.execute(
        "SELECT id, quantity, inventory_received_quantity FROM product_orders WHERE id = ?",
        (order_id,),
    ).fetchone()
    if order is None:
        return
    received = int(order["inventory_received_quantity"] or 0) + int(added_quantity)
    order_quantity = int(order["quantity"] or 0)
    if received <= 0:
        status = "未入库"
    elif received < order_quantity:
        status = "部分入库"
    else:
        status = "已入库"
    conn.execute(
        """
        UPDATE product_orders
        SET inventory_received_quantity = ?, inventory_status = ?, updated_at = ?
        WHERE id = ?
        """,
        (received, status, now, order_id),
    )


def create_inventory_transaction(
    conn,
    transaction_type,
    manual_id,
    quantity,
    from_location_id=None,
    to_location_id=None,
    related_order_type="",
    related_order_id="",
    related_order_no="",
    customer="",
    remark="",
):
    now = datetime.utcnow().isoformat(timespec="seconds")
    quantity = int(quantity)
    if transaction_type == "in":
        if not to_location_id:
            raise ValueError("请选择入库库位")
        update_inventory_balance(conn, manual_id, to_location_id, quantity)
        update_order_inventory_status(conn, parse_optional_int(related_order_id), quantity, now)
    elif transaction_type == "out":
        if not from_location_id:
            raise ValueError("请选择出库库位")
        update_inventory_balance(conn, manual_id, from_location_id, -quantity)
    elif transaction_type == "adjust":
        if not to_location_id:
            raise ValueError("请选择调整库位")
        update_inventory_balance(conn, manual_id, to_location_id, quantity)
    else:
        raise ValueError("暂不支持该库存类型")
    conn.execute(
        """
        INSERT INTO inventory_transactions (
            transaction_no, type, manual_id, quantity, from_location_id, to_location_id,
            related_order_type, related_order_id, related_order_no, customer,
            operator, remark, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            unique_inventory_transaction_no(conn),
            transaction_type,
            manual_id,
            quantity,
            from_location_id,
            to_location_id,
            related_order_type,
            str(related_order_id or ""),
            related_order_no,
            customer,
            current_admin_username(),
            remark,
            now,
        ),
    )


SHIPMENT_INVENTORY_RELATED_TYPE = "销售发货"


def shipment_inventory_related_id(shipment_id):
    return f"shipment:{int(shipment_id)}"


def shipment_inventory_deducted_quantity(conn, shipment_id):
    row = conn.execute(
        """
        SELECT COALESCE(SUM(quantity), 0) AS total
        FROM inventory_transactions
        WHERE type = 'out'
          AND related_order_type = ?
          AND related_order_id = ?
        """,
        (SHIPMENT_INVENTORY_RELATED_TYPE, shipment_inventory_related_id(shipment_id)),
    ).fetchone()
    return int(row["total"] or 0)


def validate_shipment_inventory(conn, requested_items, extra_available_by_manual=None):
    """Validate total stock for order/quantity pairs before creating shipments."""
    requested_by_order = {}
    for order_id, quantity in requested_items:
        requested_by_order[int(order_id)] = requested_by_order.get(int(order_id), 0) + int(quantity)
    if not requested_by_order:
        return

    placeholders = ",".join("?" for _ in requested_by_order)
    rows = conn.execute(
        f"""
        SELECT product_orders.id AS order_id,
               product_orders.manual_id,
               manuals.product_name,
               manuals.drawing_no
        FROM product_orders
        JOIN manuals ON manuals.id = product_orders.manual_id
        WHERE product_orders.id IN ({placeholders})
        """,
        list(requested_by_order),
    ).fetchall()
    rows_by_order = {row["order_id"]: row for row in rows}
    requested_by_manual = {}
    products_by_manual = {}
    for order_id, quantity in requested_by_order.items():
        order = rows_by_order.get(order_id)
        if order is None:
            raise ValueError("关联订单或产品不存在，无法扣减库存")
        manual_id = int(order["manual_id"])
        requested_by_manual[manual_id] = requested_by_manual.get(manual_id, 0) + quantity
        products_by_manual[manual_id] = order

    extra_available_by_manual = extra_available_by_manual or {}
    shortages = []
    for manual_id, required in requested_by_manual.items():
        available = inventory_total_for_manual(conn, manual_id) + int(extra_available_by_manual.get(manual_id, 0))
        if available >= required:
            continue
        product = products_by_manual[manual_id]
        product_label = product["product_name"] or "未命名产品"
        if product["drawing_no"]:
            product_label += f"（{product['drawing_no']}）"
        shortages.append(f"{product_label}：库存 {available}，需要 {required}")
    if shortages:
        raise ValueError(f"库存不足：{'；'.join(shortages)}。请先调整库存后再出库")


def deduct_inventory_for_shipment(conn, shipment_id, order_id, quantity):
    order = conn.execute(
        """
        SELECT product_orders.manual_id,
               product_orders.order_no,
               product_orders.customer AS order_customer,
               manuals.customer AS product_customer,
               manuals.default_location_id
        FROM product_orders
        JOIN manuals ON manuals.id = product_orders.manual_id
        WHERE product_orders.id = ?
        """,
        (order_id,),
    ).fetchone()
    if order is None:
        raise ValueError("关联订单或产品不存在，无法扣减库存")

    remaining = int(quantity)
    locations = conn.execute(
        """
        SELECT location_id, quantity
        FROM inventory_balances
        WHERE manual_id = ? AND quantity > 0
        ORDER BY CASE WHEN location_id = ? THEN 0 ELSE 1 END, location_id ASC
        """,
        (order["manual_id"], order["default_location_id"] or -1),
    ).fetchall()
    if sum(int(row["quantity"] or 0) for row in locations) < remaining:
        raise ValueError("库存不足，请先调整库存后再出库")

    for location in locations:
        if remaining <= 0:
            break
        outbound_quantity = min(remaining, int(location["quantity"] or 0))
        if outbound_quantity <= 0:
            continue
        create_inventory_transaction(
            conn,
            "out",
            order["manual_id"],
            outbound_quantity,
            from_location_id=location["location_id"],
            related_order_type=SHIPMENT_INVENTORY_RELATED_TYPE,
            related_order_id=shipment_inventory_related_id(shipment_id),
            related_order_no=order["order_no"] or "",
            customer=order["order_customer"] or order["product_customer"] or "",
            remark="发货记录自动扣减库存",
        )
        remaining -= outbound_quantity


def reverse_shipment_inventory_deduction(conn, shipment_id):
    transactions = conn.execute(
        """
        SELECT id, manual_id, quantity, from_location_id
        FROM inventory_transactions
        WHERE type = 'out'
          AND related_order_type = ?
          AND related_order_id = ?
        ORDER BY id DESC
        """,
        (SHIPMENT_INVENTORY_RELATED_TYPE, shipment_inventory_related_id(shipment_id)),
    ).fetchall()
    for transaction in transactions:
        update_inventory_balance(
            conn,
            transaction["manual_id"],
            transaction["from_location_id"],
            int(transaction["quantity"] or 0),
        )
    if transactions:
        conn.execute(
            """
            DELETE FROM inventory_transactions
            WHERE type = 'out'
              AND related_order_type = ?
              AND related_order_id = ?
            """,
            (SHIPMENT_INVENTORY_RELATED_TYPE, shipment_inventory_related_id(shipment_id)),
        )


ASSEMBLY_INVENTORY_RELATED_TYPE = "组装发货"


def assembly_inventory_related_id(item_id):
    return f"assembly-item:{parse_positive_int(item_id, '组装发货明细 ID')}"


def deduct_inventory_allow_shortage(
    conn, item_id, manual_id, quantity, customer, reference
):
    item_id = parse_positive_int(item_id, "组装发货明细 ID")
    manual_id = parse_positive_int(manual_id, "配件 ID")
    quantity = parse_positive_int(quantity, "配件实际发货数量")
    manual = conn.execute(
        """
        SELECT id, default_location_id
        FROM manuals
        WHERE id = ?
        """,
        (manual_id,),
    ).fetchone()
    if manual is None:
        raise ValueError("配件不存在，无法扣减库存")

    locations = conn.execute(
        """
        SELECT location_id, quantity
        FROM inventory_balances
        WHERE manual_id = ? AND quantity > 0
        ORDER BY CASE WHEN location_id = ? THEN 0 ELSE 1 END, location_id ASC
        """,
        (manual_id, manual["default_location_id"] or -1),
    ).fetchall()
    remaining = quantity
    deducted = 0
    for location in locations:
        if remaining <= 0:
            break
        outbound_quantity = min(remaining, int(location["quantity"] or 0))
        if outbound_quantity <= 0:
            continue
        create_inventory_transaction(
            conn,
            "out",
            manual_id,
            outbound_quantity,
            from_location_id=location["location_id"],
            related_order_type=ASSEMBLY_INVENTORY_RELATED_TYPE,
            related_order_id=assembly_inventory_related_id(item_id),
            related_order_no=str(reference or "").strip(),
            customer=str(customer or "").strip(),
            remark="组装发货自动扣减库存",
        )
        remaining -= outbound_quantity
        deducted += outbound_quantity
    return deducted


def reverse_assembly_inventory_deduction(conn, item_id):
    related_id = assembly_inventory_related_id(item_id)
    transactions = conn.execute(
        """
        SELECT id, manual_id, quantity, from_location_id
        FROM inventory_transactions
        WHERE type = 'out'
          AND related_order_type = ?
          AND related_order_id = ?
        ORDER BY id DESC
        """,
        (ASSEMBLY_INVENTORY_RELATED_TYPE, related_id),
    ).fetchall()
    restored = 0
    for transaction in transactions:
        quantity = int(transaction["quantity"] or 0)
        update_inventory_balance(
            conn,
            transaction["manual_id"],
            transaction["from_location_id"],
            quantity,
        )
        restored += quantity
    if transactions:
        conn.execute(
            """
            DELETE FROM inventory_transactions
            WHERE type = 'out'
              AND related_order_type = ?
              AND related_order_id = ?
            """,
            (ASSEMBLY_INVENTORY_RELATED_TYPE, related_id),
        )
    return restored


def reverse_assembly_shipment_batch(conn, batch_id):
    batch_id = parse_positive_int(batch_id, "组装发货批次 ID")
    items = conn.execute(
        """
        SELECT id
        FROM assembly_shipment_items
        WHERE batch_id = ?
        ORDER BY id ASC
        """,
        (batch_id,),
    ).fetchall()
    affected_order_ids = {
        int(row["order_id"])
        for row in conn.execute(
            """
            SELECT DISTINCT assembly_shipment_allocations.order_id
            FROM assembly_shipment_allocations
            JOIN assembly_shipment_items
              ON assembly_shipment_items.id = assembly_shipment_allocations.item_id
            WHERE assembly_shipment_items.batch_id = ?
              AND assembly_shipment_allocations.order_id IS NOT NULL
            """,
            (batch_id,),
        ).fetchall()
    }
    for item in items:
        reverse_assembly_inventory_deduction(conn, item["id"])
    conn.execute(
        """
        DELETE FROM assembly_shipment_allocations
        WHERE item_id IN (
            SELECT id FROM assembly_shipment_items WHERE batch_id = ?
        )
        """,
        (batch_id,),
    )
    conn.execute(
        "DELETE FROM assembly_shipment_items WHERE batch_id = ?", (batch_id,)
    )
    return affected_order_ids


def _insert_assembly_shipment_allocation(
    conn, item_id, order_id, quantity, created_at
):
    conn.execute(
        """
        INSERT INTO assembly_shipment_allocations (
            item_id, order_id, quantity, created_at
        ) VALUES (?, ?, ?, ?)
        """,
        (item_id, order_id, quantity, created_at),
    )


def _validated_assembly_preview_header(preview, shipped_at):
    if not isinstance(preview, dict) or not isinstance(preview.get("items"), list):
        raise ValueError("组装发货预览格式无效")
    customer = str(preview.get("customer") or "").strip()
    assembly_drawing_no = str(preview.get("assembly_drawing_no") or "").strip()
    set_quantity = parse_positive_int(preview.get("set_quantity"), "组装件套数")
    shipped_at = str(shipped_at or "").strip()
    if not customer:
        raise ValueError("请选择客户")
    if not assembly_drawing_no:
        raise ValueError("请选择组装件图号")
    if not shipped_at:
        raise ValueError("请填写发货时间")
    if not preview["items"]:
        raise ValueError("该组装件没有有效配件")
    return customer, assembly_drawing_no, set_quantity, shipped_at


def _price_snapshot_values(unit_price_minor, currency, recorded_by, recorded_at):
    has_price = unit_price_minor is not None
    return {
        "unit_price_minor": unit_price_minor,
        "currency": str(currency or "CNY"),
        "price_recorded_by": str(recorded_by or "") if has_price else "",
        "price_recorded_at": str(recorded_at or "") if has_price else "",
    }


def current_product_price_snapshot(conn, manual_id, recorded_by, recorded_at):
    manual = conn.execute(
        """
        SELECT unit_price_minor, currency
        FROM manuals
        WHERE id = ?
        """,
        (parse_positive_int(manual_id, "产品 ID"),),
    ).fetchone()
    if manual is None:
        raise ValueError("产品不存在，无法记录发货价格")
    return _price_snapshot_values(
        manual["unit_price_minor"],
        manual["currency"],
        recorded_by,
        recorded_at,
    )


def _current_product_price_snapshots(conn, manual_ids, recorded_by, recorded_at):
    manual_ids = list(dict.fromkeys(manual_ids))
    if not manual_ids:
        return {}
    placeholders = ",".join("?" for _ in manual_ids)
    rows = conn.execute(
        f"""
        SELECT id, unit_price_minor, currency
        FROM manuals
        WHERE id IN ({placeholders})
        """,
        manual_ids,
    ).fetchall()
    snapshots = {
        int(row["id"]): _price_snapshot_values(
            row["unit_price_minor"],
            row["currency"],
            recorded_by,
            recorded_at,
        )
        for row in rows
    }
    if len(snapshots) != len(manual_ids):
        raise ValueError("产品不存在，无法记录发货价格")
    return snapshots


def _save_assembly_shipment_items(
    conn,
    batch_id,
    preview,
    now,
    recorded_by,
    existing_price_snapshots=None,
):
    customer = str(preview["customer"] or "").strip()
    assembly_drawing_no = str(preview["assembly_drawing_no"] or "").strip()
    manual_ids = [
        parse_positive_int(preview_item.get("manual_id"), "配件 ID")
        for preview_item in preview["items"]
    ]
    existing_price_snapshots = existing_price_snapshots or {}
    introduced_manual_ids = [
        manual_id
        for manual_id in manual_ids
        if manual_id not in existing_price_snapshots
    ]
    current_price_snapshots = _current_product_price_snapshots(
        conn,
        introduced_manual_ids,
        recorded_by,
        now,
    )
    affected_order_ids = set()
    for preview_item, manual_id in zip(preview["items"], manual_ids):
        quantity_per_set = parse_positive_int(
            preview_item.get("quantity_per_set"), "每套用量"
        )
        calculated_quantity = parse_positive_int(
            preview_item.get("calculated_quantity"), "配件计算数量"
        )
        shipped_quantity = parse_positive_int(
            preview_item.get("shipped_quantity"), "配件实际发货数量"
        )
        expected_deducted = max(
            0, int(preview_item.get("inventory_deducted_quantity") or 0)
        )
        expected_shortage = max(
            0, int(preview_item.get("inventory_shortage_quantity") or 0)
        )
        if expected_deducted + expected_shortage != shipped_quantity:
            raise ValueError("组装发货库存预览无效")
        allocations = preview_item.get("allocations")
        if not isinstance(allocations, list) or not allocations:
            raise ValueError("组装发货订单分配无效")
        allocation_total = sum(
            parse_positive_int(allocation.get("quantity"), "订单分配数量")
            for allocation in allocations
        )
        if allocation_total != shipped_quantity:
            raise ValueError("组装发货订单分配数量无效")
        if manual_id in existing_price_snapshots:
            price_snapshot = existing_price_snapshots[manual_id]
        else:
            price_snapshot = current_price_snapshots[manual_id]

        item_id = conn.execute(
            """
            INSERT INTO assembly_shipment_items (
                batch_id, manual_id, drawing_no, product_name,
                quantity_per_set, calculated_quantity, shipped_quantity,
                inventory_deducted_quantity, inventory_shortage_quantity,
                unit_price_minor, currency, price_recorded_by, price_recorded_at,
                created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                batch_id,
                manual_id,
                str(preview_item.get("drawing_no") or ""),
                str(preview_item.get("product_name") or ""),
                quantity_per_set,
                calculated_quantity,
                shipped_quantity,
                expected_deducted,
                expected_shortage,
                price_snapshot["unit_price_minor"],
                price_snapshot["currency"],
                price_snapshot["price_recorded_by"],
                price_snapshot["price_recorded_at"],
                now,
                now,
            ),
        ).lastrowid
        for allocation in allocations:
            order_id = allocation.get("order_id")
            if order_id is not None:
                order_id = parse_positive_int(order_id, "订单 ID")
                affected_order_ids.add(order_id)
            _insert_assembly_shipment_allocation(
                conn,
                item_id,
                order_id,
                parse_positive_int(allocation.get("quantity"), "订单分配数量"),
                now,
            )
        actual_deducted = deduct_inventory_allow_shortage(
            conn,
            item_id,
            manual_id,
            shipped_quantity,
            customer,
            assembly_drawing_no,
        )
        if actual_deducted != expected_deducted:
            raise ValueError("订单或库存状态已变化，请按最新结果重新确认")
    return affected_order_ids


def save_assembly_shipment(conn, preview, shipped_at, logistics_no, created_by):
    customer, assembly_drawing_no, set_quantity, shipped_at = (
        _validated_assembly_preview_header(preview, shipped_at)
    )

    now = datetime.utcnow().isoformat(timespec="seconds")
    batch_id = conn.execute(
        """
        INSERT INTO assembly_shipment_batches (
            customer, assembly_drawing_no, set_quantity, shipped_at,
            logistics_no, created_by, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            customer,
            assembly_drawing_no,
            set_quantity,
            shipped_at,
            str(logistics_no or "").strip(),
            str(created_by or "").strip(),
            now,
            now,
        ),
    ).lastrowid
    affected_order_ids = _save_assembly_shipment_items(
        conn,
        batch_id,
        preview,
        now,
        created_by,
    )
    for order_id in sorted(affected_order_ids):
        sync_order_shipment_summary(conn, order_id, updated_at=now)
    return int(batch_id)


def replace_assembly_shipment(
    conn,
    batch_id,
    preview,
    shipped_at,
    logistics_no,
    old_order_ids,
    existing_price_snapshots=None,
    recorded_by="",
):
    customer, assembly_drawing_no, set_quantity, shipped_at = (
        _validated_assembly_preview_header(preview, shipped_at)
    )
    now = datetime.utcnow().isoformat(timespec="seconds")
    conn.execute(
        """
        UPDATE assembly_shipment_batches
        SET customer = ?, assembly_drawing_no = ?, set_quantity = ?,
            shipped_at = ?, logistics_no = ?, updated_at = ?
        WHERE id = ?
        """,
        (
            customer,
            assembly_drawing_no,
            set_quantity,
            shipped_at,
            str(logistics_no or "").strip(),
            now,
            batch_id,
        ),
    )
    new_order_ids = _save_assembly_shipment_items(
        conn,
        batch_id,
        preview,
        now,
        recorded_by,
        existing_price_snapshots=existing_price_snapshots,
    )
    for order_id in sorted(set(old_order_ids) | set(new_order_ids)):
        sync_order_shipment_summary(conn, order_id, updated_at=now)
    return int(batch_id)


def shipment_plan_box_count(planned_quantity, pack_quantity):
    pack_count = parse_positive_number(pack_quantity)
    if pack_count <= 0:
        return ""
    try:
        planned_count = float(planned_quantity or 0)
    except (TypeError, ValueError):
        planned_count = 0
    if planned_count <= 0:
        return "0"
    return str(math.ceil(planned_count / pack_count))


def fetch_shipment_plan_detail(conn, plan_id):
    plan = conn.execute(
        f"""
        SELECT shipment_plans.*,
               COALESCE(summary.item_count, 0) AS item_count,
               COALESCE(summary.total_planned_quantity, 0) AS total_planned_quantity,
               COALESCE(summary.total_shipped_quantity, 0) AS total_shipped_quantity
        FROM shipment_plans
        LEFT JOIN ({shipment_plan_summary_subquery()}) AS summary
            ON summary.plan_id = shipment_plans.id
        WHERE shipment_plans.id = ?
        """,
        (plan_id,),
    ).fetchone()
    if plan is None:
        return None
    plan = dict(plan)
    rows = conn.execute(
        f"""
        SELECT shipment_plan_items.*,
               product_orders.order_no,
               product_orders.manual_id,
               product_orders.quantity AS order_quantity,
               product_orders.planned_ship_at AS order_planned_ship_at,
               product_orders.customer AS order_customer,
               manuals.drawing_no,
               manuals.product_name,
               manuals.customer AS product_customer,
               manuals.pack_quantity,
               manuals.pack_carton_size,
               manuals.pack_weight,
               product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS unshipped_quantity
        FROM shipment_plan_items
        JOIN product_orders ON product_orders.id = shipment_plan_items.order_id
        JOIN manuals ON manuals.id = product_orders.manual_id
        LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments ON shipments.order_id = product_orders.id
        WHERE shipment_plan_items.plan_id = ?
        ORDER BY shipment_plan_items.id ASC
        """,
        (plan_id,),
    ).fetchall()
    items = []
    total_boxes = 0
    has_box_total = False
    for row in rows:
        item = dict(row)
        item["remaining_plan_quantity"] = max(
            0,
            int(item["planned_quantity"] or 0) - int(item["shipped_quantity"] or 0),
        )
        item["max_ship_quantity"] = min(
            int(item["remaining_plan_quantity"] or 0),
            int(item["unshipped_quantity"] or 0),
        )
        item["box_count"] = shipment_plan_box_count(
            item["planned_quantity"],
            item["pack_quantity"],
        )
        if item["box_count"]:
            has_box_total = True
            total_boxes += int(item["box_count"])
        items.append(item)
    plan["plan_items"] = items
    plan["total_boxes"] = total_boxes if has_box_total else ""
    return plan


def build_shipment_plan_workbook(plan):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "计划发货清单"

    teal_fill = PatternFill("solid", fgColor="166C70")
    pale_fill = PatternFill("solid", fgColor="EAF4F6")
    border_color = "8FA5B0"
    thin_border = Border(
        left=Side(style="thin", color=border_color),
        right=Side(style="thin", color=border_color),
        top=Side(style="thin", color=border_color),
        bottom=Side(style="thin", color=border_color),
    )
    title_font = Font(bold=True, size=20, color="111827")
    label_font = Font(bold=True, color="334155")
    white_font = Font(bold=True, color="FFFFFF")
    value_font = Font(bold=True, color="111827")

    worksheet.merge_cells("A1:J1")
    worksheet["A1"] = "计划发货清单"
    worksheet["A1"].font = title_font
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 32

    worksheet.merge_cells("A2:J2")
    worksheet["A2"] = plan["plan_no"]
    worksheet["A2"].font = Font(bold=True, size=14, color="166C70")
    worksheet["A2"].alignment = Alignment(horizontal="center", vertical="center")

    info_rows = [
        ("客户", plan["customer"] or "未指定客户", "计划发货日期", plan["planned_ship_at"] or "-"),
        ("状态", plan["status"] or "-", "生成时间", plan["created_at"][:10] if plan["created_at"] else "-"),
        ("产品项数", plan["item_count"], "计划数量", plan["total_planned_quantity"]),
        ("已发数量", plan["total_shipped_quantity"], "预计箱数", plan["total_boxes"] or "-"),
    ]
    row_index = 4
    for left_label, left_value, right_label, right_value in info_rows:
        worksheet.cell(row=row_index, column=1, value=left_label)
        worksheet.cell(row=row_index, column=2, value=left_value)
        worksheet.cell(row=row_index, column=5, value=right_label)
        worksheet.cell(row=row_index, column=6, value=right_value)
        for col in (1, 5):
            cell = worksheet.cell(row=row_index, column=col)
            cell.font = label_font
            cell.fill = pale_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for start_col, end_col in ((2, 4), (6, 10)):
            worksheet.merge_cells(
                start_row=row_index,
                start_column=start_col,
                end_row=row_index,
                end_column=end_col,
            )
            cell = worksheet.cell(row=row_index, column=start_col)
            cell.font = value_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")
            for col in range(start_col, end_col + 1):
                worksheet.cell(row=row_index, column=col).border = thin_border
        row_index += 1

    headers = ["序号", "订单号", "图号", "产品名称", "计划数量", "一箱数量", "箱数", "箱子尺寸", "一箱重量", "备注"]
    header_row = 9
    for col_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=header_row, column=col_index, value=header)
        cell.fill = teal_fill
        cell.font = white_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[header_row].height = 24

    for offset, item in enumerate(plan["plan_items"], start=1):
        row = header_row + offset
        values = [
            offset,
            item["order_no"] or "-",
            item["drawing_no"] or "-",
            item["product_name"] or "-",
            item["planned_quantity"],
            item["pack_quantity"] or "-",
            item["box_count"] or "-",
            item["pack_carton_size"] or "-",
            item["pack_weight"] or "-",
            "",
        ]
        for col_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row, column=col_index, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center" if col_index in {1, 5, 6, 7} else "left",
                vertical="center",
                wrap_text=True,
            )
        worksheet.row_dimensions[row].height = 30

    sign_row = header_row + len(plan["plan_items"]) + 3
    worksheet.cell(row=sign_row, column=1, value="制单：")
    worksheet.cell(row=sign_row, column=4, value="仓库确认：")
    worksheet.cell(row=sign_row, column=7, value="审核：")
    for col in (1, 4, 7):
        worksheet.cell(row=sign_row, column=col).font = label_font
        worksheet.cell(row=sign_row, column=col).alignment = Alignment(vertical="center")

    widths = [6, 15, 17, 24, 9, 9, 7, 15, 12, 16]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width

    worksheet.freeze_panes = "A10"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.3
    worksheet.page_margins.right = 0.3
    worksheet.page_margins.top = 0.45
    worksheet.page_margins.bottom = 0.45
    return workbook


def fetch_shipped_orders(
    conn, query="", selected_customer="", shipped_at="", include_prices=False
):
    sql = get_shipped_orders_query(include_prices=include_prices)
    params = []
    conditions = []

    if query:
        like = f"%{query}%"
        conditions.append(
            """
            (
                product_order_shipments.shipped_at LIKE ?
                OR product_orders.order_no LIKE ?
                OR product_orders.customer LIKE ?
                OR manuals.drawing_no LIKE ?
                OR manuals.product_name LIKE ?
                OR manuals.customer LIKE ?
                OR manuals.supplier LIKE ?
            )
            """
        )
        params.extend([like, like, like, like, like, like, like])

    if selected_customer:
        conditions.append(
            """
            (
                product_orders.customer = ?
                OR manuals.customer = ?
            )
            """
        )
        params.extend([selected_customer, selected_customer])

    if shipped_at:
        conditions.append("product_order_shipments.shipped_at LIKE ?")
        params.append(f"{shipped_at}%")

    if conditions:
        sql += " WHERE " + " AND ".join(conditions)

    sql += " ORDER BY product_order_shipments.shipped_at DESC, product_order_shipments.id DESC"
    return conn.execute(sql, params).fetchall()


def _fetch_assembly_shipment_batches_by_ids(
    conn, batch_ids, batches=None, include_prices=False
):
    batch_ids = [int(batch_id) for batch_id in batch_ids]
    if not batch_ids:
        return []
    placeholders = ",".join("?" for _ in batch_ids)
    if batches is None:
        batches = conn.execute(
            f"""
            SELECT id, customer, assembly_drawing_no, set_quantity, shipped_at,
                   logistics_no, created_by, created_at, updated_at
            FROM assembly_shipment_batches
            WHERE id IN ({placeholders})
            ORDER BY shipped_at DESC, id DESC
            """,
            batch_ids,
        ).fetchall()
    batches = [dict(row) for row in batches]
    item_price_columns = ""
    if include_prices:
        item_price_columns = """,
                unit_price_minor, currency, price_recorded_by, price_recorded_at,
                unit_price_minor * shipped_quantity AS line_total_minor
        """
    items = [
        dict(row)
        for row in conn.execute(
            f"""
            SELECT id, batch_id, manual_id, drawing_no, product_name,
                   quantity_per_set, calculated_quantity, shipped_quantity,
                   inventory_deducted_quantity, inventory_shortage_quantity,
                   created_at, updated_at
                   {item_price_columns}
            FROM assembly_shipment_items
            WHERE batch_id IN ({placeholders})
            ORDER BY batch_id ASC, id ASC
            """,
            batch_ids,
        ).fetchall()
    ]
    item_ids = [item["id"] for item in items]
    allocations = []
    if item_ids:
        allocations = [
            dict(row)
            for row in conn.execute(
                f"""
                SELECT assembly_shipment_allocations.id,
                       assembly_shipment_allocations.item_id,
                       assembly_shipment_allocations.order_id,
                       assembly_shipment_allocations.quantity,
                       assembly_shipment_allocations.created_at,
                       COALESCE(product_orders.order_no, '') AS order_no
                FROM assembly_shipment_allocations
                JOIN assembly_shipment_items
                  ON assembly_shipment_items.id = assembly_shipment_allocations.item_id
                LEFT JOIN product_orders
                  ON product_orders.id = assembly_shipment_allocations.order_id
                WHERE assembly_shipment_items.batch_id IN ({placeholders})
                ORDER BY assembly_shipment_allocations.item_id ASC,
                         assembly_shipment_allocations.id ASC
                """,
                batch_ids,
            ).fetchall()
        ]
    images = [
        dict(row)
        for row in conn.execute(
            f"""
            SELECT id, batch_id, filename, original_filename, content_type,
                   file_size, uploaded_at, uploaded_ip, uploaded_user_agent
            FROM assembly_shipment_images
            WHERE batch_id IN ({placeholders})
            ORDER BY batch_id ASC, uploaded_at DESC, id DESC
            """,
            batch_ids,
        ).fetchall()
    ]

    allocations_by_item = {item_id: [] for item_id in item_ids}
    for allocation in allocations:
        allocations_by_item.setdefault(allocation["item_id"], []).append(allocation)
    items_by_batch = {batch_id: [] for batch_id in batch_ids}
    for item in items:
        item["allocations"] = allocations_by_item.get(item["id"], [])
        item["no_order_quantity"] = sum(
            int(allocation["quantity"] or 0)
            for allocation in item["allocations"]
            if allocation["order_id"] is None
        )
        items_by_batch.setdefault(item["batch_id"], []).append(item)
    images_by_batch = {batch_id: [] for batch_id in batch_ids}
    for image in images:
        images_by_batch.setdefault(image["batch_id"], []).append(image)
    for batch in batches:
        batch["items"] = items_by_batch.get(batch["id"], [])
        batch["images"] = images_by_batch.get(batch["id"], [])
    return batches


def fetch_assembly_shipment_batches(
    conn, query="", selected_customer="", shipped_at="", include_prices=False
):
    sql = """
        SELECT id, customer, assembly_drawing_no, set_quantity, shipped_at,
               logistics_no, created_by, created_at, updated_at
        FROM assembly_shipment_batches
    """
    params = []
    conditions = []
    if query:
        like = f"%{query}%"
        conditions.append(
            """
            (
                assembly_shipment_batches.shipped_at LIKE ?
                OR assembly_shipment_batches.customer LIKE ?
                OR assembly_shipment_batches.assembly_drawing_no LIKE ?
                OR EXISTS (
                    SELECT 1
                    FROM assembly_shipment_items
                    WHERE assembly_shipment_items.batch_id = assembly_shipment_batches.id
                      AND (
                          assembly_shipment_items.drawing_no LIKE ?
                          OR assembly_shipment_items.product_name LIKE ?
                      )
                )
            )
            """
        )
        params.extend([like, like, like, like, like])
    if selected_customer:
        conditions.append("assembly_shipment_batches.customer = ?")
        params.append(selected_customer)
    if shipped_at:
        conditions.append("assembly_shipment_batches.shipped_at LIKE ?")
        params.append(f"{shipped_at}%")
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += " ORDER BY assembly_shipment_batches.shipped_at DESC, assembly_shipment_batches.id DESC"
    batch_rows = conn.execute(sql, params).fetchall()
    batch_ids = [row["id"] for row in batch_rows]
    return _fetch_assembly_shipment_batches_by_ids(
        conn,
        batch_ids,
        batches=batch_rows,
        include_prices=include_prices,
    )


def fetch_shipped_orders_by_ids(conn, shipment_ids, include_prices=False):
    if not shipment_ids:
        return []
    placeholders = ",".join("?" for _ in shipment_ids)
    sql = f"""
        {get_shipped_orders_query(include_prices=include_prices)}
        WHERE product_order_shipments.id IN ({placeholders})
        ORDER BY product_order_shipments.shipped_at DESC, product_order_shipments.id DESC
    """
    return conn.execute(sql, shipment_ids).fetchall()


def fetch_shipment_by_id(conn, shipment_id, include_prices=False):
    return conn.execute(
        f"""
        {get_shipped_orders_query(include_prices=include_prices)}
        WHERE product_order_shipments.id = ?
        """,
        (shipment_id,),
    ).fetchone()


def fetch_shipment_by_token(conn, token):
    return conn.execute(
        f"""
        {get_shipped_orders_query()}
        WHERE product_order_shipments.signature_token = ?
        """,
        (token,),
    ).fetchone()


def fetch_shipment_by_photo_token(conn, token):
    return conn.execute(
        f"""
        {get_shipped_orders_query()}
        WHERE product_order_shipments.photo_upload_token = ?
        """,
        (token,),
    ).fetchone()


FINANCE_SOURCE_TYPES = frozenset({"ordinary", "assembly_item"})
FINANCE_STATUS_LABELS = {
    "pending": "待开票",
    "invoiced": "已开票",
    "paid": "已收款",
    "void": "已作废",
}


def finance_claim_key(source_type, source_id):
    if source_type not in FINANCE_SOURCE_TYPES:
        raise ValueError("发货记录来源无效")
    try:
        normalized_id = int(source_id)
    except (TypeError, ValueError) as error:
        raise ValueError("发货记录来源无效") from error
    if normalized_id <= 0:
        raise ValueError("发货记录来源无效")
    return f"{source_type}:{normalized_id}"


def parse_finance_source_refs(raw_refs):
    refs = []
    for raw_ref in raw_refs or ():
        match = re.fullmatch(r"(ordinary|assembly_item):([1-9][0-9]*)", str(raw_ref))
        if not match:
            raise ValueError("请选择有效的发货记录")
        refs.append((match.group(1), int(match.group(2))))
    if not refs:
        raise ValueError("请至少选择一条发货记录")
    if len(set(refs)) != len(refs):
        raise ValueError("不能重复选择同一发货记录")
    return refs


def _finance_source_select(source_type):
    if source_type == "ordinary":
        return """
            SELECT 'ordinary' AS source_type,
                   product_order_shipments.id AS source_id,
                   COALESCE(
                       NULLIF(TRIM(product_orders.customer), ''),
                       TRIM(manuals.customer)
                   ) AS customer_name,
                   product_orders.order_no AS order_no,
                   NULL AS assembly_batch_id,
                   '' AS assembly_drawing_no,
                   product_order_shipments.shipped_at AS shipped_at,
                   manuals.drawing_no AS drawing_no,
                   manuals.product_name AS product_name,
                   product_order_shipments.shipped_quantity AS quantity,
                   product_order_shipments.unit_price_minor AS unit_price_minor,
                   product_order_shipments.currency AS currency,
                   product_order_shipments.unit_price_minor
                       * product_order_shipments.shipped_quantity AS line_total_minor
            FROM product_order_shipments
            JOIN product_orders
              ON product_orders.id = product_order_shipments.order_id
            JOIN manuals
              ON manuals.id = product_orders.manual_id
        """
    if source_type == "assembly_item":
        return """
            SELECT 'assembly_item' AS source_type,
                   assembly_shipment_items.id AS source_id,
                   TRIM(assembly_shipment_batches.customer) AS customer_name,
                   COALESCE((
                       SELECT GROUP_CONCAT(product_orders.order_no, ' / ')
                       FROM assembly_shipment_allocations
                       JOIN product_orders
                         ON product_orders.id = assembly_shipment_allocations.order_id
                       WHERE assembly_shipment_allocations.item_id = assembly_shipment_items.id
                   ), '') AS order_no,
                   assembly_shipment_batches.id AS assembly_batch_id,
                   assembly_shipment_batches.assembly_drawing_no AS assembly_drawing_no,
                   assembly_shipment_batches.shipped_at AS shipped_at,
                   assembly_shipment_items.drawing_no AS drawing_no,
                   assembly_shipment_items.product_name AS product_name,
                   assembly_shipment_items.shipped_quantity AS quantity,
                   assembly_shipment_items.unit_price_minor AS unit_price_minor,
                   assembly_shipment_items.currency AS currency,
                   assembly_shipment_items.unit_price_minor
                       * assembly_shipment_items.shipped_quantity AS line_total_minor
            FROM assembly_shipment_items
            JOIN assembly_shipment_batches
              ON assembly_shipment_batches.id = assembly_shipment_items.batch_id
        """
    raise ValueError("发货记录来源无效")


def _fetch_finance_source(conn, source_type, source_id):
    claim_key = finance_claim_key(source_type, source_id)
    row = conn.execute(
        f"""
        SELECT finance_source.*,
               finance_invoice_items.invoice_id AS claimed_invoice_id
        FROM ({_finance_source_select(source_type)}) AS finance_source
        LEFT JOIN finance_invoice_items
          ON finance_invoice_items.active_claim_key = ?
        WHERE finance_source.source_id = ?
        """,
        (claim_key, int(source_id)),
    ).fetchone()
    return dict(row) if row else None


def fetch_available_finance_sources(conn, customer_name, currency=""):
    customer_name = str(customer_name or "").strip()
    normalized_currency = ""
    if currency:
        normalized_currency = normalize_currency(currency)
    sources = []
    for source_type in ("ordinary", "assembly_item"):
        params = [customer_name]
        currency_condition = ""
        if normalized_currency:
            currency_condition = " AND finance_source.currency = ?"
            params.append(normalized_currency)
        rows = conn.execute(
            f"""
            SELECT finance_source.*
            FROM ({_finance_source_select(source_type)}) AS finance_source
            WHERE finance_source.customer_name = ?
              AND finance_source.unit_price_minor IS NOT NULL
              {currency_condition}
              AND NOT EXISTS (
                  SELECT 1
                  FROM finance_invoice_items
                  WHERE active_claim_key = (
                      finance_source.source_type || ':' || finance_source.source_id
                  )
              )
            ORDER BY finance_source.shipped_at DESC, finance_source.source_id DESC
            """,
            params,
        ).fetchall()
        sources.extend(dict(row) for row in rows)
    sources.sort(
        key=lambda row: (
            str(row["shipped_at"] or ""),
            row["source_type"],
            int(row["source_id"]),
        ),
        reverse=True,
    )
    return sources


def _normalize_finance_source_refs(source_refs, require_nonempty=True):
    refs = []
    for reference in source_refs or ():
        if isinstance(reference, dict):
            source_type = reference.get("source_type")
            source_id = reference.get("source_id")
        else:
            try:
                source_type, source_id = reference
            except (TypeError, ValueError) as error:
                raise ValueError("发货记录来源无效") from error
        claim_key = finance_claim_key(str(source_type or ""), source_id)
        refs.append((str(source_type), int(source_id), claim_key))
    if require_nonempty and not refs:
        raise ValueError("请至少选择一条发货记录")
    if len({claim_key for _, _, claim_key in refs}) != len(refs):
        raise ValueError("不能重复选择同一发货记录")
    return refs


@contextmanager
def _finance_write_scope(conn, operation):
    if not conn.in_transaction:
        conn.execute("BEGIN IMMEDIATE")
    savepoint = f"finance_{operation}_{uuid.uuid4().hex}"
    conn.execute(f"SAVEPOINT {savepoint}")
    try:
        yield
    except Exception:
        conn.execute(f"ROLLBACK TO SAVEPOINT {savepoint}")
        conn.execute(f"RELEASE SAVEPOINT {savepoint}")
        raise
    else:
        conn.execute(f"RELEASE SAVEPOINT {savepoint}")


def _load_finance_sources(
    conn,
    source_refs,
    *,
    expected_customer="",
    expected_currency="",
    allowed_invoice_id=None,
    require_nonempty=True,
):
    refs = _normalize_finance_source_refs(
        source_refs,
        require_nonempty=require_nonempty,
    )
    sources = []
    for source_type, source_id, _claim_key in refs:
        source = _fetch_finance_source(conn, source_type, source_id)
        if source is None:
            raise ValueError("发货记录不存在")
        if source["unit_price_minor"] is None:
            raise ValueError("该发货记录尚未记录价格，不能开票")
        try:
            source["currency"] = normalize_currency(source["currency"])
        except ValueError as error:
            raise ValueError("价格格式或币种不正确") from error
        claimed_invoice_id = source.pop("claimed_invoice_id")
        if claimed_invoice_id is not None and int(claimed_invoice_id) != int(
            allowed_invoice_id or 0
        ):
            raise ValueError("该发货记录已加入其他开票单")
        sources.append(source)

    customer_names = {str(source["customer_name"] or "").strip() for source in sources}
    currencies = {source["currency"] for source in sources}
    if (
        len(customer_names) > 1
        or len(currencies) > 1
        or (sources and expected_customer and customer_names != {expected_customer})
        or (sources and expected_currency and currencies != {expected_currency})
    ):
        raise ValueError("只能合并同一客户、同一币种的发货记录")
    return refs, sources


def _insert_finance_invoice_item(conn, invoice_id, source, created_at):
    claim_key = finance_claim_key(source["source_type"], source["source_id"])
    try:
        conn.execute(
            """
            INSERT INTO finance_invoice_items (
                invoice_id, source_type, source_id, active_claim_key,
                order_no, assembly_batch_id, assembly_drawing_no,
                shipped_at, drawing_no, product_name, quantity,
                unit_price_minor, currency, line_total_minor, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                invoice_id,
                source["source_type"],
                source["source_id"],
                claim_key,
                source["order_no"] or "",
                source["assembly_batch_id"],
                source["assembly_drawing_no"] or "",
                source["shipped_at"],
                source["drawing_no"] or "",
                source["product_name"] or "",
                int(source["quantity"]),
                int(source["unit_price_minor"]),
                source["currency"],
                int(source["unit_price_minor"]) * int(source["quantity"]),
                created_at,
            ),
        )
    except sqlite3.IntegrityError as error:
        if "active_claim_key" in str(error) or "UNIQUE constraint failed" in str(error):
            raise ValueError("该发货记录已加入其他开票单") from error
        raise


def _recalculate_finance_invoice_total(conn, invoice_id, updated_by, updated_at):
    total = conn.execute(
        """
        SELECT COALESCE(SUM(line_total_minor), 0) AS total_minor
        FROM finance_invoice_items
        WHERE invoice_id = ?
        """,
        (invoice_id,),
    ).fetchone()["total_minor"]
    conn.execute(
        """
        UPDATE finance_invoices
        SET total_minor = ?, updated_by = ?, updated_at = ?
        WHERE id = ?
        """,
        (int(total or 0), str(updated_by or ""), updated_at, invoice_id),
    )


def create_finance_invoice(conn, customer_id, source_refs, created_by):
    with _finance_write_scope(conn, "create"):
        customer = conn.execute(
            "SELECT * FROM customers WHERE id = ?",
            (customer_id,),
        ).fetchone()
        if customer is None:
            raise ValueError("客户不存在")
        snapshot = customer_invoice_snapshot(customer)
        _refs, sources = _load_finance_sources(
            conn,
            source_refs,
            expected_customer=snapshot["customer_name"],
        )
        currency = sources[0]["currency"]
        now = datetime.utcnow().isoformat(timespec="seconds")
        invoice_id = conn.execute(
            """
            INSERT INTO finance_invoices (
                customer_id, customer_name, invoice_title, tax_id,
                registered_address, registered_phone, bank_name, bank_account,
                invoice_email, currency, total_minor, status,
                created_by, updated_by, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 'pending', ?, ?, ?, ?)
            """,
            (
                customer_id,
                snapshot["customer_name"],
                snapshot["invoice_title"],
                snapshot["tax_id"],
                snapshot["registered_address"],
                snapshot["registered_phone"],
                snapshot["bank_name"],
                snapshot["bank_account"],
                snapshot["invoice_email"],
                currency,
                str(created_by or ""),
                str(created_by or ""),
                now,
                now,
            ),
        ).lastrowid
        for source in sources:
            _insert_finance_invoice_item(conn, invoice_id, source, now)
        _recalculate_finance_invoice_total(conn, invoice_id, created_by, now)
    return invoice_id


def replace_pending_invoice_items(
    conn,
    invoice_id,
    source_refs,
    updated_by,
):
    with _finance_write_scope(conn, "replace"):
        invoice = conn.execute(
            "SELECT * FROM finance_invoices WHERE id = ?",
            (invoice_id,),
        ).fetchone()
        if invoice is None:
            raise ValueError("开票单不存在")
        if invoice["status"] != "pending":
            raise ValueError("只有待开票单可以修改明细")
        refs, sources = _load_finance_sources(
            conn,
            source_refs,
            expected_customer=invoice["customer_name"],
            expected_currency=invoice["currency"],
            allowed_invoice_id=invoice_id,
            require_nonempty=False,
        )
        requested = {(source_type, source_id) for source_type, source_id, _ in refs}
        sources_by_ref = {
            (source["source_type"], int(source["source_id"])): source
            for source in sources
        }
        existing_rows = conn.execute(
            """
            SELECT id, source_type, source_id
            FROM finance_invoice_items
            WHERE invoice_id = ?
            """,
            (invoice_id,),
        ).fetchall()
        existing = {
            (row["source_type"], int(row["source_id"])): row["id"]
            for row in existing_rows
        }
        now = datetime.utcnow().isoformat(timespec="seconds")
        for reference in requested - set(existing):
            _insert_finance_invoice_item(
                conn,
                invoice_id,
                sources_by_ref[reference],
                now,
            )
        removed_ids = [
            item_id for reference, item_id in existing.items() if reference not in requested
        ]
        if removed_ids:
            placeholders = ",".join("?" for _ in removed_ids)
            conn.execute(
                f"DELETE FROM finance_invoice_items WHERE id IN ({placeholders})",
                removed_ids,
            )
        _recalculate_finance_invoice_total(conn, invoice_id, updated_by, now)


def finance_source_is_claimed(conn, source_type, source_id):
    claim_key = finance_claim_key(source_type, source_id)
    return (
        conn.execute(
            """
            SELECT 1
            FROM finance_invoice_items
            WHERE active_claim_key = ?
            LIMIT 1
            """,
            (claim_key,),
        ).fetchone()
        is not None
    )


def mark_finance_invoice_invoiced(
    conn,
    invoice_id,
    invoice_no,
    invoice_date,
    finance_remark,
    updated_by,
):
    invoice_no = str(invoice_no or "").strip()
    invoice_date = str(invoice_date or "").strip()
    if not invoice_no or not invoice_date:
        raise ValueError("请填写发票号码和开票日期")
    with _finance_write_scope(conn, "invoice"):
        invoice = conn.execute(
            "SELECT status FROM finance_invoices WHERE id = ?",
            (invoice_id,),
        ).fetchone()
        if invoice is None:
            raise ValueError("开票单不存在")
        if invoice["status"] != "pending":
            raise ValueError("只有待开票单可以登记开票")
        item_count = conn.execute(
            "SELECT COUNT(*) AS c FROM finance_invoice_items WHERE invoice_id = ?",
            (invoice_id,),
        ).fetchone()["c"]
        if not item_count:
            raise ValueError("请至少选择一条发货记录")
        now = datetime.utcnow().isoformat(timespec="seconds")
        conn.execute(
            """
            UPDATE finance_invoices
            SET status = 'invoiced', invoice_no = ?, invoice_date = ?,
                finance_remark = ?, updated_by = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                invoice_no,
                invoice_date,
                str(finance_remark or "").strip(),
                str(updated_by or ""),
                now,
                invoice_id,
            ),
        )


def mark_finance_invoice_paid(conn, invoice_id, payment_date, updated_by):
    payment_date = str(payment_date or "").strip()
    if not payment_date:
        raise ValueError("请填写收款日期")
    with _finance_write_scope(conn, "payment"):
        invoice = conn.execute(
            "SELECT status FROM finance_invoices WHERE id = ?",
            (invoice_id,),
        ).fetchone()
        if invoice is None:
            raise ValueError("开票单不存在")
        if invoice["status"] != "invoiced":
            raise ValueError("只有已开票记录可以登记收款")
        now = datetime.utcnow().isoformat(timespec="seconds")
        conn.execute(
            """
            UPDATE finance_invoices
            SET status = 'paid', payment_date = ?, updated_by = ?, updated_at = ?
            WHERE id = ?
            """,
            (payment_date, str(updated_by or ""), now, invoice_id),
        )


def reopen_finance_invoice_payment(conn, invoice_id, updated_by):
    with _finance_write_scope(conn, "reopen"):
        invoice = conn.execute(
            "SELECT status FROM finance_invoices WHERE id = ?",
            (invoice_id,),
        ).fetchone()
        if invoice is None:
            raise ValueError("开票单不存在")
        if invoice["status"] != "paid":
            raise ValueError("只有已收款记录可以撤销收款")
        now = datetime.utcnow().isoformat(timespec="seconds")
        conn.execute(
            """
            UPDATE finance_invoices
            SET status = 'invoiced', payment_date = '',
                updated_by = ?, updated_at = ?
            WHERE id = ?
            """,
            (str(updated_by or ""), now, invoice_id),
        )


def void_finance_invoice(conn, invoice_id, updated_by):
    with _finance_write_scope(conn, "void"):
        invoice = conn.execute(
            "SELECT status FROM finance_invoices WHERE id = ?",
            (invoice_id,),
        ).fetchone()
        if invoice is None:
            raise ValueError("开票单不存在")
        if invoice["status"] == "paid":
            raise ValueError("已收款记录需先撤销收款后才能作废")
        if invoice["status"] != "invoiced":
            raise ValueError("只有已开票记录可以作废")
        now = datetime.utcnow().isoformat(timespec="seconds")
        conn.execute(
            """
            UPDATE finance_invoice_items
            SET active_claim_key = NULL
            WHERE invoice_id = ?
            """,
            (invoice_id,),
        )
        conn.execute(
            """
            UPDATE finance_invoices
            SET status = 'void', voided_by = ?, voided_at = ?,
                updated_by = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                str(updated_by or ""),
                now,
                str(updated_by or ""),
                now,
                invoice_id,
            ),
        )


def delete_pending_finance_invoice(conn, invoice_id):
    with _finance_write_scope(conn, "delete"):
        invoice = conn.execute(
            "SELECT status FROM finance_invoices WHERE id = ?",
            (invoice_id,),
        ).fetchone()
        if invoice is None:
            raise ValueError("开票单不存在")
        if invoice["status"] != "pending":
            raise ValueError("只有待开票单可以删除")
        # Foreign keys are intentionally disabled for this SMB-hosted SQLite app,
        # so remove children explicitly instead of relying on ON DELETE CASCADE.
        conn.execute(
            "DELETE FROM finance_invoice_items WHERE invoice_id = ?",
            (invoice_id,),
        )
        conn.execute("DELETE FROM finance_invoices WHERE id = ?", (invoice_id,))


def fetch_finance_invoice(conn, invoice_id):
    return conn.execute(
        """
        SELECT finance_invoices.*,
               (
                   SELECT COUNT(*)
                   FROM finance_invoice_items
                   WHERE finance_invoice_items.invoice_id = finance_invoices.id
               ) AS item_count
        FROM finance_invoices
        WHERE finance_invoices.id = ?
        """,
        (invoice_id,),
    ).fetchone()


def fetch_finance_invoice_items(conn, invoice_id):
    return conn.execute(
        """
        SELECT id, invoice_id, source_type, source_id, active_claim_key,
               order_no, assembly_batch_id, assembly_drawing_no, shipped_at,
               drawing_no, product_name, quantity, unit_price_minor, currency,
               line_total_minor, created_at
        FROM finance_invoice_items
        WHERE invoice_id = ?
        ORDER BY shipped_at ASC, id ASC
        """,
        (invoice_id,),
    ).fetchall()


def get_shipment_images(conn, shipment_ids):
    if not shipment_ids:
        return {}
    placeholders = ",".join("?" for _ in shipment_ids)
    rows = conn.execute(
        f"""
        SELECT *
        FROM product_order_shipment_images
        WHERE shipment_id IN ({placeholders})
        ORDER BY uploaded_at DESC, id DESC
        """,
        shipment_ids,
    ).fetchall()
    images_by_shipment = {shipment_id: [] for shipment_id in shipment_ids}
    for row in rows:
        images_by_shipment.setdefault(row["shipment_id"], []).append(row)
    return images_by_shipment


def attach_shipment_images(conn, shipments):
    shipment_dicts = [dict(shipment) for shipment in shipments]
    images_by_shipment = get_shipment_images(conn, [shipment["id"] for shipment in shipment_dicts])
    for shipment in shipment_dicts:
        shipment["images"] = images_by_shipment.get(shipment["id"], [])
    return shipment_dicts


def get_purchase_followups_pending_purchase_at_count(conn):
    return conn.execute(
        """
        SELECT COUNT(*) AS c
        FROM purchase_followups
        WHERE TRIM(COALESCE(purchased_at, '')) = ''
        """
    ).fetchone()["c"]


def purchase_followup_sse_serializer():
    return URLSafeSerializer(app.config["SECRET_KEY"], salt="purchase-followups-sse")


def purchase_followup_pending_purchase_at_sse_token(scope="pending_purchase_at_json"):
    return purchase_followup_sse_serializer().dumps({"scope": scope})


def verify_purchase_followup_sse_token(token):
    try:
        payload = purchase_followup_sse_serializer().loads(token)
    except BadSignature:
        return ""
    return payload.get("scope", "")


def purchase_followup_pending_purchase_at_sse_url():
    token = purchase_followup_pending_purchase_at_sse_token("pending_purchase_at_json")
    return f"{configured_base_url()}{url_for('purchase_followups_pending_purchase_at_sse', token=token)}"


def purchase_followup_pending_purchase_at_plain_sse_url():
    token = purchase_followup_pending_purchase_at_sse_token("pending_purchase_at_plain")
    return f"{configured_base_url()}{url_for('purchase_followups_pending_purchase_at_plain_sse', token=token)}"


def sync_order_shipment_summary(conn, order_id, updated_at=""):
    totals = conn.execute(
        f"""
        SELECT shipped_total, last_shipped_at
        FROM ({order_shipment_summary_subquery(include_assembly=True, conn=conn)})
        WHERE order_id = ?
        """,
        (order_id,),
    ).fetchone()
    now = updated_at or datetime.utcnow().isoformat(timespec="seconds")
    conn.execute(
        """
        UPDATE product_orders
        SET shipped_quantity = ?, shipped_at = ?, updated_at = ?
        WHERE id = ?
        """,
        (
            totals["shipped_total"] if totals else 0,
            totals["last_shipped_at"] if totals else "",
            now,
            order_id,
        ),
    )


def reset_shipment_signature_state(conn, shipment_id):
    shipment = fetch_shipment_by_id(conn, shipment_id)
    if shipment is None:
        return None
    old_signature = shipment["signature_image"] or ""
    conn.execute(
        """
        UPDATE product_order_shipments
        SET signature_token = ?,
            signature_status = '未签收',
            signature_image = '',
            signed_at = '',
            signed_ip = '',
            signed_user_agent = '',
            signature_expires_at = ?
        WHERE id = ?
        """,
        (unique_signature_token(conn), signature_expires_at(), shipment_id),
    )
    if old_signature:
        (SIGNATURES_DIR / secure_filename(old_signature)).unlink(missing_ok=True)
    return fetch_shipment_by_id(conn, shipment_id)


def delete_shipment_assets(conn, shipment_id, signature_image=""):
    images = get_shipment_images(conn, [shipment_id]).get(shipment_id, [])
    conn.execute("DELETE FROM product_order_shipment_images WHERE shipment_id = ?", (shipment_id,))
    for image in images:
        (SHIPMENT_IMAGES_DIR / secure_filename(image["filename"])).unlink(missing_ok=True)
    if signature_image:
        (SIGNATURES_DIR / secure_filename(signature_image)).unlink(missing_ok=True)


def configured_base_url():
    return os.getenv("BASE_URL", "http://127.0.0.1:8006").rstrip("/")


def shipment_signature_url(shipment):
    token = shipment["signature_token"] if shipment else ""
    if not token:
        return ""
    return f"{configured_base_url()}{url_for('sign_shipment', token=token)}"


def shipment_photo_upload_url(shipment):
    token = shipment["photo_upload_token"] if shipment else ""
    if not token:
        return ""
    return f"{configured_base_url()}{url_for('upload_shipment_photos', token=token)}"


def signature_expires_at(now=None):
    current = now or datetime.utcnow()
    return (current + SIGNATURE_LINK_TTL).isoformat(timespec="seconds")


def signature_link_expired(shipment, now=None):
    expires_at = shipment["signature_expires_at"] if shipment else ""
    if not expires_at:
        return False
    try:
        expires_at_value = datetime.fromisoformat(expires_at)
    except ValueError:
        return False
    return (now or datetime.utcnow()) > expires_at_value


def unique_signature_token(conn):
    while True:
        token = secrets.token_urlsafe(32)
        exists = conn.execute(
            """
            SELECT 1
            FROM product_order_shipments
            WHERE signature_token = ?
            LIMIT 1
            """,
            (token,),
        ).fetchone()
        if not exists:
            return token


def unique_shipment_photo_token(conn):
    while True:
        token = secrets.token_urlsafe(32)
        exists = conn.execute(
            """
            SELECT 1
            FROM product_order_shipments
            WHERE photo_upload_token = ?
            LIMIT 1
            """,
            (token,),
        ).fetchone()
        if not exists:
            return token


def safe_signature_filename(shipment_id, token):
    safe_token = re.sub(r"[^A-Za-z0-9_-]", "", token)[:48]
    timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    return f"{shipment_id}_{safe_token}_{timestamp}.png"


def decode_signature_image(data_url):
    prefix = "data:image/png;base64,"
    if not data_url.startswith(prefix):
        raise ValueError("签名图片格式不正确")
    try:
        image_bytes = base64.b64decode(data_url[len(prefix):], validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ValueError("签名图片无法解析") from exc
    if not image_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ValueError("签名图片必须是 PNG 格式")
    if len(image_bytes) > 2 * 1024 * 1024:
        raise ValueError("签名图片过大")
    return image_bytes


def client_ip():
    forwarded_for = request.headers.get("X-Forwarded-For", "")
    if forwarded_for:
        return forwarded_for.split(",", 1)[0].strip()
    return request.remote_addr or ""


def smtp_configured():
    return bool(os.getenv("SMTP_HOST") and (os.getenv("SMTP_FROM") or os.getenv("SMTP_USERNAME")))


def send_delivery_note_email(recipient, shipped_orders):
    if not recipient:
        return False, "未填写客户邮箱"
    if not smtp_configured():
        return False, "未配置 SMTP 邮件发送信息"
    if not shipped_orders:
        return False, "没有可发送的发货记录"

    customer_name = shipped_orders[0]["order_customer"] or shipped_orders[0]["product_customer"] or ""
    ship_date = shipped_orders[0]["shipped_at"] or datetime.now().strftime("%Y-%m-%d")
    order_no = shipped_orders[0]["order_no"] or ""
    subject_prefix = os.getenv("SMTP_SUBJECT_PREFIX", "送货单")
    subject = f"{subject_prefix} {order_no} {ship_date}".strip()
    body = "\n".join([
        "客户您好：",
        "",
        "附件为本次发货送货单，请查收。",
        "",
        "此邮件由杰德机械（生产管理）自动发送。",
    ])

    pdf_buffer = build_shipped_orders_pdf(
        shipped_orders,
        query="",
        selected_customer=customer_name,
        shipped_at=ship_date,
    )

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = os.getenv("SMTP_FROM") or os.getenv("SMTP_USERNAME")
    message["To"] = recipient
    message["Cc"] = os.getenv("DELIVERY_NOTE_CC", "derek@nbliwan.com")
    message.set_content(body)
    message.add_attachment(
        pdf_buffer.getvalue(),
        maintype="application",
        subtype="pdf",
        filename=f"送货单-{order_no or ship_date}.pdf",
    )

    host = os.getenv("SMTP_HOST")
    port = int(os.getenv("SMTP_PORT", "587"))
    username = os.getenv("SMTP_USERNAME", "")
    password = os.getenv("SMTP_PASSWORD", "")
    use_tls = os.getenv("SMTP_USE_TLS", "1").strip().lower() not in {"0", "false", "no", "off"}
    use_ssl = os.getenv("SMTP_USE_SSL", "").strip().lower() in {"1", "true", "yes", "on"} or port == 465

    try:
        smtp_class = smtplib.SMTP_SSL if use_ssl else smtplib.SMTP
        with smtp_class(host, port, timeout=20) as smtp:
            if use_tls and not use_ssl:
                smtp.starttls()
            if username:
                smtp.login(username, password)
            smtp.send_message(message)
    except Exception as exc:
        app.logger.exception("Failed to send delivery note email")
        return False, f"邮件发送失败：{exc}"
    return True, ""


def send_delivery_note_for_shipment_ids(conn, shipment_ids):
    if not shipment_ids:
        return []
    results = []
    for shipment_id in shipment_ids:
        shipped_orders = fetch_shipped_orders_by_ids(conn, [shipment_id])
        if not shipped_orders:
            continue
        shipment = shipped_orders[0]
        order = conn.execute(
            """
            SELECT customer_email
            FROM product_orders
            WHERE id = ?
            """,
            (shipment["order_id"],),
        ).fetchone()
        recipient = order["customer_email"] if order else ""
        success, error = send_delivery_note_email(recipient, shipped_orders)
        if success:
            sent_at = datetime.utcnow().isoformat(timespec="seconds")
            conn.execute(
                """
                UPDATE product_order_shipments
                SET email_sent_at = ?
                WHERE id = ?
                """,
                (sent_at, shipment_id),
            )
            results.append((shipment_id, True, recipient))
        else:
            results.append((shipment_id, False, error))
    return results


def get_pdf_font_name():
    global PDF_FONT_NAME
    if PDF_FONT_NAME:
        return PDF_FONT_NAME
    for index, font_path in enumerate(PDF_FONT_CANDIDATES):
        if not font_path.exists():
            continue
        font_name = f"ManualWeb2Unicode{index}"
        try:
            pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
        except Exception:
            continue
        PDF_FONT_NAME = font_name
        return PDF_FONT_NAME
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        PDF_FONT_NAME = "STSong-Light"
        return PDF_FONT_NAME
    except Exception:
        pass
    PDF_FONT_NAME = "Helvetica"
    return PDF_FONT_NAME


def build_shipped_orders_workbook(shipped_orders, query, selected_customer, shipped_at):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "发货记录"
    worksheet.freeze_panes = "A4"

    title = "发货记录"
    filters = []
    if query:
        filters.append(f"搜索：{query}")
    if selected_customer:
        filters.append(f"客户：{selected_customer}")
    if shipped_at:
        filters.append(f"发货时间：{shipped_at}")
    filter_text = "；".join(filters) if filters else "全部记录"

    worksheet["A1"] = title
    worksheet["A2"] = filter_text
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet["A2"].font = Font(italic=True, color="666666")
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
    headers = ["发货时间", "发货产品", "发货数量", "订单号", "客户名称", "供应商", "产品图号", "订单下单时间", "订单计划发货时间"]
    for col_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=3, column=col_index)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F81BD")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_index = 4
    if shipped_orders:
        for shipment in shipped_orders:
            worksheet.cell(row=row_index, column=1, value=shipment["shipped_at"])
            worksheet.cell(row=row_index, column=2, value=shipment["product_name"])
            worksheet.cell(row=row_index, column=3, value=shipment["shipped_quantity"])
            worksheet.cell(row=row_index, column=4, value=shipment["order_no"])
            worksheet.cell(row=row_index, column=5, value=shipment["order_customer"] or shipment["product_customer"] or "")
            worksheet.cell(row=row_index, column=6, value=shipment["supplier"] or "")
            worksheet.cell(row=row_index, column=7, value=shipment["drawing_no"] or "")
            worksheet.cell(row=row_index, column=8, value=shipment["ordered_at"])
            worksheet.cell(row=row_index, column=9, value=shipment["planned_ship_at"])
            row_index += 1
    else:
        worksheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=9)
        worksheet["A4"] = "暂无发货记录"
        worksheet["A4"].alignment = Alignment(horizontal="center", vertical="center")

    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 26
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 18
    worksheet.column_dimensions["E"].width = 18
    worksheet.column_dimensions["F"].width = 18
    worksheet.column_dimensions["G"].width = 18
    worksheet.column_dimensions["H"].width = 16
    worksheet.column_dimensions["I"].width = 16

    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    return workbook


def build_order_import_template_workbook():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "订单导入模板"
    worksheet.freeze_panes = "A2"

    headers = ["图号", "数量", "计划发货时间"]
    for col_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col_index)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F81BD")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    example_values = ["示例图号001", 100, datetime.now().date().isoformat()]
    for col_index, value in enumerate(example_values, start=1):
        worksheet.cell(row=2, column=col_index, value=value)

    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 12
    worksheet.column_dimensions["C"].width = 18
    worksheet["C2"].number_format = "yyyy-mm-dd"

    notes_sheet = workbook.create_sheet("填写说明")
    notes = [
        "图号必须和系统产品资料里的产品图号完全一致。",
        "数量必须填写大于 0 的整数。",
        "计划发货时间可以每行填写；如果留空，导入页面需要填写默认计划发货时间。",
        "导入前可以删除示例行，只保留表头和订单明细。",
    ]
    notes_sheet["A1"] = "填写说明"
    notes_sheet["A1"].font = Font(bold=True, size=14)
    for row_index, note in enumerate(notes, start=2):
        notes_sheet.cell(row=row_index, column=1, value=note)
    notes_sheet.column_dimensions["A"].width = 72
    return workbook


class SignaturePngReader(ImageReader):
    def __init__(self, path):
        self.fileName = str(path)
        self._dataA = None
        self._width, self._height, self._data = read_png_rgb(path)
        self.mode = "RGB"

    def getSize(self):
        return self._width, self._height

    def getRGBData(self):
        return self._data

    def getTransparent(self):
        return None


class SignatureImageFlowable(Flowable):
    def __init__(self, path, width, height):
        super().__init__()
        self.reader = SignaturePngReader(path)
        self.width = width
        self.height = height

    def wrap(self, avail_width, avail_height):
        return self.width, self.height

    def draw(self):
        self.canv.drawImage(self.reader, 0, 0, width=self.width, height=self.height)


def read_png_rgb(path):
    data = Path(path).read_bytes()
    if not data.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ValueError("signature image is not PNG")

    offset = 8
    width = height = color_type = bit_depth = None
    compressed = bytearray()
    while offset + 8 <= len(data):
        length = struct.unpack(">I", data[offset:offset + 4])[0]
        chunk_type = data[offset + 4:offset + 8]
        chunk_data = data[offset + 8:offset + 8 + length]
        offset += 12 + length
        if chunk_type == b"IHDR":
            width, height, bit_depth, color_type = struct.unpack(">IIBB", chunk_data[:10])
        elif chunk_type == b"IDAT":
            compressed.extend(chunk_data)
        elif chunk_type == b"IEND":
            break

    if not width or not height or bit_depth != 8 or color_type not in {0, 2, 4, 6}:
        raise ValueError("unsupported signature PNG format")

    channels = {0: 1, 2: 3, 4: 2, 6: 4}[color_type]
    stride = width * channels
    raw = zlib.decompress(bytes(compressed))
    rows = []
    previous = bytearray(stride)
    index = 0
    for _ in range(height):
        filter_type = raw[index]
        index += 1
        row = bytearray(raw[index:index + stride])
        index += stride
        for i in range(stride):
            left = row[i - channels] if i >= channels else 0
            up = previous[i]
            up_left = previous[i - channels] if i >= channels else 0
            if filter_type == 1:
                row[i] = (row[i] + left) & 0xFF
            elif filter_type == 2:
                row[i] = (row[i] + up) & 0xFF
            elif filter_type == 3:
                row[i] = (row[i] + ((left + up) // 2)) & 0xFF
            elif filter_type == 4:
                predictor = left + up - up_left
                distances = (
                    abs(predictor - left),
                    abs(predictor - up),
                    abs(predictor - up_left),
                )
                row[i] = (row[i] + (left, up, up_left)[distances.index(min(distances))]) & 0xFF
            elif filter_type != 0:
                raise ValueError("unsupported PNG filter")
        rows.append(row)
        previous = row

    rgb = bytearray(width * height * 3)
    out = 0
    for row in rows:
        for x in range(width):
            source = x * channels
            if color_type == 0:
                r = g = b = row[source]
                a = 255
            elif color_type == 2:
                r, g, b = row[source:source + 3]
                a = 255
            elif color_type == 4:
                r = g = b = row[source]
                a = row[source + 1]
            else:
                r, g, b, a = row[source:source + 4]
            alpha = a / 255
            rgb[out] = int(r * alpha + 255 * (1 - alpha))
            rgb[out + 1] = int(g * alpha + 255 * (1 - alpha))
            rgb[out + 2] = int(b * alpha + 255 * (1 - alpha))
            out += 3
    return width, height, bytes(rgb)


def shipped_order_signature_flowable(shipped_orders, small_style):
    signed_shipments = [
        shipment
        for shipment in shipped_orders
        if shipment["signature_status"] == "已签收" and shipment["signature_image"]
    ]
    if not signed_shipments:
        return Paragraph("收货/签收：__________________", small_style)

    shipment = signed_shipments[0]
    safe_name = secure_filename(shipment["signature_image"])
    signature_path = SIGNATURES_DIR / safe_name
    if safe_name != shipment["signature_image"] or not signature_path.exists():
        return Paragraph("收货/签收：__________________", small_style)

    signed_at = format_datetime(shipment["signed_at"])
    signature_image = SignatureImageFlowable(signature_path, width=44 * mm, height=16 * mm)
    return Table(
        [
            [Paragraph("收货/签收：", small_style), signature_image],
            ["", Paragraph(f"签收时间：{signed_at or '-'}", small_style)],
        ],
        colWidths=[20 * mm, 58 * mm],
    )


def build_shipped_orders_pdf(shipped_orders, query, selected_customer, shipped_at, document_no=None):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )
    font_name = get_pdf_font_name()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DeliveryNoteTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=20,
        leading=24,
        alignment=1,
        spaceAfter=3 * mm,
    )
    info_style = ParagraphStyle(
        "DeliveryNoteInfo",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=9,
        leading=11,
        alignment=0,
    )
    cell_style = ParagraphStyle(
        "DeliveryNoteCell",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8,
        leading=9,
    )
    small_style = ParagraphStyle(
        "DeliveryNoteSmall",
        parent=cell_style,
        fontSize=7.5,
        leading=8.5,
    )

    filters = []
    if query:
        filters.append(f"搜索：{query}")
    if selected_customer:
        filters.append(f"客户：{selected_customer}")
    if shipped_at:
        filters.append(f"发货时间：{shipped_at}")
    filter_text = "；".join(filters) if filters else "全部记录"

    customer_name = selected_customer or (shipped_orders[0]["order_customer"] if shipped_orders else "")
    ship_date = shipped_at or (shipped_orders[0]["shipped_at"] if shipped_orders else "")
    document_no = document_no or f"DN-{datetime.now().strftime('%Y%m%d%H%M%S')}"

    story = [
        Paragraph("送货单", title_style),
    ]

    meta_data = [
        [Paragraph("客户名称", info_style), Paragraph(customer_name or "-", info_style), Paragraph("送货单号", info_style), Paragraph(document_no, info_style)],
        [Paragraph("发货日期", info_style), Paragraph(ship_date or "-", info_style), Paragraph("制单时间", info_style), Paragraph(datetime.now().strftime('%Y-%m-%d %H:%M:%S'), info_style)],
        [Paragraph("筛选条件", info_style), Paragraph(filter_text or "-", info_style), Paragraph("备注", info_style), Paragraph("", info_style)],
    ]
    meta_table = Table(meta_data, colWidths=[24 * mm, 52 * mm, 24 * mm, 76 * mm])
    meta_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#6B7C87")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EFF3")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E8EFF3")),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([meta_table, Spacer(1, 4 * mm)])

    data = [[
        "序号",
        "产品图号",
        "产品名称",
        "订单号",
        "发货时间",
        "发货数量",
    ]]
    for index, shipment in enumerate(shipped_orders, start=1):
        data.append([
            Paragraph(str(index), cell_style),
            Paragraph(shipment["drawing_no"] or "-", cell_style),
            Paragraph(shipment["product_name"] or "-", cell_style),
            Paragraph(shipment["order_no"] or "-", cell_style),
            Paragraph(shipment["shipped_at"] or "-", cell_style),
            Paragraph(str(shipment["shipped_quantity"] or 0), cell_style),
        ])
    if len(data) == 1:
        data.append([Paragraph("暂无发货记录", cell_style)] + [""] * 5)

    table = Table(
        data,
        colWidths=[14 * mm, 28 * mm, 58 * mm, 32 * mm, 26 * mm, 18 * mm],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("LEADING", (0, 0), (-1, -1), 10),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5C7280")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#6F8390")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F7")]),
                ("WORDWRAP", (0, 0), (-1, -1), True),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 4 * mm))
    sign_table = Table(
        [
            [
                shipped_order_signature_flowable(shipped_orders, small_style),
                Paragraph("制单：__________________", small_style),
                Paragraph("发货公司：宁波市杰德机械科技有限公司", small_style),
            ]
        ],
        colWidths=[82 * mm, 54 * mm, doc.width - 136 * mm],
    )
    sign_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("LEADING", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ALIGN", (2, 0), (2, 0), "RIGHT"),
            ]
        )
    )
    story.append(sign_table)
    doc.build(story)
    buffer.seek(0)
    return buffer


def delete_upload_file(filename):
    upload_path = MANUALS_DIR / filename
    if upload_path.exists():
        upload_path.unlink()


def delete_production_drawing_file(filename):
    drawing_path = PRODUCTION_DRAWINGS_DIR / filename
    if drawing_path.exists():
        drawing_path.unlink()


@app.before_request
def ensure_database():
    global DATABASE_READY
    if DATABASE_READY:
        if request.endpoint and request.endpoint.startswith(DISABLED_ENDPOINT_PREFIXES):
            abort(404)
        return
    with application_write_lock():
        if not DATABASE_READY:
            init_db()
            DATABASE_READY = True
    if request.endpoint and request.endpoint.startswith(DISABLED_ENDPOINT_PREFIXES):
        abort(404)


@app.route("/")
def index():
    return redirect(url_for("admin_index"))


@app.route("/admin/products")
@login_required
def products_index():
    query = request.args.get("q", "").strip()
    selected_supplier = request.args.get("supplier", "").strip()
    selected_customer = request.args.get("customer", "").strip()
    sort, direction = sort_state()

    sql = """
        SELECT manuals.id,
               manuals.drawing_no,
               manuals.product_name,
               manuals.supplier,
               manuals.customer,
               manuals.updated_at,
               manuals.remark,
               manuals.filename,
               COUNT(manual_files.id) AS file_count
        FROM manuals
        LEFT JOIN manual_files ON manual_files.manual_id = manuals.id
    """
    params = []
    conditions = []
    if query:
        conditions.append(
            """
            (
                manuals.drawing_no LIKE ?
                OR manuals.product_name LIKE ?
                OR manuals.supplier LIKE ?
                OR manuals.customer LIKE ?
                OR manuals.sku LIKE ?
                OR manuals.barcode LIKE ?
                OR manuals.qr_code LIKE ?
                OR manuals.remark LIKE ?
                OR EXISTS (
                    SELECT 1
                    FROM product_assembly_components
                    WHERE product_assembly_components.manual_id = manuals.id
                      AND product_assembly_components.assembly_drawing_no LIKE ?
                )
                OR EXISTS (
                    SELECT 1
                    FROM product_materials
                    WHERE product_materials.manual_id = manuals.id
                      AND (
                        product_materials.material LIKE ?
                        OR product_materials.thickness LIKE ?
                        OR product_materials.surface_type LIKE ?
                        OR product_materials.supplier LIKE ?
                      )
                )
            )
            """
        )
        like = f"%{query}%"
        params.extend([like, like, like, like, like, like, like, like, like, like, like, like, like])
    if selected_supplier:
        conditions.append("manuals.supplier = ?")
        params.append(selected_supplier)
    if selected_customer:
        conditions.append("manuals.customer = ?")
        params.append(selected_customer)
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += f" GROUP BY manuals.id {order_clause(sort, direction)}"

    with get_db() as conn:
        ensure_all_product_inventory_codes(conn)
        manuals = conn.execute(sql, params).fetchall()
        materials_by_manual = get_product_materials_map(conn, [manual["id"] for manual in manuals])
        assembly_components_by_manual = get_product_assembly_components_map(
            conn, [manual["id"] for manual in manuals]
        )
        suppliers = [row["supplier"] for row in get_suppliers_with_conn(conn)]
        customers = [row["customer"] for row in get_customers_with_conn(conn)]

    return render_template(
        "index.html",
        manuals=manuals,
        materials_by_manual=materials_by_manual,
        assembly_components_by_manual=assembly_components_by_manual,
        suppliers=suppliers,
        customers=customers,
        query=query,
        selected_supplier=selected_supplier,
        selected_customer=selected_customer,
        sort=sort,
        direction=direction,
    )


@app.route("/admin/products/assembly-components/batch", methods=["POST"])
@permission_required("product_edit")
def batch_set_product_assembly_components():
    redirect_args = {
        "q": request.form.get("return_q", "").strip(),
        "supplier": request.form.get("return_supplier", "").strip(),
        "customer": request.form.get("return_customer", "").strip(),
        "sort": request.form.get("return_sort", "").strip(),
        "direction": request.form.get("return_direction", "").strip(),
    }

    manual_ids = []
    try:
        for value in request.form.getlist("manual_id"):
            manual_id = int(str(value).strip())
            if manual_id <= 0:
                raise ValueError
            manual_ids.append(manual_id)
    except (TypeError, ValueError):
        flash("请选择有效的产品", "error")
        return redirect(url_for("products_index", **redirect_args))
    manual_ids = list(dict.fromkeys(manual_ids))
    if not manual_ids:
        flash("请先勾选要设置组装图号的产品", "error")
        return redirect(url_for("products_index", **redirect_args))

    assembly_drawing_no = request.form.get("assembly_drawing_no", "").strip()
    quantity_text = request.form.get("quantity_per_set", "").strip()
    try:
        quantity_per_set = int(quantity_text)
    except (TypeError, ValueError):
        quantity_per_set = 0
    if not assembly_drawing_no:
        flash("请填写组装图号", "error")
        return redirect(url_for("products_index", **redirect_args))
    if quantity_per_set <= 0 or quantity_per_set > MAX_ORDER_QUANTITY:
        flash(
            f"每套配件数量必须是 1 到 {MAX_ORDER_QUANTITY} 之间的整数",
            "error",
        )
        return redirect(url_for("products_index", **redirect_args))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        placeholders = ",".join("?" for _ in manual_ids)
        manuals = conn.execute(
            f"SELECT id, customer FROM manuals WHERE id IN ({placeholders})",
            manual_ids,
        ).fetchall()
        if len(manuals) != len(manual_ids):
            flash("请选择有效的产品", "error")
            return redirect(url_for("products_index", **redirect_args))
        customers = {str(row["customer"] or "").strip() for row in manuals}
        if "" in customers:
            flash("未设置客户的产品不能批量设置组装图号", "error")
            return redirect(url_for("products_index", **redirect_args))
        if len(customers) != 1:
            flash("只能批量设置同一客户的产品", "error")
            return redirect(url_for("products_index", **redirect_args))

        component_rows = conn.execute(
            f"""
            SELECT id, manual_id, assembly_drawing_no
            FROM product_assembly_components
            WHERE manual_id IN ({placeholders})
            ORDER BY id ASC
            """,
            manual_ids,
        ).fetchall()
        existing_component_ids = {}
        for row in component_rows:
            key = (row["manual_id"], str(row["assembly_drawing_no"] or "").casefold())
            existing_component_ids.setdefault(key, row["id"])
        requested_component_key = assembly_drawing_no.casefold()

        for manual_id in manual_ids:
            existing_id = existing_component_ids.get(
                (manual_id, requested_component_key)
            )
            if existing_id is not None:
                conn.execute(
                    """
                    UPDATE product_assembly_components
                    SET assembly_drawing_no = ?, quantity_per_set = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        assembly_drawing_no,
                        quantity_per_set,
                        now,
                        existing_id,
                    ),
                )
            else:
                next_sort_order = conn.execute(
                    """
                    SELECT COALESCE(MAX(sort_order), -1) + 1
                    FROM product_assembly_components
                    WHERE manual_id = ?
                    """,
                    (manual_id,),
                ).fetchone()[0]
                conn.execute(
                    """
                    INSERT INTO product_assembly_components (
                        manual_id, assembly_drawing_no, quantity_per_set,
                        sort_order, created_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (
                        manual_id,
                        assembly_drawing_no,
                        quantity_per_set,
                        next_sort_order,
                        now,
                        now,
                    ),
                )
        conn.execute(
            f"""
            UPDATE manuals
            SET updated_by = ?, updated_at = ?
            WHERE id IN ({placeholders})
            """,
            [current_admin_username(), now, *manual_ids],
        )

    flash(
        f"已为 {len(manual_ids)} 个产品设置组装图号 {assembly_drawing_no}，每套 {quantity_per_set} 个",
        "success",
    )
    return redirect(url_for("products_index", **redirect_args))


@app.route("/dashboard")
@login_required
def dashboard():
    with get_db() as conn:
        order_shipment_summary = order_shipment_summary_subquery(
            include_assembly=True, conn=conn
        )
        stats = {
            "products": conn.execute("SELECT COUNT(*) AS c FROM manuals").fetchone()["c"],
            "customers": conn.execute("SELECT COUNT(*) AS c FROM customers").fetchone()["c"],
            "common_infos": conn.execute("SELECT COUNT(*) AS c FROM common_infos").fetchone()["c"],
            "purchase_followups": conn.execute("SELECT COUNT(*) AS c FROM purchase_followups").fetchone()["c"],
            "powder_coating_records": conn.execute("SELECT COUNT(*) AS c FROM powder_coating_records").fetchone()["c"],
            "carton_purchases": conn.execute("SELECT COUNT(*) AS c FROM carton_purchases").fetchone()["c"],
            "arrival_records": conn.execute("SELECT COUNT(*) AS c FROM arrival_records").fetchone()["c"],
            "purchase_followups_pending_purchase_at": conn.execute(
                """
                SELECT COUNT(*) AS c
                FROM purchase_followups
                WHERE TRIM(COALESCE(purchased_at, '')) = ''
                """
            ).fetchone()["c"],
            "orders": conn.execute("SELECT COUNT(*) AS c FROM product_orders").fetchone()["c"],
            "unshipped": conn.execute(
                f"""
                SELECT COALESCE(SUM(product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0)), 0) AS c
                FROM product_orders
                LEFT JOIN ({order_shipment_summary}) AS shipments
                  ON shipments.order_id = product_orders.id
                WHERE product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) > 0
                """
            ).fetchone()["c"],
            "shipped": conn.execute(
                f"SELECT COALESCE(SUM(shipped_total), 0) AS c FROM ({order_shipment_summary})"
            ).fetchone()["c"],
            "production_followups": conn.execute(
                "SELECT COUNT(*) AS c FROM production_followups"
            ).fetchone()["c"],
        }
        recent_orders = conn.execute(
            """
            SELECT product_orders.order_no, product_orders.ordered_at,
                   product_orders.quantity, manuals.drawing_no, manuals.product_name
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            ORDER BY product_orders.created_at DESC, product_orders.id DESC
            LIMIT 6
            """
        ).fetchall()
        recent_shipments = conn.execute(
            """
            SELECT product_order_shipments.shipped_at,
                   product_order_shipments.shipped_quantity,
                   product_orders.order_no,
                   manuals.drawing_no,
                   manuals.product_name
            FROM product_order_shipments
            JOIN product_orders ON product_orders.id = product_order_shipments.order_id
            JOIN manuals ON manuals.id = product_orders.manual_id
            ORDER BY product_order_shipments.created_at DESC, product_order_shipments.id DESC
            LIMIT 6
            """
        ).fetchall()
    return render_template(
        "dashboard.html",
        stats=stats,
        recent_orders=recent_orders,
        recent_shipments=recent_shipments,
    )


@app.route("/admin/production-followups", methods=["GET", "POST"])
@login_required
def production_followups():
    query = request.args.get("q", "").strip()
    if request.method == "POST":
        ordered_at = request.form.get("ordered_at", "").strip()
        drawing_no = request.form.get("drawing_no", "").strip()
        product_name = request.form.get("product_name", "").strip()
        drawings = selected_production_drawings()

        if not all([ordered_at, drawing_no, product_name]):
            flash("下单时间、产品图号、产品名称为必填项", "error")
            return redirect(url_for("production_followups"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        try:
            with get_db() as conn:
                batch_no = next_production_batch_no(conn, ordered_at)
                cursor = conn.execute(
                    """
                    INSERT INTO production_followups (
                        batch_no, ordered_at, drawing_no, product_name, created_by,
                        created_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        batch_no,
                        ordered_at,
                        drawing_no,
                        product_name,
                        current_admin_username(),
                        now,
                        now,
                    ),
                )
                followup_id = cursor.lastrowid
                if drawings:
                    save_production_drawing_files(conn, followup_id, drawings)
        except ValueError as error:
            flash(f"图纸文件格式不支持：{error}", "error")
            return redirect(url_for("production_followups"))

        flash("生产跟进已新增", "success")
        return redirect(url_for("production_followups"))

    with get_db() as conn:
        followups = fetch_production_followups(conn, query)
    return render_template(
        "production_followups.html",
        followups=followups,
        query=query,
        today=datetime.now().date().isoformat(),
        stage_labels=PRODUCTION_STAGE_LABELS,
        stage_can_complete=production_stage_can_complete,
        stage_can_revert=production_stage_can_revert,
        stage_waiting_text=production_stage_waiting_text,
    )


@app.route("/admin/production-followups/<int:followup_id>/<stage>", methods=["POST"])
@permission_required("production_followups_manage")
def complete_production_stage(followup_id, stage):
    column = production_stage_column(stage)
    if not column:
        abort(404)

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        followup = conn.execute(
            "SELECT * FROM production_followups WHERE id = ?",
            (followup_id,),
        ).fetchone()
        if followup is None:
            abort(404)
        if followup[column]:
            if not production_stage_can_revert(followup, stage):
                flash("请先撤回后续工序，再撤回当前工序", "error")
                return redirect(url_for("production_followups"))
            conn.execute(
                f"""
                UPDATE production_followups
                SET {column} = '', updated_at = ?
                WHERE id = ?
                """,
                (now, followup_id),
            )
            flash(f"{PRODUCTION_STAGE_LABELS[stage]}已改回未完成", "success")
        else:
            if not production_stage_can_complete(followup, stage):
                flash("请按激光、折弯、焊接的顺序完成", "error")
                return redirect(url_for("production_followups"))
            conn.execute(
                f"""
                UPDATE production_followups
                SET {column} = ?, updated_at = ?
                WHERE id = ?
                """,
                (now, now, followup_id),
            )
            flash(f"{PRODUCTION_STAGE_LABELS[stage]}已完成", "success")
    return redirect(url_for("production_followups"))


@app.route("/admin/production-followups/<int:followup_id>/delete", methods=["POST"])
@permission_required("production_followups_manage")
def delete_production_followup(followup_id):
    with get_db() as conn:
        followup = conn.execute(
            "SELECT id FROM production_followups WHERE id = ?",
            (followup_id,),
        ).fetchone()
        if followup is None:
            abort(404)
        files = conn.execute(
            """
            SELECT filename
            FROM production_followup_files
            WHERE followup_id = ?
            """,
            (followup_id,),
        ).fetchall()
        for file_row in files:
            delete_production_drawing_file(file_row["filename"])
        conn.execute("DELETE FROM production_followup_files WHERE followup_id = ?", (followup_id,))
        conn.execute("DELETE FROM production_followups WHERE id = ?", (followup_id,))

    flash("生产跟进记录已删除", "success")
    return redirect(url_for("production_followups"))


@app.route("/admin/production-followups/<int:followup_id>/process-card")
@login_required
def preview_production_process_card(followup_id):
    with get_db() as conn:
        followup = fetch_production_followup_detail(conn, followup_id)
    if followup is None:
        abort(404)
    return render_template(
        "production_process_card.html",
        followup=followup,
        stage_labels=PRODUCTION_STAGE_LABELS,
        generated_at=datetime.now().strftime("%Y-%m-%d"),
    )


@app.route("/admin/production-followups/files/<int:file_id>")
@login_required
def production_followup_file(file_id):
    with get_db() as conn:
        drawing_file = conn.execute(
            "SELECT * FROM production_followup_files WHERE id = ?",
            (file_id,),
        ).fetchone()
    if drawing_file is None:
        abort(404)
    return send_from_directory(
        PRODUCTION_DRAWINGS_DIR,
        drawing_file["filename"],
        as_attachment=False,
        download_name=drawing_file["original_filename"],
        mimetype=guess_type(drawing_file["filename"])[0] or "application/octet-stream",
    )


@app.route("/feedback", methods=["POST"])
@login_required
def submit_feedback():
    content = request.form.get("content", "").strip()
    if not content:
        flash("请填写反馈内容", "error")
        return redirect(url_for("dashboard"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        conn.execute(
            """
            INSERT INTO feedbacks (content, created_by, created_at)
            VALUES (?, ?, ?)
            """,
            (content, current_admin_username(), now),
        )

    flash("反馈已提交", "success")
    return redirect(url_for("dashboard"))


@app.route("/admin/feedbacks")
@admin_required
def admin_feedbacks():
    with get_db() as conn:
        feedbacks = conn.execute(
            """
            SELECT *
            FROM feedbacks
            ORDER BY created_at DESC, id DESC
            """
        ).fetchall()
    return render_template("feedbacks.html", feedbacks=feedbacks)


@app.route("/admin/feedbacks/<int:feedback_id>/handle", methods=["POST"])
@admin_required
def handle_feedback(feedback_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        feedback = conn.execute(
            "SELECT id FROM feedbacks WHERE id = ?",
            (feedback_id,),
        ).fetchone()
        if feedback is None:
            abort(404)
        conn.execute(
            """
            UPDATE feedbacks
            SET handled = 1, handled_by = ?, handled_at = ?
            WHERE id = ?
            """,
            (current_admin_username(), now, feedback_id),
        )
    flash("反馈已标记为已处理", "success")
    return redirect(url_for("admin_feedbacks"))


MANUAL_SAFE_COLUMNS = (
    "id",
    "product_name",
    "customer",
    "model",
    "category",
    "version",
    "remark",
    "filename",
    "original_filename",
    "created_at",
    "updated_at",
    "drawing_no",
    "supplier",
    "pack_quantity",
    "pack_carton_size",
    "pack_weight",
    "description_html",
    "file_type",
    "created_by",
    "updated_by",
    "sku",
    "barcode",
    "qr_code",
    "default_location_id",
    "min_stock",
)


def fetch_manual_by_id(conn, manual_id, include_price=False):
    columns = list(MANUAL_SAFE_COLUMNS)
    if include_price:
        columns.extend(("unit_price_minor", "currency"))
    projection = ", ".join(columns)
    return conn.execute(
        f"SELECT {projection} FROM manuals WHERE id = ?",
        (manual_id,),
    ).fetchone()


def posted_product_price(existing_unit_price_minor=None, existing_currency="CNY"):
    if not user_can_view_prices():
        return existing_unit_price_minor, existing_currency
    currency = normalize_currency(request.form.get("currency", "CNY"))
    unit_price_minor = parse_money_minor(request.form.get("unit_price", ""), currency)
    return unit_price_minor, currency


@app.route("/manual/<int:manual_id>")
def manual_detail(manual_id):
    include_price = user_can_view_prices()
    with get_db() as conn:
        ensure_product_inventory_code(conn, manual_id)
        manual = fetch_manual_by_id(conn, manual_id, include_price=include_price)
        inventory_total = inventory_total_for_manual(conn, manual_id)
        inventory_locations = location_stock_map(conn, [manual_id]).get(manual_id, [])
        default_location = None
        if manual and manual["default_location_id"]:
            default_location = conn.execute("SELECT * FROM warehouse_locations WHERE id = ?", (manual["default_location_id"],)).fetchone()
    if manual is None:
        abort(404)
    return render_template(
        "detail.html",
        manual=manual,
        inventory_code=product_inventory_code(manual),
        inventory_total=inventory_total,
        inventory_locations=inventory_locations,
        default_location=default_location,
    )


@app.route("/manual/<int:manual_id>/technical")
def manual_technical(manual_id):
    with get_db() as conn:
        manual = fetch_manual_by_id(conn, manual_id)
        files = get_manual_files(conn, manual_id)
        inspection_requirements = get_manual_inspection_requirements(conn, manual_id)
        materials = get_product_materials(conn, manual_id)
    if manual is None:
        abort(404)
    return render_template(
        "detail_technical.html",
        manual=manual,
        files=files,
        inspection_requirements=inspection_requirements,
        materials=materials,
    )


@app.route("/manual/<int:manual_id>/preview")
def manual_preview(manual_id):
    with get_db() as conn:
        manual = fetch_manual_by_id(conn, manual_id)
        files = get_manual_files(conn, manual_id)
    if manual is None:
        abort(404)
    return render_template("preview.html", manual=manual, files=files)


@app.route("/manuals/<path:filename>")
def preview_manual(filename):
    mimetype = guess_type(filename)[0] or "application/octet-stream"
    return send_from_directory(MANUALS_DIR, filename, mimetype=mimetype, max_age=3600)


@app.route("/download/<path:filename>")
def download_manual(filename):
    with get_db() as conn:
        manual_file = conn.execute(
            """
            SELECT original_filename
            FROM manual_files
            WHERE filename = ?
            UNION
            SELECT original_filename
            FROM manuals
            WHERE filename = ?
            LIMIT 1
            """,
            (filename, filename),
        ).fetchone()
    if manual_file is None:
        abort(404)

    return send_from_directory(
        MANUALS_DIR,
        filename,
        as_attachment=True,
        download_name=manual_file["original_filename"],
        mimetype=guess_type(filename)[0] or "application/octet-stream",
    )


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        user = valid_admin_login(username, password)
        if user:
            session.permanent = True
            session["admin_logged_in"] = True
            session["admin_username"] = user["username"]
            session["admin_role"] = user["role"]
            flash("登录成功", "success")
            return redirect(url_for("dashboard"))
        flash("账号或密码错误", "error")

    return render_template("login.html")


@app.route("/admin/logout")
@login_required
def admin_logout():
    session.clear()
    flash("已退出登录", "success")
    return redirect(url_for("admin_login"))


@app.route("/admin/users", methods=["GET", "POST"])
@admin_required
def admin_users():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        role = request.form.get("role", "operator").strip()
        can_manage_products = 1 if role == "admin" or request.form.get("can_manage_products") else 0
        can_manage_orders = 1 if role == "admin" or request.form.get("can_manage_orders") else 0
        can_view_orders = 1 if role == "admin" or can_manage_orders or request.form.get("can_view_orders") else 0
        can_manage_shipped = 1 if role == "admin" or request.form.get("can_manage_shipped") else 0
        can_view_shipped = 1 if role == "admin" or can_manage_shipped or request.form.get("can_view_shipped") else 0
        can_manage_customers = 1 if role == "admin" or request.form.get("can_manage_customers") else 0
        can_manage_common_info = 1 if role == "admin" or request.form.get("can_manage_common_info") else 0
        can_manage_purchase_followups = 1 if role == "admin" or request.form.get("can_manage_purchase_followups") else 0
        can_manage_powder_coating = 1 if role == "admin" or request.form.get("can_manage_powder_coating") else 0
        can_manage_carton_purchases = 1 if role == "admin" or request.form.get("can_manage_carton_purchases") else 0
        can_manage_warehouse_inventory = 1 if role == "admin" or request.form.get("can_manage_warehouse_inventory") else 0
        can_manage_production_followups = 1 if role == "admin" or request.form.get("can_manage_production_followups") else 0
        can_create_products = 1 if role == "admin" or request.form.get("can_create_products") else 0
        can_edit_products = 1 if role == "admin" or request.form.get("can_edit_products") else 0
        can_view_prices = 1 if role == "admin" or request.form.get("can_view_prices") else 0
        can_manage_finance = 1 if role == "admin" or request.form.get("can_manage_finance") else 0

        if role not in {"admin", "operator"}:
            flash("请选择有效的用户权限", "error")
            return redirect(url_for("admin_users"))
        if not username or not password:
            flash("用户名和密码为必填项", "error")
            return redirect(url_for("admin_users"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        try:
            with get_db() as conn:
                conn.execute(
                    """
                    INSERT INTO users (
                        username, password_hash, role, active,
                        can_manage_products, can_manage_orders, can_view_orders,
                        can_manage_shipped, can_view_shipped,
                        can_manage_customers, can_manage_common_info,
                        can_manage_purchase_followups, can_manage_powder_coating,
                        can_manage_carton_purchases, can_manage_warehouse_inventory,
                        can_manage_production_followups, can_create_products, can_edit_products,
                        can_view_prices, can_manage_finance,
                        created_at, updated_at
                    )
                    VALUES (?, ?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        username,
                        generate_password_hash(password),
                        role,
                        can_manage_products,
                        can_manage_orders,
                        can_view_orders,
                        can_manage_shipped,
                        can_view_shipped,
                        can_manage_customers,
                        can_manage_common_info,
                        can_manage_purchase_followups,
                        can_manage_powder_coating,
                        can_manage_carton_purchases,
                        can_manage_warehouse_inventory,
                        can_manage_production_followups,
                        can_create_products,
                        can_edit_products,
                        can_view_prices,
                        can_manage_finance,
                        now,
                        now,
                    ),
                )
        except sqlite3.IntegrityError:
            flash("该用户名已存在", "error")
            return redirect(url_for("admin_users"))

        flash("用户已新增", "success")
        return redirect(url_for("admin_users"))

    with get_db() as conn:
        users = conn.execute(
            """
            SELECT *
            FROM users
            ORDER BY role ASC, username COLLATE NOCASE ASC
            """
        ).fetchall()
    return render_template("users.html", users=users)


@app.route("/admin/users/<int:user_id>/edit", methods=["POST"])
@admin_required
def edit_user(user_id):
    role = request.form.get("role", "operator").strip()
    password = request.form.get("password", "").strip()
    can_manage_products = 1 if role == "admin" or request.form.get("can_manage_products") else 0
    can_manage_orders = 1 if role == "admin" or request.form.get("can_manage_orders") else 0
    can_view_orders = 1 if role == "admin" or can_manage_orders or request.form.get("can_view_orders") else 0
    can_manage_shipped = 1 if role == "admin" or request.form.get("can_manage_shipped") else 0
    can_view_shipped = 1 if role == "admin" or can_manage_shipped or request.form.get("can_view_shipped") else 0
    can_manage_customers = 1 if role == "admin" or request.form.get("can_manage_customers") else 0
    can_manage_common_info = 1 if role == "admin" or request.form.get("can_manage_common_info") else 0
    can_manage_purchase_followups = 1 if role == "admin" or request.form.get("can_manage_purchase_followups") else 0
    can_manage_powder_coating = 1 if role == "admin" or request.form.get("can_manage_powder_coating") else 0
    can_manage_carton_purchases = 1 if role == "admin" or request.form.get("can_manage_carton_purchases") else 0
    can_manage_warehouse_inventory = 1 if role == "admin" or request.form.get("can_manage_warehouse_inventory") else 0
    can_manage_production_followups = 1 if role == "admin" or request.form.get("can_manage_production_followups") else 0
    can_create_products = 1 if role == "admin" or request.form.get("can_create_products") else 0
    can_edit_products = 1 if role == "admin" or request.form.get("can_edit_products") else 0
    can_view_prices = 1 if role == "admin" or request.form.get("can_view_prices") else 0
    can_manage_finance = 1 if role == "admin" or request.form.get("can_manage_finance") else 0

    if role not in {"admin", "operator"}:
        flash("请选择有效的用户权限", "error")
        return redirect(url_for("admin_users"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        user = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if user is None:
            abort(404)

        if user["role"] == "admin" and role != "admin":
            admin_count = conn.execute(
                "SELECT COUNT(*) AS c FROM users WHERE role = 'admin' AND active = 1"
            ).fetchone()["c"]
            if admin_count <= 1:
                flash("至少需要保留一个管理员账号", "error")
                return redirect(url_for("admin_users"))

        if password:
            conn.execute(
                """
                UPDATE users
                SET role = ?, password_hash = ?,
                    can_manage_products = ?, can_manage_orders = ?, can_view_orders = ?,
                    can_manage_shipped = ?, can_view_shipped = ?,
                    can_manage_customers = ?, can_manage_common_info = ?,
                    can_manage_purchase_followups = ?, can_manage_powder_coating = ?,
                    can_manage_carton_purchases = ?, can_manage_warehouse_inventory = ?,
                    can_manage_production_followups = ?, can_create_products = ?, can_edit_products = ?,
                    can_view_prices = ?, can_manage_finance = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    role,
                    generate_password_hash(password),
                    can_manage_products,
                    can_manage_orders,
                    can_view_orders,
                    can_manage_shipped,
                    can_view_shipped,
                    can_manage_customers,
                    can_manage_common_info,
                    can_manage_purchase_followups,
                    can_manage_powder_coating,
                    can_manage_carton_purchases,
                    can_manage_warehouse_inventory,
                    can_manage_production_followups,
                    can_create_products,
                    can_edit_products,
                    can_view_prices,
                    can_manage_finance,
                    now,
                    user_id,
                ),
            )
        else:
            conn.execute(
                """
                UPDATE users
                SET role = ?, can_manage_products = ?, can_manage_orders = ?, can_view_orders = ?,
                    can_manage_shipped = ?, can_view_shipped = ?,
                    can_manage_customers = ?, can_manage_common_info = ?,
                    can_manage_purchase_followups = ?, can_manage_powder_coating = ?,
                    can_manage_carton_purchases = ?, can_manage_warehouse_inventory = ?,
                    can_manage_production_followups = ?, can_create_products = ?, can_edit_products = ?,
                    can_view_prices = ?, can_manage_finance = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    role,
                    can_manage_products,
                    can_manage_orders,
                    can_view_orders,
                    can_manage_shipped,
                    can_view_shipped,
                    can_manage_customers,
                    can_manage_common_info,
                    can_manage_purchase_followups,
                    can_manage_powder_coating,
                    can_manage_carton_purchases,
                    can_manage_warehouse_inventory,
                    can_manage_production_followups,
                    can_create_products,
                    can_edit_products,
                    can_view_prices,
                    can_manage_finance,
                    now,
                    user_id,
                ),
            )

    if session.get("admin_username") == user["username"]:
        session["admin_role"] = role
    flash("用户已更新", "success")
    return redirect(url_for("admin_users"))


@app.route("/admin/users/password", methods=["POST"])
@admin_required
def change_admin_password():
    current_password = request.form.get("current_password", "")
    new_password = request.form.get("new_password", "")
    confirm_password = request.form.get("confirm_password", "")

    if not current_password or not new_password or not confirm_password:
        flash("当前密码、新密码和确认密码为必填项", "error")
        return redirect(url_for("admin_users"))
    if new_password != confirm_password:
        flash("两次输入的新密码不一致", "error")
        return redirect(url_for("admin_users"))

    user = current_user()
    if not user or not check_password_hash(user["password_hash"], current_password):
        flash("当前密码不正确", "error")
        return redirect(url_for("admin_users"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        conn.execute(
            """
            UPDATE users
            SET password_hash = ?, updated_at = ?
            WHERE id = ?
            """,
            (generate_password_hash(new_password), now, user["id"]),
        )

    flash("管理员密码已修改", "success")
    return redirect(url_for("admin_users"))


@app.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@admin_required
def delete_user(user_id):
    with get_db() as conn:
        user = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if user is None:
            abort(404)

        if user["username"] == session.get("admin_username"):
            flash("不能删除当前登录账号", "error")
            return redirect(url_for("admin_users"))

        if user["role"] == "admin":
            admin_count = conn.execute(
                "SELECT COUNT(*) AS c FROM users WHERE role = 'admin' AND active = 1"
            ).fetchone()["c"]
            if admin_count <= 1:
                flash("至少需要保留一个管理员账号", "error")
                return redirect(url_for("admin_users"))

        conn.execute("DELETE FROM users WHERE id = ?", (user_id,))

    flash("用户已删除", "success")
    return redirect(url_for("admin_users"))


@app.route("/admin/customers", methods=["GET", "POST"])
@permission_required("customers")
def admin_customers():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        contact = request.form.get("contact", "").strip()
        address = request.form.get("address", "").strip()
        email = request.form.get("email", "").strip()
        phone = request.form.get("phone", "").strip()
        remark = request.form.get("remark", "").strip()

        if not name:
            flash("客户名称为必填项", "error")
            return redirect(url_for("admin_customers"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        try:
            with get_db() as conn:
                conn.execute(
                    """
                    INSERT INTO customers (
                        name, contact, address, email, phone, remark,
                        created_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (name, contact, address, email, phone, remark, now, now),
                )
        except sqlite3.IntegrityError:
            flash("该客户名称已存在", "error")
            return redirect(url_for("admin_customers"))

        flash("客户信息已新增", "success")
        return redirect(url_for("admin_customers"))

    query = request.args.get("q", "").strip()
    sql = "SELECT * FROM customers"
    params = []
    if query:
        sql += """
            WHERE name LIKE ?
               OR contact LIKE ?
               OR address LIKE ?
               OR email LIKE ?
               OR phone LIKE ?
               OR remark LIKE ?
               OR invoice_title LIKE ?
               OR tax_id LIKE ?
        """
        like = f"%{query}%"
        params.extend([like, like, like, like, like, like, like, like])
    sql += " ORDER BY name COLLATE NOCASE ASC, id DESC"
    with get_db() as conn:
        customers = conn.execute(sql, params).fetchall()
    return render_template("customers.html", customers=customers, query=query)


@app.route("/admin/customers/<int:customer_id>/edit", methods=["POST"])
@permission_required("customers")
def edit_customer(customer_id):
    name = request.form.get("name", "").strip()
    contact = request.form.get("contact", "").strip()
    address = request.form.get("address", "").strip()
    email = request.form.get("email", "").strip()
    phone = request.form.get("phone", "").strip()
    remark = request.form.get("remark", "").strip()

    if not name:
        flash("客户名称为必填项", "error")
        return redirect(url_for("admin_customers"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    try:
        with get_db() as conn:
            customer = conn.execute(
                "SELECT id FROM customers WHERE id = ?",
                (customer_id,),
            ).fetchone()
            if customer is None:
                abort(404)
            conn.execute(
                """
                UPDATE customers
                SET name = ?, contact = ?, address = ?, email = ?,
                    phone = ?, remark = ?, updated_at = ?
                WHERE id = ?
                """,
                (name, contact, address, email, phone, remark, now, customer_id),
            )
    except sqlite3.IntegrityError:
        flash("该客户名称已存在", "error")
        return redirect(url_for("admin_customers"))

    flash("客户信息已更新", "success")
    return redirect(url_for("admin_customers"))


@app.route("/admin/customers/<int:customer_id>/billing", methods=["GET", "POST"])
@login_required
def customer_billing(customer_id):
    can_manage_customers = user_has_permission("customers")
    can_manage_finance = user_has_permission("finance_manage")
    if not can_manage_customers and not can_manage_finance:
        flash("当前账号没有权限查看客户开票信息", "error")
        return redirect(url_for("admin_index"))
    if request.method == "POST" and not can_manage_customers:
        flash("当前账号没有权限修改客户开票信息", "error")
        return redirect(url_for("customer_billing", customer_id=customer_id))

    with get_db() as conn:
        customer = conn.execute(
            "SELECT * FROM customers WHERE id = ?",
            (customer_id,),
        ).fetchone()
        if customer is None:
            abort(404)

        if request.method == "POST":
            values = {
                field: request.form.get(field, "").strip()
                for field in CUSTOMER_INVOICE_FIELDS
            }
            if values["invoice_email"] and not re.fullmatch(
                r"[^\s@]+@[^\s@]+\.[^\s@]+",
                values["invoice_email"],
            ):
                flash("收票邮箱格式不正确", "error")
                return redirect(
                    url_for("customer_billing", customer_id=customer_id)
                )
            now = datetime.utcnow().isoformat(timespec="seconds")
            conn.execute(
                """
                UPDATE customers
                SET invoice_title = ?, tax_id = ?, registered_address = ?,
                    registered_phone = ?, bank_name = ?, bank_account = ?,
                    invoice_email = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    *(values[field] for field in CUSTOMER_INVOICE_FIELDS),
                    now,
                    customer_id,
                ),
            )
            flash("客户开票信息已保存", "success")
            return redirect(url_for("customer_billing", customer_id=customer_id))

    return render_template("customer_billing.html", customer=customer)


@app.route("/admin/customers/<int:customer_id>/delete", methods=["POST"])
@permission_required("customers")
def delete_customer(customer_id):
    with get_db() as conn:
        customer = conn.execute(
            "SELECT id FROM customers WHERE id = ?",
            (customer_id,),
        ).fetchone()
        if customer is None:
            abort(404)
        finance_record = conn.execute(
            "SELECT 1 FROM finance_invoices WHERE customer_id = ? LIMIT 1",
            (customer_id,),
        ).fetchone()
        if finance_record is not None:
            flash("该客户已有财务记录，不能删除", "error")
            return redirect(url_for("admin_customers"))
        conn.execute("DELETE FROM customers WHERE id = ?", (customer_id,))

    flash("客户信息已删除", "success")
    return redirect(url_for("admin_customers"))


@app.route("/admin/finance")
@permission_required("finance_manage")
def finance_invoices():
    selected_customer = request.args.get("customer", "").strip()
    selected_status = request.args.get("status", "").strip()
    invoice_no = request.args.get("invoice_no", "").strip()
    invoice_date = request.args.get("date", "").strip()
    conditions = []
    params = []
    if selected_customer:
        conditions.append("finance_invoices.customer_name = ?")
        params.append(selected_customer)
    if selected_status:
        if selected_status not in FINANCE_STATUS_LABELS:
            selected_status = ""
        else:
            conditions.append("finance_invoices.status = ?")
            params.append(selected_status)
    if invoice_no:
        conditions.append("finance_invoices.invoice_no LIKE ?")
        params.append(f"%{invoice_no}%")
    if invoice_date:
        conditions.append("finance_invoices.invoice_date = ?")
        params.append(invoice_date)

    sql = """
        SELECT finance_invoices.*,
               (
                   SELECT COUNT(*)
                   FROM finance_invoice_items
                   WHERE finance_invoice_items.invoice_id = finance_invoices.id
               ) AS item_count
        FROM finance_invoices
    """
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += " ORDER BY finance_invoices.created_at DESC, finance_invoices.id DESC"
    with get_db() as conn:
        invoices = conn.execute(sql, params).fetchall()
        customers = conn.execute(
            "SELECT id, name FROM customers ORDER BY name COLLATE NOCASE, id"
        ).fetchall()
    return render_template(
        "finance_invoices.html",
        invoices=invoices,
        customers=customers,
        status_labels=FINANCE_STATUS_LABELS,
        selected_customer=selected_customer,
        selected_status=selected_status,
        selected_invoice_no=invoice_no,
        selected_invoice_date=invoice_date,
    )


@app.route("/admin/finance/new", methods=["GET", "POST"])
@permission_required("finance_manage")
def new_finance_invoice():
    selected_customer_id = request.values.get("customer_id", "").strip()
    selected_currency = request.values.get("currency", "").strip().upper()
    if request.method == "POST":
        try:
            customer_id = int(selected_customer_id)
            if customer_id <= 0:
                raise ValueError
        except (TypeError, ValueError):
            flash("请选择客户", "error")
            return redirect(url_for("new_finance_invoice"))
        try:
            source_refs = parse_finance_source_refs(
                request.form.getlist("source_ref")
            )
            with get_db() as conn:
                invoice_id = create_finance_invoice(
                    conn,
                    customer_id,
                    source_refs,
                    current_admin_username(),
                )
        except ValueError as error:
            flash(str(error), "error")
            return redirect(
                url_for(
                    "new_finance_invoice",
                    customer_id=selected_customer_id,
                    currency=selected_currency,
                )
            )
        flash("待开票单已创建", "success")
        return redirect(url_for("finance_invoice_detail", invoice_id=invoice_id))

    if selected_currency:
        try:
            selected_currency = normalize_currency(selected_currency)
        except ValueError as error:
            flash(str(error), "error")
            selected_currency = ""
    selected_customer = None
    sources = []
    with get_db() as conn:
        customers = conn.execute(
            "SELECT id, name FROM customers ORDER BY name COLLATE NOCASE, id"
        ).fetchall()
        if selected_customer_id.isdigit() and int(selected_customer_id) > 0:
            selected_customer = conn.execute(
                "SELECT id, name FROM customers WHERE id = ?",
                (int(selected_customer_id),),
            ).fetchone()
            if selected_customer:
                sources = fetch_available_finance_sources(
                    conn,
                    selected_customer["name"],
                    selected_currency,
                )
    return render_template(
        "finance_invoice_new.html",
        customers=customers,
        selected_customer=selected_customer,
        selected_customer_id=selected_customer_id,
        selected_currency=selected_currency,
        supported_currencies=SUPPORTED_CURRENCIES,
        sources=sources,
    )


@app.route("/admin/finance/<int:invoice_id>")
@permission_required("finance_manage")
def finance_invoice_detail(invoice_id):
    with get_db() as conn:
        invoice = fetch_finance_invoice(conn, invoice_id)
        if invoice is None:
            abort(404)
        items = fetch_finance_invoice_items(conn, invoice_id)
        available_sources = []
        if invoice["status"] == "pending":
            available_sources = fetch_available_finance_sources(
                conn,
                invoice["customer_name"],
                invoice["currency"],
            )
    return render_template(
        "finance_invoice_detail.html",
        invoice=invoice,
        items=items,
        available_sources=available_sources,
        status_labels=FINANCE_STATUS_LABELS,
        default_date=datetime.now().date().isoformat(),
    )


@app.route("/admin/finance/<int:invoice_id>/items", methods=["POST"])
@permission_required("finance_manage")
def edit_finance_invoice_items(invoice_id):
    raw_refs = request.form.getlist("source_ref")
    try:
        source_refs = parse_finance_source_refs(raw_refs) if raw_refs else []
        with get_db() as conn:
            replace_pending_invoice_items(
                conn,
                invoice_id,
                source_refs,
                current_admin_username(),
            )
    except ValueError as error:
        flash(str(error), "error")
    else:
        flash("待开票明细已更新", "success")
    return redirect(url_for("finance_invoice_detail", invoice_id=invoice_id))


@app.route("/admin/finance/<int:invoice_id>/issue", methods=["POST"])
@permission_required("finance_manage")
def issue_finance_invoice(invoice_id):
    try:
        with get_db() as conn:
            mark_finance_invoice_invoiced(
                conn,
                invoice_id,
                request.form.get("invoice_no", ""),
                request.form.get("invoice_date", ""),
                request.form.get("finance_remark", ""),
                current_admin_username(),
            )
    except ValueError as error:
        flash(str(error), "error")
    else:
        flash("开票信息已登记", "success")
    return redirect(url_for("finance_invoice_detail", invoice_id=invoice_id))


@app.route("/admin/finance/<int:invoice_id>/pay", methods=["POST"])
@permission_required("finance_manage")
def pay_finance_invoice(invoice_id):
    try:
        with get_db() as conn:
            mark_finance_invoice_paid(
                conn,
                invoice_id,
                request.form.get("payment_date", ""),
                current_admin_username(),
            )
    except ValueError as error:
        flash(str(error), "error")
    else:
        flash("收款信息已登记", "success")
    return redirect(url_for("finance_invoice_detail", invoice_id=invoice_id))


@app.route("/admin/finance/<int:invoice_id>/reopen", methods=["POST"])
@permission_required("finance_manage")
def reopen_finance_invoice(invoice_id):
    try:
        with get_db() as conn:
            reopen_finance_invoice_payment(
                conn,
                invoice_id,
                current_admin_username(),
            )
    except ValueError as error:
        flash(str(error), "error")
    else:
        flash("已撤销收款状态", "success")
    return redirect(url_for("finance_invoice_detail", invoice_id=invoice_id))


@app.route("/admin/finance/<int:invoice_id>/void", methods=["POST"])
@permission_required("finance_manage")
def void_finance_invoice_route(invoice_id):
    try:
        with get_db() as conn:
            void_finance_invoice(conn, invoice_id, current_admin_username())
    except ValueError as error:
        flash(str(error), "error")
    else:
        flash("开票单已作废，关联发货明细已释放", "success")
    return redirect(url_for("finance_invoice_detail", invoice_id=invoice_id))


@app.route("/admin/finance/<int:invoice_id>/delete", methods=["POST"])
@permission_required("finance_manage")
def delete_finance_invoice(invoice_id):
    try:
        with get_db() as conn:
            delete_pending_finance_invoice(conn, invoice_id)
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("finance_invoice_detail", invoice_id=invoice_id))
    flash("待开票单已删除", "success")
    return redirect(url_for("finance_invoices"))


@app.route("/admin/common-info", methods=["GET", "POST"])
@permission_required("common_info")
def admin_common_info():
    if request.method == "POST":
        category = request.form.get("category", "").strip()
        title = request.form.get("title", "").strip()
        content = request.form.get("content", "").strip()
        remark = request.form.get("remark", "").strip()

        if not title:
            flash("标题为必填项", "error")
            return redirect(url_for("admin_common_info"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO common_infos (
                    category, title, content, remark, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (category, title, content, remark, now, now),
            )

        flash("常用信息已新增", "success")
        return redirect(url_for("admin_common_info"))

    query = request.args.get("q", "").strip()
    category = request.args.get("category", "").strip()
    sql = "SELECT * FROM common_infos"
    params = []
    conditions = []
    if query:
        conditions.append(
            """
            (
                category LIKE ?
                OR title LIKE ?
                OR content LIKE ?
                OR remark LIKE ?
            )
            """
        )
        like = f"%{query}%"
        params.extend([like, like, like, like])
    if category:
        conditions.append("category = ?")
        params.append(category)
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += " ORDER BY updated_at DESC, id DESC"
    with get_db() as conn:
        common_infos = conn.execute(sql, params).fetchall()
        categories = conn.execute(
            """
            SELECT DISTINCT category
            FROM common_infos
            WHERE category != ''
            ORDER BY category COLLATE NOCASE ASC
            """
        ).fetchall()
    return render_template(
        "common_info.html",
        common_infos=common_infos,
        categories=[row["category"] for row in categories],
        query=query,
        selected_category=category,
    )


@app.route("/admin/common-info/<int:info_id>/edit", methods=["POST"])
@permission_required("common_info")
def edit_common_info(info_id):
    category = request.form.get("category", "").strip()
    title = request.form.get("title", "").strip()
    content = request.form.get("content", "").strip()
    remark = request.form.get("remark", "").strip()

    if not title:
        flash("标题为必填项", "error")
        return redirect(url_for("admin_common_info"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        common_info = conn.execute(
            "SELECT id FROM common_infos WHERE id = ?",
            (info_id,),
        ).fetchone()
        if common_info is None:
            abort(404)
        conn.execute(
            """
            UPDATE common_infos
            SET category = ?, title = ?, content = ?, remark = ?, updated_at = ?
            WHERE id = ?
            """,
            (category, title, content, remark, now, info_id),
        )

    flash("常用信息已更新", "success")
    return redirect(url_for("admin_common_info"))


@app.route("/admin/common-info/<int:info_id>/delete", methods=["POST"])
@permission_required("common_info")
def delete_common_info(info_id):
    with get_db() as conn:
        common_info = conn.execute(
            "SELECT id FROM common_infos WHERE id = ?",
            (info_id,),
        ).fetchone()
        if common_info is None:
            abort(404)
        conn.execute("DELETE FROM common_infos WHERE id = ?", (info_id,))

    flash("常用信息已删除", "success")
    return redirect(url_for("admin_common_info"))


@app.route("/admin/purchase-followups", methods=["GET", "POST"])
@permission_required("purchase_followups")
def admin_purchase_followups():
    if request.method == "POST":
        recorded_at = request.form.get("recorded_at", "").strip()
        item_name = request.form.get("item_name", "").strip()
        remark = request.form.get("remark", "").strip()
        quantity = request.form.get("quantity", "").strip()
        purchased = 0
        purchased_at = request.form.get("purchased_at", "").strip()
        image_upload = request.files.get("image")

        if not all([recorded_at, item_name, quantity]):
            flash("入录时间、采购物品和采购数量为必填项", "error")
            return redirect(url_for("admin_purchase_followups"))
        try:
            quantity_value = int(quantity)
        except ValueError:
            flash("采购数量必须为整数", "error")
            return redirect(url_for("admin_purchase_followups"))
        if quantity_value <= 0:
            flash("采购数量必须大于 0", "error")
            return redirect(url_for("admin_purchase_followups"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        try:
            image_filename, image_original_filename = save_purchase_followup_image(image_upload)
        except ValueError:
            flash("请上传图片格式文件", "error")
            return redirect(url_for("admin_purchase_followups"))
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO purchase_followups (
                    recorded_at, item_name, remark, recorded_by, quantity,
                    purchased, purchased_at, image_filename, image_original_filename,
                    created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    recorded_at,
                    item_name,
                    remark,
                    current_admin_username(),
                    quantity_value,
                    purchased,
                    purchased_at,
                    image_filename,
                    image_original_filename,
                    now,
                    now,
                ),
            )

        flash("其他物品采购记录已新增", "success")
        return redirect(url_for("admin_purchase_followups"))

    query = request.args.get("q", "").strip()
    status = request.args.get("status", "").strip()
    show_completed = request.args.get("show_completed") == "1"
    purchase_sort_columns = {
        "recorded_at": "recorded_at",
        "item_name": "item_name COLLATE NOCASE",
        "remark": "remark COLLATE NOCASE",
        "recorded_by": "recorded_by COLLATE NOCASE",
        "quantity": "quantity",
        "image": "CASE WHEN image_filename != '' THEN 1 ELSE 0 END",
        "purchased_at": "purchased_at",
        "purchased": "purchased",
    }
    sort = request.args.get("sort", "").strip()
    direction = request.args.get("direction", "desc").strip().lower()
    if sort not in purchase_sort_columns:
        sort = ""
    if direction not in {"asc", "desc"}:
        direction = "desc"
    sql = "SELECT * FROM purchase_followups"
    params = []
    conditions = []
    if query:
        conditions.append(
            """
            (
                item_name LIKE ?
                OR remark LIKE ?
                OR recorded_by LIKE ?
            )
            """
        )
        like = f"%{query}%"
        params.extend([like, like, like])
    if status == "purchased":
        conditions.append("purchased = 1")
    elif status == "pending":
        conditions.append("purchased = 0")
    elif not show_completed:
        conditions.append("purchased = 0")
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    if sort:
        sql_direction = "ASC" if direction == "asc" else "DESC"
        sql += f" ORDER BY {purchase_sort_columns[sort]} {sql_direction}, id DESC"
    else:
        sql += " ORDER BY purchased ASC, recorded_at DESC, id DESC"
    with get_db() as conn:
        purchase_followups = conn.execute(sql, params).fetchall()
        pending_purchase_at_count = get_purchase_followups_pending_purchase_at_count(conn)
    return render_template(
        "purchase_followups.html",
        purchase_followups=purchase_followups,
        pending_purchase_at_count=pending_purchase_at_count,
        query=query,
        status=status,
        show_completed=show_completed,
        sort=sort,
        direction=direction,
        default_recorded_at=datetime.now().date().isoformat(),
    )


@app.route("/purchase-followups/pending-purchase-at/stream/<token>")
def purchase_followups_pending_purchase_at_sse(token):
    if verify_purchase_followup_sse_token(token) != "pending_purchase_at_json":
        abort(404)

    def event_stream():
        last_count = None
        while True:
            with get_db() as conn:
                count = get_purchase_followups_pending_purchase_at_count(conn)
            if count != last_count:
                payload = json.dumps(
                    {
                        "count": count,
                        "generated_at": datetime.utcnow().isoformat(timespec="seconds"),
                    },
                    ensure_ascii=False,
                )
                yield f"event: purchase_followups_pending_purchase_at\n"
                yield f"data: {payload}\n\n"
                last_count = count
            else:
                yield ": keep-alive\n\n"
            time.sleep(10)

    response = Response(event_stream(), mimetype="text/event-stream")
    response.headers["Cache-Control"] = "no-cache"
    response.headers["X-Accel-Buffering"] = "no"
    return response


@app.route("/purchase-followups/pending-purchase-at/plain-stream/<token>")
def purchase_followups_pending_purchase_at_plain_sse(token):
    if verify_purchase_followup_sse_token(token) != "pending_purchase_at_plain":
        abort(404)

    def event_stream():
        last_count = None
        while True:
            with get_db() as conn:
                count = get_purchase_followups_pending_purchase_at_count(conn)
            if count != last_count:
                yield f"data: {count}\n\n"
                last_count = count
            else:
                yield ": keep-alive\n\n"
            time.sleep(10)

    response = Response(event_stream(), mimetype="text/event-stream")
    response.headers["Cache-Control"] = "no-cache"
    response.headers["X-Accel-Buffering"] = "no"
    return response


@app.route("/admin/purchase-followups/<int:followup_id>/edit", methods=["POST"])
@permission_required("purchase_followups")
def edit_purchase_followup(followup_id):
    recorded_at = request.form.get("recorded_at", "").strip()
    item_name = request.form.get("item_name", "").strip()
    remark = request.form.get("remark", "").strip()
    quantity = request.form.get("quantity", "").strip()
    purchased_at = request.form.get("purchased_at", "").strip()
    complete = request.form.get("complete") == "1"
    image_upload = request.files.get("image")

    if not all([recorded_at, item_name, quantity]):
        flash("入录时间、采购物品和采购数量为必填项", "error")
        return redirect(url_for("admin_purchase_followups"))
    try:
        quantity_value = int(quantity)
    except ValueError:
        flash("采购数量必须为整数", "error")
        return redirect(url_for("admin_purchase_followups"))
    if quantity_value <= 0:
        flash("采购数量必须大于 0", "error")
        return redirect(url_for("admin_purchase_followups"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        followup = conn.execute(
            "SELECT id, image_filename FROM purchase_followups WHERE id = ?",
            (followup_id,),
        ).fetchone()
        if followup is None:
            abort(404)
        image_filename = followup["image_filename"] or ""
        image_original_filename = ""
        if image_upload and (image_upload.filename or image_upload.mimetype):
            try:
                image_filename, image_original_filename = save_purchase_followup_image(image_upload)
            except ValueError:
                flash("请上传图片格式文件", "error")
                return redirect(url_for("admin_purchase_followups"))
            if followup["image_filename"]:
                delete_upload_file(followup["image_filename"])
        conn.execute(
            """
            UPDATE purchase_followups
            SET recorded_at = ?, item_name = ?, remark = ?, quantity = ?,
                purchased_at = ?, purchased = CASE WHEN ? = 1 THEN 1 ELSE purchased END,
                image_filename = CASE WHEN ? != '' THEN ? ELSE image_filename END,
                image_original_filename = CASE WHEN ? != '' THEN ? ELSE image_original_filename END,
                updated_at = ?
            WHERE id = ?
            """,
            (
                recorded_at,
                item_name,
                remark,
                quantity_value,
                purchased_at,
                1 if complete else 0,
                image_filename,
                image_filename,
                image_original_filename,
                image_original_filename,
                now,
                followup_id,
            ),
        )

    flash("其他物品采购记录已更新", "success")
    return redirect(url_for("admin_purchase_followups"))


@app.route("/admin/purchase-followups/<int:followup_id>/copy", methods=["POST"])
@permission_required("purchase_followups")
def copy_purchase_followup(followup_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    recorded_at = datetime.now().date().isoformat()
    with get_db() as conn:
        source = conn.execute(
            """
            SELECT item_name, remark, quantity, image_filename, image_original_filename
            FROM purchase_followups
            WHERE id = ?
            """,
            (followup_id,),
        ).fetchone()
        if source is None:
            abort(404)

        image_filename = ""
        image_original_filename = ""
        if source["image_filename"]:
            image_filename, image_original_filename, _ = copy_upload_file(
                source["image_filename"],
                source["image_original_filename"] or source["image_filename"],
            )

        conn.execute(
            """
            INSERT INTO purchase_followups (
                recorded_at, item_name, remark, recorded_by, quantity,
                purchased, purchased_at, image_filename, image_original_filename,
                created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, 0, '', ?, ?, ?, ?)
            """,
            (
                recorded_at,
                source["item_name"],
                source["remark"] or "",
                current_admin_username(),
                source["quantity"] or 0,
                image_filename,
                image_original_filename,
                now,
                now,
            ),
        )

    flash("采购跟进记录已复制", "success")
    return redirect(url_for("admin_purchase_followups"))


@app.route("/admin/purchase-followups/<int:followup_id>/toggle", methods=["POST"])
@permission_required("purchase_followups")
def toggle_purchase_followup(followup_id):
    purchased = 1 if request.form.get("purchased") == "1" else 0
    keep_visible = request.form.get("keep_visible") == "1"
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        followup = conn.execute(
            "SELECT id FROM purchase_followups WHERE id = ?",
            (followup_id,),
        ).fetchone()
        if followup is None:
            abort(404)
        conn.execute(
            """
            UPDATE purchase_followups
            SET purchased = ?, updated_at = ?
            WHERE id = ?
            """,
            (purchased, now, followup_id),
        )
    redirect_target = request.referrer or url_for("admin_purchase_followups")
    if keep_visible:
        parts = urlsplit(redirect_target)
        query = dict(parse_qsl(parts.query, keep_blank_values=True))
        query["show_completed"] = "1"
        query.pop("status", None)
        redirect_target = urlunsplit(
            (parts.scheme, parts.netloc, parts.path, urlencode(query), parts.fragment)
        )
    return redirect(redirect_target)


@app.route("/admin/purchase-followups/<int:followup_id>/image/delete", methods=["POST"])
@permission_required("purchase_followups")
def delete_purchase_followup_image(followup_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        followup = conn.execute(
            """
            SELECT id, image_filename
            FROM purchase_followups
            WHERE id = ?
            """,
            (followup_id,),
        ).fetchone()
        if followup is None:
            abort(404)
        if followup["image_filename"]:
            delete_upload_file(followup["image_filename"])
        conn.execute(
            """
            UPDATE purchase_followups
            SET image_filename = '', image_original_filename = '', updated_at = ?
            WHERE id = ?
            """,
            (now, followup_id),
        )

    flash("图片已删除", "success")
    return redirect(request.referrer or url_for("admin_purchase_followups"))


@app.route("/admin/purchase-followups/<int:followup_id>/delete", methods=["POST"])
@permission_required("purchase_followups")
def delete_purchase_followup(followup_id):
    with get_db() as conn:
        followup = conn.execute(
            "SELECT id FROM purchase_followups WHERE id = ?",
            (followup_id,),
        ).fetchone()
        if followup is None:
            abort(404)
        conn.execute("DELETE FROM purchase_followups WHERE id = ?", (followup_id,))

    flash("其他物品采购记录已删除", "success")
    return redirect(url_for("admin_purchase_followups"))


def powder_coating_filters_from_request():
    query = request.args.get("q", "").strip()
    date_type = request.args.get("date_type", "delivered_at").strip()
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    sort = request.args.get("sort", "").strip()
    direction = request.args.get("direction", "desc").strip().lower()
    show_completed = request.args.get("show_completed") == "1"

    date_columns = {
        "delivered_at": "delivered_at",
        "received_at": "received_at",
        "created_at": "created_at",
    }
    sort_columns = {
        "delivered_at": "delivered_at",
        "received_at": "received_at",
        "product_name": "product_name COLLATE NOCASE",
        "drawing_no": "drawing_no COLLATE NOCASE",
        "supplier_name": "supplier_name COLLATE NOCASE",
        "quantity": "quantity",
        "unit_price": "unit_price",
        "total_price": "(quantity * unit_price)",
        "recorded_by": "recorded_by COLLATE NOCASE",
        "completed": "completed",
    }
    if date_type not in date_columns:
        date_type = "delivered_at"
    if sort not in sort_columns:
        sort = ""
    if direction not in {"asc", "desc"}:
        direction = "desc"

    params = []
    conditions = []
    if query:
        conditions.append(
            """
            (
                product_name LIKE ?
                OR drawing_no LIKE ?
                OR supplier_name LIKE ?
                OR remark LIKE ?
                OR recorded_by LIKE ?
            )
            """
        )
        like = f"%{query}%"
        params.extend([like, like, like, like, like])
    if start_date:
        conditions.append(f"{date_columns[date_type]} >= ?")
        params.append(start_date)
    if end_date:
        conditions.append(f"{date_columns[date_type]} <= ?")
        params.append(end_date)
    if not show_completed:
        conditions.append("completed = 0")

    if sort:
        sql_direction = "ASC" if direction == "asc" else "DESC"
        order_sql = f" ORDER BY {sort_columns[sort]} {sql_direction}, id DESC"
    else:
        order_sql = " ORDER BY delivered_at DESC, id DESC"

    return {
        "query": query,
        "date_type": date_type,
        "start_date": start_date,
        "end_date": end_date,
        "show_completed": show_completed,
        "sort": sort,
        "direction": direction,
        "conditions": conditions,
        "params": params,
        "where_sql": " WHERE " + " AND ".join(conditions) if conditions else "",
        "order_sql": order_sql,
    }


def fetch_powder_coating_records(conn, filters):
    return conn.execute(
        f"SELECT * FROM powder_coating_records{filters['where_sql']}{filters['order_sql']}",
        filters["params"],
    ).fetchall()


def fetch_powder_coating_summary(conn, filters):
    return conn.execute(
        f"""
        SELECT
            COUNT(*) AS record_count,
            COALESCE(SUM(quantity), 0) AS total_quantity,
            COALESCE(SUM(quantity * unit_price), 0) AS total_amount
        FROM powder_coating_records
        {filters['where_sql']}
        """,
        filters["params"],
    ).fetchone()


def powder_coating_filter_text(filters):
    labels = {
        "delivered_at": "送货时间",
        "received_at": "送来时间",
        "created_at": "入录时间",
    }
    parts = []
    if filters["query"]:
        parts.append(f"搜索：{filters['query']}")
    if filters["start_date"] or filters["end_date"]:
        start = filters["start_date"] or "不限"
        end = filters["end_date"] or "不限"
        parts.append(f"{labels[filters['date_type']]}：{start} 至 {end}")
    return "；".join(parts) if parts else "全部记录"


def build_powder_coating_workbook(records, summary, filters):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "喷塑对账单"
    worksheet.freeze_panes = "A5"

    worksheet["A1"] = "喷塑对账单"
    worksheet["A2"] = f"筛选条件：{powder_coating_filter_text(filters)}"
    worksheet["A3"] = f"制单时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

    headers = ["序号", "送货时间", "送来时间", "供应商", "产品名称", "图号", "数量", "单价", "金额", "备注"]
    worksheet.append([])
    worksheet.append(headers)
    header_row = 5
    for cell in worksheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155E63")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, item in enumerate(records, start=1):
        unit_price = float(item["unit_price"] or 0)
        quantity = int(item["quantity"] or 0)
        worksheet.append(
            [
                index,
                item["delivered_at"] or "",
                item["received_at"] or "",
                item["supplier_name"] or "",
                item["product_name"] or "",
                item["drawing_no"] or "",
                quantity,
                unit_price,
                quantity * unit_price,
                item["remark"] or "",
            ]
        )

    total_row = worksheet.max_row + 1
    worksheet.append(["", "", "", "", "合计", "", summary["total_quantity"] or 0, "", summary["total_amount"] or 0, ""])
    for cell in worksheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8F3F1")

    widths = [8, 14, 14, 18, 24, 20, 10, 12, 14, 30]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    for row in worksheet.iter_rows(min_row=6, max_row=worksheet.max_row, min_col=8, max_col=9):
        for cell in row:
            cell.number_format = '0.00'

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def powder_pdf_image_cell(filename, max_width=18 * mm, max_height=14 * mm):
    filename = str(filename or "").strip()
    if not filename:
        return "无图"
    image_path = MANUALS_DIR / filename
    if not image_path.exists():
        return "无图"
    try:
        reader = ImageReader(str(image_path))
        image_width, image_height = reader.getSize()
        if not image_width or not image_height:
            return "无图"
        ratio = min(max_width / image_width, max_height / image_height)
        return Image(str(image_path), width=image_width * ratio, height=image_height * ratio)
    except Exception:
        return "无图"


def build_powder_coating_purchase_order_pdf(records):
    records = list(records)
    if not records:
        abort(404)
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    font_name = get_pdf_font_name()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PowderPurchaseOrderTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=20,
        leading=24,
        alignment=1,
        spaceAfter=5 * mm,
    )
    cell_style = ParagraphStyle(
        "PowderPurchaseOrderCell",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=9,
        leading=13,
        wordWrap="CJK",
    )
    latin_style = ParagraphStyle(
        "PowderPurchaseOrderLatin",
        parent=cell_style,
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=13,
        splitLongWords=0,
    )
    small_style = ParagraphStyle(
        "PowderPurchaseOrderSmall",
        parent=cell_style,
        fontSize=8.5,
        leading=11,
    )
    supplier_names = sorted({record["supplier_name"] for record in records if record["supplier_name"]})
    delivered_dates = sorted({record["delivered_at"] for record in records if record["delivered_at"]})
    received_dates = sorted({record["received_at"] for record in records if record["received_at"]})
    supplier_text = supplier_names[0] if len(supplier_names) == 1 else "多个供应商"
    delivered_text = delivered_dates[0] if len(delivered_dates) == 1 else "多个日期"
    received_text = received_dates[0] if len(received_dates) == 1 else ("多个日期" if received_dates else "-")
    order_no = f"PC-B{datetime.now().strftime('%Y%m%d%H%M')}"
    total_quantity = sum(int(record["quantity"] or 0) for record in records)
    total_amount = sum(int(record["quantity"] or 0) * float(record["unit_price"] or 0) for record in records)

    story = [Paragraph("喷塑采购单", title_style)]
    meta = [
        [Paragraph("采购单号", cell_style), Paragraph(order_no, cell_style), Paragraph("制单时间", cell_style), Paragraph(datetime.now().strftime("%Y-%m-%d %H:%M:%S"), cell_style)],
        [Paragraph("送货时间", cell_style), Paragraph(delivered_text or "-", cell_style), Paragraph("送来时间", cell_style), Paragraph(received_text or "-", cell_style)],
        [Paragraph("供应商", cell_style), Paragraph(supplier_text or "-", cell_style), Paragraph("产品项数", cell_style), Paragraph(str(len(records)), cell_style)],
    ]
    meta_table = Table(meta, colWidths=[32 * mm, 92 * mm, 32 * mm, 92 * mm])
    meta_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#6B7C87")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EFF3")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E8EFF3")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([meta_table, Spacer(1, 6 * mm)])

    data = [["图片", "产品名称", "图号", "供应商", "数量", "单价", "金额", "送货时间", "送来时间", "备注"]]
    for record in records:
        quantity = int(record["quantity"] or 0)
        unit_price = float(record["unit_price"] or 0)
        product_name = str(record["product_name"] or "")
        drawing_no = str(record["drawing_no"] or "")
        product_style = latin_style if product_name and all(ord(char) < 128 for char in product_name) else cell_style
        drawing_style = latin_style if drawing_no and all(ord(char) < 128 for char in drawing_no) else cell_style
        data.append([
            powder_pdf_image_cell(record["image_filename"]),
            pdf_single_line_paragraph(product_name, product_style),
            pdf_single_line_paragraph(drawing_no, drawing_style),
            pdf_wrapped_paragraph(record["supplier_name"], cell_style, chunk_size=10),
            str(quantity),
            f"{unit_price:.2f}",
            f"{quantity * unit_price:.2f}",
            record["delivered_at"] or "",
            record["received_at"] or "",
            pdf_wrapped_paragraph(record["remark"], cell_style, chunk_size=14),
        ])
    data.append(["", "", "", "合计", str(total_quantity), "", f"{total_amount:.2f}", "", "", ""])
    table = Table(data, colWidths=[20 * mm, 46 * mm, 32 * mm, 30 * mm, 14 * mm, 16 * mm, 20 * mm, 21 * mm, 21 * mm, 49 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155E63")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F3F1")),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#6B7C87")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (4, 1), (6, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([table, Spacer(1, 14 * mm)])
    sign_table = Table(
        [[Paragraph("采购确认：", small_style), "", Paragraph("供应商确认：", small_style), ""]],
        colWidths=[32 * mm, 92 * mm, 32 * mm, 92 * mm],
        rowHeights=[18 * mm],
    )
    sign_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#A0ADB5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(sign_table)
    doc.build(story)
    buffer.seek(0)
    return buffer


def build_powder_coating_pdf(records, summary, filters):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=9 * mm,
        bottomMargin=9 * mm,
    )
    font_name = get_pdf_font_name()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PowderStatementTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=22,
        alignment=1,
        spaceAfter=3 * mm,
    )
    info_style = ParagraphStyle(
        "PowderStatementInfo",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8.5,
        leading=10,
    )
    cell_style = ParagraphStyle(
        "PowderStatementCell",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=7.5,
        leading=8.5,
        wordWrap="CJK",
    )
    story = [Paragraph("喷塑对账单", title_style)]
    meta_data = [
        [Paragraph("筛选条件", info_style), Paragraph(powder_coating_filter_text(filters), info_style), Paragraph("制单时间", info_style), Paragraph(datetime.now().strftime("%Y-%m-%d %H:%M:%S"), info_style)],
        [Paragraph("记录数", info_style), Paragraph(str(summary["record_count"] or 0), info_style), Paragraph("合计金额", info_style), Paragraph(f"{float(summary['total_amount'] or 0):.2f}", info_style)],
    ]
    meta_table = Table(meta_data, colWidths=[22 * mm, 138 * mm, 22 * mm, 84 * mm])
    meta_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#6B7C87")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EFF3")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E8EFF3")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.extend([meta_table, Spacer(1, 4 * mm)])

    data = [["序号", "图片", "送货时间", "送来时间", "供应商", "产品名称", "图号", "数量", "单价", "金额", "备注"]]
    for index, item in enumerate(records, start=1):
        quantity = int(item["quantity"] or 0)
        unit_price = float(item["unit_price"] or 0)
        data.append(
            [
                str(index),
                powder_pdf_image_cell(item["image_filename"], max_width=14 * mm, max_height=12 * mm),
                item["delivered_at"] or "",
                item["received_at"] or "",
                Paragraph(item["supplier_name"] or "", cell_style),
                Paragraph(item["product_name"] or "", cell_style),
                Paragraph(item["drawing_no"] or "", cell_style),
                str(quantity),
                f"{unit_price:.2f}",
                f"{quantity * unit_price:.2f}",
                Paragraph(item["remark"] or "", cell_style),
            ]
        )
    data.append(["", "", "", "", "", "合计", "", str(summary["total_quantity"] or 0), "", f"{float(summary['total_amount'] or 0):.2f}", ""])

    table = Table(data, colWidths=[8 * mm, 16 * mm, 19 * mm, 19 * mm, 25 * mm, 36 * mm, 27 * mm, 13 * mm, 15 * mm, 18 * mm, 70 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155E63")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F3F1")),
                ("FONTNAME", (0, -1), (-1, -1), font_name),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#8AA0A8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("ALIGN", (7, 1), (9, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    buffer.seek(0)
    return buffer


@app.route("/admin/powder-coating", methods=["GET", "POST"])
@permission_required("powder_coating")
def admin_powder_coating():
    if request.method == "POST":
        product_id = request.form.get("product_id", "").strip()
        quantity = request.form.get("quantity", "").strip()
        supplier_name = request.form.get("supplier_name", "").strip()
        delivered_at = request.form.get("delivered_at", "").strip()
        received_at = request.form.get("received_at", "").strip()
        remark = request.form.get("remark", "").strip()

        if not all([product_id, delivered_at]):
            flash("产品和送货时间为必填项", "error")
            return redirect(url_for("admin_powder_coating"))
        try:
            product_id_value = int(product_id)
            quantity_value = parse_positive_int(quantity, "数量") if quantity else 0
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("admin_powder_coating"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        with get_db() as conn:
            product = conn.execute(
                """
                SELECT *
                FROM powder_coating_products
                WHERE id = ?
                """,
                (product_id_value,),
            ).fetchone()
            if product is None:
                flash("请选择产品库里的喷塑产品", "error")
                return redirect(url_for("admin_powder_coating"))
            selected_supplier_name = supplier_name or product["supplier_name"] or ""

            conn.execute(
                """
                INSERT INTO powder_coating_records (
                    product_id, product_name, drawing_no, supplier_name, unit_price, quantity,
                    delivered_at, received_at, remark, recorded_by,
                    image_filename, image_original_filename, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    product["id"],
                    product["product_name"],
                    product["drawing_no"] or "",
                    selected_supplier_name,
                    product["unit_price"] or 0,
                    quantity_value,
                    delivered_at,
                    received_at,
                    remark,
                    current_admin_username(),
                    product["image_filename"] or "",
                    product["image_original_filename"] or "",
                    now,
                    now,
                ),
            )

        flash("喷塑记录已新增", "success")
        return redirect(url_for("admin_powder_coating"))

    filters = powder_coating_filters_from_request()

    with get_db() as conn:
        products = conn.execute(
            """
            SELECT *
            FROM powder_coating_products
            ORDER BY product_name COLLATE NOCASE ASC, drawing_no COLLATE NOCASE ASC, id DESC
            """
        ).fetchall()
        suppliers = conn.execute(
            """
            SELECT *
            FROM powder_coating_suppliers
            ORDER BY name COLLATE NOCASE ASC, id DESC
            """
        ).fetchall()
        records = fetch_powder_coating_records(conn, filters)
        summary = fetch_powder_coating_summary(conn, filters)

    return render_template(
        "powder_coating.html",
        products=products,
        suppliers=suppliers,
        supplier_names=[supplier["name"] for supplier in suppliers],
        records=records,
        summary=summary,
        query=filters["query"],
        date_type=filters["date_type"],
        start_date=filters["start_date"],
        end_date=filters["end_date"],
        show_completed=filters["show_completed"],
        sort=filters["sort"],
        direction=filters["direction"],
        default_delivered_at=datetime.now().date().isoformat(),
    )


@app.route("/admin/powder-coating/export")
@permission_required("powder_coating")
def export_powder_coating_statement():
    export_format = request.args.get("format", "pdf").strip().lower()
    filters = powder_coating_filters_from_request()
    with get_db() as conn:
        records = fetch_powder_coating_records(conn, filters)
        summary = fetch_powder_coating_summary(conn, filters)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if export_format == "xlsx":
        buffer = build_powder_coating_workbook(records, summary, filters)
        return send_file(
            buffer,
            as_attachment=True,
            download_name=f"powder-coating-statement-{timestamp}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    buffer = build_powder_coating_pdf(records, summary, filters)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"powder-coating-statement-{timestamp}.pdf",
        mimetype="application/pdf",
    )


@app.route("/admin/powder-coating/purchase-order", methods=["POST"])
@permission_required("powder_coating")
def selected_powder_coating_purchase_order_pdf():
    record_ids = []
    for value in request.form.getlist("record_id"):
        try:
            record_ids.append(int(value))
        except ValueError:
            continue
    record_ids = list(dict.fromkeys(record_ids))
    if not record_ids:
        flash("请先选择喷塑记录", "error")
        return redirect(url_for("admin_powder_coating"))
    placeholders = ",".join("?" for _ in record_ids)
    with get_db() as conn:
        records = conn.execute(
            f"""
            SELECT *
            FROM powder_coating_records
            WHERE id IN ({placeholders})
            ORDER BY delivered_at DESC, id DESC
            """,
            record_ids,
        ).fetchall()
    if not records:
        abort(404)
    buffer = build_powder_coating_purchase_order_pdf(records)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return send_file(
        buffer,
        as_attachment=False,
        download_name=f"powder-coating-purchase-order-{timestamp}.pdf",
        mimetype="application/pdf",
    )


@app.route("/admin/powder-coating/suppliers", methods=["POST"])
@permission_required("powder_coating")
def create_powder_coating_supplier():
    name = request.form.get("name", "").strip()
    contact = request.form.get("contact", "").strip()
    phone = request.form.get("phone", "").strip()
    remark = request.form.get("remark", "").strip()
    if not name:
        flash("供应商名称为必填项", "error")
        return redirect(url_for("admin_powder_coating"))
    now = datetime.utcnow().isoformat(timespec="seconds")
    try:
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO powder_coating_suppliers (
                    name, contact, phone, remark, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (name, contact, phone, remark, now, now),
            )
    except sqlite3.IntegrityError:
        flash("这个供应商已经存在", "error")
        return redirect(url_for("admin_powder_coating"))
    flash("供应商已新增", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/suppliers/<int:supplier_id>/edit", methods=["POST"])
@permission_required("powder_coating")
def edit_powder_coating_supplier(supplier_id):
    name = request.form.get("name", "").strip()
    contact = request.form.get("contact", "").strip()
    phone = request.form.get("phone", "").strip()
    remark = request.form.get("remark", "").strip()
    if not name:
        flash("供应商名称为必填项", "error")
        return redirect(url_for("admin_powder_coating"))
    now = datetime.utcnow().isoformat(timespec="seconds")
    try:
        with get_db() as conn:
            supplier = conn.execute(
                "SELECT name FROM powder_coating_suppliers WHERE id = ?",
                (supplier_id,),
            ).fetchone()
            if supplier is None:
                abort(404)
            old_name = supplier["name"] or ""
            conn.execute(
                """
                UPDATE powder_coating_suppliers
                SET name = ?, contact = ?, phone = ?, remark = ?, updated_at = ?
                WHERE id = ?
                """,
                (name, contact, phone, remark, now, supplier_id),
            )
            if old_name and old_name != name:
                conn.execute(
                    """
                    UPDATE powder_coating_records
                    SET supplier_name = ?, updated_at = ?
                    WHERE supplier_name = ?
                    """,
                    (name, now, old_name),
                )
                conn.execute(
                    """
                    UPDATE powder_coating_products
                    SET supplier_name = ?, updated_at = ?
                    WHERE supplier_name = ?
                    """,
                    (name, now, old_name),
                )
    except sqlite3.IntegrityError:
        flash("这个供应商名称已经存在", "error")
        return redirect(url_for("admin_powder_coating"))
    flash("供应商已更新", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/suppliers/<int:supplier_id>/delete", methods=["POST"])
@permission_required("powder_coating")
def delete_powder_coating_supplier(supplier_id):
    with get_db() as conn:
        supplier = conn.execute(
            "SELECT name FROM powder_coating_suppliers WHERE id = ?",
            (supplier_id,),
        ).fetchone()
        if supplier is None:
            abort(404)
        used_count = conn.execute(
            "SELECT COUNT(*) AS c FROM powder_coating_records WHERE supplier_name = ?",
            (supplier["name"] or "",),
        ).fetchone()["c"]
        product_count = conn.execute(
            "SELECT COUNT(*) AS c FROM powder_coating_products WHERE supplier_name = ?",
            (supplier["name"] or "",),
        ).fetchone()["c"]
        if used_count or product_count:
            flash("这个供应商已有喷塑产品或记录，不能删除", "error")
            return redirect(url_for("admin_powder_coating"))
        conn.execute("DELETE FROM powder_coating_suppliers WHERE id = ?", (supplier_id,))
    flash("供应商已删除", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/products", methods=["POST"])
@permission_required("powder_coating")
def create_powder_coating_product():
    product_name = request.form.get("product_name", "").strip()
    drawing_no = request.form.get("drawing_no", "").strip()
    supplier_name = request.form.get("supplier_name", "").strip()
    unit_price = request.form.get("unit_price", "").strip()
    image_upload = request.files.get("image")

    if not product_name:
        flash("产品名称为必填项", "error")
        return redirect(url_for("admin_powder_coating"))
    try:
        unit_price_value = parse_nonnegative_price(unit_price or "0")
        image_filename, image_original_filename = save_powder_coating_image(image_upload)
    except ValueError as error:
        flash(str(error) if str(error) else "请上传图片格式文件", "error")
        return redirect(url_for("admin_powder_coating"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        conn.execute(
            """
            INSERT INTO powder_coating_products (
                product_name, drawing_no, supplier_name, unit_price, image_filename,
                image_original_filename, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                product_name,
                drawing_no,
                supplier_name,
                unit_price_value,
                image_filename,
                image_original_filename,
                now,
                now,
            ),
        )

    flash("喷塑产品已加入产品库", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/products/<int:product_id>/edit", methods=["POST"])
@permission_required("powder_coating")
def edit_powder_coating_product(product_id):
    product_name = request.form.get("product_name", "").strip()
    drawing_no = request.form.get("drawing_no", "").strip()
    supplier_name = request.form.get("supplier_name", "").strip()
    unit_price = request.form.get("unit_price", "").strip()
    image_upload = request.files.get("image")

    if not product_name:
        flash("产品名称为必填项", "error")
        return redirect(url_for("admin_powder_coating"))
    try:
        unit_price_value = parse_nonnegative_price(unit_price or "0")
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("admin_powder_coating"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        product = conn.execute(
            "SELECT id, image_filename FROM powder_coating_products WHERE id = ?",
            (product_id,),
        ).fetchone()
        if product is None:
            abort(404)

        image_filename = ""
        image_original_filename = ""
        if image_upload and (image_upload.filename or image_upload.mimetype):
            try:
                image_filename, image_original_filename = save_powder_coating_image(image_upload)
            except ValueError:
                flash("请上传图片格式文件", "error")
                return redirect(url_for("admin_powder_coating"))
            if product["image_filename"]:
                delete_upload_file(product["image_filename"])

        conn.execute(
            """
            UPDATE powder_coating_products
            SET product_name = ?, drawing_no = ?, supplier_name = ?, unit_price = ?,
                image_filename = CASE WHEN ? != '' THEN ? ELSE image_filename END,
                image_original_filename = CASE WHEN ? != '' THEN ? ELSE image_original_filename END,
                updated_at = ?
            WHERE id = ?
            """,
            (
                product_name,
                drawing_no,
                supplier_name,
                unit_price_value,
                image_filename,
                image_filename,
                image_original_filename,
                image_original_filename,
                now,
                product_id,
            ),
        )

    flash("喷塑产品已更新", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/products/<int:product_id>/image/delete", methods=["POST"])
@permission_required("powder_coating")
def delete_powder_coating_product_image(product_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        product = conn.execute(
            "SELECT id, image_filename FROM powder_coating_products WHERE id = ?",
            (product_id,),
        ).fetchone()
        if product is None:
            abort(404)
        if product["image_filename"]:
            delete_upload_file(product["image_filename"])
        conn.execute(
            """
            UPDATE powder_coating_products
            SET image_filename = '', image_original_filename = '', updated_at = ?
            WHERE id = ?
            """,
            (now, product_id),
        )

    flash("产品图片已删除", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/products/<int:product_id>/copy", methods=["POST"])
@permission_required("powder_coating")
def copy_powder_coating_product(product_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        product = conn.execute(
            "SELECT * FROM powder_coating_products WHERE id = ?",
            (product_id,),
        ).fetchone()
        if product is None:
            abort(404)
        image_filename = ""
        image_original_filename = ""
        if product["image_filename"]:
            try:
                image_filename, image_original_filename, _ = copy_upload_file(
                    product["image_filename"],
                    product["image_original_filename"] or product["image_filename"],
                )
            except FileNotFoundError:
                image_filename = ""
                image_original_filename = ""
        conn.execute(
            """
            INSERT INTO powder_coating_products (
                product_name, drawing_no, supplier_name, unit_price, image_filename,
                image_original_filename, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f"{product['product_name']} - 副本",
                product["drawing_no"] or "",
                product["supplier_name"] or "",
                product["unit_price"] or 0,
                image_filename,
                image_original_filename,
                now,
                now,
            ),
        )
    flash("喷塑产品已复制", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/products/<int:product_id>/delete", methods=["POST"])
@permission_required("powder_coating")
def delete_powder_coating_product(product_id):
    with get_db() as conn:
        product = conn.execute(
            "SELECT id, image_filename FROM powder_coating_products WHERE id = ?",
            (product_id,),
        ).fetchone()
        if product is None:
            abort(404)
        record_count = conn.execute(
            "SELECT COUNT(*) AS c FROM powder_coating_records WHERE product_id = ?",
            (product_id,),
        ).fetchone()["c"]
        if record_count:
            flash("这个产品已有喷塑记录，不能删除", "error")
            return redirect(url_for("admin_powder_coating"))
        if product["image_filename"]:
            delete_upload_file(product["image_filename"])
        conn.execute("DELETE FROM powder_coating_products WHERE id = ?", (product_id,))

    flash("喷塑产品已删除", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/<int:record_id>/edit", methods=["POST"])
@permission_required("powder_coating")
def edit_powder_coating_record(record_id):
    product_name = request.form.get("product_name", "").strip()
    drawing_no = request.form.get("drawing_no", "").strip()
    supplier_name = request.form.get("supplier_name", "").strip()
    unit_price = request.form.get("unit_price", "").strip()
    quantity = request.form.get("quantity", "").strip()
    delivered_at = request.form.get("delivered_at", "").strip()
    received_at = request.form.get("received_at", "").strip()
    remark = request.form.get("remark", "").strip()

    if not all([product_name, delivered_at]):
        flash("产品名称和送货时间为必填项", "error")
        return redirect(url_for("admin_powder_coating"))
    try:
        quantity_value = parse_positive_int(quantity, "数量") if quantity else 0
        unit_price_value = parse_nonnegative_price(unit_price or "0")
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("admin_powder_coating"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        record = conn.execute(
            "SELECT id FROM powder_coating_records WHERE id = ?",
            (record_id,),
        ).fetchone()
        if record is None:
            abort(404)
        conn.execute(
            """
            UPDATE powder_coating_records
            SET product_name = ?, drawing_no = ?, supplier_name = ?, unit_price = ?, quantity = ?,
                delivered_at = ?, received_at = ?, remark = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                product_name,
                drawing_no,
                supplier_name,
                unit_price_value,
                quantity_value,
                delivered_at,
                received_at,
                remark,
                now,
                record_id,
            ),
        )

    flash("喷塑记录已更新", "success")
    return redirect(request.referrer or url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/<int:record_id>/copy", methods=["POST"])
@permission_required("powder_coating")
def copy_powder_coating_record(record_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    delivered_at = datetime.now().date().isoformat()
    with get_db() as conn:
        source = conn.execute(
            """
            SELECT product_id, product_name, drawing_no, unit_price, quantity,
                   supplier_name, received_at, remark, image_filename, image_original_filename
            FROM powder_coating_records
            WHERE id = ?
            """,
            (record_id,),
        ).fetchone()
        if source is None:
            abort(404)
        conn.execute(
            """
            INSERT INTO powder_coating_records (
                product_id, product_name, drawing_no, supplier_name, unit_price, quantity,
                delivered_at, received_at, remark, recorded_by,
                image_filename, image_original_filename, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, '', ?, ?, ?, ?, ?, ?)
            """,
            (
                source["product_id"],
                source["product_name"],
                source["drawing_no"] or "",
                source["supplier_name"] or "",
                source["unit_price"] or 0,
                source["quantity"] or 0,
                delivered_at,
                source["remark"] or "",
                current_admin_username(),
                source["image_filename"] or "",
                source["image_original_filename"] or "",
                now,
                now,
            ),
        )

    flash("喷塑记录已复制", "success")
    return redirect(url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/<int:record_id>/complete", methods=["POST"])
@permission_required("powder_coating")
def toggle_powder_coating_completed(record_id):
    completed = 1 if request.form.get("completed") == "1" else 0
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        record = conn.execute(
            "SELECT id FROM powder_coating_records WHERE id = ?",
            (record_id,),
        ).fetchone()
        if record is None:
            abort(404)
        conn.execute(
            """
            UPDATE powder_coating_records
            SET completed = ?, updated_at = ?
            WHERE id = ?
            """,
            (completed, now, record_id),
        )
    flash("喷塑记录已完成" if completed else "喷塑记录已取消完成", "success")
    return redirect(request.referrer or url_for("admin_powder_coating"))


@app.route("/admin/powder-coating/<int:record_id>/delete", methods=["POST"])
@permission_required("powder_coating")
def delete_powder_coating_record(record_id):
    with get_db() as conn:
        record = conn.execute(
            "SELECT id FROM powder_coating_records WHERE id = ?",
            (record_id,),
        ).fetchone()
        if record is None:
            abort(404)
        conn.execute("DELETE FROM powder_coating_records WHERE id = ?", (record_id,))

    flash("喷塑记录已删除", "success")
    return redirect(url_for("admin_powder_coating"))


def carton_purchase_filters_from_request():
    query = request.args.get("q", "").strip()
    date_type = request.args.get("date_type", "ordered_at").strip()
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    sort = request.args.get("sort", "").strip()
    direction = request.args.get("direction", "desc").strip().lower()
    show_completed = request.args.get("show_completed") == "1"

    date_columns = {
        "ordered_at": "ordered_at",
        "received_at": "received_at",
        "created_at": "created_at",
    }
    sort_columns = {
        "ordered_at": "ordered_at",
        "print_mark": "print_mark COLLATE NOCASE",
        "supplier_name": "supplier_name COLLATE NOCASE",
        "board_type": "board_type COLLATE NOCASE",
        "quantity": "quantity",
        "unit_price": "unit_price",
        "total_price": "(quantity * unit_price)",
        "carton_size": "carton_size COLLATE NOCASE",
        "received_at": "received_at",
        "remark": "remark COLLATE NOCASE",
        "recorded_by": "recorded_by COLLATE NOCASE",
        "completed": "completed",
    }
    if date_type not in date_columns:
        date_type = "ordered_at"
    if sort not in sort_columns:
        sort = ""
    if direction not in {"asc", "desc"}:
        direction = "desc"

    params = []
    conditions = []
    if query:
        conditions.append(
            """
            (
                print_mark LIKE ?
                OR carton_size LIKE ?
                OR supplier_name LIKE ?
                OR board_type LIKE ?
                OR remark LIKE ?
                OR recorded_by LIKE ?
            )
            """
        )
        like = f"%{query}%"
        params.extend([like, like, like, like, like, like])
    if start_date:
        conditions.append(f"{date_columns[date_type]} >= ?")
        params.append(start_date)
    if end_date:
        conditions.append(f"{date_columns[date_type]} <= ?")
        params.append(end_date)
    if not show_completed:
        conditions.append("completed = 0")

    if sort:
        sql_direction = "ASC" if direction == "asc" else "DESC"
        order_sql = f" ORDER BY {sort_columns[sort]} {sql_direction}, id DESC"
    else:
        order_sql = " ORDER BY ordered_at DESC, id DESC"

    return {
        "query": query,
        "date_type": date_type,
        "start_date": start_date,
        "end_date": end_date,
        "show_completed": show_completed,
        "sort": sort,
        "direction": direction,
        "where_sql": " WHERE " + " AND ".join(conditions) if conditions else "",
        "order_sql": order_sql,
        "params": params,
    }


def fetch_carton_purchase_records(conn, filters):
    return conn.execute(
        f"SELECT * FROM carton_purchases{filters['where_sql']}{filters['order_sql']}",
        filters["params"],
    ).fetchall()


def fetch_carton_purchase_summary(conn, filters):
    return conn.execute(
        f"""
        SELECT COUNT(*) AS record_count,
               COALESCE(SUM(quantity), 0) AS total_quantity,
               COALESCE(SUM(quantity * unit_price), 0) AS total_amount
        FROM carton_purchases
        {filters['where_sql']}
        """,
        filters["params"],
    ).fetchone()


def carton_purchase_filter_text(filters):
    labels = {
        "ordered_at": "下单时间",
        "received_at": "送来时间",
        "created_at": "录入时间",
    }
    parts = []
    if filters["query"]:
        parts.append(f"搜索：{filters['query']}")
    if filters["start_date"] or filters["end_date"]:
        start = filters["start_date"] or "不限"
        end = filters["end_date"] or "不限"
        parts.append(f"{labels[filters['date_type']]}：{start} 至 {end}")
    return "；".join(parts) if parts else "全部记录"


def build_carton_purchase_statement_pdf(records, summary, filters):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=9 * mm,
        bottomMargin=9 * mm,
    )
    font_name = get_pdf_font_name()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CartonStatementTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=22,
        alignment=1,
        spaceAfter=3 * mm,
    )
    info_style = ParagraphStyle(
        "CartonStatementInfo",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8.5,
        leading=10,
    )
    cell_style = ParagraphStyle(
        "CartonStatementCell",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=7.5,
        leading=8.5,
        wordWrap="CJK",
    )
    story = [Paragraph("纸箱采购对账单", title_style)]
    meta_data = [
        [Paragraph("筛选条件", info_style), Paragraph(carton_purchase_filter_text(filters), info_style), Paragraph("制单时间", info_style), Paragraph(datetime.now().strftime("%Y-%m-%d %H:%M:%S"), info_style)],
        [Paragraph("记录数", info_style), Paragraph(str(summary["record_count"] or 0), info_style), Paragraph("合计金额", info_style), Paragraph(f"{float(summary['total_amount'] or 0):.2f}", info_style)],
    ]
    meta_table = Table(meta_data, colWidths=[22 * mm, 138 * mm, 22 * mm, 84 * mm])
    meta_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#6B7C87")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EFF3")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E8EFF3")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.extend([meta_table, Spacer(1, 4 * mm)])

    data = [["序号", "下单时间", "送来时间", "供应商", "印刷唛头", "纸板类型", "箱子尺寸CM", "数量", "单价", "金额", "备注"]]
    for index, item in enumerate(records, start=1):
        quantity = int(item["quantity"] or 0)
        unit_price = float(item["unit_price"] or 0)
        data.append(
            [
                str(index),
                item["ordered_at"] or "",
                item["received_at"] or "",
                Paragraph(item["supplier_name"] or "", cell_style),
                Paragraph(item["print_mark"] or "", cell_style),
                Paragraph(item["board_type"] or "", cell_style),
                Paragraph(item["carton_size"] or "", cell_style),
                str(quantity),
                f"{unit_price:.2f}",
                f"{quantity * unit_price:.2f}",
                Paragraph(item["remark"] or "", cell_style),
            ]
        )
    data.append(["", "", "", "", "合计", "", "", str(summary["total_quantity"] or 0), "", f"{float(summary['total_amount'] or 0):.2f}", ""])

    table = Table(data, colWidths=[9 * mm, 19 * mm, 19 * mm, 24 * mm, 42 * mm, 22 * mm, 28 * mm, 13 * mm, 15 * mm, 18 * mm, 57 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155E63")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F3F1")),
                ("FONTNAME", (0, -1), (-1, -1), font_name),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#8AA0A8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("ALIGN", (7, 1), (9, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    buffer.seek(0)
    return buffer


def build_carton_purchase_order_pdf(records):
    if hasattr(records, "keys"):
        records = [records]
    records = list(records)
    if not records:
        abort(404)
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )
    font_name = get_pdf_font_name()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CartonPurchaseOrderTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=20,
        leading=24,
        alignment=1,
        spaceAfter=5 * mm,
    )
    cell_style = ParagraphStyle(
        "CartonPurchaseOrderCell",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8.5,
        leading=11,
        wordWrap="CJK",
    )
    mark_style = ParagraphStyle(
        "CartonPurchaseOrderMark",
        parent=cell_style,
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=11,
        splitLongWords=0,
        spaceBefore=1,
        spaceAfter=1,
    )
    mark_cjk_style = ParagraphStyle(
        "CartonPurchaseOrderMarkCJK",
        parent=cell_style,
        fontSize=9.5,
        leading=11,
        splitLongWords=0,
        spaceBefore=1,
        spaceAfter=1,
    )
    size_style = ParagraphStyle(
        "CartonPurchaseOrderSize",
        parent=cell_style,
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=11,
        splitLongWords=0,
    )
    small_style = ParagraphStyle(
        "CartonPurchaseOrderSmall",
        parent=cell_style,
        fontSize=8.5,
        leading=11,
    )
    first_record = records[0]
    supplier_names = sorted({record["supplier_name"] for record in records if record["supplier_name"]})
    ordered_dates = sorted({record["ordered_at"] for record in records if record["ordered_at"]})
    received_dates = sorted({record["received_at"] for record in records if record["received_at"]})
    order_no = f"CT-{first_record['id']:05d}" if len(records) == 1 else f"CT-B{datetime.now().strftime('%Y%m%d%H%M')}"
    supplier_text = supplier_names[0] if len(supplier_names) == 1 else "多个供应商"
    ordered_text = ordered_dates[0] if len(ordered_dates) == 1 else "多个日期"
    received_text = received_dates[0] if len(received_dates) == 1 else ("多个日期" if received_dates else "-")
    story = [Paragraph("纸箱采购单", title_style)]
    meta = [
        [Paragraph("采购单号", cell_style), Paragraph(order_no, cell_style), Paragraph("制单时间", cell_style), Paragraph(datetime.now().strftime("%Y-%m-%d %H:%M:%S"), cell_style)],
        [Paragraph("下单时间", cell_style), Paragraph(ordered_text or "-", cell_style), Paragraph("送来时间", cell_style), Paragraph(received_text or "-", cell_style)],
        [Paragraph("供应商", cell_style), Paragraph(supplier_text or "-", cell_style), Paragraph("产品项数", cell_style), Paragraph(str(len(records)), cell_style)],
    ]
    meta_table = Table(meta, colWidths=[32 * mm, 92 * mm, 32 * mm, 92 * mm])
    meta_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#6B7C87")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EFF3")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E8EFF3")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.extend([meta_table, Spacer(1, 6 * mm)])

    details = [
        ["麦头", "纸板类型", "尺寸mm", "纸箱数量", "单价", "总价", "备注"],
    ]
    total_quantity = 0
    total_amount = 0
    for record in records:
        quantity = int(record["quantity"] or 0)
        unit_price = float(record["unit_price"] or 0)
        total_price = quantity * unit_price
        total_quantity += quantity
        total_amount += total_price
        carton_size = (record["carton_size"] or "").replace("×", "x")
        details.append(
            [
                pdf_single_line_paragraph(
                    record["print_mark"],
                    mark_cjk_style if any(ord(char) > 127 for char in str(record["print_mark"] or "")) else mark_style,
                ),
                pdf_wrapped_paragraph(record["board_type"], cell_style, chunk_size=6),
                pdf_single_line_paragraph(carton_size, size_style),
                str(quantity),
                f"{unit_price:.2f}",
                f"{total_price:.2f}",
                pdf_wrapped_paragraph(record["remark"], cell_style, chunk_size=12),
            ]
        )
    details.append(["", "", "合计", str(total_quantity), "", f"{total_amount:.2f}", ""])
    detail_table = Table(details, colWidths=[68 * mm, 14 * mm, 38 * mm, 18 * mm, 16 * mm, 18 * mm, 58 * mm], repeatRows=1, hAlign="LEFT")
    detail_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155E63")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F3F1")),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#6B7C87")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
                ("ALIGN", (3, 1), (5, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([detail_table, Spacer(1, 16 * mm)])

    sign_table = Table(
        [
            [Paragraph("采购确认：", small_style), "", Paragraph("供应商确认：", small_style), ""],
        ],
        colWidths=[32 * mm, 92 * mm, 32 * mm, 92 * mm],
        rowHeights=[18 * mm],
    )
    sign_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#A0ADB5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(sign_table)
    doc.build(story)
    buffer.seek(0)
    return buffer


@app.route("/admin/carton-purchases", methods=["GET", "POST"])
@permission_required("carton_purchases")
def admin_carton_purchases():
    if request.method == "POST":
        ordered_at = request.form.get("ordered_at", "").strip()
        received_at = request.form.get("received_at", "").strip()
        shared_supplier_name = request.form.get("supplier_name", "").strip()

        product_ids = [value.strip() for value in request.form.getlist("product_id")]
        quantities = [value.strip() for value in request.form.getlist("quantity")]
        board_types = [value.strip() for value in request.form.getlist("board_type")]
        remarks = [value.strip() for value in request.form.getlist("remark")]

        if not ordered_at:
            flash("下单时间为必填项", "error")
            return redirect(url_for("admin_carton_purchases"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        created_count = 0
        batch_supplier_name = shared_supplier_name
        with get_db() as conn:
            for product_id, quantity, board_type, remark in zip_longest(
                product_ids, quantities, board_types, remarks, fillvalue=""
            ):
                if not any([product_id, quantity, remark]):
                    continue
                if not product_id or not quantity:
                    flash("每一行都需要选择纸箱产品并填写数量", "error")
                    return redirect(url_for("admin_carton_purchases"))
                try:
                    product_id_value = int(product_id)
                    quantity_value = parse_positive_int(quantity, "数量")
                except ValueError as error:
                    flash(str(error), "error")
                    return redirect(url_for("admin_carton_purchases"))

                product = conn.execute(
                    "SELECT * FROM carton_products WHERE id = ?",
                    (product_id_value,),
                ).fetchone()
                if product is None:
                    flash("请选择纸箱产品库里的产品", "error")
                    return redirect(url_for("admin_carton_purchases"))
                selected_board_type = board_type or product["board_type"] or ""
                product_supplier_name = product["supplier_name"] or ""
                if not batch_supplier_name:
                    batch_supplier_name = product_supplier_name
                if batch_supplier_name and product_supplier_name and product_supplier_name != batch_supplier_name:
                    flash("一次新增纸箱采购只能选择同一个供应商的产品", "error")
                    return redirect(url_for("admin_carton_purchases"))
                selected_supplier_name = batch_supplier_name or product_supplier_name
                unit_price_value = calculate_carton_unit_price(
                    product["carton_length"] or 0,
                    product["carton_width"] or 0,
                    product["carton_height"] or 0,
                    carton_board_price(product, selected_board_type),
                )
                size_text = carton_size_text(product["carton_length"], product["carton_width"], product["carton_height"]) or product["carton_size"] or ""
                conn.execute(
                    """
                    INSERT INTO carton_purchases (
                        product_id, ordered_at, print_mark, supplier_name, board_type, quantity, unit_price, carton_size, received_at,
                        carton_length, carton_width, carton_height, remark, recorded_by, created_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        product["id"],
                        ordered_at,
                        product["print_mark"] or "",
                        selected_supplier_name,
                        selected_board_type,
                        quantity_value,
                        unit_price_value,
                        size_text,
                        received_at,
                        product["carton_length"] or 0,
                        product["carton_width"] or 0,
                        product["carton_height"] or 0,
                        remark,
                        current_admin_username(),
                        now,
                        now,
                    ),
                )
                created_count += 1

        if not created_count:
            flash("请至少选择一个纸箱产品", "error")
            return redirect(url_for("admin_carton_purchases"))

        flash(f"纸箱采购记录已新增：{created_count} 条", "success")
        return redirect(url_for("admin_carton_purchases"))

    filters = carton_purchase_filters_from_request()
    with get_db() as conn:
        products = conn.execute(
            """
            SELECT *
            FROM carton_products
            ORDER BY print_mark COLLATE NOCASE ASC, carton_size COLLATE NOCASE ASC, id DESC
            """
        ).fetchall()
        suppliers = conn.execute(
            """
            SELECT *
            FROM carton_suppliers
            ORDER BY name COLLATE NOCASE ASC, id DESC
            """
        ).fetchall()
        records = fetch_carton_purchase_records(conn, filters)
        summary = fetch_carton_purchase_summary(conn, filters)

    return render_template(
        "carton_purchases.html",
        products=products,
        suppliers=suppliers,
        supplier_names=[supplier["name"] for supplier in suppliers],
        carton_board_type_options=CARTON_BOARD_TYPE_OPTIONS,
        records=records,
        summary=summary,
        query=filters["query"],
        date_type=filters["date_type"],
        start_date=filters["start_date"],
        end_date=filters["end_date"],
        show_completed=filters["show_completed"],
        sort=filters["sort"],
        direction=filters["direction"],
        default_ordered_at=datetime.now().date().isoformat(),
    )


@app.route("/admin/carton-purchases/statement")
@permission_required("carton_purchases")
def carton_purchase_statement_pdf():
    filters = carton_purchase_filters_from_request()
    with get_db() as conn:
        records = fetch_carton_purchase_records(conn, filters)
        summary = fetch_carton_purchase_summary(conn, filters)

    buffer = build_carton_purchase_statement_pdf(records, summary, filters)
    as_attachment = request.args.get("download") == "1"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return send_file(
        buffer,
        as_attachment=as_attachment,
        download_name=f"carton-purchase-statement-{timestamp}.pdf",
        mimetype="application/pdf",
    )


@app.route("/admin/carton-suppliers", methods=["POST"])
@permission_required("carton_purchases")
def create_carton_supplier():
    name = request.form.get("name", "").strip()
    contact = request.form.get("contact", "").strip()
    phone = request.form.get("phone", "").strip()
    remark = request.form.get("remark", "").strip()
    if not name:
        flash("供应商名称为必填项", "error")
        return redirect(url_for("admin_carton_purchases"))
    now = datetime.utcnow().isoformat(timespec="seconds")
    try:
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO carton_suppliers (
                    name, contact, phone, remark, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (name, contact, phone, remark, now, now),
            )
    except sqlite3.IntegrityError:
        flash("这个供应商已经存在", "error")
        return redirect(url_for("admin_carton_purchases"))
    flash("纸箱供应商已新增", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/carton-suppliers/<int:supplier_id>/edit", methods=["POST"])
@permission_required("carton_purchases")
def edit_carton_supplier(supplier_id):
    name = request.form.get("name", "").strip()
    contact = request.form.get("contact", "").strip()
    phone = request.form.get("phone", "").strip()
    remark = request.form.get("remark", "").strip()
    if not name:
        flash("供应商名称为必填项", "error")
        return redirect(url_for("admin_carton_purchases"))
    now = datetime.utcnow().isoformat(timespec="seconds")
    try:
        with get_db() as conn:
            supplier = conn.execute(
                "SELECT name FROM carton_suppliers WHERE id = ?",
                (supplier_id,),
            ).fetchone()
            if supplier is None:
                abort(404)
            old_name = supplier["name"] or ""
            conn.execute(
                """
                UPDATE carton_suppliers
                SET name = ?, contact = ?, phone = ?, remark = ?, updated_at = ?
                WHERE id = ?
                """,
                (name, contact, phone, remark, now, supplier_id),
            )
            if old_name and old_name != name:
                conn.execute(
                    """
                    UPDATE carton_purchases
                    SET supplier_name = ?, updated_at = ?
                    WHERE supplier_name = ?
                    """,
                    (name, now, old_name),
                )
    except sqlite3.IntegrityError:
        flash("这个供应商名称已经存在", "error")
        return redirect(url_for("admin_carton_purchases"))
    flash("纸箱供应商已更新", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/carton-suppliers/<int:supplier_id>/delete", methods=["POST"])
@permission_required("carton_purchases")
def delete_carton_supplier(supplier_id):
    with get_db() as conn:
        supplier = conn.execute(
            "SELECT name FROM carton_suppliers WHERE id = ?",
            (supplier_id,),
        ).fetchone()
        if supplier is None:
            abort(404)
        used_count = conn.execute(
            "SELECT COUNT(*) AS c FROM carton_purchases WHERE supplier_name = ?",
            (supplier["name"] or "",),
        ).fetchone()["c"]
        if used_count:
            flash("这个供应商已有纸箱采购记录，不能删除", "error")
            return redirect(url_for("admin_carton_purchases"))
        conn.execute("DELETE FROM carton_suppliers WHERE id = ?", (supplier_id,))
    flash("纸箱供应商已删除", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/carton-products", methods=["POST"])
@permission_required("carton_purchases")
def create_carton_product():
    print_mark = request.form.get("print_mark", "").strip()
    supplier_name = request.form.get("supplier_name", "").strip()
    board_type = request.form.get("board_type", "").strip()
    carton_length = request.form.get("carton_length", "").strip()
    carton_width = request.form.get("carton_width", "").strip()
    carton_height = request.form.get("carton_height", "").strip()
    board_price_high = request.form.get("board_price_high", "").strip()
    board_price_middle = request.form.get("board_price_middle", "").strip()
    board_price_low = request.form.get("board_price_low", "").strip()
    remark = request.form.get("remark", "").strip()
    if not all([print_mark, carton_length, carton_width, carton_height]):
        flash("印刷唛头、长、宽、高为必填项", "error")
        return redirect(url_for("admin_carton_purchases"))
    try:
        carton_length_value = parse_nonnegative_number(carton_length, "长")
        carton_width_value = parse_nonnegative_number(carton_width, "宽")
        carton_height_value = parse_nonnegative_number(carton_height, "高")
        board_price_high_value = parse_nonnegative_price(board_price_high or "0")
        board_price_middle_value = parse_nonnegative_price(board_price_middle or "0")
        board_price_low_value = parse_nonnegative_price(board_price_low or "0")
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("admin_carton_purchases"))
    carton_size = carton_size_text(carton_length_value, carton_width_value, carton_height_value)
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        conn.execute(
            """
            INSERT INTO carton_products (
                print_mark, supplier_name, board_type, carton_size, carton_length, carton_width, carton_height,
                board_price_high, board_price_middle, board_price_low,
                remark, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                print_mark,
                supplier_name,
                board_type,
                carton_size,
                carton_length_value,
                carton_width_value,
                carton_height_value,
                board_price_high_value,
                board_price_middle_value,
                board_price_low_value,
                remark,
                now,
                now,
            ),
        )
    flash("纸箱产品已加入产品库", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/carton-products/<int:product_id>/edit", methods=["POST"])
@permission_required("carton_purchases")
def edit_carton_product(product_id):
    print_mark = request.form.get("print_mark", "").strip()
    supplier_name = request.form.get("supplier_name", "").strip()
    board_type = request.form.get("board_type", "").strip()
    carton_length = request.form.get("carton_length", "").strip()
    carton_width = request.form.get("carton_width", "").strip()
    carton_height = request.form.get("carton_height", "").strip()
    board_price_high = request.form.get("board_price_high", "").strip()
    board_price_middle = request.form.get("board_price_middle", "").strip()
    board_price_low = request.form.get("board_price_low", "").strip()
    remark = request.form.get("remark", "").strip()
    if not all([print_mark, carton_length, carton_width, carton_height]):
        flash("印刷唛头、长、宽、高为必填项", "error")
        return redirect(url_for("admin_carton_purchases"))
    try:
        carton_length_value = parse_nonnegative_number(carton_length, "长")
        carton_width_value = parse_nonnegative_number(carton_width, "宽")
        carton_height_value = parse_nonnegative_number(carton_height, "高")
        board_price_high_value = parse_nonnegative_price(board_price_high or "0")
        board_price_middle_value = parse_nonnegative_price(board_price_middle or "0")
        board_price_low_value = parse_nonnegative_price(board_price_low or "0")
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("admin_carton_purchases"))
    carton_size = carton_size_text(carton_length_value, carton_width_value, carton_height_value)
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        product = conn.execute(
            "SELECT id FROM carton_products WHERE id = ?",
            (product_id,),
        ).fetchone()
        if product is None:
            abort(404)
        conn.execute(
            """
            UPDATE carton_products
            SET print_mark = ?, supplier_name = ?, board_type = ?, carton_size = ?,
                carton_length = ?, carton_width = ?, carton_height = ?,
                board_price_high = ?, board_price_middle = ?, board_price_low = ?,
                remark = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                print_mark,
                supplier_name,
                board_type,
                carton_size,
                carton_length_value,
                carton_width_value,
                carton_height_value,
                board_price_high_value,
                board_price_middle_value,
                board_price_low_value,
                remark,
                now,
                product_id,
            ),
        )
    flash("纸箱产品已更新", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/carton-products/<int:product_id>/delete", methods=["POST"])
@permission_required("carton_purchases")
def delete_carton_product(product_id):
    with get_db() as conn:
        product = conn.execute(
            "SELECT id FROM carton_products WHERE id = ?",
            (product_id,),
        ).fetchone()
        if product is None:
            abort(404)
        used_count = conn.execute(
            "SELECT COUNT(*) AS c FROM carton_purchases WHERE product_id = ?",
            (product_id,),
        ).fetchone()["c"]
        if used_count:
            flash("这个纸箱产品已有采购记录，不能删除", "error")
            return redirect(url_for("admin_carton_purchases"))
        conn.execute("DELETE FROM carton_products WHERE id = ?", (product_id,))
    flash("纸箱产品已删除", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/carton-products/<int:product_id>/copy", methods=["POST"])
@permission_required("carton_purchases")
def copy_carton_product(product_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        product = conn.execute(
            "SELECT * FROM carton_products WHERE id = ?",
            (product_id,),
        ).fetchone()
        if product is None:
            abort(404)
        conn.execute(
            """
            INSERT INTO carton_products (
                print_mark, supplier_name, board_type, carton_size,
                carton_length, carton_width, carton_height,
                board_price_high, board_price_middle, board_price_low,
                remark, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f"{product['print_mark']} - 副本",
                product["supplier_name"] or "",
                product["board_type"] or "",
                product["carton_size"] or "",
                product["carton_length"] or 0,
                product["carton_width"] or 0,
                product["carton_height"] or 0,
                product["board_price_high"] or 0,
                product["board_price_middle"] or 0,
                product["board_price_low"] or 0,
                product["remark"] or "",
                now,
                now,
            ),
        )
    flash("纸箱产品已复制", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/carton-purchases/<int:record_id>/purchase-order")
@permission_required("carton_purchases")
def carton_purchase_order_pdf(record_id):
    with get_db() as conn:
        record = conn.execute(
            "SELECT * FROM carton_purchases WHERE id = ?",
            (record_id,),
        ).fetchone()
    if record is None:
        abort(404)
    buffer = build_carton_purchase_order_pdf(record)
    as_attachment = request.args.get("download") == "1"
    return send_file(
        buffer,
        as_attachment=as_attachment,
        download_name=f"carton-purchase-order-{record_id}.pdf",
        mimetype="application/pdf",
    )


@app.route("/admin/carton-purchases/purchase-order", methods=["POST"])
@permission_required("carton_purchases")
def selected_carton_purchase_order_pdf():
    record_ids = []
    for value in request.form.getlist("record_id"):
        try:
            record_ids.append(int(value))
        except (TypeError, ValueError):
            continue
    if not record_ids:
        flash("请先选择要预览采购单的纸箱采购记录", "error")
        return redirect(url_for("admin_carton_purchases"))
    placeholders = ",".join("?" for _ in record_ids)
    with get_db() as conn:
        records = conn.execute(
            f"""
            SELECT *
            FROM carton_purchases
            WHERE id IN ({placeholders})
            ORDER BY ordered_at ASC, id ASC
            """,
            record_ids,
        ).fetchall()
    if not records:
        abort(404)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if request.form.get("download") == "1":
        buffer = build_carton_purchase_order_pdf(records)
        return send_file(
            buffer,
            as_attachment=True,
            download_name=f"carton-purchase-order-selected-{timestamp}.pdf",
            mimetype="application/pdf",
        )
    total_quantity = sum(int(record["quantity"] or 0) for record in records)
    supplier_names = sorted({record["supplier_name"] for record in records if record["supplier_name"]})
    ordered_dates = sorted({record["ordered_at"] for record in records if record["ordered_at"]})
    return render_template(
        "carton_purchase_order_preview.html",
        records=records,
        total_quantity=total_quantity,
        supplier_text=supplier_names[0] if len(supplier_names) == 1 else ("多个供应商" if supplier_names else "-"),
        ordered_text=ordered_dates[0] if len(ordered_dates) == 1 else ("多个日期" if ordered_dates else "-"),
        generated_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


@app.route("/admin/carton-purchases/<int:record_id>/edit", methods=["POST"])
@permission_required("carton_purchases")
def edit_carton_purchase(record_id):
    ordered_at = request.form.get("ordered_at", "").strip()
    print_mark = request.form.get("print_mark", "").strip()
    supplier_name = request.form.get("supplier_name", "").strip()
    board_type = request.form.get("board_type", "").strip()
    quantity = request.form.get("quantity", "").strip()
    unit_price = request.form.get("unit_price", "").strip()
    carton_length = request.form.get("carton_length", "").strip()
    carton_width = request.form.get("carton_width", "").strip()
    carton_height = request.form.get("carton_height", "").strip()
    received_at = request.form.get("received_at", "").strip()
    remark = request.form.get("remark", "").strip()

    if not all([ordered_at, print_mark, quantity, carton_length, carton_width, carton_height]):
        flash("下单时间、印刷唛头、数量、长、宽、高为必填项", "error")
        return redirect(url_for("admin_carton_purchases"))
    try:
        quantity_value = parse_positive_int(quantity, "数量")
        carton_length_value = parse_nonnegative_number(carton_length, "长")
        carton_width_value = parse_nonnegative_number(carton_width, "宽")
        carton_height_value = parse_nonnegative_number(carton_height, "高")
        unit_price_value = parse_nonnegative_price(unit_price or "0")
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("admin_carton_purchases"))
    carton_size = carton_size_text(carton_length_value, carton_width_value, carton_height_value)

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        record = conn.execute(
            "SELECT id FROM carton_purchases WHERE id = ?",
            (record_id,),
        ).fetchone()
        if record is None:
            abort(404)
        conn.execute(
            """
            UPDATE carton_purchases
            SET ordered_at = ?, print_mark = ?, supplier_name = ?, board_type = ?, quantity = ?, unit_price = ?, carton_size = ?,
                carton_length = ?, carton_width = ?, carton_height = ?,
                received_at = ?, remark = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                ordered_at,
                print_mark,
                supplier_name,
                board_type,
                quantity_value,
                unit_price_value,
                carton_size,
                carton_length_value,
                carton_width_value,
                carton_height_value,
                received_at,
                remark,
                now,
                record_id,
            ),
        )

    flash("纸箱采购记录已更新", "success")
    return redirect(request.referrer or url_for("admin_carton_purchases"))


@app.route("/admin/carton-purchases/<int:record_id>/copy", methods=["POST"])
@permission_required("carton_purchases")
def copy_carton_purchase(record_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    ordered_at = datetime.now().date().isoformat()
    with get_db() as conn:
        source = conn.execute(
            """
            SELECT product_id, print_mark, supplier_name, board_type, quantity, unit_price, carton_size,
                   carton_length, carton_width, carton_height, remark
            FROM carton_purchases
            WHERE id = ?
            """,
            (record_id,),
        ).fetchone()
        if source is None:
            abort(404)
        conn.execute(
            """
            INSERT INTO carton_purchases (
                product_id, ordered_at, print_mark, supplier_name, board_type, quantity, unit_price, carton_size, received_at,
                carton_length, carton_width, carton_height, remark, recorded_by, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, '', ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                source["product_id"],
                ordered_at,
                source["print_mark"] or "",
                source["supplier_name"] or "",
                source["board_type"] or "",
                source["quantity"] or 0,
                source["unit_price"] or 0,
                source["carton_size"] or "",
                source["carton_length"] or 0,
                source["carton_width"] or 0,
                source["carton_height"] or 0,
                source["remark"] or "",
                current_admin_username(),
                now,
                now,
            ),
        )

    flash("纸箱采购记录已复制", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/carton-purchases/<int:record_id>/complete", methods=["POST"])
@permission_required("carton_purchases")
def toggle_carton_purchase_completed(record_id):
    completed = 1 if request.form.get("completed") == "1" else 0
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        record = conn.execute(
            "SELECT id FROM carton_purchases WHERE id = ?",
            (record_id,),
        ).fetchone()
        if record is None:
            abort(404)
        conn.execute(
            """
            UPDATE carton_purchases
            SET completed = ?, updated_at = ?
            WHERE id = ?
            """,
            (completed, now, record_id),
        )
    flash("纸箱采购记录已完成" if completed else "纸箱采购记录已取消完成", "success")
    return redirect(request.referrer or url_for("admin_carton_purchases"))


@app.route("/admin/carton-purchases/<int:record_id>/delete", methods=["POST"])
@permission_required("carton_purchases")
def delete_carton_purchase(record_id):
    with get_db() as conn:
        record = conn.execute(
            "SELECT id FROM carton_purchases WHERE id = ?",
            (record_id,),
        ).fetchone()
        if record is None:
            abort(404)
        conn.execute("DELETE FROM carton_purchases WHERE id = ?", (record_id,))

    flash("纸箱采购记录已删除", "success")
    return redirect(url_for("admin_carton_purchases"))


def arrival_record_filters_from_request():
    query = request.args.get("q", "").strip()
    month = request.args.get("month", "").strip()
    show_older = request.args.get("show_older") == "1"
    sort = request.args.get("sort", "").strip()
    direction = request.args.get("direction", "desc").strip().lower()

    sort_columns = {
        "arrived_at": "arrived_at",
        "item_name": "item_name COLLATE NOCASE",
        "spec": "spec COLLATE NOCASE",
        "quantity": "quantity",
        "unit_price": "unit_price",
        "total_price": "(quantity * unit_price)",
        "supplier_name": "supplier_name COLLATE NOCASE",
    }
    if sort not in sort_columns:
        sort = ""
    if direction not in {"asc", "desc"}:
        direction = "desc"

    params = []
    conditions = []
    default_start_date = (datetime.now().date() - timedelta(days=30)).isoformat()
    month_start = ""
    month_end = ""
    if query:
        conditions.append(
            """
            (
                item_name LIKE ?
                OR spec LIKE ?
                OR supplier_name LIKE ?
                OR remark LIKE ?
                OR recorded_by LIKE ?
            )
            """
        )
        like = f"%{query}%"
        params.extend([like, like, like, like, like])
    if re.fullmatch(r"\d{4}-\d{2}", month):
        month_start = f"{month}-01"
        parsed_month = datetime.strptime(month_start, "%Y-%m-%d").date()
        if parsed_month.month == 12:
            next_month = parsed_month.replace(year=parsed_month.year + 1, month=1)
        else:
            next_month = parsed_month.replace(month=parsed_month.month + 1)
        month_end = next_month.isoformat()
        conditions.append("arrived_at >= ?")
        conditions.append("arrived_at < ?")
        params.extend([month_start, month_end])
    elif not show_older:
        conditions.append("arrived_at >= ?")
        params.append(default_start_date)

    if sort:
        sql_direction = "ASC" if direction == "asc" else "DESC"
        order_sql = f" ORDER BY {sort_columns[sort]} {sql_direction}, id DESC"
    else:
        order_sql = " ORDER BY arrived_at DESC, id DESC"

    return {
        "query": query,
        "month": month,
        "show_older": show_older,
        "default_start_date": default_start_date,
        "month_start": month_start,
        "month_end": month_end,
        "sort": sort,
        "direction": direction,
        "where_sql": " WHERE " + " AND ".join(conditions) if conditions else "",
        "order_sql": order_sql,
        "params": params,
    }


@app.route("/admin/arrival-records", methods=["GET", "POST"])
@permission_required("carton_purchases")
def admin_arrival_records():
    if request.method == "POST":
        arrived_at = request.form.get("arrived_at", "").strip()
        item_name = request.form.get("item_name", "").strip()
        spec = request.form.get("spec", "").strip()
        quantity = request.form.get("quantity", "").strip()
        unit_price = request.form.get("unit_price", "").strip()
        supplier_name = request.form.get("supplier_name", "").strip()
        remark = request.form.get("remark", "").strip()

        if not all([arrived_at, item_name, quantity]):
            flash("日期、物品名称和数量为必填项", "error")
            return redirect(url_for("admin_arrival_records"))
        try:
            quantity_value = parse_positive_int(quantity, "数量")
            unit_price_value = parse_nonnegative_price(unit_price or "0")
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("admin_arrival_records"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO arrival_records (
                    arrived_at, item_name, spec, quantity, unit_price,
                    supplier_name, remark, recorded_by, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    arrived_at,
                    item_name,
                    spec,
                    quantity_value,
                    unit_price_value,
                    supplier_name,
                    remark,
                    current_admin_username(),
                    now,
                    now,
                ),
            )
        flash("到货记录已新增", "success")
        return redirect(url_for("admin_arrival_records"))

    filters = arrival_record_filters_from_request()
    with get_db() as conn:
        suppliers = conn.execute(
            """
            SELECT *
            FROM arrival_suppliers
            ORDER BY name COLLATE NOCASE ASC, id DESC
            """
        ).fetchall()
        records = conn.execute(
            f"SELECT * FROM arrival_records{filters['where_sql']}{filters['order_sql']}",
            filters["params"],
        ).fetchall()
        summary = conn.execute(
            f"""
            SELECT COUNT(*) AS record_count,
                   COALESCE(SUM(quantity), 0) AS total_quantity,
                   COALESCE(SUM(quantity * unit_price), 0) AS total_amount
            FROM arrival_records
            {filters['where_sql']}
            """,
            filters["params"],
        ).fetchone()

    return render_template(
        "arrival_records.html",
        suppliers=suppliers,
        supplier_names=[supplier["name"] for supplier in suppliers],
        records=records,
        summary=summary,
        query=filters["query"],
        month=filters["month"],
        show_older=filters["show_older"],
        default_start_date=filters["default_start_date"],
        sort=filters["sort"],
        direction=filters["direction"],
        default_arrived_at=datetime.now().date().isoformat(),
    )


@app.route("/admin/arrival-records/suppliers", methods=["POST"])
@permission_required("carton_purchases")
def create_arrival_supplier():
    name = request.form.get("name", "").strip()
    contact = request.form.get("contact", "").strip()
    phone = request.form.get("phone", "").strip()
    remark = request.form.get("remark", "").strip()
    if not name:
        flash("供应商名称为必填项", "error")
        return redirect(url_for("admin_arrival_records"))
    now = datetime.utcnow().isoformat(timespec="seconds")
    try:
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO arrival_suppliers (
                    name, contact, phone, remark, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (name, contact, phone, remark, now, now),
            )
    except sqlite3.IntegrityError:
        flash("这个供应商已经存在", "error")
        return redirect(url_for("admin_arrival_records"))
    flash("供应商已新增", "success")
    return redirect(url_for("admin_arrival_records"))


@app.route("/admin/arrival-records/suppliers/<int:supplier_id>/edit", methods=["POST"])
@permission_required("carton_purchases")
def edit_arrival_supplier(supplier_id):
    name = request.form.get("name", "").strip()
    contact = request.form.get("contact", "").strip()
    phone = request.form.get("phone", "").strip()
    remark = request.form.get("remark", "").strip()
    if not name:
        flash("供应商名称为必填项", "error")
        return redirect(url_for("admin_arrival_records"))
    now = datetime.utcnow().isoformat(timespec="seconds")
    try:
        with get_db() as conn:
            supplier = conn.execute(
                "SELECT name FROM arrival_suppliers WHERE id = ?",
                (supplier_id,),
            ).fetchone()
            if supplier is None:
                abort(404)
            old_name = supplier["name"] or ""
            conn.execute(
                """
                UPDATE arrival_suppliers
                SET name = ?, contact = ?, phone = ?, remark = ?, updated_at = ?
                WHERE id = ?
                """,
                (name, contact, phone, remark, now, supplier_id),
            )
            if old_name and old_name != name:
                conn.execute(
                    """
                    UPDATE arrival_records
                    SET supplier_name = ?, updated_at = ?
                    WHERE supplier_name = ?
                    """,
                    (name, now, old_name),
                )
    except sqlite3.IntegrityError:
        flash("这个供应商名称已经存在", "error")
        return redirect(url_for("admin_arrival_records"))
    flash("供应商已更新", "success")
    return redirect(url_for("admin_arrival_records"))


@app.route("/admin/arrival-records/suppliers/<int:supplier_id>/delete", methods=["POST"])
@permission_required("carton_purchases")
def delete_arrival_supplier(supplier_id):
    with get_db() as conn:
        supplier = conn.execute(
            "SELECT name FROM arrival_suppliers WHERE id = ?",
            (supplier_id,),
        ).fetchone()
        if supplier is None:
            abort(404)
        used_count = conn.execute(
            "SELECT COUNT(*) AS c FROM arrival_records WHERE supplier_name = ?",
            (supplier["name"] or "",),
        ).fetchone()["c"]
        if used_count:
            flash("这个供应商已有到货记录，不能删除", "error")
            return redirect(url_for("admin_arrival_records"))
        conn.execute("DELETE FROM arrival_suppliers WHERE id = ?", (supplier_id,))
    flash("供应商已删除", "success")
    return redirect(url_for("admin_arrival_records"))


@app.route("/admin/arrival-records/<int:record_id>/edit", methods=["POST"])
@permission_required("carton_purchases")
def edit_arrival_record(record_id):
    arrived_at = request.form.get("arrived_at", "").strip()
    item_name = request.form.get("item_name", "").strip()
    spec = request.form.get("spec", "").strip()
    quantity = request.form.get("quantity", "").strip()
    unit_price = request.form.get("unit_price", "").strip()
    supplier_name = request.form.get("supplier_name", "").strip()
    remark = request.form.get("remark", "").strip()

    if not all([arrived_at, item_name, quantity]):
        flash("日期、物品名称和数量为必填项", "error")
        return redirect(url_for("admin_arrival_records"))
    try:
        quantity_value = parse_positive_int(quantity, "数量")
        unit_price_value = parse_nonnegative_price(unit_price or "0")
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("admin_arrival_records"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        record = conn.execute(
            "SELECT id FROM arrival_records WHERE id = ?",
            (record_id,),
        ).fetchone()
        if record is None:
            abort(404)
        conn.execute(
            """
            UPDATE arrival_records
            SET arrived_at = ?, item_name = ?, spec = ?, quantity = ?,
                unit_price = ?, supplier_name = ?, remark = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                arrived_at,
                item_name,
                spec,
                quantity_value,
                unit_price_value,
                supplier_name,
                remark,
                now,
                record_id,
            ),
        )
    flash("到货记录已更新", "success")
    return redirect(request.referrer or url_for("admin_arrival_records"))


@app.route("/admin/arrival-records/<int:record_id>/copy", methods=["POST"])
@permission_required("carton_purchases")
def copy_arrival_record(record_id):
    now = datetime.utcnow().isoformat(timespec="seconds")
    arrived_at = datetime.now().date().isoformat()
    with get_db() as conn:
        source = conn.execute(
            """
            SELECT item_name, spec, quantity, unit_price, supplier_name, remark
            FROM arrival_records
            WHERE id = ?
            """,
            (record_id,),
        ).fetchone()
        if source is None:
            abort(404)
        conn.execute(
            """
            INSERT INTO arrival_records (
                arrived_at, item_name, spec, quantity, unit_price,
                supplier_name, remark, recorded_by, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                arrived_at,
                source["item_name"] or "",
                source["spec"] or "",
                source["quantity"] or 0,
                source["unit_price"] or 0,
                source["supplier_name"] or "",
                source["remark"] or "",
                current_admin_username(),
                now,
                now,
            ),
        )
    flash("到货记录已复制", "success")
    return redirect(url_for("admin_arrival_records"))


@app.route("/admin/arrival-records/<int:record_id>/delete", methods=["POST"])
@permission_required("carton_purchases")
def delete_arrival_record(record_id):
    with get_db() as conn:
        record = conn.execute(
            "SELECT id FROM arrival_records WHERE id = ?",
            (record_id,),
        ).fetchone()
        if record is None:
            abort(404)
        conn.execute("DELETE FROM arrival_records WHERE id = ?", (record_id,))
    flash("到货记录已删除", "success")
    return redirect(url_for("admin_arrival_records"))


INBOUND_TYPES = ["生产入库", "采购入库", "退货入库", "库存调整"]
OUTBOUND_TYPES = ["销售发货", "生产领料", "样品", "报废", "其他"]


def inventory_overview_filters():
    query = request.args.get("q", "").strip()
    status = request.args.get("status", "").strip()
    if status not in {"", "low", "zero", "positive"}:
        status = ""
    return query, status


def inventory_product_rows(conn, query="", status=""):
    ensure_all_product_inventory_codes(conn)
    manual_projection = ", ".join(
        f"manuals.{column}" for column in MANUAL_SAFE_COLUMNS
    )
    rows = conn.execute(
        f"""
        SELECT {manual_projection},
               COALESCE(SUM(inventory_balances.quantity), 0) AS total_stock
        FROM manuals
        LEFT JOIN inventory_balances ON inventory_balances.manual_id = manuals.id
        WHERE (
            ? = ''
            OR manuals.product_name LIKE ?
            OR manuals.drawing_no LIKE ?
            OR manuals.sku LIKE ?
            OR manuals.barcode LIKE ?
            OR manuals.qr_code LIKE ?
            OR manuals.remark LIKE ?
        )
        GROUP BY manuals.id
        ORDER BY manuals.product_name COLLATE NOCASE ASC, manuals.id ASC
        """,
        (query, f"%{query}%", f"%{query}%", f"%{query}%", f"%{query}%", f"%{query}%", f"%{query}%"),
    ).fetchall()
    filtered = []
    for row in rows:
        total = int(row["total_stock"] or 0)
        min_stock = int(row["min_stock"] or 0)
        if status == "low" and not (min_stock > 0 and total < min_stock):
            continue
        if status == "zero" and total != 0:
            continue
        if status == "positive" and total <= 0:
            continue
        filtered.append(row)
    return filtered


def location_stock_map(conn, manual_ids):
    if not manual_ids:
        return {}
    placeholders = ",".join("?" for _ in manual_ids)
    rows = conn.execute(
        f"""
        SELECT inventory_balances.manual_id, inventory_balances.quantity,
               warehouse_locations.name, warehouse_locations.code
        FROM inventory_balances
        JOIN warehouse_locations ON warehouse_locations.id = inventory_balances.location_id
        WHERE inventory_balances.manual_id IN ({placeholders})
          AND inventory_balances.quantity != 0
        ORDER BY warehouse_locations.code COLLATE NOCASE ASC, warehouse_locations.id ASC
        """,
        manual_ids,
    ).fetchall()
    result = {manual_id: [] for manual_id in manual_ids}
    for row in rows:
        result.setdefault(row["manual_id"], []).append(row)
    return result


def inventory_status_label(total_stock, min_stock):
    total_stock = int(total_stock or 0)
    min_stock = int(min_stock or 0)
    if total_stock <= 0:
        return "无库存"
    if min_stock > 0 and total_stock < min_stock:
        return "库存不足"
    return "正常"


def fetch_inventory_product_by_code(conn, code):
    code = (code or "").strip()
    if not code:
        return None
    ensure_all_product_inventory_codes(conn)
    projection = ", ".join(MANUAL_SAFE_COLUMNS)
    return conn.execute(
        f"""
        SELECT {projection}
        FROM manuals
        WHERE sku = ? OR barcode = ? OR qr_code = ?
        LIMIT 1
        """,
        (code, code, code),
    ).fetchone()


def fetch_recent_inventory_transactions(conn, manual_id, limit=8):
    return conn.execute(
        """
        SELECT inventory_transactions.*,
               from_loc.name AS from_location_name,
               to_loc.name AS to_location_name
        FROM inventory_transactions
        LEFT JOIN warehouse_locations AS from_loc ON from_loc.id = inventory_transactions.from_location_id
        LEFT JOIN warehouse_locations AS to_loc ON to_loc.id = inventory_transactions.to_location_id
        WHERE inventory_transactions.manual_id = ?
        ORDER BY inventory_transactions.created_at DESC, inventory_transactions.id DESC
        LIMIT ?
        """,
        (manual_id, limit),
    ).fetchall()


@app.route("/admin/inventory")
@permission_required("warehouse_inventory")
def admin_inventory_overview():
    query, status = inventory_overview_filters()
    with get_db() as conn:
        products = inventory_product_rows(conn, query, status)
        locations_by_manual = location_stock_map(conn, [row["id"] for row in products])
    return render_template(
        "inventory_overview.html",
        products=products,
        locations_by_manual=locations_by_manual,
        query=query,
        status=status,
        status_label=inventory_status_label,
        inventory_code=product_inventory_code,
    )


@app.route("/admin/inventory/export.xlsx")
@permission_required("warehouse_inventory")
def export_inventory_overview():
    query, status = inventory_overview_filters()
    with get_db() as conn:
        products = inventory_product_rows(conn, query, status)
        locations_by_manual = location_stock_map(conn, [row["id"] for row in products])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存总览"
    sheet.append(["产品名称", "产品图号", "SKU/库存编码", "条码", "总库存", "各库位库存", "最低库存", "库存状态"])
    for row in products:
        location_text = "；".join(f"{item['name']}({item['code']}): {item['quantity']}" for item in locations_by_manual.get(row["id"], []))
        total_stock = int(row["total_stock"] or 0)
        min_stock = int(row["min_stock"] or 0)
        sheet.append([
            row["product_name"],
            row["drawing_no"],
            product_inventory_code(row),
            row["barcode"] or "",
            total_stock,
            location_text,
            min_stock,
            inventory_status_label(total_stock, min_stock),
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8F3F1")
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"inventory-overview-{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/admin/inventory/scan")
@permission_required("warehouse_inventory")
def inventory_scan():
    return render_template("inventory_scan.html", mode="query")


@app.route("/admin/inventory/api/product")
@permission_required("warehouse_inventory")
def inventory_api_product():
    code = request.args.get("code", "").strip()
    with get_db() as conn:
        product = fetch_inventory_product_by_code(conn, code)
        if product is None:
            return jsonify({"error": "未找到产品"}), 404
        total_stock = inventory_total_for_manual(conn, product["id"])
        location_rows = location_stock_map(conn, [product["id"]]).get(product["id"], [])
        recent = fetch_recent_inventory_transactions(conn, product["id"], 6)
    image_url = url_for("preview_manual", filename=product["filename"]) if product["filename"] else ""
    return jsonify({
        "product": {
            "id": product["id"],
            "product_name": product["product_name"],
            "drawing_no": product["drawing_no"] or "",
            "sku": product["sku"] or "",
            "barcode": product["barcode"] or "",
            "inventory_code": product_inventory_code(product),
            "spec": product["pack_carton_size"] or product["remark"] or "",
            "image_url": image_url,
            "total_stock": total_stock,
            "min_stock": int(product["min_stock"] or 0),
            "default_location_id": product["default_location_id"] or "",
            "locations": [
                {"name": row["name"], "code": row["code"], "quantity": row["quantity"]}
                for row in location_rows
            ],
            "recent_transactions": [
                {
                    "transaction_no": item["transaction_no"],
                    "type": item["type"],
                    "quantity": item["quantity"],
                    "from_location": item["from_location_name"] or "",
                    "to_location": item["to_location_name"] or "",
                    "remark": item["remark"] or "",
                    "created_at": item["created_at"],
                }
                for item in recent
            ],
        }
    })


def inventory_form_context(selected_manual_id=None, selected_order_id=None, selected_order_no=""):
    with get_db() as conn:
        ensure_all_product_inventory_codes(conn)
        projection = ", ".join(MANUAL_SAFE_COLUMNS)
        products = conn.execute(
            f"""
            SELECT {projection}
            FROM manuals
            ORDER BY product_name COLLATE NOCASE ASC, id ASC
            """
        ).fetchall()
        locations = conn.execute(
            """
            SELECT *
            FROM warehouse_locations
            WHERE enabled = 1
            ORDER BY code COLLATE NOCASE ASC, id ASC
            """
        ).fetchall()
        if not locations:
            default_location = get_or_create_default_location(conn)
            locations = [default_location]
        order = None
        if selected_order_id:
            order = conn.execute(
                """
                SELECT product_orders.*, manuals.product_name, manuals.drawing_no
                FROM product_orders
                JOIN manuals ON manuals.id = product_orders.manual_id
                WHERE product_orders.id = ?
                """,
                (selected_order_id,),
            ).fetchone()
    return {
        "products": products,
        "locations": locations,
        "selected_manual_id": selected_manual_id,
        "selected_order_id": selected_order_id or "",
        "selected_order_no": selected_order_no,
        "order": order,
        "inventory_code": product_inventory_code,
    }


@app.route("/admin/inventory/inbound", methods=["GET", "POST"])
@permission_required("warehouse_inventory")
def inventory_inbound():
    if request.method == "POST":
        manual_id = parse_optional_int(request.form.get("manual_id"))
        quantity_text = request.form.get("quantity", "").strip()
        location_id = parse_optional_int(request.form.get("location_id"))
        stock_type = request.form.get("stock_type", "").strip()
        related_order_id = request.form.get("related_order_id", "").strip()
        related_order_no = request.form.get("related_order_no", "").strip()
        remark = request.form.get("remark", "").strip()
        if not manual_id or not location_id or not stock_type:
            flash("请选择产品、库位和入库类型", "error")
            return redirect(request.referrer or url_for("inventory_inbound"))
        try:
            quantity = parse_positive_int(quantity_text, "入库数量")
            with get_db() as conn:
                manual = conn.execute("SELECT id FROM manuals WHERE id = ?", (manual_id,)).fetchone()
                location = conn.execute("SELECT id FROM warehouse_locations WHERE id = ? AND enabled = 1", (location_id,)).fetchone()
                if manual is None or location is None:
                    flash("产品或库位不存在", "error")
                    return redirect(url_for("inventory_inbound"))
                create_inventory_transaction(
                    conn,
                    "in",
                    manual_id,
                    quantity,
                    to_location_id=location_id,
                    related_order_type=stock_type,
                    related_order_id=related_order_id,
                    related_order_no=related_order_no,
                    remark=remark,
                )
        except ValueError as error:
            flash(str(error), "error")
            return redirect(request.referrer or url_for("inventory_inbound"))
        flash("入库成功，库存流水已生成", "success")
        return redirect(url_for("inventory_scan", code=related_order_no))

    selected_order_id = parse_optional_int(request.args.get("order_id"))
    selected_manual_id = parse_optional_int(request.args.get("manual_id"))
    selected_order_no = request.args.get("order_no", "").strip()
    if selected_order_id and not selected_manual_id:
        with get_db() as conn:
            order = conn.execute("SELECT manual_id, order_no FROM product_orders WHERE id = ?", (selected_order_id,)).fetchone()
            if order:
                selected_manual_id = order["manual_id"]
                selected_order_no = selected_order_no or order["order_no"]
    return render_template("inventory_inbound.html", stock_types=INBOUND_TYPES, **inventory_form_context(selected_manual_id, selected_order_id, selected_order_no))


@app.route("/admin/inventory/outbound", methods=["GET", "POST"])
@permission_required("warehouse_inventory")
def inventory_outbound():
    if request.method == "POST":
        manual_id = parse_optional_int(request.form.get("manual_id"))
        quantity_text = request.form.get("quantity", "").strip()
        location_id = parse_optional_int(request.form.get("location_id"))
        stock_type = request.form.get("stock_type", "").strip()
        related_order_no = request.form.get("related_order_no", "").strip()
        customer = request.form.get("customer", "").strip()
        remark = request.form.get("remark", "").strip()
        if not manual_id or not location_id or not stock_type:
            flash("请选择产品、库位和出库类型", "error")
            return redirect(request.referrer or url_for("inventory_outbound"))
        try:
            quantity = parse_positive_int(quantity_text, "出库数量")
            with get_db() as conn:
                create_inventory_transaction(
                    conn,
                    "out",
                    manual_id,
                    quantity,
                    from_location_id=location_id,
                    related_order_type=stock_type,
                    related_order_no=related_order_no,
                    customer=customer,
                    remark=remark,
                )
        except ValueError as error:
            flash(str(error), "error")
            return redirect(request.referrer or url_for("inventory_outbound"))
        flash("出库成功，库存流水已生成", "success")
        return redirect(url_for("admin_inventory_overview"))

    selected_manual_id = parse_optional_int(request.args.get("manual_id"))
    return render_template("inventory_outbound.html", stock_types=OUTBOUND_TYPES, **inventory_form_context(selected_manual_id))


@app.route("/admin/inventory/adjust", methods=["GET", "POST"])
@permission_required("warehouse_inventory")
def inventory_adjust():
    if request.method == "POST":
        manual_id = parse_optional_int(request.form.get("manual_id"))
        location_id = parse_optional_int(request.form.get("location_id"))
        counted_quantity_text = request.form.get("counted_quantity", "").strip()
        remark = request.form.get("remark", "").strip()
        if not manual_id or not location_id:
            flash("请选择产品和调整库位", "error")
            return redirect(request.referrer or url_for("inventory_adjust"))
        try:
            counted_quantity = int(counted_quantity_text)
        except ValueError:
            flash("盘点数量必须为整数", "error")
            return redirect(request.referrer or url_for("inventory_adjust"))
        if counted_quantity < 0:
            flash("盘点数量不能小于 0", "error")
            return redirect(request.referrer or url_for("inventory_adjust"))
        try:
            with get_db() as conn:
                manual = conn.execute("SELECT id FROM manuals WHERE id = ?", (manual_id,)).fetchone()
                location = conn.execute("SELECT id FROM warehouse_locations WHERE id = ? AND enabled = 1", (location_id,)).fetchone()
                if manual is None or location is None:
                    flash("产品或库位不存在", "error")
                    return redirect(url_for("inventory_adjust"))
                current_quantity = inventory_balance_for_location(conn, manual_id, location_id)
                delta = counted_quantity - current_quantity
                if delta == 0:
                    flash("盘点数量和当前库存一致，无需调整", "success")
                    return redirect(url_for("admin_inventory_overview"))
                create_inventory_transaction(
                    conn,
                    "adjust",
                    manual_id,
                    delta,
                    to_location_id=location_id,
                    related_order_type="库存调整",
                    remark=remark or f"盘点调整：{current_quantity} -> {counted_quantity}",
                )
        except ValueError as error:
            flash(str(error), "error")
            return redirect(request.referrer or url_for("inventory_adjust"))
        flash("库存调整成功，调整流水已生成", "success")
        return redirect(url_for("admin_inventory_overview"))

    selected_manual_id = parse_optional_int(request.args.get("manual_id"))
    return render_template("inventory_adjust.html", **inventory_form_context(selected_manual_id))


@app.route("/admin/inventory/locations", methods=["GET", "POST"])
@permission_required("warehouse_inventory")
def inventory_locations():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        code = request.form.get("code", "").strip()
        remark = request.form.get("remark", "").strip()
        enabled = 1 if request.form.get("enabled", "1") else 0
        if not name or not code:
            flash("库位名称和库位编码为必填项", "error")
            return redirect(url_for("inventory_locations"))
        now = datetime.utcnow().isoformat(timespec="seconds")
        try:
            with get_db() as conn:
                conn.execute(
                    """
                    INSERT INTO warehouse_locations (name, code, remark, enabled, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (name, code, remark, enabled, now, now),
                )
        except sqlite3.IntegrityError:
            flash("库位编码已经存在", "error")
            return redirect(url_for("inventory_locations"))
        flash("库位已新增", "success")
        return redirect(url_for("inventory_locations"))
    with get_db() as conn:
        locations = conn.execute("SELECT * FROM warehouse_locations ORDER BY enabled DESC, code COLLATE NOCASE ASC, id ASC").fetchall()
    return render_template("inventory_locations.html", locations=locations)


@app.route("/admin/inventory/locations/<int:location_id>/edit", methods=["POST"])
@permission_required("warehouse_inventory")
def edit_inventory_location(location_id):
    name = request.form.get("name", "").strip()
    code = request.form.get("code", "").strip()
    remark = request.form.get("remark", "").strip()
    enabled = 1 if request.form.get("enabled") else 0
    if not name or not code:
        flash("库位名称和库位编码为必填项", "error")
        return redirect(url_for("inventory_locations"))
    now = datetime.utcnow().isoformat(timespec="seconds")
    try:
        with get_db() as conn:
            location = conn.execute("SELECT id FROM warehouse_locations WHERE id = ?", (location_id,)).fetchone()
            if location is None:
                abort(404)
            conn.execute(
                """
                UPDATE warehouse_locations
                SET name = ?, code = ?, remark = ?, enabled = ?, updated_at = ?
                WHERE id = ?
                """,
                (name, code, remark, enabled, now, location_id),
            )
    except sqlite3.IntegrityError:
        flash("库位编码已经存在", "error")
        return redirect(url_for("inventory_locations"))
    flash("库位已更新", "success")
    return redirect(url_for("inventory_locations"))


@app.route("/admin/inventory/transactions")
@permission_required("warehouse_inventory")
def inventory_transactions():
    query = request.args.get("q", "").strip()
    params = []
    where = ""
    if query:
        like = f"%{query}%"
        where = """
            WHERE inventory_transactions.transaction_no LIKE ?
               OR manuals.product_name LIKE ?
               OR manuals.sku LIKE ?
               OR manuals.qr_code LIKE ?
               OR inventory_transactions.related_order_no LIKE ?
               OR inventory_transactions.remark LIKE ?
               OR inventory_transactions.operator LIKE ?
        """
        params = [like, like, like, like, like, like, like]
    with get_db() as conn:
        rows = conn.execute(
            f"""
            SELECT inventory_transactions.*,
                   manuals.product_name, manuals.drawing_no, manuals.sku, manuals.qr_code,
                   from_loc.name AS from_location_name,
                   to_loc.name AS to_location_name
            FROM inventory_transactions
            JOIN manuals ON manuals.id = inventory_transactions.manual_id
            LEFT JOIN warehouse_locations AS from_loc ON from_loc.id = inventory_transactions.from_location_id
            LEFT JOIN warehouse_locations AS to_loc ON to_loc.id = inventory_transactions.to_location_id
            {where}
            ORDER BY inventory_transactions.created_at DESC, inventory_transactions.id DESC
            LIMIT 300
            """,
            params,
        ).fetchall()
    return render_template("inventory_transactions.html", transactions=rows, query=query)


@app.route("/admin/inventory/products/<int:manual_id>/label")
@permission_required("warehouse_inventory")
def inventory_product_label(manual_id):
    with get_db() as conn:
        ensure_product_inventory_code(conn, manual_id)
        manual_projection = ", ".join(
            f"manuals.{column}" for column in MANUAL_SAFE_COLUMNS
        )
        product = conn.execute(
            f"""
            SELECT {manual_projection}, warehouse_locations.name AS location_name, warehouse_locations.code AS location_code
            FROM manuals
            LEFT JOIN warehouse_locations ON warehouse_locations.id = manuals.default_location_id
            WHERE manuals.id = ?
            """,
            (manual_id,),
        ).fetchone()
    if product is None:
        abort(404)
    return render_template("inventory_labels.html", products=[product], inventory_code=product_inventory_code)


@app.route("/admin/inventory/labels", methods=["POST"])
@permission_required("warehouse_inventory")
def inventory_labels():
    manual_ids = []
    for value in request.form.getlist("manual_id"):
        parsed = parse_optional_int(value)
        if parsed:
            manual_ids.append(parsed)
    manual_ids = list(dict.fromkeys(manual_ids))
    if not manual_ids:
        flash("请先选择要打印标签的产品", "error")
        return redirect(url_for("admin_inventory_overview"))
    placeholders = ",".join("?" for _ in manual_ids)
    with get_db() as conn:
        for manual_id in manual_ids:
            ensure_product_inventory_code(conn, manual_id)
        manual_projection = ", ".join(
            f"manuals.{column}" for column in MANUAL_SAFE_COLUMNS
        )
        products = conn.execute(
            f"""
            SELECT {manual_projection}, warehouse_locations.name AS location_name, warehouse_locations.code AS location_code
            FROM manuals
            LEFT JOIN warehouse_locations ON warehouse_locations.id = manuals.default_location_id
            WHERE manuals.id IN ({placeholders})
            ORDER BY manuals.product_name COLLATE NOCASE ASC, manuals.id ASC
            """,
            manual_ids,
        ).fetchall()
    return render_template("inventory_labels.html", products=products, inventory_code=product_inventory_code)


@app.route("/admin")
@login_required
def admin_index():
    if not user_can_access_admin_modules():
        flash("当前账号没有权限访问后台模块", "error")
        return redirect(url_for("dashboard"))
    with get_db() as conn:
        product_count = conn.execute(
            "SELECT COUNT(*) AS c FROM manuals"
        ).fetchone()["c"]
        customer_count = conn.execute(
            "SELECT COUNT(*) AS c FROM customers"
        ).fetchone()["c"]
        common_info_count = conn.execute(
            "SELECT COUNT(*) AS c FROM common_infos"
        ).fetchone()["c"]
        purchase_followup_count = conn.execute(
            "SELECT COUNT(*) AS c FROM purchase_followups"
        ).fetchone()["c"]
        powder_coating_count = conn.execute(
            "SELECT COUNT(*) AS c FROM powder_coating_records"
        ).fetchone()["c"]
        carton_purchase_count = conn.execute(
            "SELECT COUNT(*) AS c FROM carton_purchases"
        ).fetchone()["c"]
        arrival_record_count = conn.execute(
            "SELECT COUNT(*) AS c FROM arrival_records"
        ).fetchone()["c"]
        feedback_count = conn.execute(
            "SELECT COUNT(*) AS c FROM feedbacks"
        ).fetchone()["c"]
        suppliers = [row["supplier"] for row in get_suppliers_with_conn(conn)]
        customers = get_customer_name_options(conn)
        material_options = get_product_material_options(conn)
        inventory_locations = conn.execute(
            """
            SELECT *
            FROM warehouse_locations
            WHERE enabled = 1
            ORDER BY code COLLATE NOCASE ASC, id ASC
            """
        ).fetchall()
    return render_template(
        "admin.html",
        product_count=product_count,
        customer_count=customer_count,
        common_info_count=common_info_count,
        purchase_followup_count=purchase_followup_count,
        powder_coating_count=powder_coating_count,
        carton_purchase_count=carton_purchase_count,
        arrival_record_count=arrival_record_count,
        feedback_count=feedback_count,
        suppliers=suppliers,
        customers=customers,
        inventory_locations=inventory_locations,
        materials=[],
        material_options=material_options,
        inspection_requirements=[],
        supported_currencies=SUPPORTED_CURRENCIES,
    )


@app.route("/admin/orders")
@permission_required("orders_view")
def admin_orders():
    query = request.args.get("q", "").strip()
    selected_customer = request.args.get("customer", "").strip()
    order_sort_columns = {
        "order_no": "product_orders.order_no COLLATE NOCASE",
        "assembly_drawing_no": "product_orders.assembly_drawing_no COLLATE NOCASE",
        "assembly_set_quantity": "product_orders.assembly_set_quantity",
        "ordered_at": "product_orders.ordered_at",
        "drawing_no": "manuals.drawing_no COLLATE NOCASE",
        "product_name": "manuals.product_name COLLATE NOCASE",
        "customer": "manuals.customer COLLATE NOCASE",
        "supplier": "manuals.supplier COLLATE NOCASE",
        "quantity": "product_orders.quantity",
        "unshipped_quantity": "unshipped_quantity",
        "planned_ship_at": "product_orders.planned_ship_at",
        "material_stock_status": "COALESCE(inventory_totals.total_quantity, 0)",
        "remark": "product_orders.remark COLLATE NOCASE",
        "carton_status": "product_orders.carton_status COLLATE NOCASE",
        "inventory_status": "product_orders.inventory_status COLLATE NOCASE",
        "recent_ship_status": "product_orders.recent_ship_status COLLATE NOCASE",
        "shipped_quantity": "shipped_quantity_total",
        "shipped_at": "shipped_at_display",
    }
    sort = request.args.get("sort", "ordered_at").strip()
    direction = request.args.get("direction", "desc").strip().lower()
    if sort not in order_sort_columns:
        sort = "ordered_at"
    if direction not in {"asc", "desc"}:
        direction = "desc"
    sql_direction = "ASC" if direction == "asc" else "DESC"
    sql = """
        SELECT product_orders.*,
               manuals.drawing_no,
               manuals.product_name,
               manuals.customer AS product_customer,
               manuals.supplier AS supplier,
               COALESCE(inventory_totals.total_quantity, 0) AS material_inventory_total,
               COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS shipped_quantity_total,
               COALESCE(shipments.last_shipped_at, product_orders.shipped_at, '') AS shipped_at_display,
               product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS unshipped_quantity,
               CASE
                   WHEN product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) <= 0
                       THEN 1
                   ELSE 0
               END AS completion_rank
        FROM product_orders
        JOIN manuals ON manuals.id = product_orders.manual_id
        LEFT JOIN (__ORDER_SHIPMENT_SUMMARY__) AS shipments
          ON shipments.order_id = product_orders.id
        LEFT JOIN (
            SELECT manual_id, COALESCE(SUM(quantity), 0) AS total_quantity
            FROM inventory_balances
            GROUP BY manual_id
        ) AS inventory_totals ON inventory_totals.manual_id = product_orders.manual_id
    """
    params = []
    conditions = []
    if query:
        conditions.append(
            """
            (
                product_orders.order_no LIKE ?
                OR product_orders.assembly_drawing_no LIKE ?
                OR product_orders.customer LIKE ?
                OR manuals.customer LIKE ?
                OR manuals.supplier LIKE ?
                OR manuals.drawing_no LIKE ?
                OR manuals.product_name LIKE ?
                OR product_orders.material_stock_status LIKE ?
                OR product_orders.remark LIKE ?
                OR product_orders.carton_status LIKE ?
                OR product_orders.inventory_status LIKE ?
                OR product_orders.recent_ship_status LIKE ?
            )
            """
        )
        like = f"%{query}%"
        params.extend([like, like, like, like, like, like, like, like, like, like, like, like])
    if selected_customer:
        conditions.append(
            """
            (
                product_orders.customer = ?
                OR manuals.customer = ?
            )
            """
        )
        params.extend([selected_customer, selected_customer])
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += (
        f" ORDER BY completion_rank ASC, "
        f"CASE WHEN completion_rank = 0 THEN {order_sort_columns[sort]} END {sql_direction}, "
        "CASE WHEN completion_rank = 1 THEN product_orders.ordered_at END ASC, "
        "product_orders.ordered_at DESC, product_orders.id DESC"
    )

    with get_db() as conn:
        sql = sql.replace(
            "__ORDER_SHIPMENT_SUMMARY__",
            order_shipment_summary_subquery(include_assembly=True, conn=conn),
        )
        orders = conn.execute(sql, params).fetchall()
        customers = get_order_customer_options(conn)
        user_options = get_active_user_options(conn)

    today = datetime.now().date()
    warning_until = today + timedelta(days=10)

    return render_template(
        "orders.html",
        orders=orders,
        query=query,
        customers=customers,
        selected_customer=selected_customer,
        user_options=user_options,
        sort=sort,
        direction=direction,
        today=today.isoformat(),
        warning_until=warning_until.isoformat(),
    )


@app.route("/admin/orders/shipment-plans", methods=["POST"])
@permission_required("orders_manage")
def create_shipment_plan_from_orders():
    order_ids = []
    for value in request.form.getlist("order_id"):
        try:
            order_ids.append(int(value))
        except ValueError:
            continue
    order_ids = list(dict.fromkeys(order_ids))
    planned_ship_at = request.form.get("planned_ship_at", "").strip() or datetime.now().date().isoformat()
    if not order_ids:
        flash("请先选择要生成计划发货清单的订单", "error")
        return redirect(url_for("admin_orders"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        placeholders = ",".join("?" for _ in order_ids)
        rows = conn.execute(
            f"""
            SELECT product_orders.id,
                   product_orders.order_no,
                   product_orders.customer AS order_customer,
                   manuals.customer AS product_customer,
                   product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS unshipped_quantity
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments ON shipments.order_id = product_orders.id
            WHERE product_orders.id IN ({placeholders})
            ORDER BY product_orders.planned_ship_at ASC, product_orders.id ASC
            """,
            order_ids,
        ).fetchall()
        items = [
            row
            for row in rows
            if int(row["unshipped_quantity"] or 0) > 0
        ]
        if not items:
            flash("选择的订单都已发完，不能生成计划发货清单", "error")
            return redirect(url_for("admin_orders"))
        customers = sorted({
            (row["order_customer"] or row["product_customer"] or "").strip()
            for row in items
            if (row["order_customer"] or row["product_customer"] or "").strip()
        })
        customer_text = customers[0] if len(customers) == 1 else ("多个客户" if customers else "")
        plan_no = unique_shipment_plan_no(conn)
        cursor = conn.execute(
            """
            INSERT INTO shipment_plans (
                plan_no, planned_ship_at, customer, status, remark,
                created_by, created_at, updated_at
            )
            VALUES (?, ?, ?, '待发货', '', ?, ?, ?)
            """,
            (plan_no, planned_ship_at, customer_text, current_admin_username(), now, now),
        )
        plan_id = cursor.lastrowid
        conn.executemany(
            """
            INSERT INTO shipment_plan_items (
                plan_id, order_id, planned_quantity, shipped_quantity,
                created_at, updated_at
            )
            VALUES (?, ?, ?, 0, ?, ?)
            """,
            [
                (plan_id, row["id"], int(row["unshipped_quantity"] or 0), now, now)
                for row in items
            ],
        )

    flash(f"计划发货清单已生成：{plan_no}", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/orders/carton-purchases", methods=["POST"])
@permission_required("orders_manage")
def create_carton_purchases_from_orders():
    order_ids = []
    for value in request.form.getlist("order_id"):
        try:
            order_ids.append(int(value))
        except ValueError:
            continue
    order_ids = list(dict.fromkeys(order_ids))
    if not order_ids:
        flash("请先选择要生成纸箱采购清单的订单", "error")
        return redirect(url_for("admin_orders"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        placeholders = ",".join("?" for _ in order_ids)
        rows = conn.execute(
            f"""
            SELECT product_orders.id,
                   product_orders.order_no,
                   product_orders.quantity,
                   manuals.drawing_no,
                   manuals.product_name,
                   manuals.supplier,
                   manuals.pack_quantity,
                   manuals.pack_carton_size,
                   product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS unshipped_quantity
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments ON shipments.order_id = product_orders.id
            WHERE product_orders.id IN ({placeholders})
            ORDER BY product_orders.planned_ship_at ASC, product_orders.id ASC
            """,
            order_ids,
        ).fetchall()
        items = [row for row in rows if int(row["unshipped_quantity"] or 0) > 0]
        if not items:
            flash("选择的订单都已发完，不能生成纸箱采购清单", "error")
            return redirect(url_for("admin_orders"))

        created_count = 0
        for row in items:
            order_quantity = int(row["quantity"] or 0)
            pack_count = parse_positive_number(row["pack_quantity"])
            if pack_count > 0:
                carton_quantity = max(1, math.ceil(order_quantity / pack_count))
                pack_note = f"一箱数量：{row['pack_quantity'] or '-'}"
            else:
                carton_quantity = order_quantity
                pack_note = "一箱数量未填写，纸箱数量请人工确认"
            carton_length, carton_width, carton_height = parse_carton_size_dimensions(row["pack_carton_size"])
            carton_size = carton_size_text(carton_length, carton_width, carton_height) or (row["pack_carton_size"] or "")
            drawing_no = (row["drawing_no"] or row["order_no"] or "").strip()
            pack_mark_quantity = pack_quantity_pcs_text(row["pack_quantity"])
            print_mark = f"{drawing_no}-{pack_mark_quantity}" if drawing_no and pack_mark_quantity else (drawing_no or pack_mark_quantity)
            remark_parts = [
                f"订单号：{row['order_no']}",
                f"订单数量：{order_quantity}",
                pack_note,
            ]
            if row["pack_carton_size"]:
                remark_parts.append(f"箱子尺寸：{row['pack_carton_size']}")
            conn.execute(
                """
                INSERT INTO carton_purchases (
                    product_id, ordered_at, print_mark, supplier_name, board_type,
                    quantity, unit_price, carton_size, received_at,
                    carton_length, carton_width, carton_height,
                    remark, recorded_by, created_at, updated_at
                )
                VALUES (NULL, ?, ?, ?, '', ?, 0, ?, '', ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    datetime.now().date().isoformat(),
                    print_mark,
                    row["supplier"] or "",
                    carton_quantity,
                    carton_size,
                    carton_length,
                    carton_width,
                    carton_height,
                    " / ".join(remark_parts),
                    current_admin_username(),
                    now,
                    now,
                ),
            )
            created_count += 1

    flash(f"纸箱采购清单已生成：{created_count} 条", "success")
    return redirect(url_for("admin_carton_purchases"))


@app.route("/admin/shipped-orders")
@permission_required("shipped_view")
def shipped_orders():
    query = request.args.get("q", "").strip()
    selected_customer = request.args.get("customer", "").strip()
    shipped_at = request.args.get("shipped_at", "").strip()
    include_prices = user_can_view_prices()

    with get_db() as conn:
        shipped_orders = fetch_shipped_orders(
            conn,
            query=query,
            selected_customer=selected_customer,
            shipped_at=shipped_at,
            include_prices=include_prices,
        )
        shipped_orders = attach_shipment_images(conn, shipped_orders)
        assembly_batches = fetch_assembly_shipment_batches(
            conn,
            query=query,
            selected_customer=selected_customer,
            shipped_at=shipped_at,
            include_prices=include_prices,
        )
        customers = get_shipment_customer_options(conn)
        unshipped_orders = get_unshipped_order_options(conn)
        shipment_plans = fetch_open_shipment_plans(conn)

    return render_template(
        "shipped_orders.html",
        shipped_orders=shipped_orders,
        assembly_batches=assembly_batches,
        unshipped_orders=unshipped_orders,
        shipment_plans=shipment_plans,
        query=query,
        customers=customers,
        selected_customer=selected_customer,
        selected_shipped_at=shipped_at,
        default_shipped_at=datetime.now().date().isoformat(),
    )


@app.route("/admin/shipped-orders/assembly-options")
@permission_required("shipped_view")
def assembly_shipment_options():
    customer = request.args.get("customer", "")
    try:
        with get_db() as conn:
            drawing_numbers = get_assembly_options_for_customer(conn, customer)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    return jsonify({"assembly_drawing_numbers": drawing_numbers})


@app.route("/admin/shipped-orders/assembly-preview", methods=["POST"])
@permission_required("shipped_manage")
def assembly_shipment_preview():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"error": "请求数据格式无效"}), 400
    try:
        with get_db() as conn:
            preview = build_assembly_shipment_preview(
                conn,
                payload.get("customer"),
                payload.get("assembly_drawing_no"),
                payload.get("set_quantity"),
                payload.get("overrides"),
            )
    except AssemblyDefinitionNotFound as error:
        return jsonify({"error": str(error)}), 404
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    return jsonify(preview)


def _parse_assembly_shipment_item_overrides(form):
    manual_ids = form.getlist("manual_id")
    shipped_quantities = form.getlist("shipped_quantity")
    if not manual_ids or not shipped_quantities:
        raise ValueError("请提交完整的组装发货配件明细")
    overrides = {}
    for manual_id, shipped_quantity in zip_longest(
        manual_ids, shipped_quantities, fillvalue=""
    ):
        if not str(manual_id or "").strip() or not str(shipped_quantity or "").strip():
            raise ValueError("组装发货配件 ID 和数量必须一一对应")
        manual_id = parse_positive_int(manual_id, "配件 ID")
        if manual_id in overrides:
            raise ValueError("同一配件不能重复提交")
        overrides[manual_id] = parse_positive_int(
            shipped_quantity, "配件实际发货数量"
        )
    return overrides


def _assembly_edit_preview(conn, batch, overrides=None):
    definition = get_assembly_definition(
        conn, batch["customer"], batch["assembly_drawing_no"]
    )
    definition_ids = {int(row["manual_id"]) for row in definition}
    if overrides is None:
        old_quantities = {
            int(item["manual_id"]): int(item["shipped_quantity"])
            for item in batch["items"]
        }
        overrides = {
            manual_id: quantity
            for manual_id, quantity in old_quantities.items()
            if manual_id in definition_ids
        }
    return build_assembly_shipment_preview(
        conn,
        batch["customer"],
        batch["assembly_drawing_no"],
        batch["set_quantity"],
        overrides,
        exclude_batch_id=batch["id"],
    )


@app.route("/admin/shipped-orders/assembly/new", methods=["POST"])
@permission_required("shipped_manage")
def create_assembly_shipment_from_shipped_page():
    customer = request.form.get("customer", "")
    assembly_drawing_no = request.form.get("assembly_drawing_no", "")
    set_quantity = request.form.get("set_quantity", "")
    shipped_at = request.form.get("shipped_at", "").strip()
    logistics_no = request.form.get("logistics_no", "").strip()
    submitted_token = request.form.get("preview_token", "").strip()
    if not shipped_at:
        return jsonify({"error": "请填写发货时间"}), 400
    if not submitted_token:
        return jsonify({"error": "缺少组装发货预览标识"}), 400

    staged_images = []
    transaction_committed = False
    try:
        staged_images = stage_assembly_shipment_images(
            selected_shipment_images()
        )
        overrides = _parse_assembly_shipment_item_overrides(request.form)
        with get_db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            preview = build_assembly_shipment_preview(
                conn,
                customer,
                assembly_drawing_no,
                set_quantity,
                overrides,
            )
            definition_ids = {item["manual_id"] for item in preview["items"]}
            if set(overrides) != definition_ids:
                raise ValueError("必须提交组装件定义中的全部配件且每项仅提交一次")
            if submitted_token != preview["preview_token"]:
                return (
                    jsonify(
                        {
                            "error": "订单或库存状态已变化，请按最新结果重新确认",
                            "preview": preview,
                        }
                    ),
                    409,
                )
            if preview["warnings"] and request.form.get("confirm_warnings") != "1":
                return (
                    jsonify(
                        {
                            "error": "请确认当前组装发货警告后再保存",
                            "preview": preview,
                        }
                    ),
                    409,
                )
            batch_id = save_assembly_shipment(
                conn,
                preview,
                shipped_at,
                logistics_no,
                current_admin_username(),
            )
            if staged_images:
                save_assembly_shipment_images(conn, batch_id, staged_images)
            conn.commit()
            transaction_committed = True
    except AssemblyDefinitionNotFound as error:
        return jsonify({"error": str(error)}), 400
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    except sqlite3.DatabaseError as error:
        if not transaction_committed:
            return jsonify({"error": "组装发货保存失败，所有更改已回滚"}), 500
        app.logger.error("组装发货已提交，但数据库连接退出失败：%s", error)
    except OSError as error:
        if not transaction_committed:
            return jsonify({"error": "组装发货图片保存失败，所有更改已回滚"}), 500
        app.logger.error("组装发货已提交，但请求退出失败：%s", error)
    except Exception as error:
        if not transaction_committed:
            raise
        app.logger.error("组装发货已提交，但请求退出异常：%s", error)
    finally:
        cleanup_assembly_shipment_image_stages(
            staged_images, remove_final=not transaction_committed
        )

    return (
        jsonify(
            {
                "batch_id": batch_id,
                "redirect_url": url_for("shipped_orders"),
            }
        ),
        201,
    )


@app.route(
    "/admin/shipped-orders/assembly/<int:batch_id>/edit",
    methods=["GET", "POST"],
)
@permission_required("shipped_manage")
def edit_assembly_shipment(batch_id):
    if request.method == "GET":
        with get_db() as conn:
            batches = _fetch_assembly_shipment_batches_by_ids(conn, [batch_id])
            if not batches:
                abort(404)
            batch = batches[0]
            try:
                preview = _assembly_edit_preview(conn, batch)
            except (AssemblyDefinitionNotFound, ValueError) as error:
                return str(error), 409
        return render_template(
            "assembly_shipment_edit.html", batch=batch, preview=preview
        )

    if request.is_json:
        payload = request.get_json(silent=True)
        if not isinstance(payload, dict):
            return jsonify({"error": "请求数据格式无效"}), 400
        try:
            with get_db() as conn:
                batches = _fetch_assembly_shipment_batches_by_ids(conn, [batch_id])
                if not batches:
                    abort(404)
                preview = _assembly_edit_preview(
                    conn, batches[0], payload.get("overrides")
                )
        except AssemblyDefinitionNotFound as error:
            return jsonify({"error": str(error)}), 404
        except ValueError as error:
            return jsonify({"error": str(error)}), 400
        return jsonify(preview)

    shipped_at = request.form.get("shipped_at", "").strip()
    logistics_no = request.form.get("logistics_no", "").strip()
    submitted_token = request.form.get("preview_token", "").strip()
    if not shipped_at:
        return jsonify({"error": "请填写发货时间"}), 400
    if not submitted_token:
        return jsonify({"error": "缺少组装发货预览标识"}), 400
    staged_images = []
    transaction_committed = False
    try:
        staged_images = stage_assembly_shipment_images(
            selected_shipment_images()
        )
        overrides = _parse_assembly_shipment_item_overrides(request.form)
        with get_db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            batches = _fetch_assembly_shipment_batches_by_ids(
                conn,
                [batch_id],
                include_prices=True,
            )
            if not batches:
                abort(404)
            batch = batches[0]
            preview = _assembly_edit_preview(conn, batch, overrides)
            definition_ids = {item["manual_id"] for item in preview["items"]}
            if set(overrides) != definition_ids:
                raise ValueError("必须提交组装件定义中的全部配件且每项仅提交一次")
            if submitted_token != preview["preview_token"]:
                return (
                    jsonify(
                        {
                            "error": "订单或库存状态已变化，请按最新结果重新确认",
                            "preview": preview,
                        }
                    ),
                    409,
                )
            if preview["warnings"] and request.form.get("confirm_warnings") != "1":
                return (
                    jsonify(
                        {
                            "error": "请确认当前组装发货警告后再保存",
                            "preview": preview,
                        }
                    ),
                    409,
                )
            existing_price_snapshots = {
                int(item["manual_id"]): {
                    "unit_price_minor": item["unit_price_minor"],
                    "currency": item["currency"],
                    "price_recorded_by": item["price_recorded_by"],
                    "price_recorded_at": item["price_recorded_at"],
                }
                for item in batch["items"]
            }
            old_order_ids = reverse_assembly_shipment_batch(conn, batch_id)
            replace_assembly_shipment(
                conn,
                batch_id,
                preview,
                shipped_at,
                logistics_no,
                old_order_ids,
                existing_price_snapshots=existing_price_snapshots,
                recorded_by=current_admin_username(),
            )
            if staged_images:
                save_assembly_shipment_images(conn, batch_id, staged_images)
            conn.commit()
            transaction_committed = True
    except AssemblyDefinitionNotFound as error:
        return jsonify({"error": str(error)}), 400
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    except sqlite3.DatabaseError as error:
        if not transaction_committed:
            return jsonify({"error": "组装发货修改失败，所有更改已回滚"}), 500
        app.logger.error("组装发货修改已提交，但数据库连接退出失败：%s", error)
    except OSError as error:
        if not transaction_committed:
            return jsonify({"error": "组装发货图片保存失败，所有更改已回滚"}), 500
        app.logger.error("组装发货修改已提交，但请求退出失败：%s", error)
    except Exception as error:
        if not transaction_committed:
            raise
        app.logger.error("组装发货修改已提交，但请求退出异常：%s", error)
    finally:
        cleanup_assembly_shipment_image_stages(
            staged_images, remove_final=not transaction_committed
        )
    return (
        jsonify(
            {"batch_id": batch_id, "redirect_url": url_for("shipped_orders")}
        ),
        201,
    )


@app.route(
    "/admin/shipped-orders/assembly/<int:batch_id>/delete", methods=["POST"]
)
@permission_required("shipped_manage")
def delete_assembly_shipment(batch_id):
    image_filenames = []
    try:
        with get_db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            batches = _fetch_assembly_shipment_batches_by_ids(conn, [batch_id])
            if not batches:
                abort(404)
            image_filenames = [
                image["filename"] for image in batches[0]["images"]
            ]
            old_order_ids = reverse_assembly_shipment_batch(conn, batch_id)
            conn.execute(
                "DELETE FROM assembly_shipment_images WHERE batch_id = ?",
                (batch_id,),
            )
            conn.execute(
                "DELETE FROM assembly_shipment_batches WHERE id = ?", (batch_id,)
            )
            for order_id in sorted(old_order_ids):
                sync_order_shipment_summary(conn, order_id)
    except sqlite3.DatabaseError:
        return jsonify({"error": "组装发货删除失败，所有更改已回滚"}), 500
    delete_assembly_shipment_image_files(batch_id, image_filenames)
    flash("组装发货批次已删除", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/new", methods=["POST"])
@permission_required("shipped_manage")
def create_shipment_from_shipped_page():
    shipped_at = request.form.get("shipped_at", "").strip()
    logistics_no = request.form.get("logistics_no", "").strip()
    order_ids = request.form.getlist("order_id")
    shipped_quantities = request.form.getlist("shipped_quantity")

    if not shipped_at:
        flash("请填写发货时间", "error")
        return redirect(url_for("shipped_orders"))

    requested_items = []
    for order_id, shipped_quantity in zip_longest(order_ids, shipped_quantities, fillvalue=""):
        order_id = str(order_id or "").strip()
        shipped_quantity = str(shipped_quantity or "").strip()
        if not order_id and not shipped_quantity:
            continue
        if not order_id or not shipped_quantity:
            flash("请选择订单，并填写发货数量", "error")
            return redirect(url_for("shipped_orders"))
        try:
            order_id_value = int(order_id)
            shipped_quantity_value = int(shipped_quantity)
        except ValueError:
            flash("发货数量必须为整数", "error")
            return redirect(url_for("shipped_orders"))
        if shipped_quantity_value <= 0:
            flash("发货数量必须大于 0", "error")
            return redirect(url_for("shipped_orders"))
        requested_items.append((order_id_value, shipped_quantity_value))

    if not requested_items:
        flash("请选择订单，并填写发货数量和发货时间", "error")
        return redirect(url_for("shipped_orders"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    shipment_ids = []
    image_files = selected_shipment_images()
    with get_db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        unique_order_ids = list(dict.fromkeys(order_id for order_id, _ in requested_items))
        placeholders = ",".join("?" for _ in unique_order_ids)
        rows = conn.execute(
            f"""
            SELECT product_orders.*,
                   product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS unshipped_quantity
            FROM product_orders
            LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments ON shipments.order_id = product_orders.id
            WHERE product_orders.id IN ({placeholders})
            """,
            unique_order_ids,
        ).fetchall()
        orders_by_id = {row["id"]: row for row in rows}
        remaining_by_order = {
            row["id"]: int(row["unshipped_quantity"] or 0)
            for row in rows
        }
        for order_id_value, shipped_quantity_value in requested_items:
            order = orders_by_id.get(order_id_value)
            if order is None:
                flash("请选择有效的订单", "error")
                return redirect(url_for("shipped_orders"))
            if shipped_quantity_value > remaining_by_order.get(order_id_value, 0):
                flash("发货数量不能大于未发数量", "error")
                return redirect(url_for("shipped_orders"))
            remaining_by_order[order_id_value] -= shipped_quantity_value

        try:
            validate_shipment_inventory(conn, requested_items)
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("shipped_orders"))

        for order_id_value, shipped_quantity_value in requested_items:
            order = orders_by_id[order_id_value]
            price_snapshot = current_product_price_snapshot(
                conn,
                order["manual_id"],
                current_admin_username(),
                now,
            )
            signature_token = unique_signature_token(conn)
            photo_upload_token = unique_shipment_photo_token(conn)
            signature_expires = signature_expires_at()
            cursor = conn.execute(
                """
                INSERT INTO product_order_shipments (
                    order_id, shipped_quantity, shipped_at, created_at, email_sent_at, logistics_no,
                    signature_token, signature_status, signature_expires_at, photo_upload_token,
                    unit_price_minor, currency, price_recorded_by, price_recorded_at
                )
                VALUES (?, ?, ?, ?, '', ?, ?, '未签收', ?, ?, ?, ?, ?, ?)
                """,
                (
                    order_id_value,
                    shipped_quantity_value,
                    shipped_at,
                    now,
                    logistics_no,
                    signature_token,
                    signature_expires,
                    photo_upload_token,
                    price_snapshot["unit_price_minor"],
                    price_snapshot["currency"],
                    price_snapshot["price_recorded_by"],
                    price_snapshot["price_recorded_at"],
                ),
            )
            shipment_id = cursor.lastrowid
            shipment_ids.append(shipment_id)
            deduct_inventory_for_shipment(
                conn,
                shipment_id,
                order_id_value,
                shipped_quantity_value,
            )
            if image_files:
                for file_storage in image_files:
                    try:
                        file_storage.stream.seek(0)
                    except Exception:
                        pass
                try:
                    save_shipment_images(conn, shipment_id, image_files)
                except ValueError as error:
                    conn.rollback()
                    flash(f"发货图片格式不支持：{error}", "error")
                    return redirect(url_for("shipped_orders"))
            sync_order_shipment_summary(conn, order_id_value, updated_at=now)
    flash(f"已新增 {len(shipment_ids)} 条发货记录", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/plans/<int:plan_id>/quantities", methods=["POST"])
@permission_required("shipped_manage")
def update_shipment_plan_quantities(plan_id):
    item_ids = request.form.getlist("plan_item_id")
    quantities = request.form.getlist("planned_quantity")
    updates = {}
    for item_id, quantity in zip_longest(item_ids, quantities, fillvalue=""):
        try:
            item_id_value = int(str(item_id).strip())
            quantity_value = int(str(quantity or "0").strip() or "0")
        except ValueError:
            flash("计划数量必须为整数", "error")
            return redirect(url_for("shipped_orders"))
        if quantity_value < 0:
            flash("计划数量不能小于 0", "error")
            return redirect(url_for("shipped_orders"))
        updates[item_id_value] = quantity_value

    if not updates:
        flash("没有可保存的计划数量", "error")
        return redirect(url_for("shipped_orders"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        plan = conn.execute(
            "SELECT * FROM shipment_plans WHERE id = ?",
            (plan_id,),
        ).fetchone()
        if plan is None:
            flash("计划发货清单不存在", "error")
            return redirect(url_for("shipped_orders"))
        if plan["status"] == "已完成":
            flash("已完成的计划发货清单不能修改数量", "error")
            return redirect(url_for("shipped_orders"))

        rows = conn.execute(
            """
            SELECT id, shipped_quantity
            FROM shipment_plan_items
            WHERE plan_id = ?
            """,
            (plan_id,),
        ).fetchall()
        items_by_id = {row["id"]: row for row in rows}
        for item_id, quantity_value in updates.items():
            item = items_by_id.get(item_id)
            if item is None:
                flash("计划明细不存在，请刷新后重试", "error")
                return redirect(url_for("shipped_orders"))
            shipped_quantity = int(item["shipped_quantity"] or 0)
            if quantity_value < shipped_quantity:
                flash("计划数量不能小于已发数量", "error")
                return redirect(url_for("shipped_orders"))

        conn.executemany(
            """
            UPDATE shipment_plan_items
            SET planned_quantity = ?, updated_at = ?
            WHERE id = ? AND plan_id = ?
            """,
            [
                (quantity_value, now, item_id, plan_id)
                for item_id, quantity_value in updates.items()
            ],
        )
        sync_shipment_plan_status(conn, plan_id, updated_at=now)

    flash("计划发货数量已保存", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/plans/<int:plan_id>/preview")
@permission_required("shipped_view")
def preview_shipment_plan(plan_id):
    with get_db() as conn:
        plan = fetch_shipment_plan_detail(conn, plan_id)
    if plan is None:
        flash("计划发货清单不存在", "error")
        return redirect(url_for("shipped_orders"))
    return render_template("shipment_plan_preview.html", plan=plan)


@app.route("/admin/shipped-orders/plans/<int:plan_id>/inspection-report")
@permission_required("shipped_view")
def preview_inspection_report(plan_id):
    with get_db() as conn:
        plan = fetch_shipment_plan_detail(conn, plan_id)
        if plan is not None:
            inspection_sections = build_shipment_plan_inspection_sections(conn, plan)
        else:
            inspection_sections = []
    if plan is None:
        flash("计划发货清单不存在", "error")
        return redirect(url_for("shipped_orders"))
    report_no = f"IR-{plan['plan_no']}"
    return render_template(
        "inspection_report_preview.html",
        plan=plan,
        report_no=report_no,
        inspection_sections=inspection_sections,
        generated_at=datetime.now().strftime("%Y-%m-%d"),
    )


@app.route("/admin/shipped-orders/plans/<int:plan_id>/excel")
@permission_required("shipped_view")
def export_shipment_plan_excel(plan_id):
    with get_db() as conn:
        plan = fetch_shipment_plan_detail(conn, plan_id)
    if plan is None:
        flash("计划发货清单不存在", "error")
        return redirect(url_for("shipped_orders"))

    workbook = build_shipment_plan_workbook(plan)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    safe_plan_no = re.sub(r"[^A-Za-z0-9_-]+", "-", plan["plan_no"] or f"shipment-plan-{plan_id}").strip("-")
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"计划发货清单-{safe_plan_no or plan_id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/admin/shipped-orders/plans/<int:plan_id>/delete", methods=["POST"])
@permission_required("shipped_manage")
def delete_shipment_plan(plan_id):
    with get_db() as conn:
        plan = conn.execute(
            "SELECT * FROM shipment_plans WHERE id = ?",
            (plan_id,),
        ).fetchone()
        if plan is None:
            flash("计划发货清单不存在", "error")
            return redirect(url_for("shipped_orders"))
        reports = conn.execute(
            "SELECT filename FROM shipment_plan_inspection_reports WHERE plan_id = ?",
            (plan_id,),
        ).fetchall()
        for report in reports:
            (INSPECTION_REPORTS_DIR / secure_filename(report["filename"])).unlink(missing_ok=True)
        conn.execute("DELETE FROM shipment_plan_inspection_reports WHERE plan_id = ?", (plan_id,))
        conn.execute("DELETE FROM shipment_plan_items WHERE plan_id = ?", (plan_id,))
        conn.execute("DELETE FROM shipment_plans WHERE id = ?", (plan_id,))

    flash("计划发货清单已删除，已生成的发货记录保留在已发货物清单中", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/plans/<int:plan_id>/inspection-reports", methods=["POST"])
@permission_required("shipped_manage")
def upload_shipment_plan_inspection_reports(plan_id):
    files = selected_inspection_report_files()
    if not files:
        flash("请选择要上传的检验报告文件", "error")
        return redirect(url_for("shipped_orders"))

    with get_db() as conn:
        plan = conn.execute(
            "SELECT * FROM shipment_plans WHERE id = ?",
            (plan_id,),
        ).fetchone()
        if plan is None:
            flash("计划发货清单不存在", "error")
            return redirect(url_for("shipped_orders"))
        try:
            saved = save_inspection_report_files(conn, plan_id, files)
        except ValueError as error:
            flash(f"检验报告文件格式不支持：{error}", "error")
            return redirect(url_for("shipped_orders"))

    flash(f"已上传 {len(saved)} 个检验报告文件", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/plans/<int:plan_id>/approve", methods=["POST"])
@permission_required("shipped_manage")
def approve_shipment_plan(plan_id):
    shipped_at = request.form.get("shipped_at", "").strip()
    logistics_no = request.form.get("logistics_no", "").strip()
    item_ids = request.form.getlist("plan_item_id")
    quantities = request.form.getlist("shipped_quantity")

    if not shipped_at:
        flash("请填写发货时间", "error")
        return redirect(url_for("shipped_orders"))

    requested = {}
    for item_id, quantity in zip_longest(item_ids, quantities, fillvalue=""):
        try:
            item_id_value = int(str(item_id).strip())
            quantity_value = int(str(quantity or "0").strip() or "0")
        except ValueError:
            flash("实际发货数量必须为整数", "error")
            return redirect(url_for("shipped_orders"))
        if quantity_value < 0:
            flash("实际发货数量不能小于 0", "error")
            return redirect(url_for("shipped_orders"))
        requested[item_id_value] = quantity_value

    if not any(quantity > 0 for quantity in requested.values()):
        flash("请至少填写一项实际发货数量", "error")
        return redirect(url_for("shipped_orders"))

    now = datetime.utcnow().isoformat(timespec="seconds")
    shipment_ids = []
    image_files = selected_shipment_images()
    with get_db() as conn:
        plan = conn.execute(
            "SELECT * FROM shipment_plans WHERE id = ?",
            (plan_id,),
        ).fetchone()
        if plan is None:
            flash("计划发货清单不存在", "error")
            return redirect(url_for("shipped_orders"))
        if plan["status"] == "已完成":
            flash("该计划发货清单已完成", "error")
            return redirect(url_for("shipped_orders"))

        rows = conn.execute(
            f"""
            SELECT shipment_plan_items.*,
                   product_orders.quantity - COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS unshipped_quantity
            FROM shipment_plan_items
            JOIN product_orders ON product_orders.id = shipment_plan_items.order_id
            LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments ON shipments.order_id = product_orders.id
            WHERE shipment_plan_items.plan_id = ?
            ORDER BY shipment_plan_items.id ASC
            """,
            (plan_id,),
        ).fetchall()
        items_by_id = {row["id"]: row for row in rows}
        for item_id, quantity_value in requested.items():
            if quantity_value <= 0:
                continue
            item = items_by_id.get(item_id)
            if item is None:
                flash("计划明细不存在，请刷新后重试", "error")
                return redirect(url_for("shipped_orders"))
            remaining_plan_quantity = int(item["planned_quantity"] or 0) - int(item["shipped_quantity"] or 0)
            max_quantity = min(remaining_plan_quantity, int(item["unshipped_quantity"] or 0))
            if quantity_value > max_quantity:
                flash("实际发货数量不能超过计划剩余数量或订单未发数量", "error")
                return redirect(url_for("shipped_orders"))

        shipment_requests = [
            (items_by_id[item_id]["order_id"], quantity_value)
            for item_id, quantity_value in requested.items()
            if quantity_value > 0
        ]
        try:
            validate_shipment_inventory(conn, shipment_requests)
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("shipped_orders"))

        for item_id, quantity_value in requested.items():
            if quantity_value <= 0:
                continue
            item = items_by_id[item_id]
            signature_token = unique_signature_token(conn)
            photo_upload_token = unique_shipment_photo_token(conn)
            signature_expires = signature_expires_at()
            cursor = conn.execute(
                """
                INSERT INTO product_order_shipments (
                    order_id, shipped_quantity, shipped_at, created_at, email_sent_at, logistics_no,
                    signature_token, signature_status, signature_expires_at, photo_upload_token
                )
                VALUES (?, ?, ?, ?, '', ?, ?, '未签收', ?, ?)
                """,
                (
                    item["order_id"],
                    quantity_value,
                    shipped_at,
                    now,
                    logistics_no,
                    signature_token,
                    signature_expires,
                    photo_upload_token,
                ),
            )
            shipment_id = cursor.lastrowid
            shipment_ids.append(shipment_id)
            deduct_inventory_for_shipment(
                conn,
                shipment_id,
                item["order_id"],
                quantity_value,
            )
            if image_files:
                for file_storage in image_files:
                    try:
                        file_storage.stream.seek(0)
                    except Exception:
                        pass
                try:
                    save_shipment_images(conn, shipment_id, image_files)
                except ValueError as error:
                    conn.rollback()
                    flash(f"发货图片格式不支持：{error}", "error")
                    return redirect(url_for("shipped_orders"))
            conn.execute(
                """
                UPDATE shipment_plan_items
                SET shipped_quantity = shipped_quantity + ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (quantity_value, now, item_id),
            )
            sync_order_shipment_summary(conn, item["order_id"], updated_at=now)

        conn.execute(
            """
            UPDATE shipment_plans
            SET reviewed_by = ?, reviewed_at = ?, updated_at = ?
            WHERE id = ?
            """,
            (current_admin_username(), now, now, plan_id),
        )
        status = "已完成"
        conn.execute(
            """
            UPDATE shipment_plans
            SET status = ?, updated_at = ?
            WHERE id = ?
            """,
            (status, now, plan_id),
        )

    flash(f"计划发货清单已审核，状态：{status}", "success")
    if shipment_ids:
        flash(f"已生成 {len(shipment_ids)} 条发货记录", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/<int:shipment_id>/signature-link", methods=["POST"])
@permission_required("shipped_manage")
def generate_shipment_signature_link(shipment_id):
    with get_db() as conn:
        shipment = fetch_shipment_by_id(conn, shipment_id)
        if shipment is None:
            flash("发货记录不存在", "error")
            return redirect(url_for("shipped_orders"))
        already_signed = shipment["signature_status"] == "已签收" or bool(shipment["signed_at"])
        if already_signed:
            flash("该发货记录已签收，不能重新生成签字链接", "error")
            return redirect(url_for("shipped_orders"))
        should_regenerate = not shipment["signature_token"] or signature_link_expired(shipment)
        if should_regenerate:
            conn.execute(
                """
                UPDATE product_order_shipments
                SET signature_token = ?,
                    signature_expires_at = ?,
                    signature_status = CASE
                        WHEN signature_status = '已签收' THEN signature_status
                        ELSE '未签收'
                    END
                WHERE id = ?
                """,
                (unique_signature_token(conn), signature_expires_at(), shipment_id),
            )
        shipment = fetch_shipment_by_id(conn, shipment_id)

    flash(f"签字链接：{shipment_signature_url(shipment)}", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/<int:shipment_id>/photo-link", methods=["POST"])
@permission_required("shipped_manage")
def generate_shipment_photo_link(shipment_id):
    with get_db() as conn:
        shipment = fetch_shipment_by_id(conn, shipment_id)
        if shipment is None:
            flash("发货记录不存在", "error")
            return redirect(url_for("shipped_orders"))
        if not shipment["photo_upload_token"]:
            conn.execute(
                """
                UPDATE product_order_shipments
                SET photo_upload_token = ?
                WHERE id = ?
                """,
                (unique_shipment_photo_token(conn), shipment_id),
            )
        shipment = fetch_shipment_by_id(conn, shipment_id)

    flash(f"发货图片上传链接：{shipment_photo_upload_url(shipment)}", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/<int:shipment_id>/remark", methods=["POST"])
@permission_required("shipped_manage")
def update_shipment_remark(shipment_id):
    remark = request.form.get("remark", "").strip()
    with get_db() as conn:
        shipment = fetch_shipment_by_id(conn, shipment_id)
        if shipment is None:
            flash("发货记录不存在", "error")
            return redirect(url_for("shipped_orders"))
        conn.execute(
            """
            UPDATE product_order_shipments
            SET remark = ?
            WHERE id = ?
            """,
            (remark, shipment_id),
        )
    flash("发货备注已保存", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/<int:shipment_id>/images", methods=["POST"])
@permission_required("shipped_manage")
def upload_shipment_images_admin(shipment_id):
    files = selected_shipment_images()
    if not files:
        flash("请选择要上传的发货图片", "error")
        return redirect(url_for("shipped_orders"))

    with get_db() as conn:
        shipment = fetch_shipment_by_id(conn, shipment_id)
        if shipment is None:
            flash("发货记录不存在", "error")
            return redirect(url_for("shipped_orders"))
        try:
            saved = save_shipment_images(conn, shipment_id, files)
        except ValueError as error:
            flash(f"发货图片格式不支持：{error}", "error")
            return redirect(url_for("shipped_orders"))

    flash(f"已上传 {len(saved)} 张发货图片", "success")
    return redirect(url_for("shipped_orders"))


@app.route(
    "/admin/shipped-orders/prices/<source_type>/<int:source_id>",
    methods=["POST"],
)
@permission_required("shipped_manage")
def backfill_shipment_price(source_type, source_id):
    if not user_can_view_prices():
        flash("当前账号没有权限查看或补录价格", "error")
        return redirect(url_for("shipped_orders"))

    sources = {
        "ordinary": ("product_order_shipments", "发货记录"),
        "assembly_item": ("assembly_shipment_items", "组装发货配件"),
    }
    source = sources.get(source_type)
    if source is None:
        flash("价格补录来源无效", "error")
        return redirect(url_for("shipped_orders"))

    try:
        currency = normalize_currency(request.form.get("currency", "CNY"))
        unit_price_minor = parse_money_minor(
            request.form.get("unit_price", ""),
            currency,
        )
        if unit_price_minor is None:
            raise ValueError("请填写发货单价")
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("shipped_orders"))

    table_name, source_label = source
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute(
            f"SELECT id, unit_price_minor FROM {table_name} WHERE id = ?",
            (source_id,),
        ).fetchone()
        if row is None:
            flash(f"{source_label}不存在", "error")
            return redirect(url_for("shipped_orders"))
        if row["unit_price_minor"] is not None:
            flash(f"{source_label}已经记录价格，不能覆盖", "error")
            return redirect(url_for("shipped_orders"))
        conn.execute(
            f"""
            UPDATE {table_name}
            SET unit_price_minor = ?,
                currency = ?,
                price_recorded_by = ?,
                price_recorded_at = ?
            WHERE id = ? AND unit_price_minor IS NULL
            """,
            (
                unit_price_minor,
                currency,
                current_admin_username(),
                now,
                source_id,
            ),
        )

    flash("历史发货价格已补录", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/admin/shipped-orders/<int:shipment_id>/edit", methods=["GET", "POST"])
@permission_required("shipped_manage")
def edit_shipment(shipment_id):
    with get_db() as conn:
        shipment = fetch_shipment_by_id(
            conn,
            shipment_id,
            include_prices=user_can_view_prices(),
        )
        if shipment is None:
            flash("发货记录不存在", "error")
            return redirect(url_for("shipped_orders"))

        if request.method == "POST":
            shipped_quantity = request.form.get("shipped_quantity", "").strip()
            shipped_at = request.form.get("shipped_at", "").strip()
            logistics_no = request.form.get("logistics_no", "").strip()

            if not all([shipped_quantity, shipped_at]):
                flash("发货数量和发货时间为必填项", "error")
                return redirect(url_for("edit_shipment", shipment_id=shipment_id))

            try:
                shipped_quantity_value = int(shipped_quantity)
            except ValueError:
                flash("发货数量必须为整数", "error")
                return redirect(url_for("edit_shipment", shipment_id=shipment_id))

            if shipped_quantity_value <= 0:
                flash("发货数量必须大于 0", "error")
                return redirect(url_for("edit_shipment", shipment_id=shipment_id))

            order = conn.execute(
                """
                SELECT quantity
                FROM product_orders
                WHERE id = ?
                """,
                (shipment["order_id"],),
            ).fetchone()
            combined_total = conn.execute(
                f"""
                SELECT shipped_total
                FROM ({order_shipment_summary_subquery(include_assembly=True, conn=conn)})
                WHERE order_id = ?
                """,
                (shipment["order_id"],),
            ).fetchone()
            if order is None:
                flash("关联订单不存在", "error")
                return redirect(url_for("shipped_orders"))
            other_total = max(
                0,
                int(combined_total["shipped_total"] if combined_total else 0)
                - int(shipment["shipped_quantity"] or 0),
            )
            if shipped_quantity_value + other_total > order["quantity"]:
                flash("修改后的发货数量超过了订单总数量", "error")
                return redirect(url_for("edit_shipment", shipment_id=shipment_id))

            tracked_inventory_quantity = shipment_inventory_deducted_quantity(conn, shipment_id)
            try:
                validate_shipment_inventory(
                    conn,
                    [(shipment["order_id"], shipped_quantity_value)],
                    {shipment["manual_id"]: tracked_inventory_quantity},
                )
            except ValueError as error:
                flash(str(error), "error")
                return redirect(url_for("edit_shipment", shipment_id=shipment_id))

            now = datetime.utcnow().isoformat(timespec="seconds")
            reverse_shipment_inventory_deduction(conn, shipment_id)
            conn.execute(
                """
                UPDATE product_order_shipments
                SET shipped_quantity = ?,
                    shipped_at = ?,
                    logistics_no = ?
                WHERE id = ?
                """,
                (shipped_quantity_value, shipped_at, logistics_no, shipment_id),
            )
            deduct_inventory_for_shipment(
                conn,
                shipment_id,
                shipment["order_id"],
                shipped_quantity_value,
            )
            signature_reset = shipment["signature_status"] == "已签收" or bool(shipment["signed_at"])
            if signature_reset:
                reset_shipment_signature_state(conn, shipment_id)
            sync_order_shipment_summary(conn, shipment["order_id"], updated_at=now)
            flash("发货记录已更新", "success")
            if signature_reset:
                flash("该发货记录原已签收，修改后已重置签收状态并生成新的签字链接", "success")
            return redirect(url_for("shipped_orders"))

    return render_template("shipment_edit.html", shipment=shipment)


@app.route("/admin/shipped-orders/<int:shipment_id>/delete", methods=["POST"])
@permission_required("shipped_manage")
def delete_shipment(shipment_id):
    with get_db() as conn:
        shipment = fetch_shipment_by_id(conn, shipment_id)
        if shipment is None:
            flash("发货记录不存在", "error")
            return redirect(url_for("shipped_orders"))
        reverse_shipment_inventory_deduction(conn, shipment_id)
        delete_shipment_assets(conn, shipment_id, signature_image=shipment["signature_image"] or "")
        conn.execute("DELETE FROM product_order_shipments WHERE id = ?", (shipment_id,))
        sync_order_shipment_summary(conn, shipment["order_id"])

    flash("发货记录已删除", "success")
    return redirect(url_for("shipped_orders"))


@app.route("/signature-image/<path:filename>")
@login_required
def signature_image(filename):
    safe_name = secure_filename(filename)
    if safe_name != filename or not safe_name.lower().endswith(".png"):
        abort(404)
    return send_from_directory(SIGNATURES_DIR, safe_name)


def shipment_image_file_response(safe_name):
    suffix = Path(safe_name).suffix.lower()
    safe_inline = suffix in SHIPMENT_RASTER_SUFFIX_FORMATS
    response = send_from_directory(
        SHIPMENT_IMAGES_DIR,
        safe_name,
        as_attachment=not safe_inline,
        mimetype=None if safe_inline else "application/octet-stream",
    )
    response.headers["X-Content-Type-Options"] = "nosniff"
    if not safe_inline:
        response.headers["Content-Security-Policy"] = (
            "sandbox; default-src 'none'; frame-ancestors 'none'"
        )
    return response


@app.route("/shipment-image/<path:filename>")
@permission_required("shipped_view")
def shipment_image(filename):
    safe_name = secure_filename(filename)
    if safe_name != filename or Path(safe_name).suffix.lower() not in IMAGE_EXTENSIONS:
        abort(404)
    return shipment_image_file_response(safe_name)


@app.route("/inspection-report-file/<path:filename>")
@login_required
def inspection_report_file(filename):
    safe_name = secure_filename(filename)
    if safe_name != filename or Path(safe_name).suffix.lower() not in ALLOWED_EXTENSIONS:
        abort(404)
    with get_db() as conn:
        report = conn.execute(
            """
            SELECT original_filename
            FROM shipment_plan_inspection_reports
            WHERE filename = ?
            """,
            (safe_name,),
        ).fetchone()
    if report is None:
        abort(404)
    return send_from_directory(
        INSPECTION_REPORTS_DIR,
        safe_name,
        mimetype=guess_type(safe_name)[0] or "application/octet-stream",
    )


@app.route("/shipment-photos/<token>", methods=["GET", "POST"])
def upload_shipment_photos(token):
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,}", token or ""):
        return render_template("shipment_photos.html", invalid=True), 404

    with get_db() as conn:
        shipment = fetch_shipment_by_photo_token(conn, token)
        if shipment is None:
            return render_template("shipment_photos.html", invalid=True), 404

        error = ""
        uploaded_count = 0
        if request.method == "POST":
            files = selected_shipment_images()
            if not files:
                error = "请选择要上传的发货图片"
            else:
                try:
                    uploaded_count = len(save_shipment_images(conn, shipment["id"], files))
                except ValueError as upload_error:
                    error = f"图片格式不支持：{upload_error}"
        images = get_shipment_images(conn, [shipment["id"]]).get(shipment["id"], [])

    return render_template(
        "shipment_photos.html",
        shipment=shipment,
        images=images,
        error=error,
        uploaded_count=uploaded_count,
    )


@app.route("/shipment-photos/<token>/image/<path:filename>")
def shipment_photo_public_image(token, filename):
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,}", token or ""):
        abort(404)
    safe_name = secure_filename(filename)
    if safe_name != filename or Path(safe_name).suffix.lower() not in IMAGE_EXTENSIONS:
        abort(404)
    with get_db() as conn:
        shipment = fetch_shipment_by_photo_token(conn, token)
        if shipment is None:
            abort(404)
        exists = conn.execute(
            """
            SELECT 1
            FROM product_order_shipment_images
            WHERE shipment_id = ?
              AND filename = ?
            LIMIT 1
            """,
            (shipment["id"], safe_name),
        ).fetchone()
        if not exists:
            abort(404)
    return shipment_image_file_response(safe_name)


@app.route("/sign/<token>", methods=["GET", "POST"])
def sign_shipment(token):
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,}", token or ""):
        return render_template("sign.html", invalid=True), 404

    with get_db() as conn:
        shipment = fetch_shipment_by_token(conn, token)
        if shipment is None:
            return render_template("sign.html", invalid=True), 404

        already_signed = shipment["signature_status"] == "已签收" or bool(shipment["signed_at"])
        if not already_signed and signature_link_expired(shipment):
            return render_template("sign.html", invalid=True, expired=True), 410
        if request.method == "POST":
            if already_signed:
                return redirect(url_for("sign_shipment", token=token))
            signature_data = request.form.get("signature_data", "")
            try:
                image_bytes = decode_signature_image(signature_data)
            except ValueError as error:
                return render_template("sign.html", shipment=shipment, error=str(error)), 400

            filename = safe_signature_filename(shipment["id"], token)
            signature_path = SIGNATURES_DIR / filename
            SIGNATURES_DIR.mkdir(parents=True, exist_ok=True)
            signature_path.write_bytes(image_bytes)

            signed_at = datetime.utcnow().isoformat(timespec="seconds")
            cursor = conn.execute(
                """
                UPDATE product_order_shipments
                SET signature_status = '已签收',
                    signature_image = ?,
                    signed_at = ?,
                    signed_ip = ?,
                    signed_user_agent = ?
                WHERE id = ?
                  AND signature_status != '已签收'
                  AND signed_at = ''
                """,
                (
                    filename,
                    signed_at,
                    client_ip(),
                    request.headers.get("User-Agent", "")[:500],
                    shipment["id"],
                ),
            )
            if cursor.rowcount == 0:
                signature_path.unlink(missing_ok=True)
                return redirect(url_for("sign_shipment", token=token))
            shipment = fetch_shipment_by_token(conn, token)
            return render_template("sign.html", shipment=shipment, signed=True)

    return render_template("sign.html", shipment=shipment, signed=already_signed)


@app.route("/admin/shipped-orders/export", methods=["GET", "POST"])
@permission_required("shipped_view")
def export_shipped_orders():
    if request.method == "POST":
        shipment_ids = []
        for value in request.form.getlist("shipment_id"):
            try:
                shipment_ids.append(int(value))
            except ValueError:
                continue

        query = request.form.get("q", "").strip()
        selected_customer = request.form.get("customer", "").strip()
        shipped_at = request.form.get("shipped_at", "").strip()

        if not shipment_ids:
            flash("请选择要导出的发货记录", "error")
            return redirect(url_for("shipped_orders", q=query, customer=selected_customer, shipped_at=shipped_at))

        with get_db() as conn:
            shipped_orders = fetch_shipped_orders_by_ids(conn, shipment_ids)

        document_no = f"DN-{datetime.now().strftime('%Y%m%d%H%M%S')}"
        buffer = build_shipped_orders_pdf(shipped_orders, query, selected_customer, shipped_at, document_no=document_no)
        filename = f"{document_no}.pdf"
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype="application/pdf",
        )

    export_format = request.args.get("format", "xlsx").strip().lower()
    query = request.args.get("q", "").strip()
    selected_customer = request.args.get("customer", "").strip()
    shipped_at = request.args.get("shipped_at", "").strip()

    with get_db() as conn:
        shipped_orders = fetch_shipped_orders(
            conn,
            query=query,
            selected_customer=selected_customer,
            shipped_at=shipped_at,
        )

    if export_format == "pdf":
        document_no = f"DN-{datetime.now().strftime('%Y%m%d%H%M%S')}"
        buffer = build_shipped_orders_pdf(shipped_orders, query, selected_customer, shipped_at, document_no=document_no)
        filename = f"{document_no}.pdf"
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype="application/pdf",
        )

    workbook = build_shipped_orders_workbook(shipped_orders, query, selected_customer, shipped_at)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = "发货记录.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/admin/orders/new", methods=["GET", "POST"])
@permission_required("orders_manage")
def new_order():
    if request.method == "POST":
        order_no = request.form.get("order_no", "").strip()
        ordered_at = request.form.get("ordered_at", "").strip()
        selected_customer = request.form.get("customer", "").strip()
        assembly_drawing_no = request.form.get("assembly_drawing_no", "").strip()
        assembly_set_quantity_text = request.form.get(
            "assembly_set_quantity", ""
        ).strip()
        try:
            assembly_set_quantity = int(assembly_set_quantity_text or 0)
        except ValueError:
            flash("组装数量必须大于 0", "error")
            return redirect(url_for("new_order"))
        if bool(assembly_drawing_no) != bool(assembly_set_quantity_text):
            flash("组装图号和组装数量必须同时填写", "error")
            return redirect(url_for("new_order"))
        if assembly_drawing_no and assembly_set_quantity <= 0:
            flash("组装数量必须大于 0", "error")
            return redirect(url_for("new_order"))
        if assembly_set_quantity > MAX_ORDER_QUANTITY:
            flash(f"组装数量不能超过 {MAX_ORDER_QUANTITY}", "error")
            return redirect(url_for("new_order"))
        manual_ids = request.form.getlist("manual_id")
        quantities = request.form.getlist("quantity")
        planned_ship_dates = request.form.getlist("planned_ship_at")
        material_stock_statuses = request.form.getlist("material_stock_status")
        remarks = request.form.getlist("remark")
        carton_statuses = request.form.getlist("carton_status")
        recent_ship_statuses = request.form.getlist("recent_ship_status")

        if not all([order_no, ordered_at]):
            flash("订单号、下单时间为必填项", "error")
            return redirect(url_for("new_order"))
        if not selected_customer:
            flash("请选择客户", "error")
            return redirect(url_for("new_order"))

        items = []
        for manual_id, quantity, planned_ship_at, material_stock_status, remark, carton_status, recent_ship_status in zip_longest(
            manual_ids,
            quantities,
            planned_ship_dates,
            material_stock_statuses,
            remarks,
            carton_statuses,
            recent_ship_statuses,
            fillvalue="",
        ):
            manual_id = manual_id.strip()
            quantity = quantity.strip()
            planned_ship_at = planned_ship_at.strip()
            material_stock_status = material_stock_status.strip()
            remark = remark.strip()
            carton_status = carton_status.strip()
            recent_ship_status = recent_ship_status.strip()
            if not any([manual_id, quantity, planned_ship_at, material_stock_status, remark, carton_status, recent_ship_status]):
                continue
            if not manual_id or not quantity or not planned_ship_at:
                flash("每一行都需要选择产品、填写订单数量和计划发货时间", "error")
                return redirect(url_for("new_order"))
            try:
                manual_id_value = int(manual_id)
                quantity_value = int(quantity)
            except ValueError:
                flash("订单数量必须为整数", "error")
                return redirect(url_for("new_order"))
            if quantity_value <= 0:
                flash("订单数量必须大于 0", "error")
                return redirect(url_for("new_order"))
            if quantity_value > MAX_ORDER_QUANTITY:
                flash(f"订单数量不能超过 {MAX_ORDER_QUANTITY}", "error")
                return redirect(url_for("new_order"))
            items.append((manual_id_value, quantity_value, planned_ship_at, material_stock_status, remark, carton_status, recent_ship_status))

        if not items:
            flash("至少需要添加一个产品", "error")
            return redirect(url_for("new_order"))

        now = datetime.utcnow().isoformat(timespec="seconds")
        with get_db() as conn:
            if assembly_drawing_no:
                assembly_options = get_assembly_options_for_customer(
                    conn, selected_customer
                )
                canonical_assembly_drawing_no = next(
                    (
                        option
                        for option in assembly_options
                        if option.casefold() == assembly_drawing_no.casefold()
                    ),
                    None,
                )
                if canonical_assembly_drawing_no is None:
                    flash("请选择该客户有效的组装图号", "error")
                    return redirect(url_for("new_order"))
                assembly_drawing_no = canonical_assembly_drawing_no
            manual_rows = conn.execute(
                f"""
                SELECT id, customer
                FROM manuals
                WHERE id IN ({",".join("?" for _ in items)})
                """,
                [item[0] for item in items],
            ).fetchall()
            manuals_by_id = {row["id"]: row for row in manual_rows}
            if len(manuals_by_id) != len({item[0] for item in items}):
                flash("请选择有效的产品", "error")
                return redirect(url_for("new_order"))
            if selected_customer and any(
                (manuals_by_id[item[0]]["customer"] or "") != selected_customer
                for item in items
            ):
                flash("订单产品必须属于所选客户", "error")
                return redirect(url_for("new_order"))
            customer_emails_by_name = {
                row["name"]: row["email"]
                for row in conn.execute(
                    """
                    SELECT name, email
                    FROM customers
                    WHERE name IS NOT NULL AND name != ''
                    """
                ).fetchall()
            }

            conn.executemany(
                """
                INSERT INTO product_orders (
                    manual_id, order_no, ordered_at, quantity, customer,
                    assembly_drawing_no, assembly_set_quantity,
                    customer_email, planned_ship_at, material_stock_status, material_stock_assignee, remark,
                    carton_status, carton_assignee, recent_ship_status, shipped_quantity, shipped_at,
                    created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, '', ?, ?)
                """,
                [
                    (
                        manual_id,
                        order_no,
                        ordered_at,
                        quantity,
                        manuals_by_id[manual_id]["customer"] or "",
                        assembly_drawing_no,
                        assembly_set_quantity,
                        customer_emails_by_name.get(manuals_by_id[manual_id]["customer"] or "", ""),
                        planned_ship_at,
                        material_stock_status,
                        "",
                        remark,
                        carton_status,
                        "",
                        recent_ship_status,
                        now,
                        now,
                    )
                    for manual_id, quantity, planned_ship_at, material_stock_status, remark, carton_status, recent_ship_status in items
                ],
            )
        flash("产品订单已新增", "success")
        return redirect(url_for("admin_orders"))

    with get_db() as conn:
        products = get_product_options(conn)
        product_customers = sorted({
            product["customer"]
            for product in products
            if product["customer"]
        }, key=str.lower)
    return render_template(
        "order_new.html",
        products=products,
        product_customers=product_customers,
        default_ordered_at=datetime.now().date().isoformat(),
    )


@app.route("/admin/orders/assembly-options")
@permission_required("orders_manage")
def order_assembly_options():
    customer = request.args.get("customer", "")
    try:
        with get_db() as conn:
            drawing_numbers = get_assembly_options_for_customer(conn, customer)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    return jsonify({"assembly_drawing_numbers": drawing_numbers})


@app.route("/admin/orders/assembly-definition")
@permission_required("orders_manage")
def order_assembly_definition():
    customer = request.args.get("customer", "")
    requested_drawing_no = request.args.get("assembly_drawing_no", "")
    try:
        with get_db() as conn:
            drawing_numbers = get_assembly_options_for_customer(conn, customer)
            canonical_drawing_no = next(
                (
                    drawing_no
                    for drawing_no in drawing_numbers
                    if drawing_no.casefold() == requested_drawing_no.strip().casefold()
                ),
                None,
            )
            if not requested_drawing_no.strip():
                raise ValueError("请选择组装图号")
            if canonical_drawing_no is None:
                return jsonify({"error": "未找到该客户的组装图号配置"}), 404
            definition = get_assembly_definition(
                conn, customer, canonical_drawing_no
            )
    except ValueError as error:
        return jsonify({"error": str(error)}), 400

    return jsonify(
        {
            "assembly_drawing_no": canonical_drawing_no,
            "items": [
                {
                    "manual_id": row["manual_id"],
                    "drawing_no": row["drawing_no"],
                    "product_name": row["product_name"],
                    "quantity_per_set": row["quantity_per_set"],
                }
                for row in definition
            ],
        }
    )


@app.route("/admin/orders/import", methods=["GET", "POST"])
@permission_required("orders_manage")
def import_orders():
    if request.method == "POST":
        order_no = request.form.get("order_no", "").strip()
        ordered_at = request.form.get("ordered_at", "").strip()
        planned_ship_at = request.form.get("planned_ship_at", "").strip()
        upload = request.files.get("order_file")

        if not all([order_no, ordered_at]):
            flash("订单号、下单时间为必填项", "error")
            return redirect(url_for("import_orders"))
        if upload is None or not upload.filename:
            flash("请上传包含图号和数量两列的 Excel 文件", "error")
            return redirect(url_for("import_orders"))

        import_rows, errors = parse_order_import_workbook(upload)
        if errors:
            for error in errors[:20]:
                flash(error, "error")
            if len(errors) > 20:
                flash(f"还有 {len(errors) - 20} 个错误未显示，请先修正表格后重新导入", "error")
            return redirect(url_for("import_orders"))

        with get_db() as conn:
            items, errors = match_order_import_items(conn, import_rows)
            if errors:
                for error in errors[:20]:
                    flash(error, "error")
                if len(errors) > 20:
                    flash(f"还有 {len(errors) - 20} 个错误未显示，请先修正表格后重新导入", "error")
                return redirect(url_for("import_orders"))
            if not items:
                flash("没有找到可导入的订单明细", "error")
                return redirect(url_for("import_orders"))
            if not planned_ship_at and any(not item.get("planned_ship_at") for item in import_rows):
                flash("请填写默认计划发货时间，或在 Excel 中为每一行填写计划发货时间", "error")
                return redirect(url_for("import_orders"))

            now = datetime.utcnow().isoformat(timespec="seconds")
            planned_by_manual_id = {}
            for row in import_rows:
                if not row.get("planned_ship_at"):
                    continue
                matched_items, matched_errors = match_order_import_items(conn, [row])
                if matched_items and not matched_errors:
                    planned_by_manual_id[matched_items[0]["manual"]["id"]] = row["planned_ship_at"]
            customer_emails_by_name = {
                row["name"]: row["email"]
                for row in conn.execute(
                    """
                    SELECT name, email
                    FROM customers
                    WHERE name IS NOT NULL AND name != ''
                    """
                ).fetchall()
            }
            conn.executemany(
                """
                INSERT INTO product_orders (
                    manual_id, order_no, ordered_at, quantity, customer,
                    customer_email, planned_ship_at, material_stock_status, material_stock_assignee, remark,
                    carton_status, carton_assignee, shipped_quantity, shipped_at,
                    created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, '', '', '', '', '', 0, '', ?, ?)
                """,
                [
                    (
                        item["manual"]["id"],
                        order_no,
                        ordered_at,
                        item["quantity"],
                        item["manual"]["customer"] or "",
                        customer_emails_by_name.get(item["manual"]["customer"] or "", ""),
                        planned_by_manual_id.get(item["manual"]["id"]) or planned_ship_at,
                        now,
                        now,
                    )
                    for item in items
                ],
            )

        total_quantity = sum(item["quantity"] for item in items)
        flash(f"订单导入成功：{len(items)} 个产品，共 {total_quantity} 件", "success")
        return redirect(url_for("admin_orders"))

    return render_template("order_import.html")


@app.route("/admin/orders/import-template")
@permission_required("orders_manage")
def download_order_import_template():
    workbook = build_order_import_template_workbook()
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="订单导入模板.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/admin/orders/<int:order_id>/edit", methods=["GET", "POST"])
@permission_required("orders_manage")
def edit_order(order_id):
    with get_db() as conn:
        order = conn.execute(
            f"""
            SELECT product_orders.*,
                   manuals.drawing_no,
                   manuals.product_name,
                   manuals.customer AS product_customer,
                   COALESCE(shipments.shipped_total, product_orders.shipped_quantity, 0) AS shipped_quantity_total,
                   COALESCE(shipments.last_shipped_at, product_orders.shipped_at, '') AS shipped_at_display
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments
              ON shipments.order_id = product_orders.id
            WHERE product_orders.id = ?
            """,
            (order_id,),
        ).fetchone()
    if order is None:
        abort(404)

    if request.method == "POST":
        manual_id = request.form.get("manual_id", "").strip()
        order_no = request.form.get("order_no", "").strip()
        ordered_at = request.form.get("ordered_at", "").strip()
        quantity = request.form.get("quantity", "").strip()
        customer_email = request.form.get("customer_email", "").strip()
        planned_ship_at = request.form.get("planned_ship_at", "").strip()
        material_stock_status = request.form.get("material_stock_status", "").strip()
        remark = request.form.get("remark", "").strip()
        carton_status = request.form.get("carton_status", "").strip()
        recent_ship_status = request.form.get("recent_ship_status", "").strip()

        if not all([manual_id, order_no, ordered_at, quantity, planned_ship_at]):
            flash("产品图号、订单号、下单时间、订单数量、计划发货时间为必填项", "error")
            return redirect(url_for("edit_order", order_id=order_id))

        try:
            manual_id_value = int(manual_id)
            quantity_value = int(quantity)
        except ValueError:
            flash("订单数量必须为整数", "error")
            return redirect(url_for("edit_order", order_id=order_id))

        if quantity_value <= 0:
            flash("订单数量必须大于 0", "error")
            return redirect(url_for("edit_order", order_id=order_id))
        if quantity_value > MAX_ORDER_QUANTITY:
            flash(f"订单数量不能超过 {MAX_ORDER_QUANTITY}", "error")
            return redirect(url_for("edit_order", order_id=order_id))

        now = datetime.utcnow().isoformat(timespec="seconds")
        with get_db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            current_order = conn.execute(
                f"""
                SELECT product_orders.manual_id,
                       product_orders.customer,
                       product_orders.assembly_drawing_no,
                       product_orders.assembly_set_quantity,
                       COALESCE(
                           shipments.shipped_total,
                           product_orders.shipped_quantity,
                           0
                       ) AS shipped_total,
                       (
                           product_orders.shipped_quantity > 0
                           OR EXISTS (
                               SELECT 1
                               FROM product_order_shipments
                               WHERE product_order_shipments.order_id = product_orders.id
                           )
                           OR EXISTS (
                               SELECT 1
                               FROM assembly_shipment_allocations
                               WHERE assembly_shipment_allocations.order_id = product_orders.id
                           )
                       ) AS has_shipment_reference
                FROM product_orders
                LEFT JOIN ({order_shipment_summary_subquery(include_assembly=True, conn=conn)}) AS shipments
                  ON shipments.order_id = product_orders.id
                WHERE product_orders.id = ?
                """,
                (order_id,),
            ).fetchone()
            if current_order is None:
                abort(404)
            manual = conn.execute(
                "SELECT id, customer FROM manuals WHERE id = ?",
                (manual_id_value,),
            ).fetchone()
            if manual is None:
                flash("请选择有效的产品图号", "error")
                return redirect(url_for("edit_order", order_id=order_id))
            if (
                current_order["assembly_drawing_no"]
                and (manual["customer"] or "") != (current_order["customer"] or "")
            ):
                flash("组装订单不能更换为其他客户的产品", "error")
                return redirect(url_for("edit_order", order_id=order_id))
            if current_order["has_shipment_reference"] and (
                manual_id_value != current_order["manual_id"]
                or (manual["customer"] or "") != (current_order["customer"] or "")
            ):
                flash("已有发货记录的订单不能更换产品或客户", "error")
                return redirect(url_for("edit_order", order_id=order_id))
            shipped_total = int(current_order["shipped_total"] or 0)
            if quantity_value < shipped_total:
                flash(f"订单数量不能小于累计已发数量 {shipped_total}", "error")
                return redirect(url_for("edit_order", order_id=order_id))

            conn.execute(
                """
                UPDATE product_orders
                SET manual_id = ?, order_no = ?, ordered_at = ?, quantity = ?,
                    customer = ?, customer_email = ?, planned_ship_at = ?, material_stock_status = ?,
                    material_stock_assignee = ?, remark = ?, carton_status = ?, carton_assignee = ?,
                    recent_ship_status = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    manual_id_value,
                    order_no,
                    ordered_at,
                    quantity_value,
                    manual["customer"] or "",
                    customer_email,
                    planned_ship_at,
                    material_stock_status,
                    "",
                    remark,
                    carton_status,
                    "",
                    recent_ship_status,
                    now,
                    order_id,
                ),
            )
        flash("产品订单已更新", "success")
        return redirect(url_for("admin_orders"))

    with get_db() as conn:
        products = get_product_options(conn)

    return render_template(
        "order_edit.html",
        order=order,
        products=products,
    )


@app.route("/admin/orders/<int:order_id>/carton", methods=["GET", "POST"])
@permission_required("orders_manage")
def edit_order_carton(order_id):
    with get_db() as conn:
        order = conn.execute(
            """
            SELECT product_orders.*,
                   manuals.drawing_no,
                   manuals.product_name
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            WHERE product_orders.id = ?
            """,
            (order_id,),
        ).fetchone()
    if order is None:
        abort(404)

    if request.method == "POST":
        carton_status = request.form.get("carton_status", "").strip()
        now = datetime.utcnow().isoformat(timespec="seconds")
        with get_db() as conn:
            conn.execute(
                """
                UPDATE product_orders
                SET carton_status = ?, updated_at = ?
                WHERE id = ?
                """,
                (carton_status, now, order_id),
            )
        flash("纸箱情况已更新", "success")
        return redirect(url_for("admin_orders"))

    return render_template("order_carton_edit.html", order=order)


@app.route("/admin/orders/<int:order_id>/material", methods=["GET", "POST"])
@permission_required("orders_manage")
def edit_order_material(order_id):
    with get_db() as conn:
        order = conn.execute(
            """
            SELECT product_orders.*,
                   manuals.drawing_no,
                   manuals.product_name
            FROM product_orders
            JOIN manuals ON manuals.id = product_orders.manual_id
            WHERE product_orders.id = ?
            """,
            (order_id,),
        ).fetchone()
    if order is None:
        abort(404)

    if request.method == "POST":
        material_stock_status = request.form.get("material_stock_status", "").strip()
        now = datetime.utcnow().isoformat(timespec="seconds")
        with get_db() as conn:
            conn.execute(
                """
                UPDATE product_orders
                SET material_stock_status = ?, updated_at = ?
                WHERE id = ?
                """,
                (material_stock_status, now, order_id),
            )
        flash("材料库存情况已更新", "success")
        return redirect(url_for("admin_orders"))

    return render_template("order_material_edit.html", order=order)


@app.route("/admin/orders/<int:order_id>/material-status", methods=["POST"])
@permission_required("orders_manage")
def toggle_order_material_status(order_id):
    material_stock_status = "1" if request.form.get("checked") == "1" else ""
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        order = conn.execute(
            "SELECT id FROM product_orders WHERE id = ?",
            (order_id,),
        ).fetchone()
        if order is None:
            abort(404)
        conn.execute(
            """
            UPDATE product_orders
            SET material_stock_status = ?, updated_at = ?
            WHERE id = ?
            """,
            (material_stock_status, now, order_id),
        )
    return redirect(request.referrer or url_for("admin_orders"))


@app.route("/admin/orders/<int:order_id>/carton-status", methods=["POST"])
@permission_required("orders_manage")
def toggle_order_carton_status(order_id):
    carton_status = "1" if request.form.get("checked") == "1" else ""
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        order = conn.execute(
            "SELECT id FROM product_orders WHERE id = ?",
            (order_id,),
        ).fetchone()
        if order is None:
            abort(404)
        conn.execute(
            """
            UPDATE product_orders
            SET carton_status = ?, updated_at = ?
            WHERE id = ?
            """,
            (carton_status, now, order_id),
        )
    return redirect(request.referrer or url_for("admin_orders"))


@app.route("/admin/orders/<int:order_id>/recent-ship-status", methods=["POST"])
@permission_required("orders_manage")
def toggle_order_recent_ship_status(order_id):
    recent_ship_status = "1" if request.form.get("checked") == "1" else ""
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        order = conn.execute(
            "SELECT id FROM product_orders WHERE id = ?",
            (order_id,),
        ).fetchone()
        if order is None:
            abort(404)
        conn.execute(
            """
            UPDATE product_orders
            SET recent_ship_status = ?, updated_at = ?
            WHERE id = ?
            """,
            (recent_ship_status, now, order_id),
        )
    return redirect(request.referrer or url_for("admin_orders"))


@app.route("/admin/orders/<int:order_id>/delete", methods=["POST"])
@permission_required("orders_manage")
def delete_order(order_id):
    with get_db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        order = conn.execute(
            """
            SELECT product_orders.id,
                   (
                       product_orders.shipped_quantity > 0
                       OR EXISTS (
                           SELECT 1
                           FROM product_order_shipments
                           WHERE product_order_shipments.order_id = product_orders.id
                       )
                       OR EXISTS (
                           SELECT 1
                           FROM assembly_shipment_allocations
                           WHERE assembly_shipment_allocations.order_id = product_orders.id
                       )
                   ) AS has_shipment_reference,
                   EXISTS (
                       SELECT 1
                       FROM inventory_transactions
                       WHERE inventory_transactions.related_order_id = CAST(product_orders.id AS TEXT)
                   ) AS has_inventory_reference,
                   EXISTS (
                       SELECT 1
                       FROM shipment_plan_items
                       WHERE shipment_plan_items.order_id = product_orders.id
                   ) AS has_shipment_plan_reference
            FROM product_orders
            WHERE product_orders.id = ?
            """,
            (order_id,),
        ).fetchone()
        if order is None:
            abort(404)
        if order["has_shipment_plan_reference"]:
            flash("已有计划发货记录的订单不能删除", "error")
            return redirect(url_for("admin_orders"))
        if order["has_shipment_reference"]:
            flash("已有发货记录的订单不能删除", "error")
            return redirect(url_for("admin_orders"))
        if order["has_inventory_reference"]:
            flash("已有库存流水的订单不能删除", "error")
            return redirect(url_for("admin_orders"))
        conn.execute("DELETE FROM product_orders WHERE id = ?", (order_id,))
    flash("产品订单已删除", "success")
    return redirect(url_for("admin_orders"))


@app.route("/admin/upload", methods=["POST"])
@permission_required("product_create")
def upload_manual():
    uploads = selected_uploads()

    product_name = request.form.get("product_name", "").strip()
    drawing_no = request.form.get("drawing_no", "").strip()
    supplier = request.form.get("supplier", "").strip()
    customer = request.form.get("customer", "").strip()
    pack_quantity = request.form.get("pack_quantity", "").strip()
    pack_carton_size = request.form.get("pack_carton_size", "").strip()
    pack_weight = request.form.get("pack_weight", "").strip()
    sku = request.form.get("sku", "").strip()
    barcode = request.form.get("barcode", "").strip()
    default_location_id = parse_optional_int(request.form.get("default_location_id"))
    min_stock_text = request.form.get("min_stock", "0").strip()
    model = ""
    category = ""
    remark = request.form.get("remark", "").strip()
    description_html = request.form.get("description_html", "").strip()
    materials = posted_product_materials()
    inspection_requirements = posted_inspection_requirements()
    current_user = current_admin_username()

    if not all([drawing_no, product_name]):
        flash("产品图号、产品名称为必填项", "error")
        return redirect(url_for("admin_index"))
    try:
        min_stock = int(min_stock_text or 0)
    except ValueError:
        flash("最低库存必须为整数", "error")
        return redirect(url_for("admin_index"))
    if min_stock < 0:
        flash("最低库存不能小于 0", "error")
        return redirect(url_for("admin_index"))
    try:
        unit_price_minor, currency = posted_product_price()
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("admin_index"))
    try:
        saved_files = save_uploads(uploads)
    except ValueError:
        flash("不支持该文件类型", "error")
        return redirect(url_for("admin_index"))

    if saved_files:
        filename, original_filename, file_type = saved_files[0]
    else:
        filename, original_filename, file_type = "", "", "file"
    now = datetime.utcnow().isoformat(timespec="seconds")
    version = now
    with get_db() as conn:
        cursor = conn.execute(
            """
            INSERT INTO manuals (
                drawing_no, product_name, supplier, customer, pack_quantity, pack_carton_size, pack_weight,
                sku, barcode, default_location_id, min_stock,
                model, category, version, remark,
                description_html, filename, original_filename, file_type,
                unit_price_minor, currency,
                created_by, updated_by,
                created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                drawing_no,
                product_name,
                supplier,
                customer,
                pack_quantity,
                pack_carton_size,
                pack_weight,
                sku,
                barcode,
                default_location_id,
                min_stock,
                model,
                category,
                version,
                remark,
                description_html,
                filename,
                original_filename,
                file_type,
                unit_price_minor,
                currency,
                current_user,
                current_user,
                now,
                now,
            ),
        )
        ensure_customer_exists(conn, customer, now)
        manual_id = cursor.lastrowid
        ensure_product_inventory_code(conn, manual_id)
        save_product_materials(conn, manual_id, materials)
        save_manual_inspection_requirements(conn, manual_id, inspection_requirements)
        if saved_files:
            conn.executemany(
                """
                INSERT INTO manual_files (
                    manual_id, filename, original_filename, file_type, created_at
                )
                VALUES (?, ?, ?, ?, ?)
                """,
                [
                    (manual_id, item[0], item[1], item[2], now)
                    for item in saved_files
                ],
            )
    flash("产品资料上传成功", "success")
    return redirect(url_for("admin_index"))


@app.route("/admin/editor-images", methods=["POST"])
@login_required
def upload_editor_images():
    if not (user_has_permission("product_create") or user_has_permission("product_edit")):
        return jsonify({"error": "当前账号没有权限上传编辑图片"}), 403
    images = [
        file for file in request.files.getlist("images")
        if file and (file.filename or file.mimetype)
    ]
    if not images:
        return jsonify({"error": "未选择图片"}), 400

    saved_images = []
    try:
        for image in images:
            if not is_allowed_editor_image(image):
                raise ValueError(image.filename or "")
            filename, original_filename = save_editor_image(image)
            saved_images.append(
                {
                    "filename": filename,
                    "original_filename": original_filename,
                    "url": url_for("preview_manual", filename=filename),
                }
            )
    except ValueError:
        return jsonify({"error": "不支持该图片类型"}), 400

    return jsonify({"images": saved_images})


@app.route("/admin/<int:manual_id>/edit", methods=["GET", "POST"])
@permission_required("product_edit")
def edit_manual(manual_id):
    include_price = user_can_view_prices()
    with get_db() as conn:
        manual = fetch_manual_by_id(conn, manual_id, include_price=include_price)
    if manual is None:
        abort(404)

    if request.method == "POST":
        product_name = request.form.get("product_name", "").strip()
        drawing_no = request.form.get("drawing_no", "").strip()
        supplier = request.form.get("supplier", "").strip()
        customer = request.form.get("customer", "").strip()
        pack_quantity = request.form.get("pack_quantity", "").strip()
        pack_carton_size = request.form.get("pack_carton_size", "").strip()
        pack_weight = request.form.get("pack_weight", "").strip()
        sku = request.form.get("sku", "").strip()
        barcode = request.form.get("barcode", "").strip()
        default_location_id = parse_optional_int(request.form.get("default_location_id"))
        min_stock_text = request.form.get("min_stock", "0").strip()
        remark = request.form.get("remark", "").strip()
        assembly_drawing_numbers = request.form.getlist("assembly_drawing_no")
        assembly_quantities = request.form.getlist("assembly_quantity_per_set")
        submitted_assembly_components = [
            {
                "assembly_drawing_no": drawing,
                "quantity_per_set": quantity,
            }
            for drawing, quantity in zip_longest(
                assembly_drawing_numbers, assembly_quantities, fillvalue=""
            )
        ]
        current_user = current_admin_username()

        if not all([drawing_no, product_name]):
            flash("产品图号、产品名称为必填项", "error")
            return redirect(url_for("edit_manual", manual_id=manual_id))
        try:
            min_stock = int(min_stock_text or 0)
        except ValueError:
            flash("最低库存必须为整数", "error")
            return redirect(url_for("edit_manual", manual_id=manual_id))
        if min_stock < 0:
            flash("最低库存不能小于 0", "error")
            return redirect(url_for("edit_manual", manual_id=manual_id))
        if include_price:
            try:
                unit_price_minor, currency = posted_product_price(
                    manual["unit_price_minor"], manual["currency"]
                )
            except ValueError as error:
                flash(str(error), "error")
                return render_edit_manual_page(manual, submitted_assembly_components)
        try:
            assembly_components = normalize_component_rows(
                assembly_drawing_numbers, assembly_quantities
            )
        except ValueError as error:
            flash(str(error), "error")
            return render_edit_manual_page(manual, submitted_assembly_components)
        if assembly_components and not customer:
            flash("请先选择客户，再配置所属组装件", "error")
            return render_edit_manual_page(manual, submitted_assembly_components)
        if customer != (manual["customer"] or ""):
            with get_db() as conn:
                has_assembly_order = conn.execute(
                    """
                    SELECT 1
                    FROM product_orders
                    WHERE manual_id = ?
                      AND TRIM(assembly_drawing_no) != ''
                    LIMIT 1
                    """,
                    (manual_id,),
                ).fetchone()
            if has_assembly_order:
                flash("已有组装订单的产品不能直接更换客户", "error")
                return render_edit_manual_page(
                    manual, submitted_assembly_components
                )

        now = datetime.utcnow().isoformat(timespec="seconds")
        version = now
        with get_db() as conn:
            if include_price:
                conn.execute(
                    """
                    UPDATE manuals
                    SET drawing_no = ?, product_name = ?, supplier = ?, customer = ?,
                        pack_quantity = ?, pack_carton_size = ?, pack_weight = ?,
                        sku = ?, barcode = ?, default_location_id = ?, min_stock = ?,
                        version = ?, remark = ?, unit_price_minor = ?, currency = ?,
                        updated_by = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        drawing_no,
                        product_name,
                        supplier,
                        customer,
                        pack_quantity,
                        pack_carton_size,
                        pack_weight,
                        sku,
                        barcode,
                        default_location_id,
                        min_stock,
                        version,
                        remark,
                        unit_price_minor,
                        currency,
                        current_user,
                        now,
                        manual_id,
                    ),
                )
            else:
                conn.execute(
                    """
                    UPDATE manuals
                    SET drawing_no = ?, product_name = ?, supplier = ?, customer = ?,
                        pack_quantity = ?, pack_carton_size = ?, pack_weight = ?,
                        sku = ?, barcode = ?, default_location_id = ?, min_stock = ?,
                        version = ?, remark = ?, updated_by = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        drawing_no,
                        product_name,
                        supplier,
                        customer,
                        pack_quantity,
                        pack_carton_size,
                        pack_weight,
                        sku,
                        barcode,
                        default_location_id,
                        min_stock,
                        version,
                        remark,
                        current_user,
                        now,
                        manual_id,
                    ),
                )
            ensure_product_inventory_code(conn, manual_id)
            ensure_customer_exists(conn, customer, now)
            save_product_assembly_components(conn, manual_id, assembly_components, now)
            conn.execute(
                """
                UPDATE product_orders
                SET customer = ?, updated_at = ?
                WHERE manual_id = ?
                """,
                (
                    customer,
                    now,
                    manual_id,
                ),
            )
        flash("产品资料已更新", "success")
        return redirect(url_for("admin_index"))

    return render_edit_manual_page(manual)


def render_edit_manual_page(manual, assembly_components=None):
    with get_db() as conn:
        customers = get_customer_name_options(conn)
        if assembly_components is None:
            assembly_components = get_product_assembly_components(conn, manual["id"])
        inventory_locations = conn.execute(
            """
            SELECT *
            FROM warehouse_locations
            WHERE enabled = 1
            ORDER BY code COLLATE NOCASE ASC, id ASC
            """
        ).fetchall()
    return render_template(
        "edit.html",
        manual=manual,
        suppliers=get_suppliers(),
        customers=customers,
        assembly_components=assembly_components,
        inventory_locations=inventory_locations,
        supported_currencies=SUPPORTED_CURRENCIES,
    )


@app.route("/admin/<int:manual_id>/edit/technical", methods=["GET", "POST"])
@permission_required("product_edit")
def edit_manual_technical(manual_id):
    with get_db() as conn:
        manual = fetch_manual_by_id(conn, manual_id)
    if manual is None:
        abort(404)

    if request.method == "POST":
        description_html = request.form.get("description_html", "").strip()
        materials = posted_product_materials()
        inspection_requirements = posted_inspection_requirements()
        uploads = selected_uploads()
        saved_files = []
        if uploads:
            try:
                saved_files = save_uploads(uploads)
            except ValueError:
                flash("不支持该文件类型", "error")
                return redirect(url_for("edit_manual_technical", manual_id=manual_id))

        now = datetime.utcnow().isoformat(timespec="seconds")
        with get_db() as conn:
            conn.execute(
                """
                UPDATE manuals
                SET description_html = ?, version = ?, updated_by = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    description_html,
                    now,
                    current_admin_username(),
                    now,
                    manual_id,
                ),
            )
            save_product_materials(conn, manual_id, materials)
            save_manual_inspection_requirements(conn, manual_id, inspection_requirements)
            if saved_files:
                conn.executemany(
                    """
                    INSERT INTO manual_files (
                        manual_id, filename, original_filename, file_type, created_at
                    )
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    [
                        (manual_id, item[0], item[1], item[2], now)
                        for item in saved_files
                    ],
                )
                primary = get_manual_files(conn, manual_id)[0]
                conn.execute(
                    """
                    UPDATE manuals
                    SET filename = ?, original_filename = ?, file_type = ?
                    WHERE id = ?
                    """,
                    (
                        primary["filename"],
                        primary["original_filename"],
                        primary["file_type"],
                        manual_id,
                    ),
                )
        flash("产品技术资料已更新", "success")
        return redirect(url_for("admin_index"))

    return render_edit_manual_technical_page(manual)


def render_edit_manual_technical_page(manual):
    with get_db() as conn:
        files = get_manual_files(conn, manual["id"])
        materials = get_product_materials(conn, manual["id"])
        material_options = get_product_material_options(conn)
        inspection_requirements = get_manual_inspection_requirements(conn, manual["id"])
    return render_template(
        "edit_technical.html",
        manual=manual,
        files=files,
        materials=materials,
        material_options=material_options,
        inspection_requirements=inspection_requirements,
    )


@app.route("/admin/<int:manual_id>/files/<int:file_id>/delete", methods=["POST"])
@permission_required("product_edit")
def delete_manual_file(manual_id, file_id):
    with get_db() as conn:
        manual = fetch_manual_by_id(conn, manual_id)
        if manual is None:
            abort(404)
        manual_file = conn.execute(
            "SELECT * FROM manual_files WHERE id = ? AND manual_id = ?",
            (file_id, manual_id),
        ).fetchone()
        if manual_file is None:
            abort(404)
        file_count = conn.execute(
            "SELECT COUNT(*) AS c FROM manual_files WHERE manual_id = ?",
            (manual_id,),
        ).fetchone()["c"]
        if file_count <= 1:
            flash("至少需要保留一个文件", "error")
            return redirect(url_for("edit_manual_technical", manual_id=manual_id))

        conn.execute("DELETE FROM manual_files WHERE id = ?", (file_id,))
        remaining = get_manual_files(conn, manual_id)[0]
        conn.execute(
            """
            UPDATE manuals
            SET filename = ?, original_filename = ?, file_type = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                remaining["filename"],
                remaining["original_filename"],
                remaining["file_type"],
                datetime.utcnow().isoformat(timespec="seconds"),
                manual_id,
            ),
        )

    delete_upload_file(manual_file["filename"])
    flash("附件已删除", "success")
    return redirect(url_for("edit_manual_technical", manual_id=manual_id))


@app.route("/admin/<int:manual_id>/copy", methods=["POST"])
@permission_required("product_edit")
def copy_manual(manual_id):
    current_user = current_admin_username()
    now = datetime.utcnow().isoformat(timespec="seconds")
    with get_db() as conn:
        manual = fetch_manual_by_id(conn, manual_id)
        if manual is None:
            abort(404)
        files = get_manual_files(conn, manual_id)
        materials = get_product_materials(conn, manual_id)

        copied_files = [
            copy_upload_file(file["filename"], file["original_filename"])
            for file in files
        ]
        if copied_files:
            filename, original_filename, file_type = copied_files[0]
        else:
            filename, original_filename, file_type = "", "", "file"

        cursor = conn.execute(
            """
            INSERT INTO manuals (
                drawing_no, product_name, supplier, customer, pack_quantity, pack_carton_size, pack_weight,
                model, category, version, remark,
                description_html, filename, original_filename, file_type,
                created_by, updated_by,
                created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                manual["drawing_no"],
                f"{manual['product_name']} - 副本",
                manual["supplier"],
                manual["customer"],
                manual["pack_quantity"],
                manual["pack_carton_size"],
                manual["pack_weight"],
                manual["model"],
                manual["category"],
                now,
                manual["remark"],
                manual["description_html"],
                filename,
                original_filename,
                file_type,
                current_user,
                current_user,
                now,
                now,
            ),
        )
        ensure_customer_exists(conn, manual["customer"], now)
        copied_manual_id = cursor.lastrowid
        if copied_files:
            conn.executemany(
                """
                INSERT INTO manual_files (
                    manual_id, filename, original_filename, file_type, created_at
                )
                VALUES (?, ?, ?, ?, ?)
                """,
                [
                    (copied_manual_id, item[0], item[1], item[2], now)
                    for item in copied_files
                ],
            )
        if materials:
            save_product_materials(conn, copied_manual_id, [dict(row) for row in materials])

    flash("产品资料已复制", "success")
    return redirect(url_for("edit_manual", manual_id=copied_manual_id))


@app.route("/admin/<int:manual_id>/delete", methods=["POST"])
@permission_required("product_edit")
def delete_manual(manual_id):
    with get_db() as conn:
        manual = fetch_manual_by_id(conn, manual_id)
        if manual is None:
            abort(404)
        files = get_manual_files(conn, manual_id)
        if not files and manual["filename"]:
            files = [manual]
        conn.execute("DELETE FROM product_materials WHERE manual_id = ?", (manual_id,))
        conn.execute(
            "DELETE FROM product_assembly_components WHERE manual_id = ?",
            (manual_id,),
        )
        conn.execute("DELETE FROM manual_files WHERE manual_id = ?", (manual_id,))
        conn.execute("DELETE FROM manuals WHERE id = ?", (manual_id,))

    for manual_file in files:
        delete_upload_file(manual_file["filename"])
    flash("产品资料已删除", "success")
    return redirect(url_for("admin_index"))


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=8006)
